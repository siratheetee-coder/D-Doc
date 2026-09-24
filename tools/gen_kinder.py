# -*- coding: utf-8 -*-
"""ดึงตัวบ่งชี้พัฒนาการปฐมวัยจากไฟล์ .xlsm ต้นฉบับ -> app/services/kinder.py

ต้นฉบับ: บัญชีเรียกชื่อและสมุดบันทึกพัฒนาการเด็กนักเรียนระดับปฐมวัย (อบ.1/1, 1/2, 1/3)
ดึงมาเป็นข้อความตรง ๆ ไม่แก้ถ้อยคำ (เว้นแต่ช่องว่างซ้ำ)
"""
import re
import openpyxl

SRC = r"C:\Users\Lenovo_\Downloads\Pre-elementary-25062560\แบบบันทึกสำหรับครูอนุบาล"
FILES = [("อ.1", "แบบบันทึกอนุบาล 3 ขวบ.xlsm"),
         ("อ.2", "แบบบันทึกอนุบาล1.xlsm"),
         ("อ.3", "แบบบันทึกอนุบาล2.xlsm")]
SHEETS = [("phys", ["04-ด้านร่างกาย"]),
          ("emo", ["05-ด้านอารมณ์"]),
          ("soc", ["06-ด้านสังคม"]),
          ("intel", ["07-ด้านสติปัญญา-1", "08-ด้านสติปัญญา-2"])]


def clean(t):
    t = re.sub(r"\s+", " ", str(t)).strip()
    return t


def read_items(wb, names):
    """คืน [(ลำดับ, กลุ่ม, ข้อความ)] จากชีตรายงาน (D=ลำดับ E=กลุ่ม/ข้อความ)"""
    out = []
    for nm in names:
        ws = wb[nm]
        pending = None
        for r in range(1, ws.max_row + 1):
            num = ws.cell(row=r, column=4).value
            txt = ws.cell(row=r, column=5).value
            if txt is None or not clean(txt):
                continue
            t = clean(txt)
            if num not in (None, "") and str(num).strip().isdigit():
                pending = (int(num), t)          # แถวนี้คือ ลำดับ + ชื่อกลุ่ม
            elif pending and re.match(r"^\d+\s*\.", t):
                body = re.sub(r"^\d+\s*\.\s*", "", t)   # ตัดเลขข้อหน้าข้อความ
                out.append((pending[0], pending[1], body))
                pending = None
    return out


data = {}
for level, fn in FILES:
    wb = openpyxl.load_workbook(f"{SRC}\\{fn}", data_only=True)
    data[level] = {key: read_items(wb, names) for key, names in SHEETS}
    ws = wb["10-ผู้ปกครอง"]
    home = []
    for r in range(9, 25):
        n, t = ws.cell(row=r, column=4).value, ws.cell(row=r, column=5).value
        if t and str(n).strip().isdigit():
            home.append(clean(t))
    data[level]["home"] = home

# ---- ตรวจความถูกต้องก่อนเขียนไฟล์ ----
for level in data:
    assert len(data[level]["phys"]) == 10, (level, len(data[level]["phys"]))
    assert len(data[level]["emo"]) == 10
    assert len(data[level]["soc"]) == 10
    assert len(data[level]["intel"]) == 20
    assert len(data[level]["home"]) == 10
    for key in ("phys", "emo", "soc", "intel"):
        nums = [n for n, _, _ in data[level][key]]
        assert nums == list(range(1, len(nums) + 1)), (level, key, nums)
home_all = [data[lv]["home"] for lv in data]
assert home_all[0] == home_all[1] == home_all[2], "ข้อผู้ปกครองไม่เหมือนกันทุกชั้น"

LEVEL_META = {
    "อ.1": ("อบ.1/1", 3, "ชั้นอนุบาล 3 ขวบ"),
    "อ.2": ("อบ.1/2", 4, "ชั้นอนุบาลปีที่ 1"),
    "อ.3": ("อบ.1/3", 5, "ชั้นอนุบาลปีที่ 2"),
}

L = []
w = L.append
w('# -*- coding: utf-8 -*-')
w('"""')
w('kinder.py - ตัวบ่งชี้พัฒนาการเด็กปฐมวัย (สมุดรายงานประจำตัวนักเรียน / สมุดพกอนุบาล)')
w('')
w('ที่มา: ชุด "บัญชีเรียกชื่อและสมุดบันทึกพัฒนาการเด็กนักเรียนระดับปฐมวัย"')
w('       อบ.1/1-1/3 = สมุดรายงานประจำตัวนักเรียน (สมุดพก ส่งผู้ปกครอง)')
w('       อบ.2/1-2/3 = สมุดบันทึกของครูประจำชั้น (บัญชีเรียกชื่อ + บันทึกพัฒนาการทั้งห้อง)')
w('       (ไฟล์ Excel ต้นฉบับที่โรงเรียนใช้กันทั่วไป · ดึงข้อความมาตรง ๆ ไม่ได้แต่งเอง)')
w('       สร้างไฟล์นี้ด้วยสคริปต์ gen_kinder.py อย่าพิมพ์แก้มือ ถ้าจะแก้ให้แก้ที่ต้นฉบับแล้วดึงใหม่')
w('')
w('ชั้นเรียน: อ.1 = 3 ขวบ · อ.2 = 4 ขวบ · อ.3 = 5 ขวบ (ตามที่หน้าปกของแต่ละไฟล์ระบุอายุไว้)')
w('ประเมิน 2 ภาคเรียน · 3 = ปฏิบัติได้ (ดี) · 2 = ปฏิบัติได้บางครั้ง (ปานกลาง) · 1 = ควรเสริม (ไม่ชัดเจน)')
w('"""')
w('')
w('# ระดับชั้น -> (รหัสแบบ, อายุ, ชื่อชั้นบนปก)')
w('KINDER_LEVELS = {')
for lv, (form, age, title) in LEVEL_META.items():
    w(f'    "{lv}": {{"form": "{form}", "age": {age}, "title": "{title}"}},')
w('}')
w('')
w('# ด้านพัฒนาการ (คีย์, ชื่อเต็ม, ชื่อสั้นในตารางสรุป)')
w('DOMAINS = [')
w('    ("phys", "ด้านร่างกาย", "ร่างกาย"),')
w('    ("emo", "ด้านอารมณ์ - จิตใจ", "อารมณ์ - จิตใจ"),')
w('    ("soc", "ด้านสังคม", "สังคม"),')
w('    ("intel", "ด้านสติปัญญา", "สติปัญญา"),')
w(']')
w('')
w('# ระดับผลการประเมินรายข้อ (ค่าในฐานข้อมูล -> ชื่อคอลัมน์ในเอกสาร)')
w('RATINGS = [(3, "ปฏิบัติได้"), (2, "ปฏิบัติได้บางครั้ง"), (1, "ควรเสริม")]')
w('# ระดับคุณภาพสรุปรายด้าน')
w('QUALITY = [(3, "ดี", "ปฏิบัติได้คล่องแคล่ว"), (2, "ปานกลาง", "ปฏิบัติได้บางครั้ง"),')
w('           (1, "ควรเสริม", "ไม่ชัดเจน")]')
w('')
w('# หัวข้อความเห็นครูประจำชั้น (แยกตามด้าน · ครูเขียนบรรยายรายภาคเรียน)')
w('COMMENT_TOPICS = {')
w('    "phys": ["การเจริญเติบโต", "กล้ามเนื้อใหญ่", "กล้ามเนื้อเล็ก"],')
w('    "emo": ["สุขภาพจิต", "คุณธรรม จริยธรรม", "ศิลปะ ดนตรี การเคลื่อนไหว"],')
w('    "soc": ["การอยู่ร่วมสังคม", "การช่วยเหลือตนเอง",')
w('            "การอนุรักษ์สิ่งแวดล้อม ประเพณี วัฒนธรรมไทย"],')
w('    "intel": ["การใช้ประสาททั้งห้า", "การคิดฟังอ่านพูดเขียน",')
w('              "การจำแนกเปรียบเทียบจำนวน", "ตำแหน่ง ทิศทาง และเวลา"],')
w('}')
w('')
w('# พฤติกรรมของนักเรียนขณะอยู่ที่บ้าน (ผู้ปกครองเป็นผู้ประเมิน · เหมือนกันทุกชั้น)')
w('HOME_ITEMS = [')
for t in home_all[0]:
    w(f'    {t!r},')
w(']')
w('')
w('# ตัวบ่งชี้รายชั้น: ระดับชั้น -> ด้าน -> [(ลำดับ, กลุ่มพฤติกรรม, ข้อความ)]')
w('INDICATORS = {')
for lv in ("อ.1", "อ.2", "อ.3"):
    w(f'    "{lv}": {{')
    for key, _ in SHEETS:
        w(f'        "{key}": [')
        for n, grp, body in data[lv][key]:
            w(f'            ({n}, {grp!r}, {body!r}),')
        w('        ],')
    w('    },')
w('}')
w('')
w('')
w('def items_for(level, domain):')
w('    """ตัวบ่งชี้ของชั้นนี้ในด้านนี้ · ชั้นที่ไม่ใช่อนุบาลคืนลิสต์ว่าง"""')
w('    return INDICATORS.get((level or "").strip(), {}).get(domain, [])')
w('')
w('')
w('def is_kinder(level) -> bool:')
w('    """ชั้นนี้เป็นระดับปฐมวัยไหม"""')
w('    return (level or "").strip() in KINDER_LEVELS')
w('')
w('')
w('def code_of(domain, seq) -> str:')
w('    """คีย์ที่ใช้เก็บผลรายข้อในฐานข้อมูล เช่น phys:3"""')
w('    return f"{domain}:{seq}"')
w('')
w('')
w('def domain_average(values) -> int:')
w('    """สรุประดับคุณภาพรายด้านจากผลรายข้อ (ปัดเข้าใกล้สุด · ไม่มีข้อมูล = 0)"""')
w('    vals = [v for v in values if v]')
w('    if not vals:')
w('        return 0')
w('    avg = sum(vals) / len(vals)')
w('    return 3 if avg >= 2.5 else 2 if avg >= 1.5 else 1')
w('')

out = "app/services/kinder.py"
open(out, "w", encoding="utf-8").write("\n".join(L))
print("เขียน", out, "·", sum(len(data[lv][k]) for lv in data for k in ("phys", "emo", "soc", "intel")), "ตัวบ่งชี้")
