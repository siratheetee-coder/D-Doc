# -*- coding: utf-8 -*-
"""
demo_seed.py - สร้างโรงเรียนสมมติสำหรับถ่ายคลิปสาธิต (ข้อมูลครบ 6 งาน)

ใช้บนเครื่องตัวเองเท่านั้น (ไม่ใช่เซิร์ฟเวอร์จริง) - ข้อมูลทุกอย่างเป็นของสมมติ
ไม่ตรงกับบุคคล ร้านค้า หรือโรงเรียนจริง

    python -m app.services.demo_seed            # สร้างครั้งแรก (ถามรหัสผ่าน)
    python -m app.services.demo_seed --reset    # ลบโรงเรียนเดโมเดิมแล้วสร้างใหม่ให้เหมือนตั้งต้น

บัญชีที่ได้ (รหัสผ่านเดียวกันทั้ง 3 บัญชี):
    demoschool      ไอดีหลักของโรงเรียน (เจ้าหน้าที่/หัวหน้างาน เห็นทุกงาน)
    teacher1.demo   บัญชีครู (ส่งแผนการสอน/ใบลา/ขอไปราชการ)
    director.demo   บัญชี ผอ. (อนุมัติเอกสาร)

ความปลอดภัย:
  - แตะเฉพาะโรงเรียนที่ slug = "demo" เท่านั้น โรงเรียนอื่นไม่ถูกลบหรือแก้
  - ไม่ยอมรันถ้าโฟลเดอร์ข้อมูลอยู่ใต้ /opt/ (เซิร์ฟเวอร์จริง)
  - รหัสผ่านไม่ถูกเก็บไว้ในโค้ด: พิมพ์ตอนรัน หรือส่งผ่าน env DEMO_PASSWORD
"""
import json
import sys
from datetime import date, datetime, timedelta

DEMO_SLUG = "demo"
DEMO_SCHOOL = "โรงเรียนบ้านตัวอย่างวิทยา"
OWNER_USER = "demoschool"          # "demo" เฉย ๆ มักชนกับบัญชีทดสอบเดิม


# ------------------------------------------------------------------ helpers
def _d(y, m, d):
    return datetime(y, m, d)


def _weekdays(start: date, end: date):
    cur = start
    while cur <= end:
        if cur.weekday() < 5:
            yield cur
        cur += timedelta(days=1)


def _fake_docx(title: str) -> bytes:
    """ไฟล์ Word ตัวอย่าง (แผนการสอน) สร้างในหน่วยความจำ"""
    import io
    from docx import Document
    doc = Document()
    doc.add_heading(title, 1)
    doc.add_paragraph("แผนการจัดการเรียนรู้ (ตัวอย่างสำหรับสาธิตระบบ)")
    for h in ("จุดประสงค์การเรียนรู้", "สาระสำคัญ", "กิจกรรมการเรียนรู้", "การวัดและประเมินผล"):
        doc.add_heading(h, 2)
        doc.add_paragraph("........................................................................")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ accounts
def find_demo_tenant():
    from app.accounts import acc_session, Tenant
    db = acc_session()
    try:
        t = db.query(Tenant).filter_by(slug=DEMO_SLUG).first()
        return t.id if t else None
    finally:
        db.close()


def create_accounts(password: str) -> dict:
    from app.accounts import acc_session, provision_tenant, Tenant, Account, password_problem
    from app.modules import ALL_MODULES_CSV
    bad = password_problem(password, OWNER_USER)
    if bad:
        raise ValueError(bad)
    tid = provision_tenant(DEMO_SCHOOL, DEMO_SLUG, OWNER_USER, password, expiry_date=None,
                           max_users=10, must_change=False, plan="member")
    db = acc_session()
    try:
        t = db.get(Tenant, tid)
        t.teacher_code = DEMO_SLUG
        t.policy_accepted_at = datetime.now()
        t.last_active_at = datetime.now()
        db.commit()
    finally:
        db.close()
    return {"tenant_id": tid}


def create_person_accounts(tid: int, password: str, teacher_pid: int, director_pid: int) -> list:
    from app.accounts import acc_session, add_teacher_account, Account
    from app.modules import ALL_MODULES_CSV
    out = []
    r = add_teacher_account(tid, teacher_pid, "teacher1", password, display_name="ครูสมใจ ใจดี")
    if r.get("error"):
        raise ValueError(r["error"])
    out.append(r["username"])
    r = add_teacher_account(tid, director_pid, "director", password, display_name="ผอ.วิชัย รักเรียน")
    if r.get("error"):
        raise ValueError(r["error"])
    out.append(r["username"])
    db = acc_session()
    try:
        acc = db.query(Account).filter_by(username=r["username"]).first()
        acc.is_director = True
        acc.modules = ALL_MODULES_CSV
        db.commit()
    finally:
        db.close()
    return out


# ------------------------------------------------------------------ school data
PEOPLE = [
    # (ชื่อ, ตำแหน่ง, ประเภท, วิทยฐานะ)
    ("นายวิชัย รักเรียน", "ผู้อำนวยการโรงเรียน", "ผู้บริหาร", "ผู้อำนวยการชำนาญการพิเศษ"),
    ("นางสาวสมใจ ใจดี", "ครู", "ครู", "ครูชำนาญการ"),
    ("นายประเสริฐ มั่นคง", "ครู", "ครู", "ครูชำนาญการพิเศษ"),
    ("นางมาลี สุขสันต์", "ครู", "ครู", "ครูชำนาญการพิเศษ"),
    ("นายธนากร ศรีสว่าง", "ครู", "ครู", "ครูชำนาญการ"),
    ("นางสาวกนกพร แสงทอง", "ครู", "ครู", "ครู"),
    ("นายอนุชา พงษ์ไพร", "ครู", "ครู", "ครู"),
    ("นางสุภาพร บุญมา", "ครู", "ครู", "ครูชำนาญการ"),
    ("นางสาวปิยะนุช ทองดี", "ครูผู้ช่วย", "ครู", ""),
    ("นายชาญวิทย์ เก่งกาจ", "ครู", "ครู", "ครูชำนาญการ"),
    ("นางสาวรัตนา พึ่งบุญ", "เจ้าหน้าที่ธุรการ", "ธุรการ", ""),
    ("นายสมศักดิ์ ขยันงาน", "นักการภารโรง", "นักการ", ""),
]

VENDORS = [
    ("ร้านตัวอย่างเครื่องเขียน", "นางสาวดารุณี ค้าขาย", "0000000000011", "12 หมู่ 3 ตำบลตัวอย่าง อำเภอเมือง จังหวัดสมมติ 99999", "080-000-0001"),
    ("หจก.ตัวอย่างคอมพิวเตอร์", "นายเอกชัย ไอที", "0000000000028", "45 ถนนสมมติ ตำบลในเมือง อำเภอเมือง จังหวัดสมมติ 99999", "080-000-0002"),
    ("ร้านช่างแอร์ตัวอย่าง", "นายสมพงษ์ เย็นสบาย", "0000000000036", "7 หมู่ 1 ตำบลตัวอย่าง อำเภอเมือง จังหวัดสมมติ 99999", "080-000-0003"),
    ("ร้านตัวอย่างกีฬา", "นายวีระ แข็งแรง", "0000000000044", "88 ตลาดสมมติ อำเภอเมือง จังหวัดสมมติ 99999", "080-000-0004"),
    ("ร้านอาหารแม่ตัวอย่าง", "นางบุญเรือน อิ่มอร่อย", "0000000000052", "3 หมู่ 2 ตำบลตัวอย่าง อำเภอเมือง จังหวัดสมมติ 99999", "080-000-0005"),
]

FIRST_M = ["ด.ช.ภูมิ", "ด.ช.ธีรภัทร", "ด.ช.กันต์", "ด.ช.ปัณณวัฒน์", "ด.ช.ณัฐวุฒิ", "ด.ช.พีรพัฒน์",
           "ด.ช.ชยพล", "ด.ช.กฤษณะ", "ด.ช.อัครพล", "ด.ช.ศุภกร"]
FIRST_F = ["ด.ญ.ปุณยวีร์", "ด.ญ.ณิชา", "ด.ญ.พิมพ์ชนก", "ด.ญ.กัญญาณัฐ", "ด.ญ.ชนิดา", "ด.ญ.ธัญชนก",
           "ด.ญ.อรปรียา", "ด.ญ.สุพิชญา", "ด.ญ.เบญญาภา", "ด.ญ.ปาณิสรา"]
LAST = ["มีสุข", "ใจงาม", "ศรีทอง", "บุญเย็น", "แก้วใส", "ทองคำ", "รุ่งเรือง", "สายใจ", "พูนผล", "เพชรงาม"]


def seed_school(db, *, today: date) -> dict:
    """กรอกข้อมูลสมมติลงฐานข้อมูลของโรงเรียนเดโม คืน person id ที่ต้องใช้ผูกบัญชี"""
    from app import models as m
    from app.thai_utils import current_fiscal_year, current_academic_year
    from app.services.doc_number import suggest_doc_no, commit_doc_no
    from app.services.academic import grade_of
    from app.services.asset_utils import CATEGORY_LIFE

    now = datetime(today.year, today.month, today.day, 9, 0)
    fy = current_fiscal_year(now)
    ay = current_academic_year(now)
    ce = ay - 543                                        # ปี ค.ศ. ต้นปีการศึกษา

    def ago(days, hour=9):
        x = now - timedelta(days=days)
        return x.replace(hour=hour)

    def docno(kind, source, ref_id, subject, dt):
        no = suggest_doc_no(db, kind, fy)
        commit_doc_no(db, kind, fy, no, source=source, ref_id=ref_id, subject=subject, date=dt)
        return no

    # ---------------- ตั้งค่าโรงเรียน ----------------
    school = db.query(m.School).first() or m.School()
    school.name = DEMO_SCHOOL
    school.academic_year = ay
    school.address = "99 หมู่ 9 ตำบลตัวอย่าง อำเภอเมือง จังหวัดสมมติ 99999"
    school.district, school.province = "เมือง", "สมมติ"
    school.area_office = "สำนักงานเขตพื้นที่การศึกษาประถมศึกษาสมมติ เขต 1"
    school.director_name = PEOPLE[0][0]
    school.director_position = "ผู้อำนวยการโรงเรียน"
    school.officer_name = PEOPLE[4][0]
    school.head_officer_name = PEOPLE[2][0]
    school.finance_officer_name = PEOPLE[7][0]
    school.finance_head_name = PEOPLE[3][0]
    school.admin_officer_name = PEOPLE[10][0]
    school.academic_head_name = PEOPLE[3][0]
    school.doc_prefix = "ศธ 99999"
    db.add(school)

    persons = []
    for name, pos, ptype, rank in PEOPLE:
        p = m.Person(name=name, position=pos, person_type=ptype, rank=rank,
                     start_date=_d(2010 + len(persons) % 10, 5, 16), salary=0)
        db.add(p)
        persons.append(p)
    db.flush()
    director, teacher = persons[0], persons[1]
    teachers = persons[1:10]

    for dname in ("ฝ่ายบริหารงานวิชาการ", "ฝ่ายบริหารงบประมาณ", "ฝ่ายบริหารงานบุคคล", "ฝ่ายบริหารทั่วไป"):
        db.add(m.Department(name=dname))

    vendors = []
    for name, owner, tax, addr, phone in VENDORS:
        v = m.Vendor(name=name, owner_name=owner, tax_id=tax, address=addr, phone=phone,
                     bank_account="000-0-00000-0")
        db.add(v)
        vendors.append(v)
    db.flush()

    # ---------------- โครงการ ----------------
    projects = []
    for name, budget, resp in (("โครงการพัฒนาแหล่งเรียนรู้ในโรงเรียน", 60000, "ฝ่ายบริหารงานวิชาการ"),
                               ("โครงการส่งเสริมสุขภาพและกีฬานักเรียน", 25000, "ฝ่ายบริหารทั่วไป"),
                               ("โครงการพัฒนาระบบบริหารจัดการสำนักงาน", 40000, "ฝ่ายบริหารงบประมาณ")):
        pj = m.Project(name=name, budget=budget, plan_year=fy, responsible=resp)
        db.add(pj)
        projects.append(pj)
    db.flush()
    rpt = m.ProjectReport(
        project_id=projects[1].id, title="กิจกรรมกีฬาสีภายในโรงเรียน",
        date_start=ago(40), date_end=ago(38), location=DEMO_SCHOOL,
        responsible=PEOPLE[6][0], responsible_pos="ครู",
        principles="เพื่อส่งเสริมให้นักเรียนมีสุขภาพแข็งแรง รู้จักการทำงานเป็นทีม และมีน้ำใจนักกีฬา",
        objectives=json.dumps(["เพื่อส่งเสริมสุขภาพกายและใจของนักเรียน",
                               "เพื่อฝึกการทำงานเป็นทีมและความมีน้ำใจนักกีฬา"], ensure_ascii=False),
        target_qty=json.dumps(["นักเรียนทุกคน จำนวน 40 คน เข้าร่วมกิจกรรม"], ensure_ascii=False),
        target_qual=json.dumps(["นักเรียนร้อยละ 90 มีความพึงพอใจในระดับดีขึ้นไป"], ensure_ascii=False),
        budget_planned=15000, budget_used=12450, budget_note="เงินอุดหนุนรายหัว",
        suggestions="ควรจัดกิจกรรมต่อเนื่องทุกปี และเพิ่มชนิดกีฬาพื้นบ้าน")
    db.add(rpt)

    # ---------------- งานพัสดุ ----------------
    def proc(subject, ptype, items, vendor, status, *, case="normal", project=None, days_ago=30,
             inspect="single", dept="ฝ่ายบริหารงานวิชาการ", purpose=""):
        total = sum(q * p for _, q, _, p in items)
        pr = m.Procurement(fiscal_year=fy, subject=subject, proc_type=ptype, proc_case=case,
                           department=dept, purpose=purpose or f"เพื่อใช้ในการ{subject}",
                           project_id=project.id if project else None,
                           project_name=project.name if project else "",
                           total_amount=total, delivery_days=7, inspection_mode=inspect,
                           vendor_id=vendor.id if vendor else None, status=status,
                           request_date=ago(days_ago))
        for name, q, unit, price in items:
            pr.items.append(m.ProcurementItem(name=name, quantity=q, unit=unit, unit_price=price))
        db.add(pr)
        db.flush()
        spec = m.Committee(kind="spec", mode="single")
        spec.members.append(m.CommitteeMember(name=PEOPLE[4][0], position="ครู", role="ผู้กำหนดรายละเอียด", seq=1))
        pr.committees.append(spec)
        ins = m.Committee(kind="inspect", mode=inspect)
        if inspect == "single":
            ins.members.append(m.CommitteeMember(name=PEOPLE[5][0], position="ครู", role="ผู้ตรวจรับพัสดุ", seq=1))
        else:
            for i, (who, role) in enumerate(((PEOPLE[2][0], "ประธานกรรมการ"), (PEOPLE[5][0], "กรรมการ"),
                                             (PEOPLE[6][0], "กรรมการและเลขานุการ")), 1):
                ins.members.append(m.CommitteeMember(name=who, position="ครู", role=role, seq=i))
        pr.committees.append(ins)
        if status != "ร่าง":
            pr.memo_no = docno("memo", "procurement", pr.id, subject, pr.request_date)
            pr.command_no = docno("command", "procurement", pr.id, subject, pr.request_date)
            pr.command_date = pr.request_date
            pr.order_no = docno("purchase_order" if ptype == "ซื้อ" else "hire_order",
                                "procurement", pr.id, subject, ago(days_ago - 2))
            pr.quotation_date = ago(days_ago - 1)
            pr.order_date = ago(days_ago - 2)
            pr.delivery_due_date = ago(days_ago - 9)
        if status in ("ตรวจรับแล้ว", "เบิกจ่ายแล้ว"):
            pr.delivery_date = ago(days_ago - 7)
            pr.inspect_date = ago(days_ago - 7)
            pr.delivery_note_no = "0012"
            pr.inspect_memo_no = docno("memo", "procurement", pr.id, f"ตรวจรับ {subject}", pr.inspect_date)
        return pr

    p_office = proc("ซื้อวัสดุสำนักงาน", "ซื้อ",
                    [("กระดาษ A4 70 แกรม", 20, "รีม", 125), ("หมึกเครื่องพิมพ์", 4, "กล่อง", 390),
                     ("แฟ้มเอกสาร", 30, "เล่ม", 45), ("ปากกาลูกลื่น", 5, "กล่อง", 120)],
                    vendors[0], "เบิกจ่ายแล้ว", project=projects[2], days_ago=60,
                    dept="ฝ่ายบริหารงบประมาณ")
    p_air = proc("จ้างซ่อมเครื่องปรับอากาศห้องสมุด", "จ้าง",
                 [("ค่าซ่อมและเติมน้ำยาเครื่องปรับอากาศ 18,000 BTU", 2, "เครื่อง", 2800)],
                 vendors[2], "ตรวจรับแล้ว", project=projects[0], days_ago=25, dept="ฝ่ายบริหารทั่วไป")
    p_com = proc("ซื้อครุภัณฑ์คอมพิวเตอร์ห้องเรียนรู้", "ซื้อ",
                 [("เครื่องคอมพิวเตอร์ สำหรับงานประมวลผล", 3, "เครื่อง", 22000),
                  ("เครื่องพิมพ์เลเซอร์", 1, "เครื่อง", 8900)],
                 vendors[1], "อนุมัติ", project=projects[0], days_ago=12, inspect="committee")
    proc("ซื้อน้ำดื่มสำหรับกิจกรรมวันวิทยาศาสตร์", "ซื้อ",
         [("น้ำดื่ม ขนาด 600 มล.", 10, "แพ็ค", 65)], vendors[0], "ร่าง", case="w119t1", days_ago=3)
    proc("ซื้ออุปกรณ์กีฬา", "ซื้อ",
         [("ลูกฟุตบอล หนังเย็บ เบอร์ 4", 5, "ลูก", 650), ("ตาข่ายวอลเลย์บอล", 1, "ผืน", 1200)],
         vendors[3], "ร่าง", project=projects[1], days_ago=1, dept="ฝ่ายบริหารทั่วไป")

    for name, unit, cat, price in (("กระดาษ A4 70 แกรม", "รีม", "วัสดุสำนักงาน", 125),
                                   ("หมึกเครื่องพิมพ์", "กล่อง", "วัสดุสำนักงาน", 390),
                                   ("แฟ้มเอกสาร", "เล่ม", "วัสดุสำนักงาน", 45),
                                   ("ลูกฟุตบอล หนังเย็บ เบอร์ 4", "ลูก", "วัสดุกีฬา", 650),
                                   ("น้ำดื่ม ขนาด 600 มล.", "แพ็ค", "วัสดุงานบ้านงานครัว", 65)):
        db.add(m.ItemCatalog(name=name, unit=unit, category=cat, unit_price=price))

    # บัญชีวัสดุ: รับเข้าจากเรื่องจัดซื้อ + เบิกจ่ายบางส่วน
    for it in p_office.items:
        mat = m.MaterialItem(name=it.name, unit=it.unit, min_stock=5 if it.unit == "รีม" else 0)
        mat.txns.append(m.MaterialTxn(date=p_office.inspect_date, kind="in", qty=it.quantity,
                                      unit_price=it.unit_price, ref=f"ใบสั่งซื้อ {p_office.order_no}"))
        db.add(mat)
    db.flush()
    paper = db.query(m.MaterialItem).filter_by(name="กระดาษ A4 70 แกรม").first()
    req = m.Requisition(req_no=f"1/{fy}", date=ago(40), requester=PEOPLE[5][0],
                        department="ฝ่ายบริหารงานวิชาการ", purpose="ใช้จัดทำเอกสารประกอบการสอน", status="จ่ายแล้ว")
    req.items.append(m.RequisitionItem(material_id=paper.id, name=paper.name, unit="รีม", qty=6))
    db.add(req)
    db.flush()
    paper.txns.append(m.MaterialTxn(date=ago(40), kind="out", qty=6, ref=req.req_no, requisition_id=req.id))

    # ครุภัณฑ์
    asset_rows = [
        ("7440-001-0001/2566", "เครื่องคอมพิวเตอร์ตั้งโต๊ะ", "ครุภัณฑ์คอมพิวเตอร์", 2023, 21500, "ห้องธุรการ", "Acer Veriton"),
        ("7440-001-0002/2566", "เครื่องคอมพิวเตอร์ตั้งโต๊ะ", "ครุภัณฑ์คอมพิวเตอร์", 2023, 21500, "ห้องคอมพิวเตอร์", "Acer Veriton"),
        ("7430-002-0001/2565", "เครื่องพิมพ์เลเซอร์", "ครุภัณฑ์คอมพิวเตอร์", 2022, 7900, "ห้องธุรการ", "Brother HL"),
        ("4120-001-0001/2564", "เครื่องปรับอากาศ 18,000 BTU", "ครุภัณฑ์ไฟฟ้าและวิทยุ", 2021, 24500, "ห้องสมุด", "Daikin"),
        ("4120-001-0002/2564", "เครื่องปรับอากาศ 18,000 BTU", "ครุภัณฑ์ไฟฟ้าและวิทยุ", 2021, 24500, "ห้องสมุด", "Daikin"),
        ("5810-003-0001/2567", "เครื่องฉายภาพโปรเจคเตอร์", "ครุภัณฑ์โฆษณาและเผยแพร่", 2024, 18900, "ห้องประชุม", "Epson"),
        ("7110-006-0001/2563", "โต๊ะทำงานครู พร้อมเก้าอี้", "ครุภัณฑ์สำนักงาน", 2020, 5500, "ห้องพักครู", ""),
        ("7730-010-0001/2567", "ชุดเครื่องเสียงห้องประชุม", "ครุภัณฑ์ไฟฟ้าและวิทยุ", 2024, 32000, "ห้องประชุม", "Yamaha"),
        ("7520-013-0001/2562", "ตู้เหล็กเก็บเอกสาร 2 บาน", "ครุภัณฑ์สำนักงาน", 2019, 4900, "ห้องธุรการ", ""),
        ("6675-001-0001/2560", "ตู้เย็น 2 ประตู", "ครุภัณฑ์งานบ้านงานครัว (เครื่องจักรกล)", 2017, 12900, "โรงอาหาร", "Sharp"),
    ]
    for code, name, cat, year, cost, loc, brand in asset_rows:
        db.add(m.Asset(asset_code=code, name=name, category=cat, acquired_date=_d(year, 6, 15),
                       cost=cost, useful_life=CATEGORY_LIFE.get(cat, 5), location=loc, brand_model=brand,
                       vendor_name=vendors[1].name, funding_source="เงินอุดหนุน"))
        db.add(m.AssetNumberUsed(code=code))
    db.add(m.AssetNumberSeries(prefix="7440-001", digits=4, reset_yearly=True, append_year=True))

    db.add(m.Contract(fiscal_year=fy, contract_no=p_air.order_no, ctype="ใบสั่งจ้าง", party=vendors[2].name,
                      subject=p_air.subject, amount=p_air.total_amount, sign_date=p_air.order_date,
                      start_date=p_air.order_date, end_date=p_air.delivery_due_date,
                      warranty_end=p_air.inspect_date + timedelta(days=90),
                      status="ตรวจรับแล้ว", source="procurement", ref_id=p_air.id))
    db.add(m.Contract(fiscal_year=fy, contract_no=p_com.order_no, ctype="ใบสั่งซื้อ", party=vendors[1].name,
                      subject=p_com.subject, amount=p_com.total_amount, sign_date=p_com.order_date,
                      start_date=p_com.order_date, end_date=now + timedelta(days=5),
                      status="ระหว่างดำเนินการ", source="procurement", ref_id=p_com.id))
    for i, (name, budget, period) in enumerate((("ซื้อวัสดุสำนักงาน", 15000, "ตุลาคม"),
                                                ("ซื้อครุภัณฑ์คอมพิวเตอร์", 80000, "มกราคม"),
                                                ("จ้างซ่อมแซมครุภัณฑ์", 20000, "ตลอดปี"),
                                                ("ซื้อหนังสือเรียน", 70000, "มีนาคม")), 1):
        db.add(m.ProcurementPlan(fiscal_year=fy, seq=i, name=name, budget=budget,
                                 expected_period=f"{period} {fy - 1 if period in ('ตุลาคม',) else fy}"))

    # หนังสือเรียน
    books = [("ป.1", "ภาษาไทย", "หนังสือเรียนภาษาไทย ป.1", "สำนักพิมพ์ตัวอย่าง", 68, 20),
             ("ป.1", "คณิตศาสตร์", "หนังสือเรียนคณิตศาสตร์ ป.1", "สำนักพิมพ์ตัวอย่าง", 72, 20),
             ("ป.1", "วิทยาศาสตร์", "หนังสือเรียนวิทยาศาสตร์ ป.1", "สำนักพิมพ์สมมติ", 65, 20),
             ("ป.6", "ภาษาไทย", "หนังสือเรียนภาษาไทย ป.6", "สำนักพิมพ์ตัวอย่าง", 85, 20),
             ("ป.6", "คณิตศาสตร์", "หนังสือเรียนคณิตศาสตร์ ป.6", "สำนักพิมพ์ตัวอย่าง", 88, 20),
             ("ป.6", "ภาษาอังกฤษ", "หนังสือเรียนภาษาอังกฤษ ป.6", "สำนักพิมพ์สมมติ", 95, 20)]
    sel = []
    for lvl, subj, title, pub, price, qty in books:
        db.add(m.TextBook(year=ay, level=lvl, subject=subj, title=title, publisher=pub,
                          unit_price=price, qty_received=qty))
        sel.append({"key": f"demo{len(sel)}", "title": title, "level": lvl, "subject": subj,
                    "publisher": pub, "source_id": "", "publication": "", "price": price,
                    "qty": qty, "selected": True})
    db.add(m.TextbookPurchase(year=ay + 1, fiscal_year=fy, selection_items=json.dumps(sel, ensure_ascii=False),
                              total_budget=sum(b[4] * b[5] for b in books), delivery_days=15,
                              delivery_place=DEMO_SCHOOL, academic_head=PEOPLE[3][0],
                              vendor_name="ร้านตัวอย่างเครื่องเขียน",
                              purpose="เพื่อจัดหาหนังสือเรียนให้นักเรียนครบทุกคนก่อนเปิดภาคเรียน",
                              meet_place="ห้องประชุมโรงเรียน", meet_time="09.00 - 12.00 น."))

    # ---------------- งานการเงิน ----------------
    acc_sub = m.FinanceAccount(name="เงินอุดหนุนรายหัว", opening_balance=185000, deposit_type="bank",
                               fund_type="เงินนอกงบประมาณ")
    acc_income = m.FinanceAccount(name="เงินรายได้สถานศึกษา", opening_balance=12500, deposit_type="bank",
                                  fund_type="เงินนอกงบประมาณ")
    acc_rev = m.FinanceAccount(name="เงินรายได้แผ่นดิน (ดอกเบี้ย)", opening_balance=0, deposit_type="cash",
                               fund_type="เงินรายได้แผ่นดิน")
    db.add_all([acc_sub, acc_income, acc_rev])
    db.flush()
    item_teach = m.AccountItem(account_id=acc_sub.id, fiscal_year=fy, name="ค่าจัดการเรียนการสอน", budget=120000)
    item_book = m.AccountItem(account_id=acc_sub.id, fiscal_year=fy, name="ค่าหนังสือเรียน", budget=70000)
    db.add_all([item_teach, item_book])
    db.flush()

    txns = [
        (acc_sub, item_teach, 150, "in", 95000, "รับเงินอุดหนุนรายหัว ภาคเรียนที่ 2", "สพป.สมมติ เขต 1"),
        (acc_sub, item_book, 140, "in", 70000, "รับเงินอุดหนุนค่าหนังสือเรียน", "สพป.สมมติ เขต 1"),
        (acc_income, None, 120, "in", 3500, "รับเงินบริจาคจากผู้ปกครอง", "ใบเสร็จ 001"),
        (acc_rev, None, 30, "in", 212.45, "ดอกเบี้ยเงินฝากธนาคาร", "สมุดบัญชี"),
        (acc_income, None, 20, "out", 1200, "ค่าวัสดุตกแต่งห้องเรียน", "ใบเสร็จร้านค้า"),
    ]
    for acc, item, days, kind, amt, cat, ref in txns:
        db.add(m.FinanceTxn(account_id=acc.id, item_id=item.id if item else None, fiscal_year=fy,
                            date=ago(days), kind=kind, amount=amt, category=cat, ref=ref))
    db.flush()

    # ขอเบิกจ่ายจากเรื่องจัดซื้อวัสดุสำนักงาน -> จ่ายแล้ว
    dm = m.DisburseMemo(fiscal_year=fy, date=ago(50), subject=f"ขออนุมัติเบิกจ่ายเงิน {p_office.subject}",
                        payee=vendors[0].name, amount=p_office.total_amount, wht=0, proc_kind="จัดซื้อ",
                        budget_source="เงินอุดหนุนรายหัว", account_id=acc_sub.id, item_id=item_teach.id,
                        procurement_id=p_office.id, project_id=projects[2].id, status="จ่ายแล้ว")
    db.add(dm)
    db.flush()
    dm.memo_no = docno("memo", "finance", dm.id, dm.subject, dm.date)
    db.add(m.FinanceTxn(account_id=acc_sub.id, item_id=item_teach.id, fiscal_year=fy, date=dm.date,
                        kind="out", amount=dm.amount, category=p_office.subject, ref=dm.memo_no,
                        disburse_id=dm.id))
    db.add(m.CheckPayment(fiscal_year=fy, date=dm.date, pay_method="โอน", check_no="KTB-000123",
                          bank="ธนาคารตัวอย่าง สาขาเมือง", payee=vendors[0].name, amount=dm.amount,
                          purpose=p_office.subject, account_id=acc_sub.id, cleared=True))
    db.add(m.Receipt(fiscal_year=fy, receipt_no=f"1/{fy}", date=ago(120), kind="รับ",
                     party="ผู้ปกครองนักเรียน (ตัวอย่าง)", amount=3500, account_id=acc_income.id))
    # ขอเบิกจ่ายที่ยังรออนุมัติ (ไว้กดโชว์)
    dm2 = m.DisburseMemo(fiscal_year=fy, date=ago(2), subject=f"ขออนุมัติเบิกจ่ายเงิน {p_air.subject}",
                         payee=vendors[2].name, amount=p_air.total_amount, proc_kind="จัดจ้าง",
                         budget_source="เงินอุดหนุนรายหัว", account_id=acc_sub.id, item_id=item_teach.id,
                         procurement_id=p_air.id, project_id=projects[0].id, status="ร่าง")
    db.add(dm2)
    db.flush()
    dm2.memo_no = docno("memo", "finance", dm2.id, dm2.subject, dm2.date)

    loan = m.MoneyLoan(fiscal_year=fy, contract_no=f"1/{fy}", date=ago(10), receive_date=ago(9),
                       due_date=ago(9) + timedelta(days=15), borrower=PEOPLE[6][0], position="ครู",
                       submit_to=f"ผู้อำนวยการ{DEMO_SCHOOL}", fund_from="เงินอุดหนุนรายหัว",
                       purpose="พานักเรียนไปแข่งขันงานศิลปหัตถกรรมนักเรียน ระดับเขตพื้นที่",
                       items=json.dumps([{"name": "ค่าอาหารนักเรียนและครู", "amount": 2400},
                                         {"name": "ค่าเช่ารถ", "amount": 3000}], ensure_ascii=False),
                       amount=5400, within_days=15, account_id=acc_sub.id)
    db.add(loan)

    # ---------------- งานธุรการ ----------------
    incoming = [
        ("ศธ 99999/ว 1520", "สพป.สมมติ เขต 1", "การจัดงานศิลปหัตถกรรมนักเรียน ระดับเขตพื้นที่", PEOPLE[3][0]),
        ("ศธ 99999/ว 1488", "สพป.สมมติ เขต 1", "แจ้งจัดสรรงบประมาณเงินอุดหนุน", PEOPLE[7][0]),
        ("สส 0001/123", "สำนักงานสาธารณสุขอำเภอเมือง", "ขอความร่วมมือตรวจสุขภาพนักเรียน", PEOPLE[5][0]),
        ("ศธ 99999/ว 1410", "สพป.สมมติ เขต 1", "การประเมินผลการปฏิบัติงานครู", director.name),
        ("อบต 0002/88", "องค์การบริหารส่วนตำบลตัวอย่าง", "แจ้งโอนเงินอุดหนุนอาหารกลางวัน", PEOPLE[7][0]),
        ("ศธ 99999/ว 1377", "สพป.สมมติ เขต 1", "การอบรมเชิงปฏิบัติการครูคณิตศาสตร์", PEOPLE[2][0]),
        ("ตช 0003/45", "สถานีตำรวจภูธรตัวอย่าง", "โครงการตำรวจประสานโรงเรียน", director.name),
        ("ศธ 99999/ว 1350", "สพป.สมมติ เขต 1", "สำรวจข้อมูลนักเรียน 10 มิถุนายน", PEOPLE[10][0]),
    ]
    for i, (lno, org, subj, to) in enumerate(incoming):
        dt = ago(3 + i * 6)
        db.add(m.IncomingLetter(fiscal_year=fy, recv_no=len(incoming) - i, recv_date=dt, letter_no=lno,
                                letter_date=dt - timedelta(days=2), from_org=org, to_person=to, subject=subj,
                                action_note="มอบผู้รับผิดชอบดำเนินการ"))
    for i, (to, subj) in enumerate((("สพป.สมมติ เขต 1", "ส่งรายชื่อนักเรียนเข้าแข่งขันงานศิลปหัตถกรรม"),
                                    ("องค์การบริหารส่วนตำบลตัวอย่าง", "ขอบคุณการสนับสนุนเงินอุดหนุนอาหารกลางวัน"),
                                    ("สำนักงานสาธารณสุขอำเภอเมือง", "ยืนยันวันตรวจสุขภาพนักเรียน"))):
        dt = ago(4 + i * 9)
        lt = m.OfficialLetter(fiscal_year=fy, date=dt, subject=subj, to=f"ผู้อำนวยการ{to}" if "สพป" in to else to,
                              body="ด้วยโรงเรียนบ้านตัวอย่างวิทยา มีความประสงค์ ...\n\nจึงเรียนมาเพื่อโปรดทราบ",
                              signer_name=director.name, signer_position="ผู้อำนวยการโรงเรียนบ้านตัวอย่างวิทยา")
        db.add(lt)
        db.flush()
        lt.doc_no = docno("outgoing", "admin", lt.id, subj, dt)
        db.add(m.OutgoingLetter(fiscal_year=fy, send_no=lt.doc_no, date=dt, to_org=to, subject=subj))
    for i, subj in enumerate(("ขออนุญาตใช้ห้องประชุมจัดอบรมผู้ปกครอง", "รายงานผลการตรวจสุขภาพนักเรียน")):
        dt = ago(6 + i * 11)
        memo = m.OfficeMemo(fiscal_year=fy, date=dt, from_dept="ฝ่ายบริหารทั่วไป", to_person=f"ผู้อำนวยการ{DEMO_SCHOOL}",
                            subject=subj, body="ด้วย ...\n\nจึงเรียนมาเพื่อโปรดพิจารณา",
                            signer_name=PEOPLE[5][0], signer_position="ครู")
        db.add(memo)
        db.flush()
        memo.memo_no = docno("memo", "admin", memo.id, subj, dt)
    for i, subj in enumerate(("แต่งตั้งคณะกรรมการดำเนินงานกีฬาสีภายใน",
                              "แต่งตั้งครูเวรรักษาการณ์ประจำเดือน",
                              "แต่งตั้งคณะกรรมการประเมินผลการเรียน")):
        dt = ago(8 + i * 15)
        o = m.SchoolOrder(fiscal_year=fy, date=dt, subject=subj,
                          body="เพื่อให้การดำเนินงานเป็นไปด้วยความเรียบร้อย จึงแต่งตั้งบุคลากรดังต่อไปนี้ ...")
        db.add(o)
        db.flush()
        o.order_no = docno("command", "admin", o.id, subj, dt)

    # ---------------- งานบุคคล ----------------
    for lt_, d_ in (("sick", 60), ("personal", 45), ("vacation", 10)):
        db.add(m.LeaveEntitlement(year=ay, leave_type=lt_, days=d_))
    db.add(m.LeaveRecord(person_id=persons[2].id, year=ay, leave_type="sick", start_date=ago(35),
                         end_date=ago(34), days=2, reason="ไข้หวัดใหญ่"))
    db.add(m.LeaveRecord(person_id=persons[5].id, year=ay, leave_type="personal", start_date=ago(20),
                         end_date=ago(20), days=1, reason="ธุระส่วนตัว"))
    db.add(m.TravelRecord(person_id=persons[3].id, year=ay, subject="ประชุมผู้บริหารสถานศึกษา",
                          place="สพป.สมมติ เขต 1", start_date=ago(15), end_date=ago(15), days=1, budget=0))
    db.add(m.Decoration(person_id=director.id, name="ตริตาภรณ์ช้างเผือก (ต.ช.)", year=ay - 3))
    # คำขอที่ค้างอยู่ในสายอนุมัติ (ไว้โชว์ ผอ. กดอนุมัติ)
    db.add(m.LeaveRequest(person_id=teacher.id, leave_type="ลากิจ", start_date=(now + timedelta(days=5)).date(),
                          end_date=(now + timedelta(days=5)).date(), days=1, reason="พาบุตรไปพบแพทย์",
                          contact="080-000-0099", work_group="กลุ่มสาระภาษาไทย", status="personnel",
                          personnel_by=persons[3].id, personnel_at=ago(1), personnel_comment="ตรวจสอบวันลาแล้ว"))
    db.add(m.TravelRequest(person_id=teacher.id, subject="อบรมการจัดการเรียนรู้เชิงรุก", place="สพป.สมมติ เขต 1",
                           start_date=(now + timedelta(days=12)).date(), end_date=(now + timedelta(days=13)).date(),
                           days=2, budget=1200, budget_words="หนึ่งพันสองร้อยบาทถ้วน", purpose_type="อบรม",
                           doc_ref="ศธ 99999/ว 1377", reimburse="yes", cost_types="พาหนะ,เบี้ยเลี้ยง",
                           status="pending"))

    # ---------------- งานวิชาการ ----------------
    t1s, t1e = _d(ce, 5, 16), _d(ce, 10, 10)
    t2s, t2e = _d(ce, 11, 1), _d(ce + 1, 3, 31)
    db.add(m.AcadYearSetting(year=ay, t1_start=t1s, t1_end=t1e, t2_start=t2s, t2_end=t2e))
    by_month = {}
    for d in _weekdays(t1s.date(), min(t1e.date(), today)):
        by_month.setdefault(d.month, []).append(d.day)
    for mon, days in by_month.items():
        db.add(m.AcadCalendar(year=ay, month=mon, days_csv=",".join(map(str, days))))

    classes = []
    students_all = []
    for lvl, homeroom in (("ป.1", persons[1]), ("ป.6", persons[2])):
        k = m.AcadClass(year=ay, level=lvl, room="1", homeroom_id=homeroom.id)
        db.add(k)
        db.flush()
        for i in range(20):
            first = (FIRST_M if i % 2 == 0 else FIRST_F)[i // 2]
            sname = f"{first} {LAST[(i * 3 + len(classes)) % len(LAST)]}"
            sno = f"{(1 if lvl == 'ป.1' else 2)}{i + 1:03d}"
            st = m.Student(student_no=sno, name=sname, sex="M" if i % 2 == 0 else "F", level=lvl, room="1",
                           birthdate=_d(ce - (6 if lvl == "ป.1" else 11), 1 + i % 12, 1 + i % 27),
                           nationality="ไทย", enroll_date=_d(ce - (0 if lvl == "ป.1" else 5), 5, 16))
            db.add(st)
            db.flush()
            db.add(m.StudentMeasure(student_id=st.id, year=ay, term=1, date=_d(ce, 6, 10),
                                    weight=(21 if lvl == "ป.1" else 38) + (i % 7), height=(117 if lvl == "ป.1" else 145) + (i % 9)))
            k.students.append(m.AcadStudent(student_id=st.id, seq=i + 1, student_no=sno, name=sname, sex=st.sex))
            students_all.append(st)
        classes.append(k)
    db.flush()

    subj_def = [("ท", "ภาษาไทย", "ภาษาไทย", 200), ("ค", "คณิตศาสตร์", "คณิตศาสตร์", 200),
                ("ว", "วิทยาศาสตร์และเทคโนโลยี", "วิทยาศาสตร์และเทคโนโลยี", 120),
                ("ส", "สังคมศึกษา ศาสนา และวัฒนธรรม", "สังคมศึกษา ศาสนา และวัฒนธรรม", 80),
                ("อ", "ภาษาอังกฤษ", "ภาษาต่างประเทศ", 40)]
    periods = []
    for seq, (name, tl, brk) in enumerate((("คาบ 1", "08:30-09:30", False), ("คาบ 2", "09:30-10:30", False),
                                           ("คาบ 3", "10:30-11:30", False), ("พักกลางวัน", "11:30-12:30", True),
                                           ("คาบ 4", "12:30-13:30", False), ("คาบ 5", "13:30-14:30", False)), 1):
        pe = m.AcadPeriod(year=ay, seq=seq, name=name, time_label=tl, is_break=brk)
        db.add(pe)
        periods.append(pe)
    db.flush()
    teach_periods = [p for p in periods if not p.is_break]

    for k in classes:
        g = k.level[-1]
        subjects = []
        for idx, (pre, name, group, hours) in enumerate(subj_def):
            sj = m.AcadSubject(year=ay, level=k.level, code=f"{pre}1{g}101", name=name, learn_group=group,
                               kind="พื้นฐาน", hours=hours, mid_max=70, final_max=30, term=0, seq=idx)
            db.add(sj)
            db.flush()
            tch = teacher if (k.level == "ป.1" and idx in (0, 4)) else teachers[(idx + (0 if k.level == "ป.1" else 3)) % len(teachers)]
            db.add(m.AcadTeaching(class_id=k.id, subject_id=sj.id, teacher_id=tch.id))
            subjects.append(sj)
        for act_i, act in enumerate(("แนะแนว", "ลูกเสือ-เนตรนารี", "ชุมนุม", "กิจกรรมเพื่อสังคมและสาธารณประโยชน์")):
            db.add(m.AcadActivity(year=ay, level=k.level, code=f"ก1{g}90{act_i + 1}", name=act, hours=40, seq=act_i))
        # ตารางเรียน
        for day in range(1, 6):
            for pi, pe in enumerate(teach_periods):
                sj = subjects[(day + pi) % len(subjects)]
                note = "ลูกเสือ-เนตรนารี" if (day == 5 and pi == len(teach_periods) - 1) else ""
                db.add(m.AcadTimetable(class_id=k.id, day=day, period_id=pe.id,
                                       subject_id=None if note else sj.id, note=note))
        # คะแนน: กรอกแล้ว 3 วิชาแรก ที่เหลือเว้นไว้ให้โชว์การกรอก
        for sj in subjects[:3]:
            a1 = m.AcadAssignment(subject_id=sj.id, term=0, name="ใบงานที่ 1", max_score=20, seq=1)
            a2 = m.AcadAssignment(subject_id=sj.id, term=0, name="สอบกลางภาค", max_score=30, is_midterm=True, seq=2)
            a3 = m.AcadAssignment(subject_id=sj.id, term=0, name="ชิ้นงานกลุ่ม", max_score=20, seq=3)
            db.add_all([a1, a2, a3])
            db.flush()
            for n, st in enumerate(k.students):
                base = 0.55 + ((n * 7 + sj.id) % 40) / 100
                parts = [round(a.max_score * min(base + j * 0.03, 1)) for j, a in enumerate((a1, a2, a3))]
                for a, sc in zip((a1, a2, a3), parts):
                    db.add(m.AcadAssignmentScore(assignment_id=a.id, acad_student_id=st.id, score=sc))
                mid = float(sum(parts))
                fin = float(round(30 * min(base + 0.05, 1)))
                db.add(m.AcadScore(acad_student_id=st.id, subject_id=sj.id, term=0, score_mid=mid,
                                   score_final=fin, score=mid + fin, grade=grade_of(mid + fin)))
        # เวลาเรียนรายเดือน (เดือนที่ผ่านมาแล้ว)
        for mon, days in by_month.items():
            if mon == today.month:
                continue
            for n, st in enumerate(k.students):
                marks = ["."] * 31
                for d in days:
                    marks[d - 1] = "/"
                if n % 6 == 0 and days:
                    marks[days[len(days) // 2] - 1] = "ป"
                if n % 9 == 0 and len(days) > 3:
                    marks[days[3] - 1] = "ล"
                db.add(m.AcadAttendance(acad_student_id=st.id, month=mon, marks="".join(marks),
                                        present=marks.count("/")))

    # แผนการสอนในสายอนุมัติ: 1 รอวิชาการ + 1 อนุมัติแล้ว
    db.add(m.LessonPlan(person_id=teacher.id, year=ay, term=2, title="แผนการจัดการเรียนรู้ ภาษาไทย ป.1 หน่วยที่ 5",
                        file_blob=_fake_docx("ภาษาไทย ป.1 หน่วยที่ 5"), file_name="แผนภาษาไทย_ป1_หน่วย5.docx",
                        note="ส่งแผนหน่วยที่ 5 ครับ", status="pending", submitted_at=ago(1)))
    db.add(m.LessonPlan(person_id=teacher.id, year=ay, term=1, title="แผนการจัดการเรียนรู้ ภาษาไทย ป.1 หน่วยที่ 1",
                        file_blob=_fake_docx("ภาษาไทย ป.1 หน่วยที่ 1"), file_name="แผนภาษาไทย_ป1_หน่วย1.docx",
                        status="approved", comment="เนื้อหาครบถ้วน", submitted_at=ago(90), reviewed_at=ago(88),
                        academic_by=persons[3].id, director_by=director.id, director_at=ago(87),
                        director_comment="อนุมัติ"))
    db.add(m.ClassroomVisit(person_id=teacher.id, term=1, year=ay, subject_group="ภาษาไทย", topic="การอ่านสะกดคำ",
                            grade_level="ป.1", period="2", visit_time="09.30 น.", visit_date=ago(50).date(),
                            visitor_name=director.name, scores="5,4,5,4,4,5,4,5,4,5",
                            suggestion="ควรเพิ่มสื่อการสอนที่เป็นเกมเพื่อกระตุ้นความสนใจ"))

    # ---------------- งานอาหารกลางวัน ----------------
    prog = m.LunchProgram(year=ay, days=200, rate_per_head=36, operate_mode="hire",
                          funding_org="องค์การบริหารส่วนตำบลตัวอย่าง", lunch_officer=PEOPLE[8][0])
    for i, (lvl, n) in enumerate((("อ.2", 6), ("อ.3", 6), ("ป.1", 20), ("ป.2", 5), ("ป.3", 3))):
        prog.classes.append(m.LunchClass(seq=i, level=lvl, num_students=n))
    db.add(prog)
    db.flush()
    prog.ledger.append(m.LunchLedger(date=_d(ce, 5, 20), kind="in", amount=round(prog.budget / 2, 2),
                                     detail="รับเงินอุดหนุนอาหารกลางวัน ภาคเรียนที่ 1", ref="อบต 0002/88"))
    per_day = prog.total_students * prog.rate_per_head
    rounds_def = [(1, t1s, _d(ce, 7, 31), "จ่ายแล้ว"), (2, _d(ce, 8, 1), t1e, "จ้างแล้ว")]
    from app.routers.lunch import _sync_round_procurement, _sync_installment_ledger
    for seq, rs, re_, st in rounds_def:
        days = sum(1 for _ in _weekdays(rs.date(), re_.date()))
        rnd = m.LunchHireRound(program_id=prog.id, seq=seq, period_type="month", start_date=rs, end_date=re_,
                               days=days, vendor_id=vendors[4].id, amount=round(per_day * days, 2),
                               order_no=f"{seq}/{fy}", order_date=rs - timedelta(days=3), status=st)
        db.add(rnd)
        db.flush()
        for kind in ("tor", "control", "inspect"):
            for i, (who, role) in enumerate(((PEOPLE[3][0], "ประธานกรรมการ"), (PEOPLE[8][0], "กรรมการ"),
                                             (PEOPLE[7][0], "กรรมการและเลขานุการ")), 1):
                db.add(m.LunchCommittee(round_id=rnd.id, kind=kind, seq=i, name=who, position="ครู", role=role))
        # แบ่งงวดละ 10 วันทำการ
        wd = list(_weekdays(rs.date(), re_.date()))
        for n, i0 in enumerate(range(0, len(wd), 10), 1):
            chunk = wd[i0:i0 + 10]
            done = st == "จ่ายแล้ว" or chunk[-1] < today - timedelta(days=30)
            inst = m.LunchInstallment(round_id=rnd.id, seq=n, start_date=datetime.combine(chunk[0], datetime.min.time()),
                                      end_date=datetime.combine(chunk[-1], datetime.min.time()), days=len(chunk),
                                      amount=round(per_day * len(chunk), 2),
                                      deliver_date=datetime.combine(chunk[-1], datetime.min.time()) if done else None,
                                      inspect_date=datetime.combine(chunk[-1], datetime.min.time()) if done else None,
                                      status="จ่ายแล้ว" if done else "ร่าง")
            db.add(inst)
            db.flush()
            inst.round = rnd
            _sync_installment_ledger(db, inst)
        db.refresh(rnd)
        _sync_round_procurement(db, rnd)
    menus = [("ข้าวผัดกะเพราไก่ ไข่ต้ม", "กล้วยน้ำว้า"), ("แกงจืดเต้าหู้หมูสับ", "ส้ม"),
             ("ข้าวมันไก่", "แตงโม"), ("ต้มจับฉ่าย", "ขนมกล้วย"), ("ผัดผักรวมหมู", "มะละกอ")]
    for i, d in enumerate(_weekdays(today - timedelta(days=13), today + timedelta(days=14))):
        mn, ds = menus[i % len(menus)]
        db.add(m.LunchMenu(program_id=prog.id, date=datetime.combine(d, datetime.min.time()),
                           main=mn, dessert=ds, groups="1,2,3,4,5"))
    for st in students_all[:20]:
        ls = m.LunchStudent(program_id=prog.id, student_id=st.id, name=st.name, sex=st.sex,
                            birthdate=st.birthdate, level=st.level)
        ls.measures.append(m.LunchMeasure(term=1, date=_d(ce, 6, 10),
                                          weight=st.measures[0].weight if st.measures else 22,
                                          height=st.measures[0].height if st.measures else 118))
        db.add(ls)

    db.commit()
    return {"teacher_pid": teacher.id, "director_pid": director.id, "fiscal_year": fy, "academic_year": ay}


# ------------------------------------------------------------------ main
def _taken_usernames() -> list:
    """ชื่อผู้ใช้ของเดโมที่ชนกับบัญชีเดิม (ตรวจก่อนสร้าง กันสร้างค้างครึ่งทาง)"""
    from app.accounts import acc_session, Account
    wanted = [OWNER_USER, f"teacher1.{DEMO_SLUG}", f"director.{DEMO_SLUG}"]
    db = acc_session()
    try:
        return [u for u in wanted if db.query(Account).filter_by(username=u).first()]
    finally:
        db.close()


def build(password: str, *, reset: bool = False, today: date | None = None) -> dict:
    from app.accounts import purge_tenant
    from app.database import get_data_dir
    from app.tenancy import session_for

    if str(get_data_dir()).replace("\\", "/").startswith("/opt/"):
        raise SystemExit("สคริปต์นี้ใช้บนเครื่องตัวเองเท่านั้น ไม่รันบนเซิร์ฟเวอร์จริง")
    old = find_demo_tenant()
    if old and not reset:
        raise SystemExit("มีโรงเรียนเดโมอยู่แล้ว ถ้าต้องการสร้างใหม่ให้ใส่ --reset")
    if old:
        purge_tenant(old)
    taken = _taken_usernames()
    if taken:
        raise SystemExit("ชื่อผู้ใช้ต่อไปนี้ถูกใช้โดยโรงเรียนอื่นในเครื่องนี้แล้ว: " + ", ".join(taken)
                         + " (ไม่ได้สร้างอะไรเลย)")
    info = create_accounts(password)
    tid = info["tenant_id"]
    db = session_for(tid)
    try:
        ids = seed_school(db, today=today or date.today())
    finally:
        db.close()
    users = create_person_accounts(tid, password, ids["teacher_pid"], ids["director_pid"])
    return {"tenant_id": tid, "users": [OWNER_USER] + users, **ids}


def main():
    import getpass
    import os
    reset = "--reset" in sys.argv
    pw = os.environ.get("DEMO_PASSWORD") or ""
    if not pw:
        pw = getpass.getpass("ตั้งรหัสผ่านสำหรับบัญชีเดโม (ใช้ร่วมกันทั้ง 3 บัญชี): ")
        if pw != getpass.getpass("พิมพ์รหัสผ่านอีกครั้ง: "):
            raise SystemExit("รหัสผ่านไม่ตรงกัน")
    r = build(pw, reset=reset)
    print(f"สร้าง {DEMO_SCHOOL} เรียบร้อย (ปีงบ {r['fiscal_year']} · ปีการศึกษา {r['academic_year']})")
    print("บัญชีเข้าสู่ระบบ:")
    for u, role in zip(r["users"], ("ไอดีหลัก (เจ้าหน้าที่)", "ครู", "ผู้อำนวยการ")):
        print(f"  {u:<16} {role}")


if __name__ == "__main__":
    main()
