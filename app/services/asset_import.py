"""นำเข้าทะเบียนครุภัณฑ์เดิมจาก Excel (หน้า /assets/import)

ขั้นตอน: อัปโหลด -> หน้าตรวจ (บอกแถวที่มีปัญหา) -> ยืนยัน -> บันทึกเฉพาะแถวที่ผ่าน
- อ่านหัวคอลัมน์แบบยืดหยุ่น ใช้ไฟล์ทะเบียนเดิมของโรงเรียนได้เลย ถ้าหัวคอลัมน์ใกล้เคียงแม่แบบ
- เลขครุภัณฑ์ซ้ำกับในระบบ หรือซ้ำกันเองในไฟล์ = ไม่นำเข้าแถวนั้น (เลขครุภัณฑ์ต้องไม่ซ้ำตลอดไป)
"""
import re
import tempfile
from io import BytesIO
from pathlib import Path

from app.models import Asset, AssetNumberUsed
from app.services.asset_utils import CATEGORY_LIFE, ASSET_STATUSES
from app.thai_utils import parse_be_date, thai_date_short

FUND_TYPES = ["เงินงบประมาณ", "เงินนอกงบประมาณ", "เงินบริจาค/เงินช่วยเหลือ", "อื่น"]
METHODS = ["วิธีตลาดอิเล็กทรอนิกส์", "วิธีประกวดราคาอิเล็กทรอนิกส์", "วิธีคัดเลือก", "วิธีเฉพาะเจาะจง", "รับบริจาค"]
STATUSES = [s for s in ASSET_STATUSES if s != "จำหน่ายแล้ว"]   # ของที่จำหน่ายแล้วไม่ต้องนำเข้า

# (ฟิลด์, หัวคอลัมน์ในแม่แบบ, คำที่ใช้จับหัวคอลัมน์ของไฟล์เดิม, ความกว้าง)
COLUMNS = [
    ("asset_code", "เลขครุภัณฑ์", ["เลขครุภัณฑ์", "รหัสครุภัณฑ์", "เลขทะเบียน", "รหัส", "หมายเลข"], 22),
    ("name", "ชื่อครุภัณฑ์ *", ["ชื่อครุภัณฑ์", "ชื่อ", "รายการ", "รายละเอียด"], 34),
    ("brand_model", "ยี่ห้อ/รุ่น", ["ยี่ห้อ", "รุ่น", "แบบ"], 20),
    ("category", "ประเภท", ["ประเภท", "หมวด", "ชนิด"], 30),
    ("quantity", "จำนวน", ["จำนวน"], 9),
    ("unit", "หน่วย", ["หน่วยนับ", "หน่วย"], 9),
    ("acquired_date", "วันที่ได้มา (วว/ดด/ปปปป)", ["วันที่ได้มา", "วันที่ซื้อ", "วันที่รับ", "ได้มาเมื่อ", "วันที่", "วัน เดือน ปี"], 20),
    ("cost", "ราคาต่อหน่วย (บาท)", ["ราคาต่อหน่วย", "ราคาทุน", "ราคา", "มูลค่า"], 16),
    ("useful_life", "อายุการใช้งาน (ปี)", ["อายุการใช้งาน", "อายุ"], 12),
    ("fund_type", "ประเภทเงิน", ["ประเภทเงิน", "แหล่งเงิน", "เงิน"], 22),
    ("acquire_method", "วิธีได้มา", ["วิธีได้มา", "วิธีการได้มา", "วิธีจัดหา", "วิธี"], 24),
    ("vendor_name", "ผู้ขาย/ผู้บริจาค", ["ผู้ขาย", "ผู้รับจ้าง", "ผู้บริจาค", "ร้าน"], 26),
    ("doc_ref", "เลขที่เอกสาร", ["เลขที่เอกสาร", "ที่เอกสาร", "เลขที่ใบส่งของ", "ใบส่งของ"], 18),
    ("location", "สถานที่ตั้ง/ผู้ใช้", ["สถานที่", "ที่ตั้ง", "ห้อง", "ผู้ใช้", "ผู้รับผิดชอบ"], 22),
    ("status", "สถานะ", ["สถานะ", "สภาพ"], 12),
    ("note", "หมายเหตุ", ["หมายเหตุ"], 24),
]
MAX_ROWS = 5000


def build_template() -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "ครุภัณฑ์"
    ws.append(["กรอกตั้งแต่แถวที่ 3 จำเป็นเฉพาะชื่อครุภัณฑ์ วันที่เป็น พ.ศ. เช่น 15/05/2565 "
               "ราคาเป็นราคาต่อหน่วย ประเภท/ประเภทเงิน/วิธีได้มา/สถานะ เลือกจากรายการได้"])
    ws.append([c[1] for c in COLUMNS])
    ws.append(["7440-001-0001/2565", "เครื่องคอมพิวเตอร์โน้ตบุ๊ก", "Acer Aspire 5", "ครุภัณฑ์คอมพิวเตอร์", 1, "เครื่อง",
               "15/05/2565", 21000, 5, "เงินงบประมาณ", "วิธีเฉพาะเจาะจง", "ร้านตัวอย่างคอมพิวเตอร์", "IV65-0012",
               "ห้องธุรการ", "ใช้งาน", ""])
    ws["A1"].font = Font(bold=True, color="7A5300")
    ws["A1"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    for i, c in enumerate(COLUMNS, 1):
        cell = ws.cell(row=2, column=i)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = c[3]
    # รายการให้เลือก (เก็บในชีตซ่อน เพราะรายการประเภทยาวเกินกว่าใส่ในสูตรได้)
    lists = wb.create_sheet("รายการเลือก")
    for col, values in enumerate([list(CATEGORY_LIFE), FUND_TYPES, METHODS, STATUSES], 1):
        for r, v in enumerate(values, 1):
            lists.cell(row=r, column=col, value=v)
    lists.sheet_state = "hidden"
    for field, letter, n in (("category", "A", len(CATEGORY_LIFE)), ("fund_type", "B", len(FUND_TYPES)),
                             ("acquire_method", "C", len(METHODS)), ("status", "D", len(STATUSES))):
        idx = [c[0] for c in COLUMNS].index(field) + 1
        col = ws.cell(row=2, column=idx).column_letter
        dv = DataValidation(type="list", formula1=f"='รายการเลือก'!${letter}$1:${letter}${n}", allow_blank=True)
        dv.error, dv.showErrorMessage = "เลือกจากรายการ หรือพิมพ์ให้ตรงกับรายการ", False
        ws.add_data_validation(dv)
        dv.add(f"{col}3:{col}{MAX_ROWS}")
    ws.freeze_panes = "B3"
    path = Path(tempfile.gettempdir()) / "แม่แบบนำเข้าครุภัณฑ์.xlsx"
    wb.save(path)
    return path


def _norm(s):
    return re.sub(r"[\s*().\-/]", "", str(s or "")).lower()


def _header_map(header):
    """จับคู่หัวคอลัมน์ของไฟล์กับฟิลด์ (คำที่ยาวกว่าจับก่อน กันคำว่า 'ราคา' ไปจับ 'ราคาต่อหน่วย' ผิดช่อง)"""
    cells = [_norm(h) for h in header]
    out, used = {}, set()
    pairs = sorted(((len(k), f, k) for f, _, keys, _ in COLUMNS for k in keys), reverse=True)
    for _, field, key in pairs:
        if field in out:
            continue
        k = _norm(key)
        for i, h in enumerate(cells):
            if i not in used and h and k in h:
                out[field] = i
                used.add(i)
                break
    return out


def _num(v, integer=False):
    if v is None or str(v).strip() == "":
        return None
    try:
        n = float(str(v).replace(",", "").replace("บาท", "").strip())
    except ValueError:
        raise ValueError("ไม่ใช่ตัวเลข")
    if n < 0:
        raise ValueError("ติดลบ")
    return int(n) if integer else round(n, 2)


def _date(v):
    from datetime import datetime, date
    if v in (None, ""):
        return None
    if isinstance(v, (datetime, date)):
        d = v if isinstance(v, datetime) else datetime(v.year, v.month, v.day)
        # Excel เก็บเป็นวันที่ ค.ศ. ถ้าพิมพ์ปี พ.ศ. ลงช่องวันที่ Excel ปีจะเกิน 2400
        return d.replace(year=d.year - 543) if d.year > 2400 else d
    d = parse_be_date(str(v).strip())
    if d is None:
        raise ValueError("อ่านวันที่ไม่ได้")
    return d


def _pick(value, options, default):
    v = str(value or "").strip()
    if not v:
        return default, None
    for o in options:
        if v == o or _norm(v) == _norm(o):
            return o, None
    for o in options:
        if _norm(v) in _norm(o) or _norm(o) in _norm(v):
            return o, None
    return default, f'"{v}" ไม่ตรงรายการ ใช้ "{default}" แทน'


def find_header_row(ws):
    """หาแถวหัวตารางใน 10 แถวแรก (แถวที่มีคำว่า ชื่อ/รายการ และจับคู่ได้หลายคอลัมน์ที่สุด)"""
    best, best_n = None, 1
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        m = _header_map(row or [])
        if "name" in m and len(m) > best_n:
            best, best_n = (r, m), len(m)
    return best


def parse(content: bytes, db) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    ws = wb["ครุภัณฑ์"] if "ครุภัณฑ์" in wb.sheetnames else wb.worksheets[0]
    found = find_header_row(ws)
    if not found:
        raise ValueError("ไม่พบหัวตาราง ต้องมีคอลัมน์ชื่อครุภัณฑ์ (หรือใช้แม่แบบของระบบ)")
    head_row, cmap = found
    used_codes = {c for (c,) in db.query(Asset.asset_code).filter(Asset.asset_code != "")}
    used_codes |= {c for (c,) in db.query(AssetNumberUsed.code)}
    seen, rows = set(), []
    for n, row in enumerate(ws.iter_rows(min_row=head_row + 1, values_only=True), head_row + 1):
        if len(rows) >= MAX_ROWS:
            break
        get = lambda f: (row[cmap[f]] if f in cmap and cmap[f] < len(row) else None)  # noqa: E731
        if not any(v not in (None, "") for v in row or []):
            continue
        r = {"line": n, "errors": [], "warnings": []}
        r["name"] = str(get("name") or "").strip()
        r["asset_code"] = str(get("asset_code") or "").strip()
        for f in ("brand_model", "unit", "vendor_name", "doc_ref", "location", "note"):
            r[f] = str(get(f) or "").strip()
        r["unit"] = r["unit"] or "หน่วย"
        for f, integer, label in (("quantity", False, "จำนวน"), ("cost", False, "ราคา"), ("useful_life", True, "อายุการใช้งาน")):
            try:
                r[f] = _num(get(f), integer)
            except ValueError as e:
                r[f] = None
                r["errors"].append(f"{label}: {e}")
        r["quantity"] = r["quantity"] or 1
        r["cost"] = r["cost"] or 0
        try:
            d = _date(get("acquired_date"))
            r["acquired_date"] = d.strftime("%Y-%m-%d") if d else ""
            r["acquired_th"] = thai_date_short(d) if d else ""
        except ValueError as e:
            r["acquired_date"] = ""
            r["errors"].append(f"วันที่ได้มา: {e}")
        for f, options, default in (("category", list(CATEGORY_LIFE), "ครุภัณฑ์สำนักงาน"),
                                    ("fund_type", FUND_TYPES, "เงินงบประมาณ"),
                                    ("acquire_method", METHODS, "วิธีเฉพาะเจาะจง"),
                                    ("status", STATUSES, "ใช้งาน")):
            r[f], warn = _pick(get(f), options, default)
            if warn:
                r["warnings"].append(warn)
        if not r["useful_life"]:
            r["useful_life"] = CATEGORY_LIFE.get(r["category"], 10)
        if not r["name"]:
            r["errors"].append("ไม่มีชื่อครุภัณฑ์")
        if r["asset_code"]:
            if r["asset_code"] in used_codes:
                r["errors"].append("เลขครุภัณฑ์นี้มีในระบบแล้ว")
            elif r["asset_code"] in seen:
                r["errors"].append("เลขครุภัณฑ์ซ้ำกับแถวก่อนหน้าในไฟล์")
            seen.add(r["asset_code"])
        else:
            r["warnings"].append("ไม่มีเลขครุภัณฑ์ (ใส่ภายหลังได้)")
        if not r["acquired_date"] and not r["errors"]:
            r["warnings"].append("ไม่มีวันที่ได้มา ค่าเสื่อมราคาจะคำนวณไม่ได้")
        rows.append(r)
    ok = [r for r in rows if not r["errors"]]
    return {"rows": rows, "ok": len(ok), "bad": len(rows) - len(ok),
            "total_cost": round(sum(r["cost"] * r["quantity"] for r in ok), 2),
            "mapped": [c[1] for c in COLUMNS if c[0] in cmap], "missing": [c[1] for c in COLUMNS if c[0] not in cmap]}


def save(db, rows) -> int:
    """บันทึกเฉพาะแถวที่ไม่มีข้อผิดพลาด (ตรวจเลขซ้ำอีกครั้งตอนบันทึก กันมีคนเพิ่มเลขเดียวกันระหว่างรอยืนยัน)"""
    from datetime import datetime
    from app.services.asset_numbering import lock_numbers, manual_number
    lock_numbers(db)
    n = 0
    for r in rows:
        if r.get("errors") or not (r.get("name") or "").strip():
            continue
        a = Asset(name=r["name"].strip()[:300], asset_code=(r.get("asset_code") or "").strip()[:100],
                  brand_model=r.get("brand_model", "")[:200], category=r.get("category") or "ครุภัณฑ์สำนักงาน",
                  quantity=float(r.get("quantity") or 1), unit=r.get("unit") or "หน่วย",
                  acquired_date=datetime.strptime(r["acquired_date"], "%Y-%m-%d") if r.get("acquired_date") else None,
                  cost=float(r.get("cost") or 0), useful_life=int(r.get("useful_life") or 10), salvage_value=1.0,
                  fund_type=r.get("fund_type") or "เงินงบประมาณ", acquire_method=r.get("acquire_method") or "วิธีเฉพาะเจาะจง",
                  vendor_name=r.get("vendor_name", "")[:200], doc_ref=r.get("doc_ref", "")[:100],
                  location=r.get("location", "")[:200], status=r.get("status") or "ใช้งาน",
                  note=(r.get("note", "") + (" (นำเข้าจาก Excel)" if r.get("note") else "นำเข้าจาก Excel"))[:300])
        if a.category not in CATEGORY_LIFE or a.fund_type not in FUND_TYPES or a.status not in STATUSES:
            continue
        manual_number(db, a.asset_code)          # กันเลขซ้ำ (โยน ValueError ถ้าซ้ำ)
        db.add(a)
        n += 1
    return n
