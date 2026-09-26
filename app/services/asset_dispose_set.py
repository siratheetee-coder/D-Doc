# -*- coding: utf-8 -*-
"""
asset_dispose_set.py - ชุดเอกสารจำหน่ายพัสดุประจำปี (ต่อจากการตรวจสอบพัสดุประจำปี)

อ้างระเบียบกระทรวงการคลังว่าด้วยการจัดซื้อจัดจ้างและการบริหารพัสดุภาครัฐ พ.ศ. 2560
หมวด 9 ส่วนที่ 3 (ตรวจสอบพัสดุ ข้อ 213) และส่วนที่ 4 (จำหน่ายพัสดุ ข้อ 214-218)
พ.ร.บ.การจัดซื้อจัดจ้างฯ พ.ศ. 2560 มาตรา 112 และคำสั่ง สพฐ. ที่ 1340/2560
ลงวันที่ 24 สิงหาคม 2560 เรื่อง มอบอำนาจเกี่ยวกับการจัดซื้อจัดจ้างและการบริหารพัสดุภาครัฐ

ถ้อยคำและลำดับเอกสารอิงแบบฟอร์มของสำนักงานเขตพื้นที่การศึกษาประถมศึกษาสุรินทร์
(ฉบับปรับตามระเบียบ พ.ศ. 2560) ซึ่งเป็นเอกสารราชการที่เผยแพร่ให้โรงเรียนใช้

ลำดับเอกสาร
  ขั้น 1 สอบหาข้อเท็จจริง : บันทึกขอแต่งตั้ง -> คำสั่งแต่งตั้ง -> รายงานผล -> บันทึกถ้อยคำ
  ขั้น 2 ขออนุมัติจำหน่าย : บันทึกขอจำหน่ายพัสดุ (เสนอ 3 คณะกรรมการ) -> คำสั่งแต่งตั้ง
  ขั้น 3 ดำเนินการ        : (ก) เฉพาะเจาะจง - หนังสือเชิญเสนอราคา/ใบเสนอราคา/ประเมินราคากลาง/
                                 รายงานผลการขาย
                            (ข) ทอดตลาด   - คำสั่ง กก.ทอดตลาด/ประกาศ/บัญชีพัสดุ/บัญชีคุม/รายงานผล
                            (ค) ทำลาย     - รายงานผลการทำลาย
  ขั้น 4 ปิดเรื่อง        : หนังสือแจ้ง สตง.ภูมิภาค -> หนังสือนำส่งเงินรายได้
"""
import json

from docx import Document
from docx.shared import Cm

from app.database import get_data_dir
from app.services.doc_page import set_a4, tidy
from app.thai_utils import thai_date, bahttext
from app.services.build_templates import (
    _font, _krut_and_title, _krut_center, _p, _p_runs, _sign_table, _set_cell, _hr,
    _repeat_header_row, _no_split_row, _no_borders, _fixed_cols,
)

_BLANK = "............................"
_DOT = "................................................................"

# ---- ข้อความอ้างระเบียบที่ใช้ซ้ำหลายฉบับ -------------------------------------
_DELEG = ("ซึ่งได้รับมอบอำนาจจากเลขาธิการคณะกรรมการการศึกษาขั้นพื้นฐาน ตามคำสั่งสำนักงาน"
          "คณะกรรมการการศึกษาขั้นพื้นฐาน ที่ 1340/2560 สั่ง ณ วันที่ 24 สิงหาคม 2560 "
          "เรื่อง มอบอำนาจการสั่งซื้อสั่งจ้าง และการดำเนินการตามระเบียบกระทรวงการคลัง"
          "ว่าด้วยการจัดซื้อจัดจ้างและการบริหารพัสดุภาครัฐ พ.ศ. 2560")

_REG = ("ระเบียบกระทรวงการคลังว่าด้วยการจัดซื้อจัดจ้างและการบริหารพัสดุภาครัฐ พ.ศ. 2560")

# วงเงินขายโดยวิธีเฉพาะเจาะจงตามข้อ 215 (1) (ก) · ฉบับ 2560 ใช้ 500,000 บาท
SELL_SPECIFIC_LIMIT = 500_000.0
# วงเงินที่หัวหน้าหน่วยงานของรัฐอนุมัติจำหน่ายเป็นสูญได้เอง ตามข้อ 217 (1)
WRITE_OFF_LIMIT = 1_000_000.0
# ต้องดำเนินการจำหน่ายให้เสร็จภายใน 60 วันนับถัดจากวันที่หัวหน้าหน่วยงานของรัฐสั่งการ
DISPOSE_DEADLINE_DAYS = 60

# วิธีจำหน่ายพัสดุตามระเบียบฯ ข้อ 215 (ใช้เลือกรายรายการ)
DISPOSE_ACTIONS = ["ขาย", "แลกเปลี่ยน", "โอน", "แปรสภาพ", "ทำลาย", "จำหน่ายเป็นสูญ"]

# วิธีจำหน่าย -> เลขข้อที่ต้องอ้างในเอกสาร
ACTION_CLAUSE = {
    "ขาย": "ข้อ 215 (1)",
    "แลกเปลี่ยน": "ข้อ 215 (2)",
    "โอน": "ข้อ 215 (3)",
    "แปรสภาพ": "ข้อ 215 (4)",
    "ทำลาย": "ข้อ 215 (4)",
    "จำหน่ายเป็นสูญ": "ข้อ 217 (1)",
}

_DISPOSE_WAYS = [
    ("ขาย", "ข้อ 215 (1)",
     "ให้ดำเนินการขายโดยวิธีทอดตลาดก่อน แต่ถ้าขายโดยวิธีทอดตลาดแล้วไม่ได้ผลดี "
     "ให้นำวิธีที่กำหนดเกี่ยวกับการซื้อมาใช้โดยอนุโลม เว้นแต่ (ก) การขายพัสดุครั้งหนึ่ง"
     "ซึ่งมีราคาซื้อหรือได้มารวมกันไม่เกิน 500,000 บาท จะขายโดยวิธีเฉพาะเจาะจง"
     "โดยการเจรจาตกลงราคากันโดยไม่ต้องทอดตลาดก่อนก็ได้ (ข) การขายให้แก่หน่วยงานของรัฐ "
     "หรือองค์การสถานสาธารณกุศลตามมาตรา 47 (7) แห่งประมวลรัษฎากร และ (ค) การขายอุปกรณ์"
     "อิเล็กทรอนิกส์ให้แก่เจ้าหน้าที่ของรัฐที่หน่วยงานของรัฐมอบให้ไว้ใช้งานในหน้าที่"),
    ("แลกเปลี่ยน", "ข้อ 215 (2)", "ให้ดำเนินการตามวิธีแลกเปลี่ยนที่กำหนดไว้ในระเบียบฯ"),
    ("โอน", "ข้อ 215 (3)",
     "ให้โอนแก่หน่วยงานของรัฐ หรือองค์การสถานสาธารณกุศลตามมาตรา 47 (7) "
     "แห่งประมวลรัษฎากร ทั้งนี้ ให้มีหลักฐานการส่งมอบไว้ต่อกันด้วย"),
    ("แปรสภาพหรือทำลาย", "ข้อ 215 (4)",
     "แปรสภาพหรือทำลายตามหลักเกณฑ์และวิธีการที่หน่วยงานของรัฐกำหนด"),
]


# ------------------------------------------------------------------ helper
def _safe(text: str) -> str:
    for ch in '<>:"/\\|?*\n\r\t':
        text = text.replace(ch, "_")
    return text.strip()[:80]


def _save(doc, name: str) -> str:
    tidy(doc)
    out_dir = get_data_dir() / "documents"
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / (_safe(name) + ".docx")
    doc.save(str(out_path))
    return str(out_path)


def _new():
    doc = Document(); set_a4(doc); _font(doc)
    return doc


def _portrait_section(doc):
    """กลับมาเป็นหน้าตั้ง (A4 + ระยะขอบมาตรฐานราชการ) หลังจากมี section แนวนอน"""
    from docx.enum.section import WD_SECTION, WD_ORIENT
    from app.services.doc_page import (A4_W, A4_H, MARGIN_TOP, MARGIN_BOTTOM,
                                       MARGIN_LEFT, MARGIN_RIGHT)
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = A4_W, A4_H
    sec.top_margin, sec.bottom_margin = MARGIN_TOP, MARGIN_BOTTOM
    sec.left_margin, sec.right_margin = MARGIN_LEFT, MARGIN_RIGHT
    return sec


def _break(doc) -> None:
    """ขึ้นหน้าใหม่ก่อนฉบับถัดไป · ฉบับแรกของชุดไม่ต้องขึ้น (กันหน้าแรกว่าง)

    ถ้าฉบับก่อนหน้าเป็นหน้านอน (เช่น บัญชีคุมขายทอดตลาด) ต้องเปิด section
    หน้าตั้งใหม่ ไม่งั้นฉบับถัดไปจะพิมพ์เป็นหน้านอนติดมาด้วย
    """
    if not (doc.paragraphs or doc.tables):
        return
    from docx.enum.section import WD_ORIENT
    if doc.sections and doc.sections[-1].orientation == WD_ORIENT.LANDSCAPE:
        _portrait_section(doc)
        return
    doc.add_page_break()


def _d(dt) -> str:
    """วันที่แบบมี พ.ศ. (รูปแบบที่ชุดเอกสารนี้ใช้ทั้งชุด)"""
    return thai_date(dt).replace(" 25", " พ.ศ. 25", 1) if dt else _BLANK


def _office(school) -> str:
    return "  ".join(p for p in [(school.name or "").strip(),
                                 (school.address or "").strip()] if p)


def _director_line(school) -> str:
    name = (school.name or "").strip()
    return "ผู้อำนวยการ" + name if name.startswith("โรงเรียน") else "ผู้อำนวยการโรงเรียน"


def _sname(school) -> str:
    return (school.name or "โรงเรียน").strip()


def _memo_header(doc, school, subject, doc_no, date_txt, *, to=None, encl=None):
    _krut_and_title(doc)
    _p_runs(doc, [("ส่วนราชการ  ", True), (_office(school), False)])
    _p_runs(doc, [("ที่  ", True), (doc_no or _BLANK, False),
                  ("\t", False), ("วันที่ ", True), (date_txt, False)], tab_cm=8)
    _p_runs(doc, [("เรื่อง  ", True), (subject, False)])
    _p_runs(doc, [("เรียน  ", True), (to or _director_line(school), False)])
    if encl:
        _p_runs(doc, [("สิ่งที่ส่งมาด้วย  ", True), (encl, False)])
    _hr(doc)


def _letter_header(doc, school, doc_no, date_txt, subject, to, encl=None):
    """หัวหนังสือราชการภายนอก (ครุฑกลาง + ที่ ซ้าย / ที่อยู่ ขวา)"""
    _krut_center(doc, height_cm=2.0)
    t = doc.add_table(rows=1, cols=2)
    _no_borders(t); t.autofit = False
    widths = [Cm(8.0), Cm(8.5)]
    _fixed_cols(t, widths)
    _set_cell(t.rows[0].cells[0], "ที่  " + (doc_no or _BLANK), align="left", size=16)
    addr = (getattr(school, "address", "") or "").strip()
    _set_cell(t.rows[0].cells[1], _sname(school) + (("\n" + addr) if addr else ""),
              align="right", size=16)
    for c, w in zip(t.rows[0].cells, widths):
        c.width = w
    _p(doc, date_txt, align="center", before=4, after=6)
    _p_runs(doc, [("เรื่อง  ", True), (subject, False)])
    _p_runs(doc, [("เรียน  ", True), (to, False)])
    if encl:
        _p_runs(doc, [("สิ่งที่ส่งมาด้วย  ", True), (encl, False)])
    _p(doc, "", after=4)


def _landscape_section(doc):
    """เปิด section ใหม่แบบหน้านอน (ใช้กับบัญชีคุมที่คอลัมน์เยอะ)"""
    from app.services.asset_audit_doc import _landscape_section as _ls
    return _ls(doc)


def _members(dp, key) -> list:
    """รายชื่อกรรมการชุดที่ระบุ [{name, position, role}]"""
    try:
        data = json.loads(dp.members or "{}")
    except Exception:
        data = {}
    return [m for m in (data.get(key) or []) if (m.get("name") or "").strip()]


def _member_lines(doc, members, *, n_blank=3):
    """รายชื่อกรรมการเป็นตารางไร้เส้นขอบ ให้ ชื่อ/ตำแหน่ง/บทบาท ตรงคอลัมน์กัน"""
    data = list(members) if members else [None] * n_blank
    widths = [Cm(1.0), Cm(6.6), Cm(4.7), Cm(4.2)]        # รวม 16.5 ซม. = พื้นที่พิมพ์ A4
    t = doc.add_table(rows=len(data), cols=4)
    _no_borders(t)
    for i, (row, m) in enumerate(zip(t.rows, data), 1):
        if m:
            name = (m.get("name") or "").strip()
            pos = "ตำแหน่ง " + ((m.get("position") or "ครู").strip())
            role = (m.get("role") or "กรรมการ").strip()
        else:
            name, pos = _BLANK, "ตำแหน่ง .................."
            role = "ประธานกรรมการ" if i == 1 else "กรรมการ"
        for c, v, w in zip(row.cells, [f"{i}.", name, pos, role], widths):
            _set_cell(c, v, size=16, align="left")
            c.width = w


def _committee_block(doc, dp, key, title):
    """หัวข้อชุดกรรมการ + รายชื่อ (ใช้ในบันทึกขอจำหน่าย/คำสั่งแต่งตั้งที่มี 3 ชุด)"""
    _p(doc, title, bold=True, indent=1.25, before=2, after=1)
    _member_lines(doc, _members(dp, key))


def _sign_committee(doc, dp, key, *, n=3):
    """บล็อกลงนามกรรมการเรียงลงมา (ประธาน + กรรมการ)"""
    roles = ["ประธานกรรมการ", "กรรมการ", "กรรมการ"]
    mem = _members(dp, key)[:n]
    while len(mem) < n:
        mem.append({"name": "", "role": roles[min(len(mem), 2)]})
    for i, m in enumerate(mem):
        name = (m.get("name") or "").strip()
        role = (m.get("role") or roles[min(i, 2)]).strip()
        _sign_table(doc, [[
            (f"ลงชื่อ ...................................... {role}", "center"),
            (f"( {name} )" if name else f"( {_BLANK} )", "center"),
        ]], after=2, keep=(i == 0))


def _director_sign(doc, school, *, lines=()):
    for txt in lines:
        _p(doc, txt, indent=1.25, after=1)
    _p(doc, "", after=8)
    _sign_table(doc, [[
        ("ลงชื่อ ......................................", "center"),
        (f"( {(school.director_name or '').strip() or _BLANK} )", "center"),
        (_director_line(school), "center"),
    ]])


def _officer_sign(doc, school, *, head=False):
    nm = ((school.head_officer_name if head else school.officer_name) or "").strip()
    _sign_table(doc, [[
        ("ลงชื่อ ......................................", "center"),
        (f"( {nm or _BLANK} )", "center"),
        ("หัวหน้าเจ้าหน้าที่" if head else "เจ้าหน้าที่", "center"),
    ]])


def _sao_name(dp) -> str:
    """ชื่อสำนักงานตรวจเงินแผ่นดินที่ต้องส่งรายงาน (แนวทางฉบับ 2563 ใช้ระดับจังหวัด)"""
    kind = (dp.sao_kind or "จังหวัด").strip()
    region = (dp.sao_region or "").strip() or ".........."
    if kind.startswith("ภูมิภาค"):
        return f"สำนักงานตรวจเงินแผ่นดินภูมิภาคที่ {region}"
    return f"สำนักงานตรวจเงินแผ่นดินจังหวัด{region}"


def _items(dp, action=None):
    """รายการในสำนวน (กรองตามวิธีจำหน่ายได้)

    เรียงตามเลขครุภัณฑ์เหมือนที่แสดงบนหน้าจอ เพื่อให้เลขลำดับในทุกฉบับตรงกัน
    ข้ามรายการที่ครุภัณฑ์ถูกลบไปแล้ว (asset_id ค้าง) กันเอกสารพัง
    """
    rows = sorted((it for it in dp.items if it.asset is not None),
                  key=lambda it: ((it.asset.asset_code or "~"), it.id))
    if action:
        rows = [it for it in rows if (it.action or "ขาย") == action]
    return rows


# ------------------------------------------------------- ตารางรายการที่ขอจำหน่าย
_ITEM_W = [Cm(1.1), Cm(3.0), Cm(4.0), Cm(1.6), Cm(2.3), Cm(2.3), Cm(2.2)]


def _item_table(doc, rows, *, price_col=False):
    """ตารางรายการครุภัณฑ์ที่ขอจำหน่าย · price_col=True เพิ่มช่องราคากลาง"""
    heads = ["ที่", "เลขครุภัณฑ์", "รายการ", "จำนวน", "ราคาทุน (บาท)", "วันที่ได้มา", "วิธีจำหน่าย"]
    widths = list(_ITEM_W)
    if price_col:
        heads[-1] = "ราคากลาง (บาท)"
        heads.append("วิธีจำหน่าย")
        widths = [Cm(1.0), Cm(2.5), Cm(3.3), Cm(1.3), Cm(2.1), Cm(2.0), Cm(2.2), Cm(2.1)]
    t = doc.add_table(rows=1, cols=len(heads))
    t.style = "Table Grid"
    t.autofit = False
    _fixed_cols(t, widths)
    hdr = t.rows[0]
    _repeat_header_row(hdr); _no_split_row(hdr)
    for c, h, w in zip(hdr.cells, heads, widths):
        _set_cell(c, h, bold=True, align="center", size=13)
        c.width = w
    total = mid = 0.0
    for i, it in enumerate(rows, 1):
        a = it.asset
        total += float(a.cost or 0)
        mid += float(it.price_mid or 0)
        row = t.add_row(); _no_split_row(row)
        vals = [str(i), a.asset_code or "-", a.name or "-",
                f"{(a.quantity or 1):g} {a.unit or 'หน่วย'}", f"{float(a.cost or 0):,.2f}",
                thai_date(a.acquired_date) if a.acquired_date else "-"]
        aligns = ["center", "left", "left", "center", "right", "center"]
        if price_col:
            vals.append(f"{float(it.price_mid or 0):,.2f}"); aligns.append("right")
        vals.append(it.action or "ขาย"); aligns.append("center")
        for c, v, al, w in zip(row.cells, vals, aligns, widths):
            _set_cell(c, v, align=al, size=13)
            c.width = w
    trow = t.add_row(); _no_split_row(trow)
    _set_cell(trow.cells[0], "รวม", bold=True, align="center", size=13)
    trow.cells[0].merge(trow.cells[3])
    _set_cell(trow.cells[4], f"{total:,.2f}", bold=True, align="right", size=13)
    if price_col:
        _set_cell(trow.cells[6], f"{mid:,.2f}", bold=True, align="right", size=13)
    for row in t.rows:                        # ย้ำความกว้าง (Word ชอบรีเซ็ต)
        for c, w in zip(row.cells, widths):
            c.width = w
    return total, mid


# ================================================================== ขั้นที่ 1
def render_fact_memo(school, dp, doc=None):
    """(1) บันทึกข้อความ ขอแต่งตั้งคณะกรรมการสอบหาข้อเท็จจริง (ข้อ 214)"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    _memo_header(doc, school, "แต่งตั้งคณะกรรมการสอบหาข้อเท็จจริง",
                 dp.fact_memo_no, _d(dp.fact_memo_date))
    _p(doc, "ตามที่คณะกรรมการตรวจสอบพัสดุประจำปี ได้รายงานผลการตรวจสอบให้ทราบ และได้สั่งการ"
            "ให้ตั้งคณะกรรมการสอบหาข้อเท็จจริง กรณีมีพัสดุชำรุด เสื่อมสภาพ หรือสูญไป หรือหมด"
            "ความจำเป็นต้องใช้ในหน่วยงานของรัฐต่อไป เพื่อประกอบการพิจารณาจำหน่ายตาม"
            f"{_REG} หมวด 9 ส่วนที่ 4 จำนวน {len(_items(dp))} รายการ นั้น",
       align="justify", indent=1.25, after=2)
    _p(doc, "เพื่อให้การดำเนินการสอบหาข้อเท็จจริงเป็นไปด้วยความเรียบร้อยและถูกต้องตาม"
            f"{_REG} ข้อ 214 {_DELEG} "
            "จึงขอแต่งตั้งบุคคลผู้มีรายนามต่อไปนี้เป็นคณะกรรมการสอบหาข้อเท็จจริง",
       align="justify", indent=1.25, after=2)
    _member_lines(doc, _members(dp, "fact"))
    _p(doc, "จึงเรียนมาเพื่อโปรดพิจารณา", indent=1.25, before=2, after=10)
    _officer_sign(doc, school, head=True)
    _director_sign(doc, school, lines=["- ชอบ", "- ดำเนินการ"])
    return _save(doc, f"ขอแต่งตั้งกรรมการสอบหาข้อเท็จจริง_ปีงบ{dp.year}") if own else doc


def render_fact_order(school, dp, doc=None):
    """(2) คำสั่งโรงเรียน แต่งตั้งคณะกรรมการสอบหาข้อเท็จจริง (ข้อ 214)"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    _krut_center(doc)
    _p(doc, f"คำสั่ง{_sname(school)}", align="center", bold=True, size=18, after=0)
    _p(doc, f"ที่ {dp.fact_order_no or _BLANK}", align="center", bold=True, after=0)
    _p(doc, "เรื่อง แต่งตั้งคณะกรรมการสอบหาข้อเท็จจริง",
       align="center", bold=True, size=17, after=0)
    _p(doc, "-----------------------------------", align="center", after=6)
    _p(doc, "ตามที่คณะกรรมการตรวจสอบพัสดุประจำปี ได้ตรวจสอบพัสดุปรากฏว่ามีพัสดุชำรุด "
            f"เสื่อมสภาพ หรือหมดความจำเป็นต้องใช้ในหน่วยงานของรัฐ จำนวน {len(_items(dp))} รายการ",
       align="justify", indent=1.25, after=2)
    _p(doc, f"อาศัยอำนาจตามความในข้อ 214 แห่ง{_REG} {_DELEG} "
            "จึงแต่งตั้งบุคคลผู้มีรายนามต่อไปนี้เป็นคณะกรรมการสอบหาข้อเท็จจริง "
            "เพื่อประกอบการพิจารณาจำหน่ายพัสดุ", align="justify", indent=1.25, after=2)
    _member_lines(doc, _members(dp, "fact"))
    _p(doc, "ให้คณะกรรมการที่ได้รับการแต่งตั้งตามคำสั่งนี้ ปฏิบัติหน้าที่ให้บังเกิดผลดีต่อ"
            f"ทางราชการอย่างสูงสุดโดยเคร่งครัด แล้วรายงานผลให้ทราบภายใน {dp.fact_days or 7} "
            "วันทำการ", align="justify", indent=1.25, after=2)
    _p(doc, f"ทั้งนี้ ตั้งแต่วันที่ {_d(dp.fact_order_date)} เป็นต้นไป", indent=1.25, after=1)
    _p(doc, f"สั่ง ณ วันที่ {_d(dp.fact_order_date)}", indent=1.25, after=12)
    _sign_table(doc, [[
        ("ลงชื่อ ......................................", "center"),
        (f"( {(school.director_name or '').strip() or _BLANK} )", "center"),
        (_director_line(school), "center"),
    ]])
    return _save(doc, f"คำสั่งแต่งตั้งกรรมการสอบหาข้อเท็จจริง_ปีงบ{dp.year}") if own else doc


def render_fact_report(school, dp, doc=None):
    """(3) บันทึกข้อความ รายงานผลการสอบหาข้อเท็จจริง"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    _memo_header(doc, school, "รายงานผลการสอบหาข้อเท็จจริง",
                 dp.fact_report_no, _d(dp.fact_report_date))
    _p(doc, "ตามที่ได้มีคำสั่งแต่งตั้งคณะกรรมการสอบหาข้อเท็จจริง กรณีมีพัสดุชำรุด เสื่อมสภาพ "
            "หรือสูญไป หรือหมดความจำเป็นต้องใช้ในหน่วยงานของรัฐต่อไป ดังที่คณะกรรมการ"
            f"ตรวจสอบพัสดุประจำปีได้รายงานผลให้ทราบแล้ว ตามคำสั่ง{_sname(school)} "
            f"ที่ {dp.fact_order_no or _BLANK} สั่ง ณ วันที่ {_d(dp.fact_order_date)} นั้น",
       align="justify", indent=1.25, after=2)
    _p(doc, f"คณะกรรมการได้เริ่มดำเนินการสอบหาข้อเท็จจริงเมื่อวันที่ {_d(dp.fact_start)} "
            f"แล้วเสร็จเมื่อวันที่ {_d(dp.fact_end)} ผลการสอบหาข้อเท็จจริงปรากฏว่า "
            + ((dp.fact_found or "").strip() or _DOT), align="justify", indent=1.25, after=2)

    rows = _items(dp)
    if rows:
        _item_table(doc, rows)
        _p(doc, "", after=3)
    causes = [it for it in rows if (it.cause or "").strip()]
    if causes:
        _p(doc, "สาเหตุที่ชำรุด เสื่อมสภาพ หรือสูญไป รายรายการ", bold=True,
           indent=1.25, after=1)
        for i, it in enumerate(causes, 1):
            _p(doc, f"{i}. {(it.asset.name or '').strip()} "
                    f"({(it.asset.asset_code or '-').strip()}) — {it.cause.strip()}",
               align="justify", indent=2.0, after=1)

    opinion = (dp.fact_opinion or "").strip() or (
        "พัสดุดังกล่าวชำรุดเสื่อมสภาพเนื่องมาจากการใช้งานตามปกติ หากนำมาซ่อมแซมจะไม่คุ้มค่า "
        "และหากใช้ในหน่วยงานของรัฐต่อไปจะสิ้นเปลืองค่าใช้จ่ายมาก จึงเห็นควรจำหน่ายตามระเบียบต่อไป")
    _p(doc, "คณะกรรมการสอบหาข้อเท็จจริงพิจารณาแล้วเห็นว่า " + opinion,
       align="justify", indent=1.25, after=2)
    if dp.fact_liable:
        _p(doc, "และเห็นว่ากรณีนี้ต้องดำเนินการหาตัวผู้รับผิดทางแพ่ง จึงเห็นควรแต่งตั้ง"
                "คณะกรรมการสอบข้อเท็จจริงความรับผิดทางละเมิดตามระเบียบสำนักนายกรัฐมนตรี"
                "ว่าด้วยหลักเกณฑ์การปฏิบัติเกี่ยวกับความรับผิดทางละเมิดของเจ้าหน้าที่ พ.ศ. 2539 "
                "ต่อไป", align="justify", indent=1.25, after=2)
    _p(doc, "จึงเรียนมาเพื่อโปรดทราบและพิจารณาสั่งการ พร้อมบันทึกนี้ได้แนบเอกสารการสอบ"
            f"ข้อเท็จจริง จำนวน {len(causes) or len(rows)} แผ่น มาด้วยแล้ว",
       align="justify", indent=1.25, after=10)
    _sign_committee(doc, dp, "fact")
    _director_sign(doc, school,
                   lines=["- ทราบ", "- มอบเจ้าหน้าที่ดำเนินการตามระเบียบฯ ต่อไป"])
    _p(doc, "วันที่ ........./............./..................", align="center", after=2)
    return _save(doc, f"รายงานผลการสอบหาข้อเท็จจริง_ปีงบ{dp.year}") if own else doc


def render_statement(school, dp, doc=None, *, item=None):
    """(4) บันทึกถ้อยคำในการสอบหาข้อเท็จจริง (1 แผ่นต่อผู้ให้ถ้อยคำ)

    ข้อมูลส่วนตัวของผู้ให้ถ้อยคำเว้นว่างให้กรอกเองในไฟล์ (ระบบไม่เก็บ)
    """
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    _p(doc, "บันทึกถ้อยคำในการสอบหาข้อเท็จจริง", align="center", bold=True, size=18, after=1)
    _p(doc, f"เขียนที่ {_sname(school)}", align="right", after=0)
    _p(doc, f"วันที่ {_BLANK}", align="right", after=6)
    what = ("พัสดุชำรุด เสื่อมสภาพ หรือหมดความจำเป็นต้องใช้ในหน่วยงานของรัฐ"
            if item is None else
            f"{(item.asset.name or '').strip()} ({(item.asset.asset_code or '-').strip()}) "
            "ชำรุดหรือใช้การไม่ได้")
    _p_runs(doc, [("เรื่อง  ", True), (f"การสอบหาข้อเท็จจริงกรณี {what}", False)])
    _hr(doc)
    _p(doc, f"ข้าพเจ้า {_BLANK} อายุ ......... ปี สัญชาติ .................. "
            "ศาสนา ..................", indent=1.25, after=1)
    _p(doc, "อาชีพ .......................................... ตำแหน่ง "
            ".......................................... อยู่บ้านเลขที่ ............ "
            "หมู่ที่ ........ ตรอก/ซอย .......................... ถนน ..........................",
       align="justify", after=1)
    _p(doc, "ตำบล/แขวง .......................... อำเภอ/เขต .......................... "
            "จังหวัด .......................... โทรศัพท์ ..........................",
       align="justify", after=2)
    _p(doc, f"ข้าพเจ้าได้รับทราบคำสั่ง{_sname(school)} ที่ {dp.fact_order_no or _BLANK} "
            f"ลงวันที่ {_d(dp.fact_order_date)} เรื่อง แต่งตั้งคณะกรรมการสอบหาข้อเท็จจริง "
            "จึงขอให้ถ้อยคำต่อคณะกรรมการสอบหาข้อเท็จจริงด้วยความสัตย์จริง ดังนี้",
       align="justify", indent=1.25, after=2)
    if item is not None and (item.cause or "").strip():
        _p(doc, item.cause.strip(), align="justify", indent=1.25, after=2)
    for _ in range(8):
        _p(doc, _DOT + _DOT[:30], align="justify", after=1)
    _p(doc, "อ่านแล้วรับรองว่าถูกต้อง จึงได้ลงลายมือชื่อไว้เป็นหลักฐาน",
       align="justify", indent=1.25, before=2, after=10)
    _sign_table(doc, [[
        ("ลงชื่อ ...................................... ผู้ให้ถ้อยคำ", "center"),
        (f"( {_BLANK} )", "center"),
    ]], after=2)
    _sign_committee(doc, dp, "fact")
    return _save(doc, f"บันทึกถ้อยคำสอบข้อเท็จจริง_ปีงบ{dp.year}") if own else doc


# ================================================================== ขั้นที่ 2
def render_dispose_request(school, dp, doc=None):
    """(5) บันทึกข้อความ ขอจำหน่ายพัสดุ + เสนอแต่งตั้ง 3 คณะกรรมการ (ข้อ 215)"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    _memo_header(doc, school, "ขอจำหน่ายพัสดุ", dp.req_memo_no, _d(dp.req_memo_date))
    _p(doc, "ตามที่ได้สั่งการให้ดำเนินการจำหน่ายพัสดุที่ชำรุด เสื่อมสภาพ สูญไป หรือหมด"
            f"ความจำเป็นต้องใช้ในหน่วยงานของรัฐ ตาม{_REG} ข้อ 215 (1) (2) (3) และ (4) "
            "ดังความละเอียดทราบแล้ว นั้น", align="justify", indent=1.25, after=2)
    _p(doc, "ได้ตรวจสอบพัสดุตามรายงานของคณะกรรมการสอบหาข้อเท็จจริง "
            f"ตามบันทึกที่ {dp.fact_report_no or _BLANK} ลงวันที่ {_d(dp.fact_report_date)} "
            "แล้ว ปรากฏว่าเป็นการชำรุดเสื่อมสภาพเนื่องมาจากการใช้งานตามปกติ หรือสูญไปตาม"
            "ธรรมชาติ หากใช้ในหน่วยงานของรัฐต่อไปจะสิ้นเปลืองค่าใช้จ่ายมาก เห็นสมควร"
            "ดำเนินการ ดังนี้", align="justify", indent=1.25, after=2)
    rows = _items(dp)
    total, _ = _item_table(doc, rows)
    _p(doc, "", after=3)
    for way, clause, _desc in _DISPOSE_WAYS:
        n = len(_items(dp, "แปรสภาพ" if way == "แปรสภาพหรือทำลาย" else way))
        if way == "แปรสภาพหรือทำลาย":
            n += len(_items(dp, "ทำลาย"))
        if n:
            _p(doc, f"โดยวิธีการ{way} จำนวน {n} รายการ ตามระเบียบฯ {clause}",
               indent=2.0, after=1)
    n_lost = len(_items(dp, "จำหน่ายเป็นสูญ"))
    if n_lost:
        _p(doc, f"จำหน่ายเป็นสูญ จำนวน {n_lost} รายการ ตามระเบียบฯ ข้อ 217 (1)",
           indent=2.0, after=1)
    _p(doc, f"รวมพัสดุที่ขอจำหน่ายทั้งสิ้น {len(rows)} รายการ ราคาทุนรวม {total:,.2f} บาท "
            f"({bahttext(total)})", align="justify", indent=1.25, before=2, after=2)
    _p(doc, "ฉะนั้น เพื่อให้การจำหน่ายพัสดุดำเนินการเป็นไปด้วยความเรียบร้อย ถูกต้อง "
            "และเป็นประโยชน์ต่อทางราชการอย่างสูงสุด จึงขอแต่งตั้งบุคคลผู้มีรายนามข้างท้ายนี้"
            "เป็นคณะกรรมการดำเนินการ", align="justify", indent=1.25, after=2)
    _committee_block(doc, dp, "price", "คณะกรรมการประเมินราคากลาง")
    _committee_block(doc, dp, "sale", "คณะกรรมการดำเนินการขาย")
    _committee_block(doc, dp, "destroy", "คณะกรรมการทำลาย")
    _p(doc, "จึงเรียนมาเพื่อโปรดพิจารณา", indent=1.25, before=2, after=10)
    _officer_sign(doc, school)
    _director_sign(doc, school, lines=["- ชอบ", "- ดำเนินการ"])
    _p(doc, "วันที่ ........./............./..................", align="center", after=2)
    return _save(doc, f"ขอจำหน่ายพัสดุ_ปีงบ{dp.year}") if own else doc


def render_committee_order(school, dp, doc=None):
    """(6) คำสั่งแต่งตั้งคณะกรรมการประเมินราคากลาง คณะกรรมการดำเนินการขาย และคณะกรรมการทำลาย"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    _krut_center(doc)
    _p(doc, f"คำสั่ง{_sname(school)}", align="center", bold=True, size=18, after=0)
    _p(doc, f"ที่ {dp.order_no or _BLANK}", align="center", bold=True, after=0)
    _p(doc, "เรื่อง แต่งตั้งคณะกรรมการประเมินราคากลาง คณะกรรมการดำเนินการขาย "
            "และคณะกรรมการทำลาย", align="center", bold=True, size=17, after=0)
    _p(doc, "-----------------------------------", align="center", after=6)
    _p(doc, f"ด้วย{_sname(school)} มีความประสงค์จะทำการจำหน่ายพัสดุที่ชำรุด เสื่อมสภาพ "
            "หรือหมดความจำเป็น หรือหากใช้ในหน่วยงานของรัฐต่อไปจะสิ้นเปลืองค่าใช้จ่ายมาก "
            f"ตามที่ได้รับอนุมัติจากผู้มีอำนาจ {_DELEG} ดังนี้",
       align="justify", indent=1.25, after=2)
    for label, act in [("ขายพัสดุ", "ขาย"), ("ทำลายพัสดุ", "ทำลาย"),
                       ("โอนพัสดุ", "โอน"), ("แปรสภาพพัสดุ", "แปรสภาพ")]:
        n = len(_items(dp, act))
        if n:
            _p(doc, f"{label}\tจำนวน {n} รายการ", indent=2.0, after=1)
    _p(doc, "ฉะนั้น เพื่อให้การดำเนินการจำหน่ายพัสดุดำเนินการเป็นไปด้วยความเรียบร้อย ถูกต้อง "
            "และเป็นประโยชน์ต่อทางราชการอย่างสูงสุด จึงแต่งตั้งบุคคลต่อไปนี้เป็นคณะกรรมการ"
            "ดำเนินการ", align="justify", indent=1.25, before=2, after=2)
    _committee_block(doc, dp, "price", "คณะกรรมการประเมินราคากลาง")
    _committee_block(doc, dp, "sale", "คณะกรรมการดำเนินการขาย")
    _committee_block(doc, dp, "destroy", "คณะกรรมการทำลาย")
    _p(doc, "ให้คณะกรรมการที่ได้รับแต่งตั้งตามคำสั่งนี้ ปฏิบัติหน้าที่ให้บังเกิดผลดี"
            "ต่อทางราชการอย่างสูงสุดโดยเคร่งครัด", align="justify", indent=1.25,
       before=2, after=2)
    _p(doc, f"สั่ง ณ วันที่ {_d(dp.order_date)}", indent=1.25, after=12)
    _sign_table(doc, [[
        ("ลงชื่อ ......................................", "center"),
        (f"( {(school.director_name or '').strip() or _BLANK} )", "center"),
        (_director_line(school), "center"),
    ]])
    return _save(doc, f"คำสั่งแต่งตั้งกรรมการจำหน่ายพัสดุ_ปีงบ{dp.year}") if own else doc


# ========================================= ขั้นที่ 3 (ก) ขายโดยวิธีเฉพาะเจาะจง
def render_invite_quote(school, dp, doc=None):
    """(7) หนังสือขอเชิญเสนอราคาซื้อพัสดุ (ขายโดยวิธีเฉพาะเจาะจง) - หนังสือภายนอก"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    _letter_header(doc, school, dp.invite_no, _d(dp.invite_date),
                   "ขอเชิญเสนอราคาซื้อพัสดุ",
                   (dp.invite_to or "").strip() or _BLANK,
                   "รายละเอียดพัสดุที่จะจำหน่ายโดยการขายด้วยวิธีเฉพาะเจาะจง จำนวน 1 ชุด")
    rows = _items(dp, "ขาย")
    _p(doc, f"ด้วย{_sname(school)} มีความประสงค์จะจำหน่ายพัสดุที่ชำรุด เสื่อมสภาพ หรือหมด"
            f"ความจำเป็นต้องใช้ในหน่วยงานของรัฐ จำนวน {len(rows)} รายการ โดยการขายด้วยวิธี"
            f"เฉพาะเจาะจง ตาม{_REG} ข้อ 215 (1) รายละเอียดตามสิ่งที่ส่งมาด้วยพร้อมนี้",
       align="justify", indent=1.25, after=2)
    _p(doc, f"จึงขอเชิญ {(dp.invite_to or '').strip() or _BLANK} เข้าร่วมเสนอราคากับ"
            f"{_sname(school)} ในวันที่ {_d(dp.quote_open)} "
            f"ระหว่างเวลา {(dp.quote_time or '').strip() or _BLANK}",
       align="justify", indent=1.25, after=2)
    _p(doc, "จึงเรียนมาเพื่อทราบและเข้าร่วมเสนอราคาตามวัน เวลา และสถานที่ดังกล่าวต่อไป",
       align="justify", indent=1.25, after=8)
    _p(doc, "ขอแสดงความนับถือ", align="center", after=12)
    _sign_table(doc, [[
        ("", "center"),
        (f"( {(school.director_name or '').strip() or _BLANK} )", "center"),
        (_director_line(school), "center"),
    ]])
    _p(doc, "กลุ่มบริหารงบประมาณ", size=14, before=8, after=0)
    _p(doc, f"โทร. {(dp.contact_phone or getattr(school, 'phone', '') or '').strip() or _BLANK}",
       size=14, after=6)
    # แนบรายละเอียดพัสดุ
    _p(doc, "รายละเอียดพัสดุที่จะจำหน่ายโดยการขายด้วยวิธีเฉพาะเจาะจง",
       align="center", bold=True, size=16, before=6, after=4)
    _item_table(doc, rows)
    return _save(doc, f"ขอเชิญเสนอราคาซื้อพัสดุ_ปีงบ{dp.year}") if own else doc


def render_quote_form(school, dp, doc=None):
    """(8) ใบเสนอราคา (ให้ผู้ซื้อกรอก) - เว้นข้อมูลผู้เสนอราคาไว้กรอกเอง"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    _p(doc, "ใบเสนอราคา", align="center", bold=True, size=20, after=6)
    _p_runs(doc, [("เรียน  ", True), (_director_line(school), False)])
    _p(doc, "", after=2)
    _p(doc, f"ข้าพเจ้า {_BLANK} ห้าง/ร้าน/บริษัท/หจก. {_BLANK}", align="justify",
       indent=1.25, after=1)
    _p(doc, "ตั้งอยู่เลขที่ .................. ถนน .......................... "
            "ตำบล/แขวง .......................... อำเภอ/เขต ..........................",
       align="justify", after=1)
    _p(doc, "จังหวัด .......................... เลขประจำตัวผู้เสียภาษี "
            "..................................", align="justify", after=2)
    _p(doc, f"ข้าพเจ้าได้อ่านและเข้าใจข้อความในการขายพัสดุที่ชำรุด เสื่อมสภาพ หรือหมด"
            f"ความจำเป็นต้องใช้ในหน่วยงานของรัฐ ของ{_sname(school)} เป็นอย่างดีแล้ว "
            "จึงขอเสนอราคาซื้อพัสดุที่ทางราชการประกาศขาย ดังรายการต่อไปนี้",
       align="justify", indent=1.25, after=3)

    rows = _items(dp, "ขาย")
    heads = ["ที่", "รายการ", "จำนวน", "ราคา/หน่วย", "จำนวนเงิน", "หมายเหตุ"]
    widths = [Cm(1.1), Cm(5.6), Cm(2.0), Cm(2.6), Cm(2.8), Cm(2.4)]
    t = doc.add_table(rows=1, cols=len(heads))
    t.style = "Table Grid"; t.autofit = False
    _fixed_cols(t, widths)
    _repeat_header_row(t.rows[0]); _no_split_row(t.rows[0])
    for c, h, w in zip(t.rows[0].cells, heads, widths):
        _set_cell(c, h, bold=True, align="center", size=13)
        c.width = w
    for i, it in enumerate(rows or [None] * 5, 1):
        row = t.add_row(); _no_split_row(row)
        if it is not None:
            a = it.asset
            vals = [str(i), a.name or "-", f"{(a.quantity or 1):g} {a.unit or 'หน่วย'}", "", "", ""]
        else:
            vals = ["", "", "", "", "", ""]
        for c, v, al, w in zip(row.cells, vals,
                               ["center", "left", "center", "right", "right", "left"], widths):
            _set_cell(c, v, align=al, size=13)
            c.width = w
    trow = t.add_row(); _no_split_row(trow)
    _set_cell(trow.cells[0], "รวมเป็นเงิน  ( ..................................................... )",
              bold=True, align="right", size=13)
    trow.cells[0].merge(trow.cells[3])
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = w

    _p(doc, "", after=3)
    _p(doc, "คำเสนอนี้จะยืนอยู่เป็นระยะเวลา .......... วัน นับแต่วันที่ได้ยื่นใบเสนอราคา",
       align="justify", indent=1.25, after=1)
    _p(doc, "กำหนดรับมอบพัสดุตามรายการข้างต้น ภายใน 7 วัน นับถัดจากวันที่ข้าพเจ้าได้รับ"
            "การพิจารณาให้เป็นผู้ชนะในการเสนอราคา", align="justify", indent=1.25, after=2)
    _p(doc, f"เสนอมา ณ วันที่ {_BLANK}", align="justify", indent=1.25, after=10)
    _sign_table(doc, [[
        ("ลงชื่อ ...................................... ผู้เสนอราคา", "center"),
        (f"( {_BLANK} )", "center"),
        ("ประทับตรา (ถ้ามี)", "center"),
    ]])
    return _save(doc, f"ใบเสนอราคาซื้อพัสดุ_ปีงบ{dp.year}") if own else doc


def render_price_report(school, dp, doc=None):
    """(9) บันทึกข้อความ การประเมินราคากลางพัสดุ"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    _memo_header(doc, school, "การประเมินราคากลางพัสดุ",
                 dp.price_memo_no, _d(dp.price_memo_date))
    _p(doc, "ตามที่ได้แต่งตั้งคณะกรรมการประเมินราคากลางพัสดุที่ชำรุด เสื่อมสภาพ หรือหมด"
            "ความจำเป็นต้องใช้ในหน่วยงานของรัฐต่อไป เพื่อขายพัสดุให้กับบุคคลหรือผู้สนใจ"
            f"ที่เสนอราคาอันเป็นประโยชน์ต่อทางราชการ ตามคำสั่ง{_sname(school)} "
            f"ที่ {dp.order_no or _BLANK} ลงวันที่ {_d(dp.order_date)} "
            "ดังความละเอียดทราบแล้ว นั้น", align="justify", indent=1.25, after=2)
    rows = _items(dp, "ขาย")
    _p(doc, "คณะกรรมการประเมินราคากลางได้ร่วมกันพิจารณาแล้ว มีมติเห็นสมควรประเมินราคากลาง"
            "พัสดุ เพื่อให้คณะกรรมการดำเนินการขายใช้เป็นเกณฑ์ในการพิจารณาขาย ดังนี้",
       align="justify", indent=1.25, after=3)
    _, mid = _item_table(doc, rows, price_col=True)
    _p(doc, f"รวมราคากลางที่ประเมินทั้งสิ้น {mid:,.2f} บาท ({bahttext(mid)})",
       bold=True, indent=1.25, before=3, after=2)
    _p(doc, "จึงเรียนมาเพื่อโปรดทราบ", indent=1.25, after=10)
    _sign_committee(doc, dp, "price")
    _director_sign(doc, school, lines=[
        "- ทราบ",
        "- มอบหัวหน้าเจ้าหน้าที่ แล้วให้มอบคณะกรรมการดำเนินการขายในเวลาพิจารณาการขาย"])
    _p(doc, "วันที่ ........./............./..................", align="center", after=2)
    return _save(doc, f"ประเมินราคากลางพัสดุ_ปีงบ{dp.year}") if own else doc


def _bidder_rows(dp):
    """ผู้เสนอราคา [(ชื่อ, ราคา)] จากช่องกรอกบรรทัดละ 'ชื่อ|ราคา'"""
    out = []
    for line in (dp.bidders or "").splitlines():
        line = line.strip()
        if not line:
            continue
        if "|" in line:
            nm, _, pr = line.partition("|")
            try:
                out.append((nm.strip(), float(pr.strip().replace(",", ""))))
            except ValueError:
                out.append((nm.strip(), 0.0))
        else:
            out.append((line, 0.0))
    return out


def render_sale_report(school, dp, doc=None):
    """(10) บันทึกข้อความ รายงานผลการดำเนินการขาย (พร้อมบล็อกอนุมัติ + ลงจ่ายออกจากทะเบียน)"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    _memo_header(doc, school, "รายงานผลการดำเนินการขาย", dp.sale_memo_no, _d(dp.sale_memo_date))
    rows = _items(dp, "ขาย")
    _p(doc, f"ตามที่{_sname(school)} ได้ดำเนินการขายพัสดุที่ชำรุด เสื่อมสภาพ หรือหมด"
            f"ความจำเป็นต้องใช้ในหน่วยงานของรัฐ จำนวน {len(rows)} รายการ โดยวิธีเฉพาะเจาะจง "
            f"เมื่อวันที่ {_d(dp.sale_date)} นั้น", align="justify", indent=1.25, after=2)
    bidders = _bidder_rows(dp)
    _p(doc, "คณะกรรมการดำเนินการขายได้ร่วมกันพิจารณาราคาของผู้เสนอราคาซื้อแล้ว ปรากฏว่า"
            f"มีผู้เสนอราคาซื้อ จำนวน {len(bidders) or 1} ราย คือ",
       align="justify", indent=1.25, after=1)
    for i, (nm, pr) in enumerate(bidders or [(dp.buyer_name or _BLANK, dp.sale_total or 0)], 1):
        _p(doc, f"{i}. {nm or _BLANK}  เสนอราคาซื้อเป็นเงิน {pr:,.2f} บาท",
           indent=2.0, after=1)
    buyer = (dp.buyer_name or "").strip() or _BLANK
    total = float(dp.sale_total or 0)
    _p(doc, f"ผลปรากฏว่า {buyer} เสนอราคาอันเป็นประโยชน์ต่อทางราชการอย่างสูงสุด "
            "คณะกรรมการดำเนินการขายจึงมีมติเห็นสมควรขายพัสดุที่ชำรุด เสื่อมสภาพ หรือหมด"
            f"ความจำเป็นต้องใช้ในหน่วยงานของรัฐ ให้กับ {buyer} เป็นเงิน {total:,.2f} บาท "
            f"({bahttext(total)}) ทั้งนี้ คณะกรรมการฯ ได้นำราคาประเมินของคณะกรรมการ"
            "ประเมินราคากลางมาใช้เป็นเกณฑ์แล้ว", align="justify", indent=1.25,
       before=2, after=2)
    if (dp.buyer_address or "").strip():
        _p(doc, f"ผู้ซื้ออยู่ที่ {dp.buyer_address.strip()}"
                + (f" เลขประจำตัวผู้เสียภาษี {dp.buyer_taxid.strip()}"
                   if (dp.buyer_taxid or "").strip() else ""),
           align="justify", indent=1.25, after=2)
    _p(doc, "จึงเรียนมาเพื่อโปรดพิจารณาอนุมัติ", indent=1.25, after=10)
    _sign_committee(doc, dp, "sale")
    _p(doc, "", after=4)
    _p_runs(doc, [("เรียน  ", True), (_director_line(school), False)])
    _p(doc, f"เห็นควรอนุมัติขายพัสดุฯ ให้กับ {buyer} ในราคา {total:,.2f} บาท",
       indent=1.25, after=8)
    _officer_sign(doc, school, head=True)
    _director_sign(doc, school,
                   lines=["- อนุมัติ", "- ลงจ่ายออกจากบัญชีหรือทะเบียนพัสดุ",
                          f"- นำเงินส่งเป็นเงินรายได้{(dp.revenue_kind or 'แผ่นดิน').strip()}",
                          "- รายงาน สตง. ตามระเบียบฯ"])
    _p(doc, "วันที่ ........./............./..................", align="center", after=2)
    return _save(doc, f"รายงานผลการขายพัสดุ_ปีงบ{dp.year}") if own else doc


# ========================================= ขั้นที่ 3 (ข) ขายโดยวิธีทอดตลาด
def render_auction_order(school, dp, doc=None):
    """(11) คำสั่งแต่งตั้งคณะกรรมการขายพัสดุโดยวิธีทอดตลาด"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    _krut_center(doc)
    _p(doc, f"คำสั่ง{_sname(school)}", align="center", bold=True, size=18, after=0)
    _p(doc, f"ที่ {dp.auction_order_no or _BLANK}", align="center", bold=True, after=0)
    _p(doc, "เรื่อง แต่งตั้งคณะกรรมการขายพัสดุโดยวิธีทอดตลาด",
       align="center", bold=True, size=17, after=0)
    _p(doc, "-----------------------------------", align="center", after=6)
    rows = _items(dp, "ขาย")
    _p(doc, f"ด้วย{_sname(school)} สังกัด{(school.area_office or '').strip() or _BLANK} "
            "จะทำการประมูลด้วยวาจาในการขายพัสดุโดยวิธีทอดตลาด "
            f"จำนวน {len(rows)} รายการ (รายละเอียดตามบัญชีแนบท้ายคำสั่งนี้) "
            f"ตาม{_REG} ข้อ 215 (1) จึงแต่งตั้งคณะกรรมการขายพัสดุโดยวิธีทอดตลาด ดังนี้",
       align="justify", indent=1.25, after=2)
    _member_lines(doc, _members(dp, "auction") or _members(dp, "sale"))
    _p(doc, "ให้คณะกรรมการที่ได้รับแต่งตั้งตามคำสั่งนี้ ปฏิบัติหน้าที่โดยเคร่งครัด",
       align="justify", indent=1.25, before=2, after=2)
    _p(doc, f"สั่ง ณ วันที่ {_d(dp.auction_order_date)}", indent=1.25, after=12)
    _sign_table(doc, [[
        ("ลงชื่อ ......................................", "center"),
        (f"( {(school.director_name or '').strip() or _BLANK} )", "center"),
        (_director_line(school), "center"),
    ]])
    return _save(doc, f"คำสั่งแต่งตั้งกรรมการขายทอดตลาด_ปีงบ{dp.year}") if own else doc


def render_auction_notice(school, dp, doc=None):
    """(12) ประกาศขายพัสดุโดยวิธีทอดตลาด (เงื่อนไข 4-5 ข้อตามแบบ)"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    _krut_center(doc)
    _p(doc, f"ประกาศ{_sname(school)}", align="center", bold=True, size=18, after=0)
    _p(doc, "เรื่อง การขายพัสดุโดยวิธีทอดตลาด", align="center", bold=True, size=17, after=0)
    _p(doc, "-----------------------------------", align="center", after=6)
    rows = _items(dp, "ขาย")
    place = (dp.auction_place or "").strip() or _BLANK
    _p(doc, f"ด้วย{_sname(school)} จะทำการขายพัสดุโดยวิธีทอดตลาด จำนวน {len(rows)} รายการ "
            f"ตามรายละเอียดแนบท้ายประกาศนี้ โดยกำหนดดูพัสดุ ณ {_sname(school)} "
            f"ในวันที่ {_d(dp.view_date)} เวลา {(dp.view_time or '').strip() or _BLANK} "
            f"และกำหนดขายทอดตลาดในวันที่ {_d(dp.auction_date)} "
            f"เวลา {(dp.auction_time or '').strip() or _BLANK} ณ {place}",
       align="justify", indent=1.25, after=2)
    _p(doc, "เงื่อนไขการขายทอดตลาด มีดังนี้", bold=True, indent=1.25, after=1)
    conds = [
        "การขายทอดตลาดจะขายโดยวิธีประมูลด้วยวาจา",
        "ผู้ประมูลราคาได้จะต้องจ่ายเงินสดทันที กรณีที่พัสดุมีราคาประมูลเกินกว่า 2,000.- บาท "
        "หากผู้ประมูลได้ไม่สามารถชำระเงินสดในคราวเดียวได้ ให้วางเงินสดไว้ไม่น้อยกว่าร้อยละ 25 "
        "ของราคาที่ประมูลได้ ส่วนที่เหลือต้องชำระให้ครบถ้วนภายใน 3 วัน นับตั้งแต่วันที่ประมูลได้ "
        f"หากไม่ชำระให้ครบถ้วนภายในกำหนดเวลาดังกล่าว จะถือว่าสละสิทธิ์ {_sname(school)} "
        "จะริบเงินที่วางไว้ แล้วจะขายทอดตลาดใหม่ต่อไป",
        "การตัดสินชี้ขาดในเรื่องตัวบุคคลผู้ประมูลได้ และราคาในการขายทอดตลาดพัสดุแต่ละรายการนั้น "
        "ใช้วิธีเคาะไม้",
        f"ผู้ประมูลได้จะต้องรับพัสดุนั้นไปจาก{_sname(school)} ให้เสร็จภายใน 3 วัน "
        "นับแต่วันชำระเงินครบถ้วน หากล่วงเลยกำหนดเวลาดังกล่าว โรงเรียนจะไม่รับผิดชอบ"
        "ต่อความเสียหายอันอาจจะเกิดขึ้นได้",
    ]
    if dp.auction_fee:
        conds.append("ค่าธรรมเนียมในการโอนกรรมสิทธิ์ของพัสดุที่ขายโดยวิธีทอดตลาด "
                     "ผู้ประมูลได้เป็นผู้ชำระ")
    for i, txt in enumerate(conds, 1):
        _p(doc, f"{i}. {txt}", align="justify", indent=2.0, after=1)
    phone = (dp.contact_phone or getattr(school, "phone", "") or "").strip() or _BLANK
    _p(doc, f"ผู้สนใจติดต่อสอบถามรายละเอียดได้ที่ {_sname(school)} "
            f"โทรศัพท์หมายเลข {phone} ตั้งแต่บัดนี้เป็นต้นไป",
       align="justify", indent=1.25, before=2, after=2)
    _p(doc, f"ประกาศ ณ วันที่ {_d(dp.notice_date)}", align="center", after=12)
    _sign_table(doc, [[
        ("ลงชื่อ ......................................", "center"),
        (f"( {(school.director_name or '').strip() or _BLANK} )", "center"),
        (_director_line(school), "center"),
    ]])
    return _save(doc, f"ประกาศขายทอดตลาด_ปีงบ{dp.year}") if own else doc


def render_auction_list(school, dp, doc=None):
    """(13) บัญชีรายการพัสดุที่จะทำการขายโดยวิธีทอดตลาด (แนบท้ายประกาศ)"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    _p(doc, "บัญชีรายการพัสดุที่จะทำการขายโดยวิธีทอดตลาด",
       align="center", bold=True, size=17, after=0)
    _p(doc, f"ในวันที่ {_d(dp.auction_date)}", align="center", after=0)
    _p(doc, f"แนบท้ายประกาศ{_sname(school)} ลงวันที่ {_d(dp.notice_date)}",
       align="center", after=6)
    rows = _items(dp, "ขาย")
    heads = ["ลำดับที่", "ชื่อ ขนาด ลักษณะ", "จำนวน", "ราคากลาง (บาท)", "หมายเหตุ"]
    widths = [Cm(1.8), Cm(6.4), Cm(2.2), Cm(3.0), Cm(3.1)]
    t = doc.add_table(rows=1, cols=len(heads))
    t.style = "Table Grid"; t.autofit = False
    _fixed_cols(t, widths)
    _repeat_header_row(t.rows[0]); _no_split_row(t.rows[0])
    for c, h, w in zip(t.rows[0].cells, heads, widths):
        _set_cell(c, h, bold=True, align="center", size=13)
        c.width = w
    mid = 0.0
    for i, it in enumerate(rows, 1):
        a = it.asset
        mid += float(it.price_mid or 0)
        desc = " ".join(x for x in [(a.name or "").strip(), (a.brand_model or "").strip()] if x)
        if (a.asset_code or "").strip():
            desc += f"  (เลขครุภัณฑ์ {a.asset_code.strip()})"
        row = t.add_row(); _no_split_row(row)
        for c, v, al, w in zip(row.cells,
                               [str(i), desc, f"{(a.quantity or 1):g} {a.unit or 'หน่วย'}",
                                f"{float(it.price_mid or 0):,.2f}", ""],
                               ["center", "left", "center", "right", "left"], widths):
            _set_cell(c, v, align=al, size=13)
            c.width = w
    trow = t.add_row(); _no_split_row(trow)
    _set_cell(trow.cells[0], "รวมราคากลางทั้งสิ้น", bold=True, align="right", size=13)
    trow.cells[0].merge(trow.cells[2])
    _set_cell(trow.cells[3], f"{mid:,.2f}", bold=True, align="right", size=13)
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = w
    _p(doc, "", after=8)
    _sign_committee(doc, dp, "auction" if _members(dp, "auction") else "sale")
    return _save(doc, f"บัญชีพัสดุขายทอดตลาด_ปีงบ{dp.year}") if own else doc


def render_auction_ledger(school, dp, doc=None):
    """(14) บัญชีคุมพัสดุที่ขายโดยวิธีทอดตลาด (ครั้งที่ 1 / ครั้งที่ 2)

    8 คอลัมน์ กว้างเกินหน้ากระดาษตั้ง จึงพิมพ์เป็นหน้านอน
    """
    own = doc is None
    if own:
        doc = Document(); set_a4(doc, landscape=True); _font(doc)
    else:
        _landscape_section(doc)
    _p(doc, "บัญชีคุมพัสดุที่ขายโดยวิธีทอดตลาด", align="center", bold=True, size=17, after=0)
    _p(doc, _sname(school) + f"  ปีงบประมาณ {dp.year}", align="center", after=6)
    rows = _items(dp, "ขาย")
    widths = [Cm(1.2), Cm(6.5), Cm(3.6), Cm(3.0), Cm(3.0), Cm(3.6), Cm(3.0), Cm(2.8)]
    n_rows = max(len(rows), 8)
    t = doc.add_table(rows=2 + n_rows, cols=8)
    t.style = "Table Grid"; t.autofit = False
    _fixed_cols(t, widths)
    r0, r1 = t.rows[0], t.rows[1]
    _set_cell(r0.cells[0].merge(r1.cells[0]), "ที่", bold=True, align="center", size=13)
    _set_cell(r0.cells[1].merge(r1.cells[1]), "ชื่อพัสดุ ขนาด ลักษณะ",
              bold=True, align="center", size=13)
    _set_cell(r0.cells[2].merge(r0.cells[4]), "ขายทอดตลาดครั้งที่ 1",
              bold=True, align="center", size=13)
    _set_cell(r0.cells[5].merge(r0.cells[7]), "ขายทอดตลาดครั้งที่ 2",
              bold=True, align="center", size=13)
    subs = ["ชื่อผู้ประมูลได้", "เลขที่ใบเสร็จรับเงิน", "ลายมือชื่อผู้รับรอง"] * 2
    for i, lab in enumerate(subs, start=2):
        _set_cell(r1.cells[i], lab, bold=True, align="center", size=12)
    _repeat_header_row(r0); _repeat_header_row(r1)
    _no_split_row(r0); _no_split_row(r1)
    for n in range(n_rows):
        row = t.rows[2 + n]
        _no_split_row(row)
        it = rows[n] if n < len(rows) else None
        if it is not None:
            a = it.asset
            desc = " ".join(x for x in [(a.name or "").strip(),
                                        (a.brand_model or "").strip()] if x)
            vals = [str(n + 1), desc, (it.disposal.buyer_name or "").strip(),
                    (it.receipt_no or "").strip(), "", "", "", ""]
        else:
            vals = [""] * 8
        for c, v, w in zip(row.cells, vals, widths):
            _set_cell(c, v, align="left" if w > Cm(3) else "center", size=12)
            c.width = w
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = w
    return _save(doc, f"บัญชีคุมพัสดุขายทอดตลาด_ปีงบ{dp.year}") if own else doc


def render_auction_report(school, dp, doc=None):
    """(15) บันทึกข้อความ รายงานผลการขายพัสดุโดยวิธีทอดตลาด"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    rows = _items(dp, "ขาย")
    _memo_header(doc, school,
                 f"รายงานผลการขายพัสดุโดยวิธีทอดตลาด จำนวน {len(rows)} รายการ",
                 dp.auction_report_no, _d(dp.auction_report_date))
    _p(doc, f"ตามคำสั่ง{_sname(school)} ที่ {dp.auction_order_no or _BLANK} "
            f"สั่ง ณ วันที่ {_d(dp.auction_order_date)} ได้แต่งตั้งคณะกรรมการขายพัสดุ"
            f"โดยวิธีทอดตลาด จำนวน {len(rows)} รายการ นั้น",
       align="justify", indent=1.25, after=2)
    bidders = _bidder_rows(dp)
    mid = sum(float(it.price_mid or 0) for it in rows)
    buyer = (dp.buyer_name or "").strip() or _BLANK
    total = float(dp.sale_total or 0)
    _p(doc, "คณะกรรมการขายพัสดุโดยวิธีทอดตลาดได้ดำเนินการเมื่อวันที่ "
            f"{_d(dp.auction_date)} ตั้งแต่เวลา {(dp.auction_time or '').strip() or _BLANK} "
            f"ณ {(dp.auction_place or '').strip() or _BLANK} "
            f"โดยมีผู้เข้าร่วมประมูล จำนวน {len(bidders) or 1} ราย "
            f"โดยคณะกรรมการได้แจ้งราคากลางและเป็นราคาเริ่มต้นแก่ผู้เข้าร่วมประมูล "
            f"เป็นเงิน {mid:,.2f} บาท แล้วปรากฏว่า", align="justify", indent=1.25, after=1)
    for i, (nm, pr) in enumerate(bidders or [(buyer, total)], 1):
        _p(doc, f"{i}. {nm or _BLANK}  ประมูลราคาเป็นเงิน {pr:,.2f} บาท",
           indent=2.0, after=1)
    _p(doc, f"ซึ่งต่อมาไม่มีผู้เข้าร่วมประมูลรายอื่นเสนอราคารับซื้อสูงขึ้นอีก คณะกรรมการ"
            f"จึงเห็นว่า {buyer} เสนอราคารับซื้อในราคาที่ไม่ต่ำกว่าราคากลางที่กำหนด",
       align="justify", indent=1.25, before=2, after=2)
    _p(doc, f"จึงได้นับ 1 ถึง 3 แล้วเคาะไม้เป็นสัญญาณการตัดสินขายพัสดุ จำนวน {len(rows)} "
            f"รายการ ให้แก่ {buyer} "
            f"อยู่ที่ {(dp.buyer_address or '').strip() or _BLANK} "
            + (f"เลขประจำตัวผู้เสียภาษี {dp.buyer_taxid.strip()} "
               if (dp.buyer_taxid or "").strip() else "")
            + f"ในราคา {total:,.2f} บาท ({bahttext(total)}) "
            "ดังสำเนาหลักฐานที่แนบ", align="justify", indent=1.25, after=2)
    _p(doc, "หมายเหตุ  กรณีขายพัสดุโดยวิธีทอดตลาดแยกรายการ ให้ระบุราคากลางในแต่ละรายการ "
            "และประมูลแต่ละรายการ", size=13, indent=1.25, after=8)
    _sign_committee(doc, dp, "auction" if _members(dp, "auction") else "sale")
    _director_sign(doc, school, lines=[
        "- อนุมัติ", "- ลงจ่ายออกจากบัญชีหรือทะเบียนพัสดุ",
        f"- นำเงินส่งเป็นเงินรายได้{(dp.revenue_kind or 'แผ่นดิน').strip()}",
        "- รายงาน สตง. ตามระเบียบฯ"])
    _p(doc, "วันที่ ........./............./..................", align="center", after=2)
    return _save(doc, f"รายงานผลการขายทอดตลาด_ปีงบ{dp.year}") if own else doc


# ========================================= ขั้นที่ 3 (ค) ทำลาย
def render_destroy_report(school, dp, doc=None):
    """(16) บันทึกข้อความ รายงานผลการทำลายพัสดุของคณะกรรมการทำลาย"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    rows = _items(dp, "ทำลาย")
    _memo_header(doc, school, "รายงานผลการทำลายพัสดุของคณะกรรมการทำลาย",
                 dp.destroy_memo_no, _d(dp.destroy_memo_date))
    _p(doc, "ตามที่ได้แต่งตั้งคณะกรรมการทำลาย เพื่อทำลายพัสดุที่ไม่สามารถนำมาขาย "
            "แลกเปลี่ยน โอน หรือแปรสภาพได้ เนื่องจากไม่สามารถนำไปซ่อมแซมให้ใช้การได้ดี "
            "อีกทั้งยังเสื่อมสภาพ ซึ่งในการนี้คณะกรรมการสอบหาข้อเท็จจริงได้ตรวจสอบและร่วมกัน"
            f"พิจารณาแล้ว ตามคำสั่ง{_sname(school)} ที่ {dp.order_no or _BLANK} "
            f"ลงวันที่ {_d(dp.order_date)} ดังความละเอียดทราบแล้ว นั้น",
       align="justify", indent=1.25, after=2)
    _p(doc, f"บัดนี้ คณะกรรมการทำลายได้ดำเนินการทำลายพัสดุ จำนวน {len(rows)} รายการ "
            "ตามที่ได้อนุมัติไว้แล้วแต่ต้น เสร็จเรียบร้อยแล้ว โดยวิธีการ"
            f"{(dp.destroy_way or '').strip() or _BLANK} "
            f"เมื่อวันที่ {_d(dp.destroy_date)}", align="justify", indent=1.25, after=3)
    if rows:
        _item_table(doc, rows)
        _p(doc, "", after=3)
    _p(doc, "จึงเรียนมาเพื่อโปรดทราบ", indent=1.25, after=10)
    _sign_committee(doc, dp, "destroy")
    _p(doc, "", after=4)
    _p_runs(doc, [("เรียน  ", True), (_director_line(school), False)])
    _p(doc, "- เพื่อโปรดทราบ", indent=1.25, after=1)
    _p(doc, "- ดำเนินการตามระเบียบฯ ต่อไป", indent=1.25, after=8)
    _officer_sign(doc, school)
    _director_sign(doc, school, lines=[
        "- อนุมัติ", "- ลงจ่ายออกจากบัญชีหรือทะเบียนพัสดุ", "- รายงาน สตง. ตามระเบียบฯ"])
    _p(doc, "วันที่ ........./............./..................", align="center", after=2)
    return _save(doc, f"รายงานผลการทำลายพัสดุ_ปีงบ{dp.year}") if own else doc


# ================================================================== ขั้นที่ 4
_SAO_ENCL = [
    "คำสั่งแต่งตั้งคณะกรรมการตรวจสอบพัสดุประจำปี และสำเนาคำสั่งแต่งตั้งเจ้าหน้าที่ "
    "และหัวหน้าเจ้าหน้าที่ อย่างละ 1 ฉบับ",
    "รายงานผลการตรวจสอบพัสดุประจำปี 1 ฉบับ",
    "คำสั่งแต่งตั้งคณะกรรมการสอบหาข้อเท็จจริง 2 ฉบับ",
    "รายงานผลการสอบหาข้อเท็จจริง 2 ชุด",
    "รายงานขออนุมัติจำหน่าย พร้อมวิธีการจำหน่ายและคำสั่งอนุมัติให้จำหน่าย 2 ชุด",
    "คำสั่งแต่งตั้งคณะกรรมการจำหน่ายพัสดุ เช่น คำสั่งแต่งตั้งคณะกรรมการประเมินราคากลาง "
    "คณะกรรมการดำเนินการขาย คณะกรรมการขายพัสดุโดยวิธีทอดตลาด",
    "บัญชีรายการพัสดุชำรุด เสื่อมสภาพ สูญไป หรือไม่จำเป็นต้องใช้ในหน่วยงานของรัฐ 2 ชุด",
    "รายงานผลการจำหน่ายพัสดุ 2 ชุด",
    "สำเนาใบเสร็จรับเงินที่แสดงการนำเงินส่งเป็นรายได้ 2 ชุด",
]


def render_sao_letter(school, dp, doc=None):
    """(17) หนังสือแจ้ง สตง.ภูมิภาค เรื่อง การจำหน่ายพัสดุ (ข้อ 218) - หนังสือภายนอก"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    _letter_header(doc, school, dp.sao_no, _d(dp.sao_date), "การจำหน่ายพัสดุ",
                   "ผู้อำนวยการ" + _sao_name(dp),
                   "สำเนาเอกสารการตรวจสอบพัสดุประจำปี และการจำหน่ายพัสดุประจำปี จำนวน 1 ชุด")
    clauses = "ข้อ 213 และข้อ 215"
    if _items(dp, "จำหน่ายเป็นสูญ"):
        clauses = "ข้อ 213 ข้อ 215 และข้อ 217"
    _p(doc, f"ด้วย{_sname(school)} ได้ดำเนินการตรวจสอบพัสดุประจำปี และจำหน่ายพัสดุ"
            f"ประจำปีงบประมาณ พ.ศ. {dp.year} ตาม{_REG} {clauses} "
            f"พร้อมทั้งได้ลงจ่ายพัสดุออกจากบัญชีหรือทะเบียนพัสดุ "
            f"เมื่อวันที่ {_d(dp.written_off_date)} เรียบร้อยแล้ว "
            f"จึงขอรายงานให้ทราบตามระเบียบฯ ข้อ 218 "
            "ดังสำเนาเอกสารที่ได้ส่งมาพร้อมหนังสือนี้",
       align="justify", indent=1.25, after=2)
    _p(doc, "จึงเรียนมาเพื่อโปรดทราบ", align="justify", indent=1.25, after=8)
    _p(doc, "ขอแสดงความนับถือ", align="center", after=12)
    _sign_table(doc, [[
        ("", "center"),
        (f"( {(school.director_name or '').strip() or _BLANK} )", "center"),
        (_director_line(school), "center"),
    ]])
    _p(doc, "งานการเงินและพัสดุ", size=14, before=8, after=0)
    _p(doc, f"โทร. {(dp.contact_phone or getattr(school, 'phone', '') or '').strip() or _BLANK}",
       size=14, after=6)
    _p(doc, "หมายเหตุ  สิ่งที่ส่งมาด้วย ประกอบด้วย", bold=True, size=14, after=1)
    for i, txt in enumerate(_SAO_ENCL, 1):
        _p(doc, f"{i}. {txt}", align="justify", indent=1.25, size=14, after=1)
    return _save(doc, f"หนังสือแจ้ง สตง. การจำหน่ายพัสดุ_ปีงบ{dp.year}") if own else doc


def render_writeoff_memo(school, dp, doc=None):
    """(19) บันทึกข้อความ ขออนุมัติจำหน่ายพัสดุเป็นสูญ (ข้อ 217)

    ใช้กรณีพัสดุสูญไปโดยไม่ปรากฏตัวผู้รับผิด หรือมีตัวผู้รับผิดแต่ไม่สามารถชดใช้ได้
    หรือมีตัวพัสดุอยู่แต่ไม่สมควรดำเนินการตามข้อ 215
    """
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    rows = _items(dp, "จำหน่ายเป็นสูญ")
    total = sum(float(it.asset.cost or 0) for it in rows)
    _memo_header(doc, school, "ขออนุมัติจำหน่ายพัสดุเป็นสูญ",
                 dp.wo_memo_no, _d(dp.wo_memo_date))
    _p(doc, "ตามที่คณะกรรมการสอบหาข้อเท็จจริงได้รายงานผลการสอบหาข้อเท็จจริง "
            f"ตามบันทึกที่ {dp.fact_report_no or _BLANK} ลงวันที่ {_d(dp.fact_report_date)} "
            f"ปรากฏว่ามีพัสดุ จำนวน {len(rows)} รายการ ราคาซื้อหรือได้มารวมกัน "
            f"{total:,.2f} บาท ({bahttext(total)}) ดังความละเอียดทราบแล้ว นั้น",
       align="justify", indent=1.25, after=2)
    _p(doc, f"{_REG} ข้อ 217 กำหนดว่า ในกรณีที่พัสดุสูญไปโดยไม่ปรากฏตัวผู้รับผิด "
            "หรือมีตัวผู้รับผิดแต่ไม่สามารถชดใช้ได้ หรือมีตัวพัสดุอยู่แต่ไม่สมควรดำเนินการ"
            "ตามข้อ 215 ให้จำหน่ายพัสดุนั้นเป็นสูญ โดย (1) ถ้าพัสดุนั้นมีราคาซื้อหรือได้มา"
            "รวมกันไม่เกิน 1,000,000 บาท ให้หัวหน้าหน่วยงานของรัฐเป็นผู้พิจารณาอนุมัติ",
       align="justify", indent=1.25, after=2)
    reason = (dp.wo_reason or "").strip() or (
        "พัสดุดังกล่าวสูญไปโดยไม่ปรากฏตัวผู้รับผิด และไม่มีตัวพัสดุเหลืออยู่ "
        "จึงไม่สามารถดำเนินการจำหน่ายตามวิธีการในข้อ 215 ได้")
    _p(doc, "ข้อเท็จจริงปรากฏว่า " + reason, align="justify", indent=1.25, after=3)
    if rows:
        _item_table(doc, rows)
        _p(doc, "", after=3)
    if dp.wo_no_liable:
        _p(doc, "ทั้งนี้ คณะกรรมการสอบหาข้อเท็จจริงได้พิจารณาแล้วเห็นว่าไม่ปรากฏตัว"
                "ผู้รับผิดทางแพ่ง จึงไม่ต้องดำเนินการตามระเบียบว่าด้วยความรับผิดทางละเมิด"
                "ของเจ้าหน้าที่", align="justify", indent=1.25, after=2)
    over = total > WRITE_OFF_LIMIT
    if over:
        _p(doc, f"เนื่องจากพัสดุดังกล่าวมีราคาซื้อหรือได้มารวมกันเกิน 1,000,000 บาท "
                "การอนุมัติจำหน่ายเป็นสูญจึงอยู่ในอำนาจของกระทรวงการคลัง ตามข้อ 217 (2) (ก) "
                "จึงเห็นควรเสนอเรื่องไปยังกระทรวงการคลังเพื่อพิจารณาอนุมัติต่อไป",
           align="justify", indent=1.25, after=2)
    _p(doc, "จึงเรียนมาเพื่อโปรดพิจารณา"
            + ("อนุมัติจำหน่ายพัสดุเป็นสูญตามรายการข้างต้น และสั่งการให้ลงจ่ายพัสดุออกจาก"
               "บัญชีหรือทะเบียนพัสดุต่อไป" if not over else "และดำเนินการต่อไป"),
       align="justify", indent=1.25, after=10)
    _officer_sign(doc, school)
    lines = ["- ทราบ"]
    if not over:
        lines += ["- อนุมัติให้จำหน่ายพัสดุเป็นสูญตามรายการข้างต้น",
                  "- ลงจ่ายออกจากบัญชีหรือทะเบียนพัสดุ",
                  "- รายงาน สตง. และปลัดกระทรวงการคลัง ตามระเบียบฯ ข้อ 218"]
    else:
        lines += ["- ให้เสนอกระทรวงการคลังพิจารณาอนุมัติ"]
    _director_sign(doc, school, lines=lines)
    _p(doc, "วันที่ ........./............./..................", align="center", after=2)
    return _save(doc, f"ขออนุมัติจำหน่ายพัสดุเป็นสูญ_ปีงบ{dp.year}") if own else doc


def render_area_letter(school, dp, doc=None):
    """(20) หนังสือแจ้งสำนักงานเขตพื้นที่การศึกษา (ข้อ 218) - หนังสือภายนอก"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    area = (school.area_office or "").strip() or "สำนักงานเขตพื้นที่การศึกษา"
    sao = _sao_name(dp)
    _letter_header(doc, school, dp.area_no, _d(dp.area_date),
                   f"การตรวจสอบพัสดุประจำปี และการจำหน่ายพัสดุ ปีงบประมาณ พ.ศ. {dp.year}",
                   "ผู้อำนวยการ" + area,
                   f"1. สำเนาเอกสารหลักฐานการดำเนินการตรวจสอบพัสดุประจำปี "
                   f"และการจำหน่ายพัสดุ จำนวน 1 ชุด\n"
                   f"2. สำเนาหนังสือแจ้ง{sao} จำนวน 1 ฉบับ")
    _p(doc, f"ด้วย{_sname(school)} ได้ดำเนินการตรวจสอบพัสดุประจำปี และจำหน่ายพัสดุ"
            f"ประจำปีงบประมาณ พ.ศ. {dp.year} ตาม{_REG} ข้อ 213 และข้อ 215 "
            f"พร้อมทั้งได้ลงจ่ายพัสดุออกจากบัญชีหรือทะเบียนพัสดุ "
            f"เมื่อวันที่ {_d(dp.written_off_date)} เรียบร้อยแล้ว "
            f"ซึ่งได้จัดส่งสำเนารายงานพร้อมเอกสารที่เกี่ยวข้องให้{sao} "
            "เพื่อทราบตามระเบียบดังกล่าวแล้ว รายละเอียดตามสิ่งที่ส่งมาด้วยพร้อมนี้",
       align="justify", indent=1.25, after=2)
    _p(doc, "จึงเรียนมาเพื่อโปรดทราบ", align="justify", indent=1.25, after=8)
    _p(doc, "ขอแสดงความนับถือ", align="center", after=12)
    _sign_table(doc, [[
        ("", "center"),
        (f"( {(school.director_name or '').strip() or _BLANK} )", "center"),
        (_director_line(school), "center"),
    ]])
    _p(doc, "งานการเงินและพัสดุ", size=14, before=8, after=0)
    _p(doc, f"โทร. {(dp.contact_phone or getattr(school, 'phone', '') or '').strip() or _BLANK}",
       size=14, after=6)
    return _save(doc, f"หนังสือแจ้งเขตพื้นที่ การจำหน่ายพัสดุ_ปีงบ{dp.year}") if own else doc


def render_mof_letter(school, dp, doc=None):
    """(21) หนังสือแจ้งปลัดกระทรวงการคลัง (ข้อ 218 - เฉพาะกรณีจำหน่ายเป็นสูญ)"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    rows = _items(dp, "จำหน่ายเป็นสูญ")
    total = sum(float(it.asset.cost or 0) for it in rows)
    _letter_header(doc, school, dp.mof_no, _d(dp.mof_date),
                   "การจำหน่ายพัสดุเป็นสูญ", "ปลัดกระทรวงการคลัง",
                   "สำเนาเอกสารหลักฐานการจำหน่ายพัสดุเป็นสูญ จำนวน 1 ชุด")
    _p(doc, f"ด้วย{_sname(school)} ได้ดำเนินการตรวจสอบพัสดุประจำปีงบประมาณ พ.ศ. {dp.year} "
            f"ตาม{_REG} ข้อ 213 ปรากฏว่ามีพัสดุสูญไปโดยไม่ปรากฏตัวผู้รับผิด "
            f"จำนวน {len(rows)} รายการ ราคาซื้อหรือได้มารวมกัน {total:,.2f} บาท "
            f"({bahttext(total)}) และได้จำหน่ายพัสดุดังกล่าวเป็นสูญ ตามระเบียบฯ ข้อ 217 (1) "
            f"พร้อมทั้งลงจ่ายพัสดุออกจากบัญชีหรือทะเบียนพัสดุ "
            f"เมื่อวันที่ {_d(dp.written_off_date)} เรียบร้อยแล้ว",
       align="justify", indent=1.25, after=2)
    _p(doc, f"จึงขอรายงานให้ทราบตาม{_REG} ข้อ 218 รายละเอียดตามสิ่งที่ส่งมาด้วยพร้อมนี้",
       align="justify", indent=1.25, after=2)
    _p(doc, "จึงเรียนมาเพื่อโปรดทราบ", align="justify", indent=1.25, after=8)
    _p(doc, "ขอแสดงความนับถือ", align="center", after=12)
    _sign_table(doc, [[
        ("", "center"),
        (f"( {(school.director_name or '').strip() or _BLANK} )", "center"),
        (_director_line(school), "center"),
    ]])
    _p(doc, "งานการเงินและพัสดุ", size=14, before=8, after=0)
    _p(doc, f"โทร. {(dp.contact_phone or getattr(school, 'phone', '') or '').strip() or _BLANK}",
       size=14, after=6)
    return _save(doc, f"หนังสือแจ้งกระทรวงการคลัง จำหน่ายเป็นสูญ_ปีงบ{dp.year}") if own else doc


def render_remit_letter(school, dp, doc=None):
    """(18) หนังสือนำส่งเงินรายได้จากการจำหน่ายพัสดุ - หนังสือภายนอก"""
    own = doc is None
    doc = doc or _new()
    if not own:
        _break(doc)
    kind = (dp.revenue_kind or "แผ่นดิน").strip()
    total = float(dp.sale_total or 0)
    to = (dp.remit_to or "").strip() or (
        "ผู้อำนวยการ" + ((school.area_office or "").strip() or "สำนักงานเขตพื้นที่การศึกษา"))
    _letter_header(doc, school, dp.remit_no, _d(dp.remit_date),
                   f"การนำส่งเงินรายได้{kind}", to,
                   f"1. เงินสด จำนวน {total:,.2f} บาท\n2. ใบนำส่งเงิน จำนวน 1 ฉบับ")
    _p(doc, f"ด้วย{_sname(school)} ได้ดำเนินการจำหน่ายพัสดุประจำปีงบประมาณ พ.ศ. {dp.year} "
            f"ตาม{_REG} ข้อ 215 (1) เสร็จเรียบร้อยแล้ว "
            f"จึงขอนำเงินส่งเป็นเงินรายได้{kind} จำนวน {total:,.2f} บาท "
            f"({bahttext(total)})", align="justify", indent=1.25, after=2)
    if kind == "สถานศึกษา":
        _p(doc, "ทั้งนี้ พัสดุที่จำหน่ายดังกล่าวได้มาด้วยเงินรายได้สถานศึกษาหรือเงินบริจาค "
                "จึงนำเงินที่ได้จากการจำหน่ายเข้าเป็นเงินรายได้สถานศึกษา",
           align="justify", indent=1.25, after=2)
    _p(doc, "จึงเรียนมาเพื่อโปรดทราบและดำเนินการต่อไป", align="justify", indent=1.25, after=8)
    _p(doc, "ขอแสดงความนับถือ", align="center", after=12)
    _sign_table(doc, [[
        ("", "center"),
        (f"( {(school.director_name or '').strip() or _BLANK} )", "center"),
        (_director_line(school), "center"),
    ]])
    _p(doc, "งานงบประมาณ", size=14, before=8, after=0)
    _p(doc, f"โทร. {(dp.contact_phone or getattr(school, 'phone', '') or '').strip() or _BLANK}",
       size=14, after=6)
    return _save(doc, f"หนังสือนำส่งเงินรายได้{kind}_ปีงบ{dp.year}") if own else doc


# ================================================================== ชุดรวม
DOCS = {
    "fact_memo":      ("บันทึกขอแต่งตั้งกรรมการสอบหาข้อเท็จจริง", render_fact_memo),
    "fact_order":     ("คำสั่งแต่งตั้งกรรมการสอบหาข้อเท็จจริง", render_fact_order),
    "fact_report":    ("รายงานผลการสอบหาข้อเท็จจริง", render_fact_report),
    "statement":      ("บันทึกถ้อยคำในการสอบหาข้อเท็จจริง", render_statement),
    "request":        ("บันทึกขอจำหน่ายพัสดุ", render_dispose_request),
    "order":          ("คำสั่งแต่งตั้งกรรมการประเมินราคากลาง/ขาย/ทำลาย", render_committee_order),
    "invite":         ("หนังสือขอเชิญเสนอราคาซื้อพัสดุ", render_invite_quote),
    "quote":          ("ใบเสนอราคา", render_quote_form),
    "price":          ("บันทึกการประเมินราคากลางพัสดุ", render_price_report),
    "sale":           ("รายงานผลการดำเนินการขาย", render_sale_report),
    "auction_order":  ("คำสั่งแต่งตั้งกรรมการขายทอดตลาด", render_auction_order),
    "notice":         ("ประกาศขายพัสดุโดยวิธีทอดตลาด", render_auction_notice),
    "auction_list":   ("บัญชีรายการพัสดุที่จะขายทอดตลาด", render_auction_list),
    "auction_ledger": ("บัญชีคุมพัสดุที่ขายทอดตลาด", render_auction_ledger),
    "auction_report": ("รายงานผลการขายพัสดุโดยวิธีทอดตลาด", render_auction_report),
    "destroy":        ("รายงานผลการทำลายพัสดุ", render_destroy_report),
    "writeoff":       ("บันทึกขออนุมัติจำหน่ายพัสดุเป็นสูญ", render_writeoff_memo),
    "area":           ("หนังสือแจ้งเขตพื้นที่การศึกษา", render_area_letter),
    "sao":            ("หนังสือแจ้ง สตง. การจำหน่ายพัสดุ", render_sao_letter),
    "mof":            ("หนังสือแจ้งปลัดกระทรวงการคลัง (จำหน่ายเป็นสูญ)", render_mof_letter),
    "remit":          ("หนังสือนำส่งเงินรายได้", render_remit_letter),
}

# ชุดย่อยตามขั้นตอน (คีย์ที่ใช้กับ /bundle/{stage})
BUNDLES = {
    "fact":    ("ชุดสอบหาข้อเท็จจริง",
                ["fact_memo", "fact_order", "fact_report", "statement"]),
    "approve": ("ชุดขออนุมัติจำหน่าย", ["request", "order"]),
    "sell":    ("ชุดขายโดยวิธีเฉพาะเจาะจง", ["invite", "quote", "price", "sale"]),
    "auction": ("ชุดขายทอดตลาด",
                ["auction_order", "notice", "auction_list", "auction_ledger",
                 "auction_report"]),
    "destroy": ("ชุดทำลายพัสดุ", ["destroy"]),
    "writeoff": ("ชุดจำหน่ายเป็นสูญ", ["writeoff"]),
    # ข้อ 218 ให้ส่งสำเนารายงาน 3 ทาง (ปลัดกระทรวงการคลังเฉพาะกรณีจำหน่ายเป็นสูญ)
    "close":   ("ชุดปิดเรื่อง", ["area", "sao", "mof", "remit"]),
}


def _keys_for(dp, stage):
    """คีย์เอกสารที่ควรออกจริงในขั้นนั้น (ตัดที่ไม่เกี่ยวออก)"""
    keys = list(BUNDLES[stage][1])
    if not _items(dp):                       # สำนวนเปล่า ยังไม่มีอะไรให้ออก
        return []
    if stage == "sell" and (dp.sale_mode or "specific") != "specific":
        return []
    if stage == "auction" and (dp.sale_mode or "specific") != "auction":
        return []
    if stage == "destroy" and not (_items(dp, "ทำลาย") or _items(dp, "แปรสภาพ")):
        return []
    if stage == "writeoff" and not _items(dp, "จำหน่ายเป็นสูญ"):
        return []
    if stage == "close":
        if float(dp.sale_total or 0) <= 0:            # ไม่มีเงินเข้า ไม่ต้องมีหนังสือนำส่ง
            keys = [k for k in keys if k != "remit"]
        if not _items(dp, "จำหน่ายเป็นสูญ"):          # แจ้งคลังเฉพาะกรณีจำหน่ายเป็นสูญ
            keys = [k for k in keys if k != "mof"]
    return keys


def render_bundle(school, dp, stage) -> str:
    """ออกชุดเอกสารของขั้นตอนที่ระบุเป็นไฟล์เดียว"""
    if stage not in BUNDLES:
        raise ValueError("ไม่รู้จักขั้นตอนนี้")
    label, _ = BUNDLES[stage]
    keys = _keys_for(dp, stage)
    if not keys:
        raise ValueError(f"ยังไม่มีเอกสารที่ต้องออกใน{label} (ตรวจวิธีจำหน่ายที่เลือกไว้)")
    doc = _new()
    for k in keys:
        DOCS[k][1](school, dp, doc)
    return _save(doc, f"{label}_ปีงบ{dp.year}")


def render_full_set(school, dp) -> str:
    """ออกทั้งสำนวนเป็นไฟล์เดียว ตามลำดับขั้นตอนของระเบียบ"""
    doc = _new()
    n = 0
    for stage in ("fact", "approve", "sell", "auction", "destroy", "writeoff", "close"):
        for k in _keys_for(dp, stage):
            DOCS[k][1](school, dp, doc)
            n += 1
    if not n:
        raise ValueError("ยังไม่มีเอกสารให้ออก")
    return _save(doc, f"ชุดจำหน่ายพัสดุทั้งสำนวน_ปีงบ{dp.year}")
