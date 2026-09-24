# -*- coding: utf-8 -*-
"""
kinder_book.py - สมุดรายงานประจำตัวนักเรียน ระดับปฐมวัย (สมุดพกอนุบาล)

ทำตามเล่มจริง "บัญชีเรียกชื่อและสมุดบันทึกพัฒนาการเด็กนักเรียนระดับปฐมวัย"
แบบ อบ.1/1 (3 ขวบ) · อบ.1/2 (4 ขวบ) · อบ.1/3 (5 ขวบ) — ส่วนสมุดรายงานที่ส่งผู้ปกครอง
เรียงหน้าตามต้นฉบับ 11 หน้า

  1  ปกสมุดรายงานประจำตัวนักเรียน
  2  ข้อมูลส่วนตัวผู้เรียน (+ กรอบติดรูป)
  3  เวลามาเรียนในรอบปี + น้ำหนักและส่วนสูง
  4  รายงานผลการพัฒนาด้านร่างกาย
  5  ด้านอารมณ์ - จิตใจ
  6  ด้านสังคม
  7  ด้านสติปัญญา (ข้อ 1-10)
  8  ด้านสติปัญญา (ข้อ 11-20)
  9  ความเห็นครูประจำชั้น (ภาคเรียนที่ 1 / 2)
  10 ผู้ปกครองรายงานพฤติกรรมของนักเรียนขณะอยู่ที่บ้าน (แบบฟอร์มให้ผู้ปกครองกรอก)
  11 ผลงานที่น่าภาคภูมิใจ + สรุปการพัฒนาเด็ก + ลงนามครู/ผู้บริหาร

ตัวบ่งชี้ทั้งหมดอยู่ที่ app/services/kinder.py (ดึงจากไฟล์ต้นฉบับ ไม่ได้แต่งเอง)
ความกว้างตารางต้องไม่เกินพื้นที่พิมพ์ A4 แนวตั้ง 16.0 ซม.
"""
import json

from docx import Document
from docx.shared import Cm, Pt

from app.database import get_data_dir
from app.services.doc_page import set_a4
from app.services.acad_doc import (_p, _cell, _widths, _logo_header, _class_label,
                                   _pp6_photo_box, _pp6_central, _safe, THAI_FONT)
from app.services import kinder as kd
from app.thai_utils import thai_date, be_date_input
from app.services.academic import TH_MONTH_FULL, TH_MONTHS, count_marks

_TICK = "✓"
_DOTS = "…………………………………………………………………………………………………"


def _jload(text, default):
    try:
        v = json.loads(text or "")
        return v if isinstance(v, type(default)) else default
    except Exception:
        return default


def _notes_of(db, aid):
    from app.models import KinderNote
    n = db.query(KinderNote).filter_by(acad_student_id=aid).first()
    return {
        "comments": _jload(n.comments if n else "", {}),
        "improve": _jload(n.improve if n else "", {}),
        "summary": _jload(n.summary if n else "", {}),
        "works": _jload(n.works if n else "", []),
    }


def _results_of(db, aid):
    """{(term, code): value} ของเด็ก 1 คน"""
    from app.models import KinderResult
    return {(r.term, r.code): r.value
            for r in db.query(KinderResult).filter_by(acad_student_id=aid).all()}


def _lines(cell, n, *, size=13):
    """เติมบรรทัดจุดไข่ปลาในช่องตาราง (ให้ครู/ผู้ปกครองเขียนมือ)"""
    from docx.oxml.ns import qn
    cell.text = ""
    first = cell.paragraphs[0]
    for i in range(n):
        p = first if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(_DOTS)
        r.font.size = Pt(size)
        r.font.name = THAI_FONT
        r._element.rPr.rFonts.set(qn("w:cs"), THAI_FONT)


def _text_or_lines(cell, text, n, *, size=13):
    """มีข้อความที่ครูกรอกไว้ก็พิมพ์ลงไป ไม่มีก็เว้นบรรทัดให้เขียนมือ"""
    if (text or "").strip():
        _cell(cell, text.strip(), align="left", size=size)
    else:
        _lines(cell, n, size=size)


# ---------------------------------------------------------------- หน้า 1 ปก
def _cover(doc, school, s, meta, *, page_break):
    klass = s.klass
    added = _logo_header(doc, school, page_break=page_break, height_cm=3.0, after=8)
    if not added:
        _p(doc, "", after=0, page_break=page_break)
        _p(doc, "", after=0)
    _p(doc, f"{meta['form']}        นักเรียนเลขที่  {s.seq or '......'}", align="right",
       size=13, after=10)
    _p(doc, "สมุดรายงานประจำตัวนักเรียน", align="center", bold=True, size=26, after=6)
    _p(doc, "ระดับปฐมวัย", align="center", bold=True, size=20, after=14)
    _p(doc, f"{meta['title']}  (อายุ {meta['age']} ขวบ)"
            + (f"  ห้อง {klass.room}" if (klass.room or "").strip() else ""),
       align="center", size=16, after=4)
    _p(doc, f"ปีการศึกษา {klass.year}", align="center", size=15, after=16)
    _p(doc, f"โรงเรียน{school.name or ''}", align="center", bold=True, size=17, after=2)
    # ตำบล/อำเภอ/จังหวัด ตามแบบต้นฉบับ (ระบบไม่มีช่องตำบล -> เว้นให้เขียนเอง)
    loc = []
    for attr, prefix in (("subdistrict", "ตำบล"), ("district", "อำเภอ"), ("province", "จังหวัด")):
        v = (getattr(school, attr, "") or "").strip()
        loc.append(prefix + (v if v else " ................"))
    _p(doc, "   ".join(loc), align="center", size=14, after=2)
    if getattr(school, "area_office", ""):
        _p(doc, f"สำนักงานเขตพื้นที่การศึกษา{school.area_office}", align="center", size=14, after=18)
    else:
        _p(doc, "", after=18)
    _p(doc, f"เลขที่  {s.seq or '......'}", align="center", size=15, after=4)
    _p(doc, f"ชื่อ - นามสกุล  {s.name}", align="center", bold=True, size=18, after=4)
    homerooms = [p.name for p in (klass.homeroom, klass.co_homeroom) if p]
    _p(doc, "ชื่อครูประจำชั้น  " + (" / ".join(homerooms) if homerooms else "......................"),
       align="center", size=14, after=2)
    _p(doc, "ชื่อผู้บริหารสถานศึกษา  "
            + (getattr(school, "director_name", "") or "......................"),
       align="center", size=14, after=0)


# ---------------------------------------------------------------- หน้า 2 ข้อมูลส่วนตัว
def _personal(doc, school, s, db):
    st = _pp6_central(s, db)
    _p(doc, "ข้อมูลส่วนตัวผู้เรียน", align="center", bold=True, size=18, after=8, page_break=True)
    rows = [
        ("ชื่อ - นามสกุล", s.name),
        ("เลขประจำตัวนักเรียน", s.student_no or ""),
        ("เลขที่", str(s.seq or "")),
        ("เลขประจำตัวประชาชน", ""),
        ("วัน เดือน ปีเกิด", thai_date(st.birthdate) if (st and st.birthdate) else ""),
        ("เชื้อชาติ / สัญชาติ / ศาสนา", (getattr(st, "nationality", "") or "") if st else ""),
        ("หมู่เลือด", ""),
        ("โรคประจำตัว", ""),
    ]
    # ที่อยู่เป็นบรรทัดยาวเต็มความกว้าง (ผสาน 2 ช่อง) ไม่งั้นตกบรรทัดในคอลัมน์แคบ
    addr = [
        "ที่อยู่ตามทะเบียนบ้านเลขที่ ..............  หมู่ที่ ..........  ซอย ..............  ถนน ..............",
        "ตำบล/แขวง ....................  อำเภอ/เขต ....................",
        "จังหวัด ....................  รหัสไปรษณีย์ ..............  โทรศัพท์ ....................",
    ]
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for lab, val in rows:
        cells = t.add_row().cells
        _cell(cells[0], lab, bold=True, align="left", fill="F1F5F9")
        _cell(cells[1], val or "", align="left")
    for line in addr:
        cells = t.add_row().cells
        cells[0].merge(cells[1])
        _cell(cells[0], line, align="left", size=13)
    _widths(t, [Cm(5.2), Cm(10.8)])

    # ---- ข้อมูลพี่น้อง (ตามแบบต้นฉบับ) ----
    _p(doc, "", after=6)
    _p(doc, "ข้อมูลพี่น้อง", bold=True, size=14, after=2)
    bt = doc.add_table(rows=0, cols=1)
    bt.style = "Table Grid"
    for line in ("มีพี่น้องทั้งหมด ......... คน      พี่ชาย ......... คน      พี่สาว ......... คน",
                 "เป็นบุตรคนที่ .........          น้องชาย ......... คน      น้องสาว ......... คน",
                 "สถานภาพการสมรสของบิดามารดา ......................................................"):
        _cell(bt.add_row().cells[0], line, align="left", size=13)
    _widths(bt, [Cm(16.0)])

    _p(doc, "", after=8)
    _p(doc, "รูปของนักเรียน", bold=True, size=13, after=2)
    _pp6_photo_box(doc)


# ---------------------------------------------------------------- หน้า 3 เวลาเรียน + น้ำหนักส่วนสูง
def _attendance(doc, school, s, db):
    from app.models import AcadAttendance
    from app.services import growth
    klass = s.klass
    _p(doc, "เวลามาเรียนในรอบปี", align="center", bold=True, size=18, after=8, page_break=True)
    att = {a.month: a for a in db.query(AcadAttendance)
           .filter(AcadAttendance.acad_student_id == s.id,
                   AcadAttendance.subject_id.is_(None)).all()}
    heads = ["เดือน", "พ.ศ.", "เวลาเรียน", "มาเรียน", "ป่วย", "ลา", "ขาด"]
    t = doc.add_table(rows=1, cols=7)
    t.style = "Table Grid"
    for i, h in enumerate(heads):
        _cell(t.rows[0].cells[i], h, bold=True, fill="EDE9FE")
    tot = {"open": 0, "/": 0, "ป": 0, "ล": 0, "ข": 0}
    for m, abbr in TH_MONTHS:
        a = att.get(m)
        cm = count_marks(a.marks) if (a and a.marks) else None
        # ปีการศึกษาเริ่ม พ.ค. -> ม.ค.-มี.ค. เป็นปีถัดไป
        be = klass.year + (1 if m in (1, 2, 3) else 0)
        cells = t.add_row().cells
        _cell(cells[0], TH_MONTH_FULL.get(m, abbr), align="left")
        _cell(cells[1], be)
        if cm:
            opened = cm["/"] + cm["ป"] + cm["ล"] + cm["ข"]
            _cell(cells[2], opened or "")
            _cell(cells[3], cm["/"] or "")
            _cell(cells[4], cm["ป"] or "")
            _cell(cells[5], cm["ล"] or "")
            _cell(cells[6], cm["ข"] or "")
            tot["open"] += opened
            for k in ("/", "ป", "ล", "ข"):
                tot[k] += cm[k]
        else:
            for c in cells[2:]:
                _cell(c, "")
    rc = t.add_row().cells
    _cell(rc[0], "รวม", bold=True, fill="F1F5F9")
    _cell(rc[1], "", fill="F1F5F9")
    for i, k in enumerate(("open", "/", "ป", "ล", "ข")):
        _cell(rc[2 + i], tot[k] or "", bold=True)
    _widths(t, [Cm(3.4), Cm(1.8), Cm(2.4), Cm(2.2), Cm(2.0), Cm(2.0), Cm(2.0)])
    pct = f"{100.0 * tot['/'] / tot['open']:.1f}" if tot["open"] else ".............."
    _p(doc, f"มาเรียนร้อยละ {pct}          ย้ายออก ..............................", size=13, after=8)

    _p(doc, "น้ำหนักและส่วนสูง", bold=True, size=15, after=3)
    ms = growth.measures_for(db, s.student_id, klass.year) if s.student_id else {}
    st = _pp6_central(s, db)
    who = st or s
    # ผลการประเมินแยก น้ำหนัก/ส่วนสูง และชั่ง 2 ครั้งต่อภาคเรียน ตามแบบต้นฉบับ
    gt = doc.add_table(rows=2, cols=7)
    gt.style = "Table Grid"
    g0, g1 = gt.rows[0].cells, gt.rows[1].cells
    for i, h in enumerate(["ภาคเรียนที่", "ครั้งที่", "วันที่ประเมิน", "น้ำหนัก", "ส่วนสูง"]):
        g0[i].merge(g1[i])
        _cell(g0[i], h, bold=True, size=12, fill="EDE9FE")
    g0[5].merge(g0[6])
    _cell(g0[5], "ผลการประเมิน", bold=True, size=12, fill="EDE9FE")
    _cell(g1[5], "น้ำหนัก", bold=True, size=11, fill="F5F3FF")
    _cell(g1[6], "ส่วนสูง", bold=True, size=11, fill="F5F3FF")
    for term in (1, 2):
        for times in (1, 2):
            # ระบบเก็บการชั่งภาคเรียนละ 1 ครั้ง -> ครั้งที่ 2 เว้นให้ครูกรอกเอง
            m = ms.get(term) if times == 1 else None
            res = growth.measure_result(who, m) if m else None
            cells = gt.add_row().cells
            _cell(cells[0], term, size=12)
            _cell(cells[1], times, size=12)
            _cell(cells[2], be_date_input(m.date) if (m and m.date) else "", size=12)
            _cell(cells[3], f"{m.weight:g}" if (m and m.weight) else "", size=12)
            _cell(cells[4], f"{m.height:g}" if (m and m.height) else "", size=12)
            _cell(cells[5], (res or {}).get("wa", "") or "", size=11)
            _cell(cells[6], (res or {}).get("ha", "") or "", size=11)
    _widths(gt, [Cm(2.0), Cm(1.6), Cm(2.6), Cm(2.0), Cm(2.0), Cm(2.9), Cm(2.9)])
    _p(doc, "เกณฑ์อ้างอิง: กราฟการเจริญเติบโตของกรมอนามัย", align="center", size=12, after=0)


# ---------------------------------------------------------------- หน้า 4-8 พัฒนาการรายด้าน
def _domain_page(doc, s, db, level, domain, title, items, res):
    """1 หน้า = 1 ชุดตัวบ่งชี้ (ด้านสติปัญญาแบ่ง 2 หน้า)"""
    _p(doc, f"รายงานผลการพัฒนา{title}", align="center", bold=True, size=18, after=6,
       page_break=True)
    t = doc.add_table(rows=2, cols=8)
    t.style = "Table Grid"
    h0, h1 = t.rows[0].cells, t.rows[1].cells
    h0[0].merge(h1[0]); _cell(h0[0], "ลำดับที่", bold=True, size=12, fill="EDE9FE")
    h0[1].merge(h1[1]); _cell(h0[1], "พฤติกรรม", bold=True, size=12, fill="EDE9FE")
    h0[2].merge(h0[4]); _cell(h0[2], "ภาคเรียนที่ 1", bold=True, size=12, fill="EDE9FE")
    h0[5].merge(h0[7]); _cell(h0[5], "ภาคเรียนที่ 2", bold=True, size=12, fill="EDE9FE")
    for i, (_v, label) in enumerate(kd.RATINGS):
        _cell(h1[2 + i], label, bold=True, size=10, fill="F5F3FF")
        _cell(h1[5 + i], label, bold=True, size=10, fill="F5F3FF")
    for n, grp, text in items:
        code = kd.code_of(domain, n)
        cells = t.add_row().cells
        _cell(cells[0], n, size=12)
        # \u00a0 = ช่องว่างไม่ตัดบรรทัด กันเลขข้อโดดไปอยู่บรรทัดเดียว
        _cell(cells[1], f"{n}.\u00a0{text}\n({grp})", align="left", size=12)
        for ti, term in enumerate((1, 2)):
            v = res.get((term, code))
            for i, (val, _label) in enumerate(kd.RATINGS):
                _cell(cells[2 + ti * 3 + i], _TICK if v == val else "", size=13)
    _widths(t, [Cm(1.3), Cm(6.7), Cm(1.3), Cm(1.3), Cm(1.3), Cm(1.3), Cm(1.3), Cm(1.3)])
    _p(doc, "หมายเหตุ  "
            + " · ".join(f"{label}" for _v, label in kd.RATINGS)
            + "  ตามเกณฑ์ท้ายเล่ม", size=11, after=0)


# ---------------------------------------------------------------- หน้า 9 ความเห็นครู
def _teacher_comments(doc, s, notes):
    _p(doc, "ความเห็นครูประจำชั้น", align="center", bold=True, size=18, after=8, page_break=True)
    comments = notes["comments"]
    improve = notes["improve"]
    for term in (1, 2):
        _p(doc, f"ภาคเรียนที่ {term}", bold=True, size=15, after=3)
        t = doc.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        _cell(t.rows[0].cells[0], "พัฒนาการด้าน", bold=True, size=12, fill="EDE9FE")
        _cell(t.rows[0].cells[1], "ความเห็นของครูประจำชั้น", bold=True, size=12, fill="EDE9FE")
        for key, full, _short in kd.DOMAINS:
            cells = t.add_row().cells
            topics = " ".join(f"{i + 1}. {x}" for i, x in enumerate(kd.COMMENT_TOPICS[key]))
            _cell(cells[0], f"{full}\n{topics}", align="left", size=11)
            _text_or_lines(cells[1], comments.get(str(term), {}).get(key, ""), 2, size=12)
        cells = t.add_row().cells
        _cell(cells[0], "พฤติกรรมที่ควรส่งเสริมและพัฒนา", bold=True, align="left", size=12,
              fill="F8FAFC")
        _text_or_lines(cells[1], improve.get(str(term), ""), 3, size=12)
        _widths(t, [Cm(5.4), Cm(10.6)])
        if term == 1:
            _p(doc, "", after=8)      # เว้นวรรคระหว่าง 2 ภาคเรียน
        # ไม่ใส่ย่อหน้าว่างท้ายภาค 2: ถ้าตารางเต็มหน้าพอดี ย่อหน้าว่างจะดันไปสร้างหน้าเปล่า


# ---------------------------------------------------------------- หน้า 10 ผู้ปกครอง
def _full_class(klass):
    """ชื่อชั้นแบบเต็มสำหรับเอกสารที่ส่งผู้ปกครอง เช่น อนุบาลปีที่ 1/1"""
    meta = kd.KINDER_LEVELS.get((klass.level or "").strip())
    title = (meta or {}).get("title", klass.level or "")
    title = title.replace("ชั้น", "")
    room = (klass.room or "").strip()
    return f"{title}/{room}" if room else title


def _home_page(doc, s, klass):
    _p(doc, "ผู้ปกครองรายงานพฤติกรรมของนักเรียนขณะอยู่ที่บ้าน", align="center", bold=True,
       size=17, after=4, page_break=True)
    _p(doc, f"ชื่อ - นามสกุล  {s.name}    เลขที่ {s.seq or '....'}    ชั้น {_full_class(klass)}",
       size=13, after=3)
    _p(doc, "ให้ผู้ปกครองทำเครื่องหมาย ✓ ที่ตรงกับพฤติกรรมของนักเรียนขณะอยู่ที่บ้าน",
       size=13, after=4)
    t = doc.add_table(rows=2, cols=8)
    t.style = "Table Grid"
    h0, h1 = t.rows[0].cells, t.rows[1].cells
    h0[0].merge(h1[0]); _cell(h0[0], "ลำดับที่", bold=True, size=12, fill="EDE9FE")
    h0[1].merge(h1[1]); _cell(h0[1], "พฤติกรรมของนักเรียนขณะอยู่ที่บ้าน", bold=True, size=12,
                              fill="EDE9FE")
    h0[2].merge(h0[4]); _cell(h0[2], "ภาคเรียนที่ 1", bold=True, size=12, fill="EDE9FE")
    h0[5].merge(h0[7]); _cell(h0[5], "ภาคเรียนที่ 2", bold=True, size=12, fill="EDE9FE")
    for i, (_v, label) in enumerate(kd.RATINGS):
        _cell(h1[2 + i], label, bold=True, size=10, fill="F5F3FF")
        _cell(h1[5 + i], label, bold=True, size=10, fill="F5F3FF")
    for i, text in enumerate(kd.HOME_ITEMS, start=1):
        cells = t.add_row().cells
        _cell(cells[0], i, size=12)
        _cell(cells[1], text, align="left", size=12)
        for c in cells[2:]:
            _cell(c, "", size=12)
    _widths(t, [Cm(1.3), Cm(6.7), Cm(1.3), Cm(1.3), Cm(1.3), Cm(1.3), Cm(1.3), Cm(1.3)])
    for term in (1, 2):
        _p(doc, "", after=4)
        _p(doc, f"ความเห็นผู้ปกครองภาคเรียนที่ {term}", bold=True, size=13, after=2)
        ct = doc.add_table(rows=1, cols=1)
        ct.style = "Table Grid"
        _lines(ct.rows[0].cells[0], 3, size=12)
        _widths(ct, [Cm(16.0)])
        _p(doc, "", after=22)                       # เว้นที่ให้เซ็นชื่อจริง
        _p(doc, "ลงชื่อ ................................................ ผู้ปกครอง",
           align="right", size=13, after=1)
        _p(doc, "(................................................)", align="right", size=13, after=0)


# ---------------------------------------------------------------- หน้า 11 สรุป
def _summary_page(doc, school, s, db, notes, res, level):
    klass = s.klass
    _p(doc, "ผลงาน / กิจกรรมที่น่าภาคภูมิใจ", align="center", bold=True, size=18, after=6,
       page_break=True)
    heads = ["วัน / เดือน / ปี", "ผลงาน/ความสำเร็จที่น่าภูมิใจ", "เกียรติคุณที่ได้รับ",
             "หน่วยงานที่จัด"]
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for i, h in enumerate(heads):
        _cell(t.rows[0].cells[i], h, bold=True, size=12, fill="EDE9FE")
    works = notes["works"]
    for w in works:
        cells = t.add_row().cells
        for i, key in enumerate(("date", "work", "award", "org")):
            _cell(cells[i], (w.get(key) or ""), align="left" if i == 1 else "center", size=12)
    for _ in range(max(0, 5 - len(works))):        # เว้นบรรทัดว่างให้เขียนเพิ่ม
        cells = t.add_row().cells
        for c in cells:
            _cell(c, "", size=12)
    _widths(t, [Cm(3.0), Cm(6.4), Cm(3.3), Cm(3.3)])

    _p(doc, "", after=10)
    _p(doc, "สรุปการพัฒนาเด็ก", bold=True, size=16, after=3)
    st = doc.add_table(rows=2, cols=4)
    st.style = "Table Grid"
    h0, h1 = st.rows[0].cells, st.rows[1].cells
    h0[0].merge(h1[0]); _cell(h0[0], "พัฒนาการ", bold=True, size=12, fill="EDE9FE")
    h0[1].merge(h0[3]); _cell(h0[1], "ระดับคุณภาพ", bold=True, size=12, fill="EDE9FE")
    for i, (_v, name, _mean) in enumerate(kd.QUALITY):
        _cell(h1[1 + i], name, bold=True, size=11, fill="F5F3FF")
    over = notes["summary"]
    for key, full, _short in kd.DOMAINS:
        vals = [v for (t_, c), v in res.items() if c.split(":")[0] == key]
        auto = kd.domain_average(vals)
        val = over.get(key) or auto
        cells = st.add_row().cells
        _cell(cells[0], full, align="left", size=12)
        for i, (qv, _name, _mean) in enumerate(kd.QUALITY):
            _cell(cells[1 + i], _TICK if val == qv else "", size=13)
    _widths(st, [Cm(7.0), Cm(3.0), Cm(3.0), Cm(3.0)])

    _p(doc, "", after=6)
    _p(doc, "ระดับคุณภาพของการพัฒนามี 3 ระดับ", size=13, after=2)
    for v, name, meaning in kd.QUALITY:
        _p(doc, f"        ระดับการพัฒนา  {name}  =  {v}  หมายถึง  {meaning}", size=13, after=1)

    _p(doc, "", after=18)
    homerooms = [p.name for p in (klass.homeroom, klass.co_homeroom) if p]
    sig = doc.add_table(rows=0, cols=2)
    r1 = sig.add_row().cells
    for i, (name, role) in enumerate([
            (homerooms[0] if homerooms else "", "ครูประจำชั้น"),
            (getattr(school, "director_name", "") or "", "ผู้บริหารสถานศึกษา")]):
        _cell(r1[i], "\n\nลงชื่อ ..........................................\n"
                     f"({name or '..........................................'})\n{role}\n"
                     "………… / ……………… / …………", size=13)
    _widths(sig, [Cm(8.0), Cm(8.0)])


# ---------------------------------------------------------------- เล่มเต็ม
def _book(doc, school, s, db, *, page_break):
    """สมุดพก 1 คน (11 หน้า) ลงในเอกสารที่ให้มา"""
    level = (s.klass.level or "").strip()
    meta = kd.KINDER_LEVELS.get(level)
    if not meta:
        return False
    res = _results_of(db, s.id)
    notes = _notes_of(db, s.id)
    _cover(doc, school, s, meta, page_break=page_break)
    _personal(doc, school, s, db)
    _attendance(doc, school, s, db)
    for key, full, _short in kd.DOMAINS:
        items = kd.items_for(level, key)
        if key == "intel" and len(items) > 10:      # สติปัญญา 20 ข้อ -> แบ่ง 2 หน้าตามต้นฉบับ
            _domain_page(doc, s, db, level, key, full + " (ข้อ 1 - 10)", items[:10], res)
            _domain_page(doc, s, db, level, key, full + " (ข้อ 11 - 20)", items[10:], res)
        else:
            _domain_page(doc, s, db, level, key, full, items, res)
    _teacher_comments(doc, s, notes)
    _home_page(doc, s, s.klass)
    _summary_page(doc, school, s, db, notes, res, level)
    return True


def _new_doc():
    doc = Document()
    set_a4(doc)
    style = doc.styles["Normal"]
    style.font.name = THAI_FONT
    style.font.size = Pt(14)
    return doc


def render_kinder_book(school, student, db) -> str:
    """สมุดพกอนุบาลรายคน"""
    doc = _new_doc()
    if not _book(doc, school, student, db, page_break=False):
        _p(doc, "ชั้นนี้ไม่ใช่ระดับปฐมวัย จึงออกสมุดพกอนุบาลไม่ได้", align="center", size=16)
    path = get_data_dir() / f"สมุดพกอนุบาล_{_safe(student.name)}.docx"
    doc.save(path)
    return str(path)


def render_kinder_class(school, klass, db) -> str:
    """สมุดพกอนุบาลทั้งห้อง (คนละเล่มต่อกัน ขึ้นหน้าใหม่ทุกคน)"""
    doc = _new_doc()
    students = sorted(klass.students, key=lambda x: (x.seq or 999, x.name))
    first = True
    for s in students:
        if _book(doc, school, s, db, page_break=not first):
            first = False
    if first:
        _p(doc, "ห้องนี้ยังไม่มีนักเรียน หรือไม่ใช่ระดับปฐมวัย", align="center", size=16)
    path = get_data_dir() / f"สมุดพกอนุบาล_{_safe(_class_label(klass))}.docx"
    doc.save(path)
    return str(path)
