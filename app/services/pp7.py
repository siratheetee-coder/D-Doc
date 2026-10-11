"""PP7 grade certificate: explicit historical identity, frozen issue payload."""
from collections import defaultdict
from datetime import datetime
from app.models import AcadStudent, AcadClass, AcadSubject, AcadScore, Student
from app.thai_utils import is_secondary


def history(db, student, start, end, term=0):
    if term not in (0,1,2):
        raise ValueError('เลือกทุกภาคเรียน ภาคเรียนที่ 1 หรือภาคเรียนที่ 2')
    if not 2500 <= start <= end <= 2800 or end-start > 12:
        raise ValueError('เลือกช่วงปีการศึกษาไม่เกิน 13 ปี และปีเริ่มต้นไม่เกินปีสิ้นสุด')
    q=db.query(AcadStudent).join(AcadClass).filter(AcadClass.year.between(start,end))
    q=q.filter(AcadStudent.student_id==student.student_id) if student.student_id else q.filter(AcadStudent.id==student.id)
    records=q.order_by(AcadClass.year,AcadClass.id).all()
    warnings=[];groups=[]
    if not student.student_id:
        warnings.append('นักเรียนยังไม่เชื่อมทะเบียนกลาง จึงแสดงได้เฉพาะข้อมูลห้องนี้ ไม่จับคู่ประวัติจากชื่อ')
    for year in range(start,end+1):
        if not any(s.klass.year==year for s in records):warnings.append(f'ไม่พบประวัตินักเรียนปี {year}')
    seen=set()
    for s in records:
        key=(s.klass.year,s.klass.level)
        if key in seen:
            raise ValueError(f'พบประวัติห้องซ้ำปี {key[0]} ชั้น {key[1]} กรุณาตรวจทะเบียนก่อนออกเอกสาร')
        seen.add(key)
        subs=db.query(AcadSubject).filter_by(year=s.klass.year,level=s.klass.level).order_by(AcadSubject.term,AcadSubject.kind,AcadSubject.seq,AcadSubject.code).all()
        if term and is_secondary(s.klass.level):
            subs=[sub for sub in subs if sub.term==term]
        scores={(x.subject_id,x.term):x for x in db.query(AcadScore).filter_by(acad_student_id=s.id)}
        terms=sorted({x.term or 0 for x in subs})
        if not subs:warnings.append(f'ยังไม่มีรายวิชาปี {key[0]} ชั้น {key[1]}')
        for subject_term in terms:
            rows=[]
            for sub in subs:
                if (sub.term or 0)!=subject_term:continue
                score=scores.get((sub.id,subject_term))
                grade=(score.grade or '').strip() if score else ''
                weight=float((sub.credit if is_secondary(s.klass.level) else sub.hours) or 0)
                if not grade:warnings.append(f'ยังไม่มีผลการเรียน {key[0]} {sub.code} {sub.name}'+(f' เทอม {subject_term}' if subject_term else ' รายปี'))
                if weight<=0:warnings.append(f'ยังไม่ระบุเวลาเรียน/หน่วยกิต {key[0]} {sub.code} {sub.name}')
                rows.append(dict(code=sub.code or '',name=sub.name,kind=sub.kind or '',grade=grade,weight=weight,group=sub.learn_group or ''))
            groups.append(dict(year=key[0],level=key[1],term=subject_term,unit='หน่วยกิต' if is_secondary(s.klass.level) else 'ชั่วโมง',rows=rows))
    if not groups:warnings.append('ไม่มีผลการเรียนในช่วงปีที่เลือก')
    return groups,warnings,records


def average(rows):
    total=points=0
    for r in rows:
        try:g=float(r['grade'])
        except (ValueError,TypeError):return None
        if not 0<=g<=4 or r['weight']<=0:return None
        total+=r['weight'];points+=g*r['weight']
    from app.services.academic import weighted_avg
    return weighted_avg([(r['grade'],r['weight']) for r in rows]) if total else None


def render(payload, output):
    import math
    import unicodedata
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from app.services.doc_page import set_a4
    from app.services.build_templates import _font, _krut_center, _no_split_row
    from app.services.cash_report import _set_cell, _p

    doc=Document();set_a4(doc);_font(doc)
    sec=doc.sections[0]
    sec.left_margin=sec.right_margin=Cm(1.2)
    sec.top_margin=sec.bottom_margin=Cm(0.9)
    groups=payload['groups']
    slots=32

    def borders(cell, **edges):
        pr=cell._tc.get_or_add_tcPr()
        b=pr.find(qn('w:tcBorders'))
        if b is None:b=OxmlElement('w:tcBorders');pr.append(b)
        for edge,style in edges.items():
            el=OxmlElement('w:'+edge)
            el.set(qn('w:val'),style);el.set(qn('w:sz'),'4');el.set(qn('w:color'),'000000')
            b.append(el)

    def table(widths):
        t=doc.add_table(rows=1,cols=len(widths));t.autofit=False
        for c,w in zip(t.columns,widths):c.width=Cm(w)
        for c,w in zip(t.rows[0].cells,widths):c.width=Cm(w)
        pr=t._tbl.tblPr
        margins=OxmlElement('w:tblCellMar')
        for edge,value in [('top',0),('bottom',0),('left',35),('right',35)]:
            el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa');margins.append(el)
        pr.append(margins)
        return t

    def fields(items):
        # Separate fixed-width value cells retain a dotted baseline after filling.
        t=table([w for label,value,lw,vw in items for w in (lw,vw)])
        for i,(label,value,lw,vw) in enumerate(items):
            _set_cell(t.cell(0,i*2),label,size=13)
            c=t.cell(0,i*2+1);_set_cell(c,str(value or ''),size=13)
            borders(c,bottom='dotted')
        _no_split_row(t.rows[0])

    def span(text):
        visible=sum(not unicodedata.combining(c) for c in text)
        return max(1,math.ceil(visible/32))

    # Reserve lines for group headings and GPA. Long subject names get merged
    # vertical cells, so they wrap without squeezing adjacent grades.
    panels=[];current=[];used=0
    def flush():
        nonlocal current,used
        if current:panels.append(current)
        current=[];used=0
    def add(text,weight='',grade='',height=1,bold=False):
        nonlocal used
        current.append((text,weight,grade,height,bold));used+=height
    for g in groups:
        title=f"ปีการศึกษา {g['year']} ชั้น {g['level']}"+(f" ภาค {g['term']}" if g['term'] else '')
        required=2+sum(span(f"{r['code']} {r['name']}") for r in g['rows'])
        if used and used+required+1>slots-2:flush()
        if used:add('')
        add(title,bold=True)
        for r in g['rows']:
            text=f"{r['code']} {r['name']}";height=span(text)
            if used+height+1>slots-2:
                flush();add(title+' (ต่อ)',bold=True)
            add(text,f"{r['weight']:g}",r['grade'] or '—',height)
        avg=average(g['rows'])
        add('ผลการเรียนเฉลี่ย','',f'{avg:.2f}' if avg is not None else '—',bold=True)
    flush()
    if not panels:panels=[[]]
    for page in range(0,len(panels),3):
        if page:doc.add_page_break()
        _p(doc,'ปพ.7',align='right',size=11,after=0)
        _krut_center(doc,height_cm=1.6)
        _p(doc,'ใบรับรองผลการเรียน',align='center',bold=True,size=18,after=2)
        _p(doc,f"เลขที่ {payload['number']}",size=13,after=1)
        fields([('ขอรับรองว่า',payload['name'],2,10),('เลขประจำตัวนักเรียน',payload['student_no'],3.4,3.2)])
        fields([('เป็นนักเรียนโรงเรียน',payload['school'].removeprefix('โรงเรียน'),3.1,15.5)])
        fields([('เลขประจำตัวประชาชน',payload['id_card'],3.5,6.5),('เกิดวันที่',payload['birthdate'],1.5,7.1)])
        fields([('ชื่อบิดา',payload['father'],1.4,7.9),('ชื่อมารดา',payload['mother'],1.6,7.7)])
        _p(doc,f"{payload['status']} มีรายวิชาและผลการเรียนช่วงปี {payload['start']}–{payload['end']} ดังนี้",size=12,after=2)

        t=table([4.5,.9,.8]*3)
        for col in range(3):
            for j,text in enumerate(['รหัส / รายวิชา','ชม./นก.','ผล']):
                _set_cell(t.cell(0,col*3+j),text,bold=True,size=11,align='center')
                borders(t.cell(0,col*3+j),top='single',bottom='single',left='single',right='single')
        for i in range(slots):
            row=t.add_row();row.height=Cm(.40);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
            _no_split_row(row)
            for c in row.cells:
                _set_cell(c,'',size=11)
                borders(c,left='single',right='single',bottom='single' if i==slots-1 else 'nil',top='nil')
        for col,panel in enumerate(panels[page:page+3]):
            pos=1
            for text,weight,grade,height,bold in panel:
                for j,value in enumerate([text,weight,grade]):
                    c=t.cell(pos,col*3+j)
                    if height>1:c=c.merge(t.cell(pos+height-1,col*3+j))
                    _set_cell(c,value,bold=bold,size=11,align='left' if j==0 else 'center')
                    c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pos+=height
        if page+3>=len(panels):
            same_unit=len({g['unit'] for g in groups})==1
            avg=average([r for g in groups for r in g['rows']]) if same_unit and not payload['warnings'] else None
            label=t.cell(slots,6).merge(t.cell(slots,7))
            _set_cell(label,'ผลการเรียนเฉลี่ยรวม',bold=True,size=11)
            _set_cell(t.cell(slots,8),f'{avg:.2f}' if avg is not None else '—',bold=True,size=11,align='center')
            for c in [label,t.cell(slots,8)]:borders(c,top='single',bottom='single')
        _p(doc,f"ออกให้ ณ วันที่ {payload['issued']}   ใบรับรองมีอายุ {payload['valid_days']} วันนับแต่วันที่ออก",size=12,after=2)
        footer=table([4,7.3,7.3])
        photo=footer.cell(0,0)
        _set_cell(photo,'ติดรูปถ่ายนักเรียน\nขนาด 2 นิ้ว\n(4 × 6 ซม.)',align='center',size=11)
        borders(photo,top='single',bottom='single',left='single',right='single')
        footer.rows[0].height=Cm(6);footer.rows[0].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        _set_cell(footer.cell(0,1),f"ลงชื่อ....................................\n({payload['registrar']})\nนายทะเบียน",align='center',size=13)
        _set_cell(footer.cell(0,2),f"ลงชื่อ....................................\n({payload['director']})\nผู้อำนวยการโรงเรียน",align='center',size=13)
        for c in footer.rows[0].cells:c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    def normalize(container):
        for p in container.paragraphs:
            p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=(1 if p._element.xpath('.//w:drawing') else
                Pt(max([r.font.size.pt if r.font.size else 11 for r in p.runs] or [11])*1.05))
            for r in p.runs:
                size=r.font.size.pt if r.font.size else 11
                r.font.name='TH Sarabun New'
                pr=r._element.get_or_add_rPr()
                cs=pr.find(qn('w:szCs'))
                if cs is None:cs=OxmlElement('w:szCs');pr.append(cs)
                cs.set(qn('w:val'),str(int(size*2)))
        for t in container.tables:
            for row in t.rows:
                for c in row.cells:normalize(c)
    normalize(doc)
    doc.save(output)
    return str(output)
