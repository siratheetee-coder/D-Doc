# -*- coding: utf-8 -*-
"""
w119_doc.py - เอกสารประกอบการเบิกจ่ายตาม ว.119 ตารางที่ 2

ตารางที่ 2 ไม่ใช่การจัดซื้อจัดจ้างตาม พ.ร.บ.ฯ จึงไม่มีรายงานขอความเห็นชอบ
เอกสารทั้งหมดเป็น "หลักฐานการเบิกจ่าย" ซึ่งต่างกันตาม 2 กรณี

  กรณีจ่ายตรงให้ผู้ขาย/ผู้รับจ้าง
    1. บันทึกขออนุมัติเบิกจ่ายเงิน (proc_alt_doc.render_w119_t2)
    2. ใบส่งของ/ใบแจ้งหนี้            <- ผู้ขายออกให้ ไม่ใช่เอกสารที่ระบบพิมพ์
    3. สำเนาแผนงาน/โครงการ
    4. ใบลงเวลาผู้เข้าร่วม (กรณีอบรม/ประชุม)
    5. ภาพถ่าย (กรณีค่าอาหาร/อาหารว่าง)

  กรณียืมเงิน
    ตอนยืม : บันทึกขออนุมัติยืมเงิน · สัญญาการยืมเงิน (แบบ 8500) · สำเนาโครงการ
    ตอนส่งใช้: บันทึกขอส่งใช้เงินยืม · หลักฐานการจ่าย · สำเนาบัตรผู้รับเงิน ·
               ใบลงเวลา · ภาพถ่าย

แบบฟอร์มที่อ้างอิง (ไม่ได้แต่งเอง)
  - ใบสำคัญรับเงินสำหรับวิทยากร = เอกสารหมายเลข 1 แนบท้ายระเบียบกระทรวงการคลัง
    ว่าด้วยค่าใช้จ่ายในการฝึกอบรม การจัดงาน และการประชุมระหว่างประเทศ พ.ศ. 2549
  - ใบสำคัญรับเงินค่าใช้จ่ายในการเดินทางมาฝึกอบรมสำหรับบุคคลภายนอก = เอกสารหมายเลข 2
  - ใบลงเวลาผู้เข้าร่วม = แบบที่ส่วนราชการใช้ทั่วไป (ที่ · ชื่อ-สกุล · หน่วยงาน ·
    ลายมือชื่อเช้า/บ่าย · ลงชื่อผู้ควบคุม)
  - บันทึกขออนุมัติยืมเงิน / สัญญา 8500 / บันทึกขอส่งใช้เงินยืม = ใช้ฟอร์มเดิม
    ของระบบ (lunch_ingredient_doc) ที่ทำตามคู่มืออาหารกลางวัน สพฐ. อยู่แล้ว
"""
import json
from datetime import datetime

from docx import Document
from docx.shared import Cm

from app.services.doc_page import set_a4
from app.database import get_data_dir
from app.thai_utils import thai_date, bahttext, parse_be_date
from app.services.book_receipt_doc import _safe
from app.services.build_templates import (
    _font, _p, _p_runs, _set_cell, _repeat_header_row, _no_split_row, _no_borders,
    _fixed_cols, _sign_table,
)

_BLANK = "................................"
_DOT = "..........................."

# วิธีจ่ายเงินของตารางที่ 2 (ตัดสินว่าจะต้องมีเอกสารยืมเงินหรือไม่)
PAY_DIRECT = "จ่ายตรงให้ผู้ขาย/ผู้รับจ้าง"
PAY_BORROW = "ยืมเงินไปจ่าย"
PAY_MODES = [PAY_DIRECT, PAY_BORROW]

# รายการค่าใช้จ่ายตามตารางที่ 2 (15 ข้อตามตัวหนังสือ) + งานฝึกอบรม/จัดงาน/ประชุม
T2_ITEMS = [
    "ค่าอาหารว่างและเครื่องดื่ม / ค่าอาหาร (ประชุมคาบเกี่ยวมื้ออาหาร)",
    "ค่าอาหารว่างและเครื่องดื่ม (หน่วยงานอื่นเข้าดูงาน)",
    "ค่าธรรมเนียมในการคืนบัตร เปลี่ยนบัตรโดยสาร",
    "ค่าธรรมเนียมอื่น ๆ ที่มิใช่ค่าธรรมเนียมธนาคาร",
    "ค่าใช้บริการอินเทอร์เน็ตของผู้เดินทางไปราชการ",
    "ค่าผ่านทางด่วนพิเศษ",
    "ค่าพาหนะ (ได้รับมอบหมายให้เดินทางไปปฏิบัติราชการ)",
    "ค่าตรวจสอบเพื่อการรับรองระบบการทำงาน",
    "ค่าระวาง",
    "ค่าตรวจร่างกาย",
    "ค่ารักษาสัตว์",
    "ค่าสาธารณูปโภค",
    "การบริจาคหรือการดำเนินการเพื่อการกุศล",
    "ค่าสมาชิกหรือค่าบำรุง",
    "ค่าตอบแทนวิทยากร",
    "ค่าใช้จ่ายในการฝึกอบรม / จัดงาน / ประชุม (ตามระเบียบฝึกอบรมฯ)",
]

# รายการที่ต้องแนบใบลงเวลา / ภาพถ่าย (ตามที่หน่วยตรวจสอบเรียกดูจริง)
_NEED_ATTENDANCE = ("อาหาร", "ฝึกอบรม", "ประชุม", "วิทยากร", "ดูงาน")
_NEED_PHOTO = ("อาหาร",)
_NEED_SPEAKER = ("วิทยากร", "ฝึกอบรม")


def _money(v) -> str:
    return f"{float(v or 0):,.2f}"


def _new():
    doc = Document(); set_a4(doc); _font(doc)
    return doc


def _save(doc, name: str) -> str:
    out = get_data_dir() / "documents"
    out.mkdir(exist_ok=True)
    path = out / (_safe(name) + ".docx")
    doc.save(str(path))
    return str(path)


def extra(proc) -> dict:
    try:
        return json.loads(proc.case_extra or "{}") or {}
    except Exception:
        return {}


def _x(ex, key, default=_BLANK) -> str:
    return (ex.get(key) or "").strip() or default


def _xd(ex, key):
    raw = (ex.get(key) or "").strip()
    return parse_be_date(raw) if raw else None


def _sname(school) -> str:
    return (school.name or "โรงเรียน").strip()


def _total(proc) -> float:
    return float(proc.total_amount or 0)


# ---------------------------------------------------------------- เช็คลิสต์เอกสาร
def checklist(proc) -> dict:
    """เอกสารที่ต้องมีสำหรับเรื่องนี้ แยกเป็น "ระบบพิมพ์ให้" กับ "ต้องแนบเอง"

    ใช้ทั้งแสดงบนหน้าจอและตัดสินว่าจะขึ้นปุ่มดาวน์โหลดใบไหน
    """
    ex = extra(proc)
    item = (ex.get("t2_item") or "").strip()
    borrow = (ex.get("pay_mode") or PAY_DIRECT).strip() == PAY_BORROW
    need_att = any(k in item for k in _NEED_ATTENDANCE)
    need_photo = any(k in item for k in _NEED_PHOTO)
    need_speaker = any(k in item for k in _NEED_SPEAKER)

    docs = [("w119t2", "บันทึกขออนุมัติเบิกจ่ายเงิน")]
    if borrow:
        docs = [("w119_borrow", "บันทึกขออนุมัติยืมเงิน + สัญญาการยืมเงิน (แบบ 8500)"),
                ("w119_repay", "บันทึกขอส่งใช้เงินยืม")] + docs
    if need_att:
        docs.append(("w119_attend", "ใบลงเวลาผู้เข้าร่วม"))
    if need_speaker:
        docs.append(("w119_speaker", "ใบสำคัญรับเงินสำหรับวิทยากร"))
    if need_att:
        docs.append(("w119_trainee", "ใบสำคัญรับเงินค่าพาหนะ/อาหาร ผู้เข้าอบรม (บุคคลภายนอก)"))

    attach = ["ใบส่งของ / ใบแจ้งหนี้ / ใบเสร็จรับเงินของผู้ขาย",
              "สำเนาแผนงาน/โครงการที่อนุมัติแล้ว"]
    if borrow:
        attach.append("สำเนาบัตรประจำตัวประชาชนของผู้รับเงิน")
    if need_photo:
        attach.append("ภาพถ่ายการจัดอาหาร/อาหารว่าง")
    return {"item": item, "borrow": borrow, "docs": docs, "attach": attach}


# ---------------------------------------------------------------- ใบลงเวลาผู้เข้าร่วม
def render_attendance(proc, school) -> str:
    """ใบลงเวลาผู้เข้าร่วม - แนบเป็นหลักฐานเมื่อเบิกค่าอาหาร/อาหารว่าง/ฝึกอบรม"""
    doc = _new()
    ex = extra(proc)
    rows = max(10, int(float(ex.get("participants") or 0) or 0))
    _p(doc, "ใบลงเวลาผู้เข้าร่วม", align="center", bold=True, size=18, after=0)
    _p(doc, (proc.subject or "").strip() or _BLANK, align="center", bold=True, size=16, after=0)
    _p(doc, f"ณ {_x(ex, 'event_place', _sname(school))}", align="center", size=15, after=0)
    ev = _xd(ex, "event_date") or proc.request_date
    _p(doc, f"วันที่ {thai_date(ev) if ev else _BLANK}  เวลา {_x(ex, 'event_time', '08.30 - 16.30 น.')}",
       align="center", size=15, after=6)

    headers = ["ที่", "ชื่อ - สกุล", "ตำแหน่ง/หน่วยงาน",
               "ลายมือชื่อ\n(ช่วงเช้า)", "ลายมือชื่อ\n(ช่วงบ่าย)", "หมายเหตุ"]
    widths = [Cm(1.0), Cm(4.6), Cm(3.6), Cm(2.8), Cm(2.8), Cm(1.7)]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    _fixed_cols(t, widths)
    _repeat_header_row(t.rows[0]); _no_split_row(t.rows[0])
    for c, h, w in zip(t.rows[0].cells, headers, widths):
        _set_cell(c, h, bold=True, align="center", size=14)
        c.width = w
    for i in range(1, rows + 1):
        r = t.add_row(); _no_split_row(r)
        for c, v, w, al in zip(r.cells, [str(i), "", "", "", "", ""], widths,
                               ["center", "left", "left", "center", "center", "center"]):
            _set_cell(c, v, align=al, size=14)
            c.width = w

    _p(doc, "", after=10)
    _sign_table(doc, [[("", "center")], [
        ("(ลงชื่อ) ...................................... ผู้ควบคุม", "center"),
        (f"( {_x(ex, 'responsible')} )", "center"),
        ("ผู้รับผิดชอบกิจกรรม", "center"),
    ]])
    return _save(doc, f"ใบลงเวลาผู้เข้าร่วม_{(proc.memo_no or proc.id)}")


# ------------------------------------------------- ใบสำคัญรับเงินสำหรับวิทยากร
def render_speaker_receipt(proc, school) -> str:
    """เอกสารหมายเลข 1 แนบท้ายระเบียบกระทรวงการคลังว่าด้วยค่าใช้จ่ายในการฝึกอบรมฯ 2549"""
    doc = _new()
    ex = extra(proc)
    hours = float(ex.get("speaker_hours") or 0)
    rate = float(ex.get("speaker_rate") or 0)
    amount = round(hours * rate, 2) or _total(proc)

    _p(doc, "ใบสำคัญรับเงินสำหรับวิทยากร", align="center", bold=True, size=18, after=8)
    _p(doc, f"ชื่อส่วนราชการผู้จัดฝึกอบรม  {_sname(school)}", size=15, after=1)
    _p(doc, f"โครงการ/หลักสูตร  {(proc.subject or '').strip() or _BLANK}", size=15, after=1)
    ev = _xd(ex, "event_date") or proc.request_date
    _p(doc, f"วันที่  {thai_date(ev) if ev else _BLANK}", size=15, after=6)

    _p(doc, f"ข้าพเจ้า {_x(ex, 'speaker_name')}  อยู่บ้านเลขที่ {_x(ex, 'speaker_addr')}",
       indent=1.25, size=15, after=1)
    _p(doc, f"ได้รับเงินจาก {_sname(school)} ดังรายละเอียดต่อไปนี้", indent=1.25, size=15, after=4)

    widths = [Cm(11.0), Cm(5.5)]
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    _fixed_cols(t, widths)
    for c, h, w in zip(t.rows[0].cells, ["รายการ", "จำนวนเงิน"], widths):
        _set_cell(c, h, bold=True, align="center", size=15)
        c.width = w
    detail = (f"ค่าสมนาคุณวิทยากร จำนวน {hours:g} ชั่วโมง × {_money(rate)} บาท"
              if hours and rate else "ค่าสมนาคุณวิทยากร")
    for vals in [(detail, _money(amount)), ("", ""), ("", ""),
                 ("รวมเป็นเงินทั้งสิ้น", _money(amount))]:
        r = t.add_row(); _no_split_row(r)
        bold = vals[0].startswith("รวม")
        for c, v, w, al in zip(r.cells, vals, widths, ["left", "right"]):
            _set_cell(c, v, align=al, size=15, bold=bold)
            c.width = w

    _p(doc, f"จำนวนเงิน  ({bahttext(amount)})", indent=1.25, size=15, before=6, after=12)
    _sign_table(doc, [[("", "center")], [
        ("(ลงชื่อ) ...................................... ผู้รับเงิน", "center"),
        (f"( {_x(ex, 'speaker_name')} )", "center"),
        ("", "center"),
        ("(ลงชื่อ) ...................................... ผู้จ่ายเงิน", "center"),
        (f"( {_x(ex, 'responsible')} )", "center"),
    ]])
    return _save(doc, f"ใบสำคัญรับเงินวิทยากร_{(proc.memo_no or proc.id)}")


# ------------------- ใบสำคัญรับเงินค่าใช้จ่ายเดินทางมาฝึกอบรม (บุคคลภายนอก)
def render_trainee_receipt(proc, school) -> str:
    """เอกสารหมายเลข 2 แนบท้ายระเบียบฯ ฝึกอบรม - จ่ายค่าอาหาร/ที่พัก/พาหนะ
    ให้ผู้เข้าอบรมที่เป็นบุคคลภายนอก โดยเซ็นรับเงินรายคนในใบเดียว"""
    doc = _new(); set_a4(doc, landscape=True); _font(doc)
    ex = extra(proc)
    rows = max(10, int(float(ex.get("participants") or 0) or 0))
    _p(doc, "ใบสำคัญรับเงินค่าใช้จ่ายในการเดินทางมาฝึกอบรมสำหรับบุคคลภายนอก",
       align="center", bold=True, size=17, after=0)
    _p(doc, f"ชื่อส่วนราชการผู้จัดฝึกอบรม  {_sname(school)}", align="center", size=15, after=0)
    _p(doc, f"โครงการ/หลักสูตร  {(proc.subject or '').strip() or _BLANK}",
       align="center", size=15, after=0)
    ev = _xd(ex, "event_date") or proc.request_date
    _p(doc, f"วันที่ {thai_date(ev) if ev else _BLANK} · ผู้เข้ารับการฝึกอบรมรวมทั้งสิ้น "
            f"{_x(ex, 'participants', '.....')} คน", align="center", size=15, after=6)

    headers = ["ลำดับ", "ชื่อ - สกุล", "ที่อยู่", "ค่าอาหาร\n(บาท)", "ค่าที่พัก\n(บาท)",
               "ค่าพาหนะ\n(บาท)", "รวมเป็นเงิน\n(บาท)", "วัน เดือน ปี\nที่รับเงิน",
               "ลายมือชื่อผู้รับเงิน"]
    widths = [Cm(1.3), Cm(4.6), Cm(4.6), Cm(2.3), Cm(2.3), Cm(2.3), Cm(2.5), Cm(2.6), Cm(4.2)]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    _fixed_cols(t, widths)
    _repeat_header_row(t.rows[0]); _no_split_row(t.rows[0])
    for c, h, w in zip(t.rows[0].cells, headers, widths):
        _set_cell(c, h, bold=True, align="center", size=13)
        c.width = w
    for i in range(1, rows + 1):
        r = t.add_row(); _no_split_row(r)
        for c, v, w in zip(r.cells, [str(i)] + [""] * 8, widths):
            _set_cell(c, v, align="center" if v else "left", size=13)
            c.width = w
    r = t.add_row(); _no_split_row(r)
    _set_cell(r.cells[2], "รวมเป็นเงินทั้งสิ้น", bold=True, align="right", size=13)
    for c, w in zip(r.cells, widths):
        c.width = w

    _p(doc, "", after=10)
    _sign_table(doc, [[("", "center")], [
        ("(ลงชื่อ) ...................................... ผู้จ่ายเงิน", "center"),
        (f"( {_x(ex, 'responsible')} )", "center"),
        ("ตำแหน่ง ......................................", "center"),
    ]])
    return _save(doc, f"ใบสำคัญรับเงินผู้เข้าอบรม_{(proc.memo_no or proc.id)}")


# ---------------------------------------------------------------- เอกสารยืมเงิน
class _LoanView:
    """ปรับข้อมูลเรื่องจัดซื้อให้อยู่ในรูปที่ฟอร์มยืมเงินเดิมของระบบใช้ได้
    (ฟอร์ม 8500 / บันทึกขออนุมัติยืมเงิน / บันทึกส่งใช้ ใน lunch_ingredient_doc)"""

    def __init__(self, proc, school):
        ex = extra(proc)
        self.id = proc.id
        self.borrower = _x(ex, "borrower", "")
        self.position = _x(ex, "borrower_pos", "")
        self.amount = _total(proc)
        self.contract_no = _x(ex, "contract_no", "")
        self.date = _xd(ex, "borrow_date") or proc.request_date
        self.receive_date = self.date
        self.due_date = _xd(ex, "due_date")
        self.within_days = int(float(ex.get("within_days") or 0) or 30)
        self.submit_to = f"ผู้อำนวยการ{_sname(school)}"
        self.fund_from = (proc.budget_source or "").strip() or "เงินอุดหนุน"
        self.purpose = (proc.subject or "").strip()
        self.returns = []


def render_borrow_set(proc, school) -> str:
    """บันทึกขออนุมัติยืมเงิน + สัญญาการยืมเงิน (แบบ 8500) ในไฟล์เดียว"""
    from app.services.lunch_ingredient_doc import render_loan_contract, render_repay_memo  # noqa
    loan = _LoanView(proc, school)
    doc = _new()
    _borrow_memo(doc, proc, school, loan)
    doc.add_page_break()
    render_loan_contract(None, school, doc, finance_loan=loan)
    return _save(doc, f"ขออนุมัติยืมเงิน_{(proc.memo_no or proc.id)}")


def _borrow_memo(doc, proc, school, loan) -> None:
    """บันทึกข้อความ ขออนุมัติยืมเงิน (ตาราง 2 ไม่ใช่การจัดซื้อจัดจ้าง จึงขอยืมเพื่อไปจ่าย)"""
    from app.services.build_templates import _krut_and_title, _hr
    total = _total(proc)
    _krut_and_title(doc)
    _p_runs(doc, [("ส่วนราชการ  ", True),
                  ("  ".join(x for x in [_sname(school), (school.address or "").strip()] if x),
                   False)])
    _p_runs(doc, [("ที่  ", True), ((proc.memo_no or "").strip() or _BLANK, False),
                  ("\t", False), ("วันที่ ", True),
                  (thai_date(loan.date) if loan.date else _DOT, False)], tab_cm=8)
    _p_runs(doc, [("เรื่อง  ", True), ("ขออนุมัติยืมเงิน", False)])
    _p_runs(doc, [("เรียน  ", True), (loan.submit_to, False)])
    _hr(doc)
    _p(doc, f"ด้วย {_sname(school)} จะดำเนินการ{loan.purpose or _BLANK} "
            f"ซึ่งเป็นค่าใช้จ่ายตามตารางที่ 2 แนบท้ายหนังสือด่วนที่สุด ที่ กค (กวจ) 0405.2/ว 119 "
            "ลงวันที่ 7 มีนาคม 2561 อันไม่ถือเป็นการจัดซื้อจัดจ้างตามพระราชบัญญัติการจัดซื้อจัดจ้าง"
            "และการบริหารพัสดุภาครัฐ พ.ศ. 2560 โดยเบิกจ่ายตามระเบียบการเงินที่เกี่ยวข้อง",
       align="justify", indent=1.25, after=2)
    _p(doc, f"ในการนี้ จึงขออนุมัติให้ {loan.borrower or _BLANK} "
            f"ยืมเงิน{loan.fund_from} จำนวน {_money(total)} บาท ({bahttext(total)}) "
            f"เพื่อเป็นค่าใช้จ่ายดังกล่าว และจะส่งใช้เงินยืมภายใน {loan.within_days} วัน "
            "นับแต่วันที่ได้รับเงิน", align="justify", indent=1.25, after=2)
    _p(doc, "จึงเรียนมาเพื่อโปรดพิจารณาอนุมัติ", align="justify", indent=1.25, after=12)
    _sign_table(doc, [[
        ("ลงชื่อ ...................................... ผู้ขอยืม", "center"),
        (f"( {loan.borrower or _BLANK} )", "center"),
    ], [
        ("ลงชื่อ ......................................", "center"),
        (f"( {(school.director_name or '').strip() or _BLANK} )", "center"),
        (loan.submit_to, "center"),
    ]])


def render_repay_set(proc, school) -> str:
    """บันทึกขอส่งใช้เงินยืม (ใช้ฟอร์มเดิมของระบบ)"""
    from app.services.lunch_ingredient_doc import render_repay_memo
    loan = _LoanView(proc, school)
    doc = _new()
    render_repay_memo(None, school, doc, finance_loan=loan)
    return _save(doc, f"ส่งใช้เงินยืม_{(proc.memo_no or proc.id)}")


RENDERERS = {
    "w119_attend": render_attendance,
    "w119_speaker": render_speaker_receipt,
    "w119_trainee": render_trainee_receipt,
    "w119_borrow": render_borrow_set,
    "w119_repay": render_repay_set,
}
