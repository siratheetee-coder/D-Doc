# -*- coding: utf-8 -*-
"""
book_tor_doc.py - ขอบเขตของงาน (TOR) สำหรับการจัดซื้อหนังสือเรียน

โครงสร้าง 10 หัวข้อตามแบบที่โรงเรียนใช้จริง (ครบกว่า TOR พัสดุทั่วไป):
  1 ความเป็นมา · 2 วัตถุประสงค์ · 3 คุณสมบัติของผู้เสนอราคา
  4 รายละเอียดคุณลักษณะเฉพาะ (ตารางแยกรายชั้น) · 5 ระยะเวลาดำเนินการ
  6 ระยะเวลา/สถานที่ส่งมอบ · 7 คณะกรรมการ · 8 เงื่อนไข · 9 วงเงิน · 10 ติดต่อสอบถาม
ถ้อยคำอ้างกฎหมายเป็นข้อความราชการสาธารณะ อิงไฟล์ตัวอย่างที่โรงเรียนใช้จริง
"""
import json

from docx import Document
from docx.shared import Cm

from app.services.doc_page import set_a4
from app.database import get_data_dir
from app.thai_utils import bahttext
from app.services.book_receipt_doc import _safe
from app.services.build_templates import (
    _font, _p, _set_cell, _repeat_header_row, _no_split_row, _no_borders, _fixed_cols,
)

_BLANK = "................................"

# คุณสมบัติของผู้เสนอราคา (ข้อกฎหมายมาตรฐาน ใช้เหมือนกันทุกโรงเรียน)
_QUALIFY = [
    "ผู้ประสงค์จะเสนอราคาต้องเป็นผู้มีอาชีพขายพัสดุที่จัดซื้อดังกล่าว",
    "ผู้ประสงค์จะเสนอราคาต้องไม่เป็นผู้ที่ถูกระบุชื่อไว้ในบัญชีรายชื่อผู้ทิ้งงานของทางราชการ "
    "และได้แจ้งเวียนชื่อแล้ว",
    "ผู้ประสงค์จะเสนอราคาต้องไม่เป็นผู้มีผลประโยชน์ร่วมกันกับผู้เสนอราคารายอื่น "
    "และ/หรือต้องไม่เป็นผู้มีผลประโยชน์ร่วมกันระหว่างผู้เสนอราคากับผู้ให้บริการตลาดกลาง "
    "ณ วันประกาศจัดซื้อ หรือไม่เป็นผู้กระทำการอันเป็นการขัดขวางการแข่งขันราคาอย่างเป็นธรรม",
    "ผู้ประสงค์จะเสนอราคาต้องไม่เป็นผู้ได้รับเอกสิทธิ์หรือความคุ้มกัน ซึ่งอาจปฏิเสธไม่ยอมขึ้นศาลไทย "
    "เว้นแต่รัฐบาลของผู้เสนอราคาได้มีคำสั่งให้สละสิทธิ์และความคุ้มกันเช่นว่านั้น",
    "ต้องเป็นผู้ปฏิบัติตามประกาศคณะกรรมการป้องกันและปราบปรามการทุจริตแห่งชาติ เรื่อง "
    "หลักเกณฑ์และวิธีการจัดทำและแสดงบัญชีรายการรับจ่ายของโครงการที่บุคคลหรือนิติบุคคล"
    "เป็นคู่สัญญากับหน่วยงานของรัฐ กล่าวคือ บุคคลหรือนิติบุคคลที่จะเข้าเป็นคู่สัญญา "
    "ต้องไม่อยู่ในฐานะเป็นผู้ไม่แสดงบัญชีรายรับรายจ่าย หรือแสดงบัญชีรายรับรายจ่าย"
    "ไม่ถูกต้องครบถ้วนในสาระสำคัญ",
    "คู่สัญญาต้องรับจ่ายเงินผ่านบัญชีธนาคาร เว้นแต่การรับจ่ายเงินแต่ละครั้งซึ่งมีมูลค่า"
    "ไม่เกินสามหมื่นบาท คู่สัญญาอาจรับจ่ายเป็นเงินสดก็ได้",
]

_DEFAULT_CONDITIONS = [
    "ผู้ซื้อสามารถเพิ่มหรือลดจำนวนและราคาได้ตามจำนวนนักเรียนที่มีอยู่จริง",
    "ผู้ขายส่งมอบสิ่งของที่ซื้อตามรายละเอียดคุณลักษณะเฉพาะที่กำหนด",
    "หากส่งมอบเกินกำหนดเวลา ผู้ขายยินยอมให้ปรับตามอัตราที่ทางราชการกำหนด",
]


def _money(v) -> str:
    return f"{float(v or 0):,.2f}"


def _lines(text, fallback):
    out = [ln.strip().lstrip("-").strip() for ln in (text or "").splitlines() if ln.strip()]
    return out or list(fallback)


def _class_table(doc, level, items):
    """ตารางรายละเอียดคุณลักษณะเฉพาะของชั้นหนึ่ง ๆ -> คืนยอดรวมของชั้น"""
    _p(doc, f"รายละเอียดคุณลักษณะเฉพาะ รายการหนังสือเรียน ชั้น{level}",
       bold=True, indent=0.6, before=6, after=2)
    headers = ["ที่", "รายการหนังสือ", "ราคา", "จำนวน", "เป็นเงิน"]
    widths = [Cm(1.0), Cm(8.2), Cm(2.0), Cm(2.0), Cm(2.6)]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    _fixed_cols(t, widths)
    _repeat_header_row(t.rows[0]); _no_split_row(t.rows[0])
    for c, h, w in zip(t.rows[0].cells, headers, widths):
        _set_cell(c, h, bold=True, align="center", size=14)
        c.width = w
    total = 0.0
    for i, it in enumerate(items, start=1):
        amount = float(it["price"]) * float(it["qty"])
        total += amount
        vals = [str(i), it["title"], _money(it["price"]), f"{it['qty']:g}", _money(amount)]
        r = t.add_row(); _no_split_row(r)
        for c, v, w, al in zip(r.cells, vals, widths,
                               ["center", "left", "right", "center", "right"]):
            _set_cell(c, v, align=al, size=14)
            c.width = w
    r = t.add_row(); _no_split_row(r)
    _set_cell(r.cells[1], "รวม", bold=True, align="right", size=14)
    _set_cell(r.cells[4], _money(total), bold=True, align="right", size=14)
    for c, w in zip(r.cells, widths):
        c.width = w
    return total


def render_book_tor(school, tp, groups) -> str:
    """tp = TextbookPurchase · groups = [(level, [{title, price, qty}]), ...] เรียงตามชั้น"""
    doc = Document(); set_a4(doc)
    _font(doc)
    sname = (school.name or "โรงเรียน").strip()
    n_items = sum(len(items) for _, items in groups)
    fy = tp.fiscal_year or tp.year

    _p(doc, "ขอบเขตของงาน (Terms of Reference : TOR)", align="center", bold=True, size=18, after=0)
    _p(doc, f"การจัดซื้อหนังสือเรียน จำนวน {n_items} รายการ ({tp.budget_source or 'เงินอุดหนุนรัฐบาล'})",
       align="center", bold=True, size=16, after=0)
    _p(doc, f"ประจำปีงบประมาณ พ.ศ. {fy}", align="center", bold=True, size=16, after=0)
    _p(doc, "--------------------------------", align="center", after=6)

    _p(doc, "1. ความเป็นมา", bold=True)
    _p(doc, f"{sname} {(school.address or '').strip()} มีความประสงค์ที่จะซื้อหนังสือเรียน "
            f"จำนวน {n_items} รายการ โดยวิธี{tp.method or 'เฉพาะเจาะจง'} "
            f"ได้รับจัดสรรงบประมาณรายจ่าย ประจำปี พ.ศ. {fy} "
            f"งบ{tp.budget_source or 'เงินอุดหนุนรัฐบาล'} โครงการสนับสนุนค่าใช้จ่ายในการจัดการศึกษา"
            "ตั้งแต่ระดับอนุบาลจนจบการศึกษาขั้นพื้นฐาน", align="justify", indent=1.25, after=2)

    _p(doc, "2. วัตถุประสงค์", bold=True)
    _p(doc, (tp.purpose or "").strip() or
            ("เพื่อเป็นสื่อการเรียนการสอนวิชาต่าง ๆ เพื่อพัฒนาการเรียนรู้และผลสัมฤทธิ์ทางการเรียน "
             f"จึงจัดซื้อหนังสือเรียนให้นักเรียน{sname} ตามจำนวนนักเรียนที่มีอยู่จริง"),
       align="justify", indent=1.25, after=2)

    _p(doc, "3. คุณสมบัติของผู้ประสงค์จะเสนอราคา", bold=True)
    for i, q in enumerate(_QUALIFY, start=1):
        _p(doc, f"3.{i} {q}", align="justify", indent=1.25, after=1)

    _p(doc, f"4. รายละเอียดคุณลักษณะเฉพาะ รายการหนังสือเรียน ปีการศึกษา {tp.year}",
       bold=True, before=4)
    grand = 0.0
    for level, items in groups:
        if items:
            grand += _class_table(doc, level or "-", items)
    if not groups:
        _p(doc, "(ยังไม่มีรายการหนังสือในทะเบียนของปีการศึกษานี้)", indent=1.25)
    _p(doc, f"รวมทั้งสิ้น {n_items} รายการ เป็นเงิน {_money(grand)} บาท ({bahttext(grand)})",
       bold=True, indent=1.25, before=4, after=2)

    _p(doc, "5. ระยะเวลาดำเนินการ", bold=True)
    _p(doc, (tp.period_text or "").strip() or _BLANK, indent=1.25, after=2)

    _p(doc, "6. ระยะเวลาและสถานที่ส่งมอบงาน", bold=True)
    _p(doc, f"กำหนดส่งมอบภายใน {tp.delivery_days or 15} วันทำการ นับถัดจากวันที่ลงนามในใบสั่งซื้อ "
            f"ณ {(tp.delivery_place or '').strip() or sname}", align="justify", indent=1.25, after=2)

    _p(doc, "7. คณะกรรมการจัดทำร่างขอบเขตของงานและผู้ตรวจรับพัสดุ", bold=True)
    try:
        members = json.loads(tp.members or "[]")
    except Exception:
        members = []
    roles = ["ประธานกรรมการ", "กรรมการ", "กรรมการและเลขานุการ"]
    if not members:
        members = [{"name": "", "position": ""} for _ in range(3)]
    mt = doc.add_table(rows=len(members), cols=3)
    _no_borders(mt)
    _fixed_cols(mt, [Cm(7.0), Cm(4.5), Cm(5.0)])
    for i, (row, m) in enumerate(zip(mt.rows, members)):
        name = (m.get("name") or "").strip() or _BLANK
        pos = (m.get("position") or "").strip() or "ครู"
        role = (m.get("role") or "").strip() or (roles[i] if i < len(roles) else "กรรมการ")
        for c, v in zip(row.cells, [f"{i + 1})  {name}", f"ตำแหน่ง  {pos}", role]):
            _set_cell(c, v, size=14, align="left")

    _p(doc, "8. เงื่อนไข", bold=True, before=4)
    for line in _lines(tp.conditions, _DEFAULT_CONDITIONS):
        _p(doc, f"- {line}", align="justify", indent=1.25, after=1)

    budget = float(tp.total_budget or 0) or grand
    ref = float(tp.price_ref or 0) or grand
    _p(doc, "9. วงเงินในการจัดหา", bold=True, before=2)
    _p(doc, f"- วงเงินงบประมาณที่ได้รับ เป็นเงิน {_money(budget)} บาท ({bahttext(budget)})",
       indent=1.25, after=1)
    _p(doc, f"- ราคากลาง เป็นเงิน {_money(ref)} บาท ({bahttext(ref)})", indent=1.25, after=2)

    _p(doc, "10. ติดต่อสอบถามรายละเอียดข้อมูลเพิ่มเติม และส่งข้อเสนอแนะ วิจารณ์ "
            "หรือแสดงความคิดเห็นได้ที่", bold=True)
    _p(doc, (tp.contact or "").strip() or
            (f"{sname} {(school.address or '').strip()} "
             f"โทรศัพท์ {(getattr(school, 'phone', '') or '').strip() or _BLANK}"),
       align="justify", indent=1.25, after=10)

    _p(doc, "ลงชื่อ..............................................ผู้จัดทำร่างขอบเขตของงาน",
       align="right", after=0)
    _p(doc, f"( {(school.officer_name or '').strip() or _BLANK} )", align="right", after=0)

    out_dir = get_data_dir() / "documents"
    out_dir.mkdir(exist_ok=True)
    path = out_dir / (_safe(f"TOR_จัดซื้อหนังสือเรียน_{tp.year}") + ".docx")
    doc.save(str(path))
    return str(path)
