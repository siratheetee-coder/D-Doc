# -*- coding: utf-8 -*-
"""
material_card.py - บัญชีวัสดุ (แบบบัญชีวัสดุของหน่วยงานรัฐ)

ledger_rows: แถวรับ-จ่าย-คงเหลือ เรียงตามวันที่ พร้อมราคาต่อหน่วยทุกแถว
  (แถวจ่ายที่ไม่ได้บันทึกราคาไว้ ใช้ราคารับเข้าครั้งล่าสุดก่อนหน้า)
render_material_cards: ไฟล์ Word แนวนอน วัสดุรายการละแผ่น ช่องตามแบบ + ข้อกำหนดท้ายแบบ
"""
from datetime import datetime

from docx import Document
from docx.shared import Cm

from app.database import get_data_dir
from app.services.build_templates import _font, _p, _set_cell
from app.services.doc_page import set_a4, set_margins
from app.thai_utils import thai_date_short

OPENING = "ยอดยกมา"

FOOTNOTES = [
    "จำนวนอย่างสูง หมายถึง จำนวนวัสดุประเภทนี้ที่เก็บรักษาไว้ได้สูงสุด",
    "จำนวนอย่างต่ำ หมายถึง จำนวนวัสดุประเภทนี้ที่คงเหลือเท่ากับจำนวนที่กำหนดแล้วต้องมีการจัดซื้อเพิ่ม",
    "กรณีรับจากคณะกรรมการตรวจรับให้กรรมการตรวจรับแต่หากมีผู้ตรวจรับเพียงคนเดียวให้ผู้ตรวจรับ",
    "กรณีจ่ายวัสดุให้ใครต้องระบุชื่อผู้รับให้ชัดเจน",
]


def _sorted(item):
    return sorted(item.txns, key=lambda x: (x.date or datetime.min, x.id or 0))


def last_in_price(item) -> float:
    """ราคารับเข้าครั้งล่าสุด (ไว้เติมให้แถวจ่าย และเป็นค่าเริ่มต้นในฟอร์ม)"""
    if item is None:
        return 0.0
    price = 0.0
    for t in _sorted(item):
        if t.kind == "in" and (t.unit_price or 0) > 0:
            price = float(t.unit_price)
    return price


def ledger_rows(item):
    """คืน (rows, ยอดคงเหลือ, ราคารับเข้าล่าสุด) rows = [{t, balance, price, opening}]"""
    rows, bal, price = [], 0.0, 0.0
    for t in _sorted(item):
        if t.kind == "in":
            bal += (t.qty or 0)
            if (t.unit_price or 0) > 0:
                price = float(t.unit_price)
            row_price = float(t.unit_price or 0) or price
        else:
            bal -= (t.qty or 0)
            row_price = float(t.unit_price or 0) or price
        rows.append({"t": t, "balance": bal, "price": row_price,
                     "opening": t.kind == "in" and (t.note or "").strip() == OPENING and not t.party})
    return rows, bal, price


def _num(v) -> str:
    return f"{v:g}" if v else ""


def _money(v) -> str:
    return f"{v:,.2f}" if v else ""


def _row(t):
    """แถวใหม่ สูงอย่างน้อย 0.9 ซม. (บรรทัดว่างพอเขียนด้วยมือ แถวที่ข้อความยาวยังขยายได้)"""
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    r = t.add_row()
    r.height, r.height_rule = Cm(0.9), WD_ROW_HEIGHT_RULE.AT_LEAST
    return r


ROWS_PER_SHEET = 8   # แถวรายการต่อแผ่น (A4 แนวนอน) แผ่นถัดไปขึ้นหัวใหม่ + แถวยกมา


def sheets(rows) -> list:
    """แบ่งแถวเป็นแผ่น: แผ่นแรก ROWS_PER_SHEET แถว แผ่นต่อไปเว้น 1 แถวบนสุดไว้ยกยอดมา"""
    out, rest = [rows[:ROWS_PER_SHEET]], rows[ROWS_PER_SHEET:]
    while rest:
        out.append(rest[:ROWS_PER_SHEET - 1])
        rest = rest[ROWS_PER_SHEET - 1:]
    return out


def _card(doc, school, item, page_no: int, rows=None, carry=None, new_page=False):
    """บัญชีวัสดุ 1 แผ่น rows = แถวของแผ่นนี้ carry = (ยอดคงเหลือ, ราคา) ยกมาจากแผ่นก่อน
    เติมบรรทัดว่างจนเต็มแผ่น ไว้เขียนต่อด้วยมือหลังพิมพ์"""
    office = (getattr(school, "affiliation", "") or getattr(school, "area_office", "") or "").strip()
    sname = (getattr(school, "name", "") or "").strip()
    if sname and not sname.startswith("โรงเรียน"):
        sname = "โรงเรียน" + sname
    title = _p(doc, "บัญชีวัสดุ", align="center", bold=True, size=20, after=0)
    if new_page:      # ขึ้นแผ่นใหม่ที่หัวเรื่อง (ย่อหน้าตัดหน้าแยกจะทิ้งบรรทัดว่างไว้บนแผ่นถัดไปจนล้น)
        title.paragraph_format.page_break_before = True
    _p(doc, f"ส่วนราชการ  {office or '..............................'}", align="right", after=0)
    _p(doc, f"หน่วยงาน  {sname or '..............................'}", align="right", after=2)
    _p(doc, f"แผ่นที่  {page_no}", after=0)
    _p(doc, f"ประเภท  {item.category or 'วัสดุทั่วไป'}      ชื่อหรือชนิดวัสดุ  {item.name}"
            f"      รหัส  {item.code or '..........'}", after=0)
    _p(doc, f"ขนาดหรือลักษณะ  {item.spec or '..............................'}"
            f"      จำนวนอย่างสูง  {_num(item.max_stock) or '..........'}", after=0)
    _p(doc, f"หน่วยที่นับ  {item.unit or ''}      ที่เก็บ  {item.location or '..............................'}"
            f"      จำนวนอย่างต่ำ  {_num(item.min_stock) or '..........'}", after=4)

    # รวม 25.0 ซม. = พื้นที่พิมพ์ A4 แนวนอน (เดิมเกินขอบ ทำให้ข้อความตัดบรรทัดจนล้นแผ่น)
    widths = [Cm(2.6), Cm(6.6), Cm(3.6), Cm(2.3), Cm(1.8), Cm(1.8), Cm(2.0), Cm(4.3)]
    head = ["วัน เดือน ปี", "รับจาก / จ่ายให้", "เลขที่เอกสาร", "ราคาต่อหน่วย", "รับ", "จ่าย", "คงเหลือ", "หมายเหตุ"]
    t = doc.add_table(rows=2, cols=8)
    t.style = "Table Grid"
    for i, (c, h, w) in enumerate(zip(t.rows[0].cells, head, widths)):
        _set_cell(c, h, size=14, align="center", bold=True); c.width = w
    sub = ["", "", "", "", "จำนวน", "จำนวน", "จำนวน", ""]
    for c, h, w in zip(t.rows[1].cells, sub, widths):
        _set_cell(c, h, size=14, align="center"); c.width = w
    for i in (0, 1, 2, 3, 7):            # หัวตารางสองชั้น รวมช่องที่ไม่มีคำว่า "จำนวน"
        t.cell(0, i).merge(t.cell(1, i))
        _set_cell(t.cell(0, i), head[i], size=14, align="center", bold=True)

    if rows is None:
        rows = ledger_rows(item)[0]
    aligns = ["center", "left", "center", "right", "center", "center", "center", "left"]
    used = len(rows)
    if carry is not None:
        vals = ["", f"ยกมาจากแผ่นที่ {page_no - 1}", "", _money(carry[1]), "", "", f"{carry[0]:g}", ""]
        for c, v, w, a in zip(_row(t).cells, vals, widths, aligns):
            _set_cell(c, v, size=14, align=a); c.width = w
        used += 1
    for r in rows:
        tx = r["t"]
        if r["opening"]:
            who = OPENING
        else:
            who = ("รับจาก " if tx.kind == "in" else "จ่ายให้ ") + (tx.party or "")
        vals = [thai_date_short(tx.date) if tx.date else "", who.strip(), tx.ref or "", _money(r["price"]),
                _num(tx.qty) if tx.kind == "in" else "", _num(tx.qty) if tx.kind == "out" else "",
                f"{r['balance']:g}", "" if r["opening"] else (tx.note or "")]
        for c, v, w, a in zip(_row(t).cells, vals, widths, aligns):
            _set_cell(c, v, size=14, align=a); c.width = w
    for _ in range(max(0, ROWS_PER_SHEET - used)):   # เว้นบรรทัดว่างจนเต็มแผ่น ไว้เขียนต่อด้วยมือ
        for c, w in zip(_row(t).cells, widths):
            _set_cell(c, "", size=14); c.width = w

    for i, n in enumerate(FOOTNOTES):      # ชิดตาราง ตัวเล็กลงนิด ให้ 8 แถวที่ตัดบรรทัด + ข้อกำหนดจบในแผ่นเดียว
        _p(doc, "•  " + n, size=13, before=6 if i == 0 else 0, after=0, line=1.0)


def render_material_cards(school, items) -> str:
    """บัญชีวัสดุรายการละแผ่น (A4 แนวนอน) คืน path ไฟล์ .docx"""
    doc = Document()
    _font(doc)                       # _font ตั้งขนาดแนวตั้ง ต้องเรียกก่อนสลับเป็นแนวนอน
    set_a4(doc, landscape=True)
    set_margins(doc, top=Cm(1.0), bottom=Cm(0.9))   # ขอบบน/ล่างแคบลงนิด ให้ 8 แถว + ข้อกำหนดจบในแผ่นเดียว
    first = True
    for item in items:
        rows = ledger_rows(item)[0]
        prev = None
        for no, chunk in enumerate(sheets(rows), 1):
            _card(doc, school, item, no, chunk, carry=(prev["balance"], prev["price"]) if prev else None,
                  new_page=not first)
            first = False
            prev = chunk[-1] if chunk else prev
    out_dir = get_data_dir() / "documents"
    out_dir.mkdir(exist_ok=True)
    name = items[0].name if len(items) == 1 else f"รวม {len(items)} รายการ"
    for ch in '<>:"/\\|?*':
        name = name.replace(ch, "_")
    from app.services.filenames import fit_name
    out_path = out_dir / fit_name(f"บัญชีวัสดุ_{name}.docx")
    doc.save(str(out_path))
    return str(out_path)
