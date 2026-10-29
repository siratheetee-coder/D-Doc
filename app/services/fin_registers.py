# -*- coding: utf-8 -*-
"""
fin_registers.py - ทะเบียนคุมเงินเฉพาะประเภท + บันทึกการรับเงินเพื่อเก็บรักษา

ยึดตามชุดแบบฟอร์มการเงินของสถานศึกษา (สพป.เพชรบุรี เขต 2) ซึ่งมีทะเบียนคุม
แยกตาม "ประเภทเงิน" ไม่ใช่แยกเป็นระบบใหม่ ระบบเราจึงผูกไว้กับ "บัญชี" ที่ครูตั้งเอง
แล้วเลือกรูปแบบฟอร์มให้อัตโนมัติจากประเภทของบัญชีนั้น (special_form)

  - บัญชีชื่อ/ประเภทเป็นเงินประกันสัญญา  -> ทะเบียนคุมเงินฝาก (เงินประกันสัญญา)
  - บัญชี fund_type = เงินรายได้แผ่นดิน  -> ทะเบียนคุมการรับและนำส่งเงินรายได้แผ่นดิน
  - บัญชี deposit_type = agency          -> สมุดคู่ฝาก (ส่วนราชการผู้เบิก)

และอีก 2 ฟอร์มที่เกาะกับหน้าเดิม
  - บันทึกการรับเงินเพื่อเก็บรักษา  -> คู่กับรายงานเงินคงเหลือประจำวัน
  - ทะเบียนคุมหลักฐานขอเบิก        -> คู่กับหน้าบันทึกขอเบิกจ่าย
"""
from datetime import datetime

from docx.shared import Cm

from app.thai_utils import thai_date, thai_date_short, bahttext
from app.services.build_templates import (
    _p, _set_cell, _repeat_header_row, _no_split_row, _fixed_cols, _sign_table,
)
from app.services.finance_forms_doc import _new, _save, _money, _grid, _row, _BLANK
from app.services.org_names import head_title, head_title_short

# ---------------------------------------------------------------- ชนิดฟอร์มพิเศษ
SPECIAL_LABEL = {
    "guarantee": "ทะเบียนคุมเงินฝาก (เงินประกันสัญญา)",
    "revenue": "ทะเบียนคุมการรับและนำส่งเงินรายได้แผ่นดิน",
    "agency": "สมุดคู่ฝาก (ส่วนราชการผู้เบิก)",
}


def special_form(account) -> str:
    """บัญชีนี้ควรพิมพ์ทะเบียนคุมแบบไหน ('' = ใช้บัญชีแยกประเภทปกติ)"""
    name = (account.name or "")
    if "ประกัน" in name:
        return "guarantee"
    if (account.fund_type or "") == "เงินรายได้แผ่นดิน" or "รายได้แผ่นดิน" in name:
        return "revenue"
    if (account.deposit_type or "") == "agency":
        return "agency"
    return ""


def _head2(doc, groups, widths, *, size=12):
    """หัวตาราง 2 ชั้น - groups: list ของ (ชื่อ, [หัวย่อย...])
    หัวย่อยว่าง = คอลัมน์เดียวที่ผสานสองแถวเข้าด้วยกัน"""
    ncol = sum(max(1, len(g[1])) for g in groups)
    t = doc.add_table(rows=2, cols=ncol)
    t.style = "Table Grid"
    _fixed_cols(t, widths)
    _repeat_header_row(t.rows[0]); _repeat_header_row(t.rows[1])
    _no_split_row(t.rows[0]); _no_split_row(t.rows[1])
    c = 0
    for label, subs in groups:
        if not subs:
            cell = t.cell(0, c).merge(t.cell(1, c))
            _set_cell(cell, label, bold=True, align="center", size=size)
            c += 1
        else:
            top = t.cell(0, c).merge(t.cell(0, c + len(subs) - 1))
            _set_cell(top, label, bold=True, align="center", size=size)
            for s in subs:
                _set_cell(t.cell(1, c), s, bold=True, align="center", size=size)
                c += 1
    for r in t.rows:
        for cell, w in zip(r.cells, widths):
            cell.width = w
    return t


def _blank_rows(t, n, widths, ncol):
    for _ in range(n):
        _row(t, [""] * ncol, widths, ["center"] * ncol)


def _sign_finance(doc, school):
    _p(doc, "", after=10)
    _sign_table(doc, [[
        ("ลงชื่อ.......................................เจ้าหน้าที่การเงิน", "center"),
        (f"( {(getattr(school, 'finance_officer_name', '') or getattr(school, 'officer_name', '') or '').strip() or '.......................................'} )", "center"),
    ], [
        ("ลงชื่อ.......................................", "center"),
        (f"( {(school.director_name or '').strip() or _BLANK} )", "center"),
        (head_title_short(school), "center"),
    ]])


def _merge_label(row, first, last, text, *, size=13, bold=True):
    """รวมช่อง first..last ของแถวเป็นช่องเดียว ใส่ข้อความกึ่งกลาง (แถว "รวม" แบบฟอร์มราชการ)"""
    cell = row.cells[first].merge(row.cells[last])
    for extra in cell.paragraphs[1:]:              # ช่องที่รวมเข้ามาทิ้งย่อหน้าว่างไว้ ทำให้แถวสูง
        extra._element.getparent().remove(extra._element)
    _set_cell(cell, text, align="center", size=size, bold=bold)
    return cell


def _school_only(school) -> str:
    """ชื่อโรงเรียนแบบตัดคำว่า "โรงเรียน" นำหน้าออก - กันข้อความซ้ำ เช่น "โรงเรียนโรงเรียนบ้าน..." """
    name = (school.name or "").strip()
    return name[len("โรงเรียน"):].strip() if name.startswith("โรงเรียน") else name


def _right_block(doc, lines, *, align="center", gap=True):
    """บล็อกข้อความชิดครึ่งขวาของหน้า โดยทุกบรรทัด "จัดตรงกันเอง" ในบล็อก
    ใช้ตารางไร้ขอบ 2 ช่อง (ซ้ายว่าง) แทนการสั่ง align=right ทีละบรรทัด
    ซึ่งทำให้บรรทัดยาวไม่เท่ากันไปเกาะขอบขวาแล้วดูเหลื่อมกัน"""
    _sign_table(doc, [[("", "center")], [(t, align) for t in lines]], gap=gap)


def _title(doc, school, title, fiscal_year, sub=""):
    _p(doc, title, align="center", bold=True, size=18, after=0)
    if sub:
        _p(doc, sub, align="center", size=14, after=0)
    _p(doc, (school.name or "").strip(), align="center", bold=True, size=16, after=0)
    _p(doc, f"ปีงบประมาณ {fiscal_year}", align="center", size=14, after=6)


# --------------------------------------- 1) บันทึกการรับเงินเพื่อเก็บรักษา
def render_safe_custody(school, rows, total, as_of) -> str:
    """rows: list ของ (ชื่อรายการ, จำนวนเงิน) - เฉพาะ "เงินสด" ที่ต้องเก็บในตู้นิรภัย
    ตามแบบฟอร์มจริง: ผอ. รับเงินไปเก็บรักษา แล้วคืนเจ้าหน้าที่การเงินในวันทำการถัดไป"""
    doc = _new()
    _p(doc, "บันทึกการรับเงินเพื่อเก็บรักษา", align="center", bold=True, size=18, after=6)
    _right_block(doc, [f"โรงเรียน{_school_only(school) or _BLANK}",
                       f"วันที่ {thai_date(as_of) if as_of else _BLANK}"], align="left", gap=False)
    _p(doc, "", after=4)
    _p(doc, "ข้าพเจ้าได้รับเงินคงเหลือตามรายการ ดังต่อไปนี้", indent=1.27, size=16, after=6)

    headers = ["รายการ", "จำนวนเงิน", "หมายเหตุ"]
    widths = [Cm(9.0), Cm(3.5), Cm(4.0)]
    t = _grid(doc, headers, widths, size=15)
    for name, amt in rows:
        _row(t, [name, _money(amt), ""], widths, ["left", "right", "left"], size=15)
    _blank_rows(t, max(0, 5 - len(rows)), widths, 3)
    _row(t, ["รวมทั้งสิ้น", _money(total), ""], widths,
         ["right", "right", "left"], size=15, bold=True)

    _p(doc, "", after=4)
    _p(doc, f"จำนวนเงิน  ({bahttext(total)})", indent=1.27, size=16, after=10)
    _p(doc, "ข้าพเจ้า จะรับผิดชอบในการเก็บรักษาเงินดังกล่าว และจะส่งคืนให้เจ้าหน้าที่การเงิน "
            "เพื่อจ่ายในวันทำการถัดไป", indent=1.27, size=16, after=10)

    _right_block(doc, ["ลงชื่อ..............................................",
                       f"( {(school.director_name or '').strip() or _BLANK} )",
                       head_title(school)])
    _p(doc, "", after=6)

    got = _money(total) if total else "..............................."
    _p(doc, f"ข้าพเจ้าได้รับเงิน {got} บาท คืนจาก"
            f"{head_title(school) or _BLANK} ในวันที่ ........................................... "
            "เพื่อจะนำไปจ่ายตามระเบียบของทางราชการ", indent=1.27, size=16, after=12)
    _right_block(doc, ["ลงชื่อ..............................................",
                       "(..............................................)",
                       "เจ้าหน้าที่การเงิน"])
    _p(doc, "", after=6)

    _p(doc, "หมายเหตุ  บันทึกการรับเงินเพื่อเก็บรักษา ให้ถือเป็นหลักฐานแทนเงินสด "
            "ซึ่งจะต้องบันทึกไว้ทุกวันในวันที่เก็บรักษาเงินสด และจะต้องบันทึกไว้ในรายงานเงินคงเหลือประจำวัน",
       size=13, after=0)
    return _save(doc, f"บันทึกการรับเงินเพื่อเก็บรักษา_{thai_date(as_of) if as_of else ''}")


# --------------------------------------- 2) ทะเบียนคุมหลักฐานขอเบิก
def render_disburse_register(school, fiscal_year, memos) -> str:
    doc = _new(landscape=True)
    _title(doc, school, "ทะเบียนคุมหลักฐานขอเบิก", fiscal_year)
    # 9 คอลัมน์ในกระดาษ A4 แนวนอน (กว้างใช้ได้ 26.7 ซม.) วันที่แบบย่อ กันตกบรรทัดในช่องแคบ
    # ความกว้างตรวจแล้วกับเอกสารจริง (Word -> PDF) ตัวเลข/วันที่/เลขที่บันทึกไม่ตกบรรทัด
    headers = ["วัน เดือน ปี", "เลขที่บันทึก", "เจ้าหนี้หรือผู้ขอเบิก", "ประเภทรายจ่าย", "จำนวนเงิน",
               "ชื่อผู้รับเงิน", "ลายมือชื่อ\nผู้รับหลักฐาน", "วัน เดือน ปี ที่ส่ง\nส่วนราชการผู้เบิก", "หมายเหตุ"]
    widths = [Cm(2.3), Cm(2.1), Cm(4.2), Cm(4.2), Cm(2.6), Cm(3.4), Cm(2.8), Cm(2.6), Cm(2.5)]
    aligns = ["center", "center", "left", "left", "right", "left", "center", "center", "left"]
    t = _grid(doc, headers, widths)
    total = 0.0
    for m in memos:
        total += float(m.amount or 0)
        _row(t, [thai_date_short(m.date) if m.date else "", m.memo_no or "", m.payee or "",
                 (m.budget_source or m.subject or "")[:100], _money(m.amount),
                 (getattr(m, "receiver_name", "") or "").strip(), "", "",
                 (getattr(m, "register_note", "") or "").strip()],
             widths, aligns)
    _blank_rows(t, max(0, 6 - len(memos)), widths, len(headers))
    _row(t, ["", "", "", "รวม", _money(total), "", "", "", ""], widths,
         ["center", "center", "left", "right", "right", "left", "center", "center", "left"], bold=True)
    _sign_finance(doc, school)
    return _save(doc, f"ทะเบียนคุมหลักฐานขอเบิก_ปีงบ{fiscal_year}")


# --------------------------------------- 3) ทะเบียนคุมเงินฝาก (เงินประกันสัญญา)
def render_guarantee_register(school, account, txns, fiscal_year) -> str:
    """แถวละ 1 รายการเงินประกันที่รับเข้ามา พร้อมวันครบกำหนดและวันที่คืนเงินให้ผู้มีสิทธิ์
    (ช่อง "การฝาก" เว้นไว้ให้กรอกวันที่นำเงินเข้าบัญชีธนาคารตามใบนำฝากจริง)"""
    doc = _new(landscape=True)
    _title(doc, school, "ทะเบียนคุมเงินนอกงบประมาณ",
           fiscal_year, sub=f"ประเภท เงินฝาก (เงินประกันสัญญา) · {account.name}")
    groups = [("ที่", []), ("รายการ", []), ("ประเภท", []),
              ("การรับ", ["วัน เดือน ปี", "ที่เอกสาร", "จำนวนเงิน"]),
              ("การฝาก", ["วัน เดือน ปี", "ที่เอกสาร", "จำนวนเงิน"]),
              ("วันครบกำหนด", []), ("วันที่เบิกจ่ายเงินคืน\nผู้มีสิทธิ์", []), ("หมายเหตุ", [])]
    widths = [Cm(1.2), Cm(3.8), Cm(2.2), Cm(2.2), Cm(2.0), Cm(2.2),
              Cm(2.2), Cm(2.0), Cm(2.2), Cm(2.2), Cm(2.4), Cm(2.0)]
    t = _head2(doc, groups, widths)
    total = 0.0
    ins = [x for x in txns if x.kind == "in"]
    for i, x in enumerate(ins, 1):
        total += float(x.amount or 0)
        _row(t, [str(i), x.note or x.category or "", "เงินประกันสัญญา",
                 thai_date(x.date) if x.date else "", x.ref or "", _money(x.amount),
                 "", "", "",
                 thai_date(x.due_date) if x.due_date else "",
                 thai_date(x.refund_date) if x.refund_date else "", ""],
             widths, ["center", "left", "center", "center", "center", "right",
                      "center", "center", "right", "center", "center", "left"])
    _blank_rows(t, max(0, 6 - len(ins)), widths, 12)
    _row(t, ["", "รวม", "", "", "", _money(total), "", "", "", "", "", ""], widths,
         ["center", "right", "center", "center", "center", "right",
          "center", "center", "right", "center", "center", "left"], bold=True)
    _sign_finance(doc, school)
    return _save(doc, f"ทะเบียนคุมเงินประกันสัญญา_ปีงบ{fiscal_year}")


# --------------------------------------- 4) ทะเบียนคุมการรับและนำส่งเงินรายได้แผ่นดิน
# 4 ช่องประเภทตามแบบทะเบียนคุมของกลุ่มตรวจสอบภายใน สพฐ. (จัดประเภทเดียวกับหน้านำส่งรายได้แผ่นดิน)
_REV_COLS = ["ดอกเบี้ยเงินฝาก", "เงินเหลือจ่ายปีเก่าส่งคืน", "ค่าขายของเบ็ดเตล็ด", "ค่าปรับและอื่น ๆ"]


def _rev_bucket(txn) -> int:
    from app.services.revenue_remit import classify
    return {"interest": 0, "old_return": 1, "sale": 2}.get(classify(txn), 3)


def render_revenue_register(school, account, txns, fiscal_year) -> str:
    doc = _new(landscape=True)
    _title(doc, school, "ทะเบียนคุมการรับและนำส่งเงินรายได้แผ่นดิน",
           fiscal_year, sub=account.name)
    groups = [("วัน เดือน ปี", []), ("ที่เอกสาร", []), ("รายการ", [])]
    for label in _REV_COLS:
        groups.append((label, ["รับ", "นำส่ง", "คงเหลือ"]))
    groups += [("รวมเงินคงเหลือ\nทุกประเภท", []), ("หมายเหตุ", [])]
    widths = [Cm(2.0), Cm(1.8), Cm(2.6)] + [Cm(1.45)] * 12 + [Cm(2.0), Cm(1.4)]
    t = _head2(doc, groups, widths)

    bal = [0.0] * len(_REV_COLS)
    ordered = sorted(txns, key=lambda x: (x.date or datetime.min, x.id))
    for x in ordered:
        b = _rev_bucket(x)
        amt = float(x.amount or 0)
        got = amt if x.kind == "in" else 0.0
        sent = amt if x.kind == "out" else 0.0     # "นำส่ง" คลัง = จ่ายออก
        bal[b] += got - sent
        cells = [""] * (3 * len(_REV_COLS))
        cells[b * 3] = _money(got) if got else ""
        cells[b * 3 + 1] = _money(sent) if sent else ""
        cells[b * 3 + 2] = _money(bal[b])
        _row(t, [thai_date(x.date) if x.date else "", x.ref or "",
                 (x.note or x.category or "")[:30]] + cells + [_money(sum(bal)), ""],
             widths, ["center", "center", "left"] + ["right"] * 12 + ["right", "left"], size=11)
    _blank_rows(t, max(0, 6 - len(ordered)), widths, 17)
    _sign_finance(doc, school)
    return _save(doc, f"ทะเบียนคุมเงินรายได้แผ่นดิน_ปีงบ{fiscal_year}")


# --------------------------------------- 5) สมุดคู่ฝาก (ส่วนราชการผู้เบิก)
def render_agency_passbook(school, account, txns, opening, fiscal_year) -> str:
    doc = _new(landscape=True)
    _title(doc, school, "สมุดคู่ฝาก", fiscal_year, sub=f"(ส่วนราชการผู้เบิก) · {account.name}")
    groups = [("วัน เดือน ปี", []), ("ที่เอกสาร\nใบนำฝาก/ใบเบิกถอน", []),
              ("จำนวนเงิน", ["ฝาก", "ถอน", "คงเหลือ"]),
              ("ลายมือชื่อ\nผู้รับฝาก", []),
              ("ลายมือชื่อผู้นำฝาก\nหรือผู้เบิกถอน", []), ("หมายเหตุ", [])]
    widths = [Cm(2.6), Cm(4.0), Cm(2.8), Cm(2.8), Cm(3.0), Cm(4.0), Cm(4.5), Cm(3.0)]
    t = _head2(doc, groups, widths)
    bal = float(opening or 0)
    _row(t, ["", "ยอดยกมา", "", "", _money(bal), "", "", ""], widths,
         ["center", "left", "right", "right", "right", "center", "center", "left"], bold=True)
    ordered = sorted(txns, key=lambda x: (x.date or datetime.min, x.id))
    for x in ordered:
        amt = float(x.amount or 0)
        bal += amt if x.kind == "in" else -amt
        _row(t, [thai_date(x.date) if x.date else "", x.ref or (x.note or "")[:28],
                 _money(amt) if x.kind == "in" else "",
                 _money(amt) if x.kind == "out" else "", _money(bal), "", "", ""],
             widths, ["center", "left", "right", "right", "right", "center", "center", "left"])
    _blank_rows(t, max(0, 6 - len(ordered)), widths, 8)
    _sign_finance(doc, school)
    return _save(doc, f"สมุดคู่ฝาก_ปีงบ{fiscal_year}")


def render_account_register(school, account, txns, opening, fiscal_year, key) -> str:
    if key == "guarantee":
        return render_guarantee_register(school, account, txns, fiscal_year)
    if key == "revenue":
        return render_revenue_register(school, account, txns, fiscal_year)
    return render_agency_passbook(school, account, txns, opening, fiscal_year)


# ============================== ทะเบียนคุมเงิน (ฟอร์มกลาง ใช้ได้ทุกบัญชี) ==============================
# คอลัมน์ตามแบบฟอร์มทะเบียนคุมเงินนอกงบประมาณที่โรงเรียนใช้จริง
#   วัน เดือน ปี | ที่เอกสาร | รายการ | รับ
#   จ่าย  [ ลูกหนี้ | ค่าตอบแทน ค่าใช้สอย ค่าวัสดุ ค่าครุภัณฑ์ อื่น ๆ | รวมจ่าย ]
#   คงเหลือ [ เงินสด | เงินฝากธนาคาร | เงินฝากส่วนราชการผู้เบิก ] | หมายเหตุ
#
# แถวแรกเป็น "ยอดยกมาจากปีงบประมาณ ก่อนหน้า" เสมอ ตามแบบฟอร์ม
_MR_GROUPS = [
    ("วัน เดือน ปี", []), ("ที่เอกสาร", []), ("รายการ", []), ("รับ", []),
    ("จ่าย", ["ลูกหนี้", "ค่าตอบแทน ค่าใช้สอย\nค่าวัสดุ ค่าครุภัณฑ์ อื่น ๆ", "รวมจ่าย"]),
    ("คงเหลือ", ["เงินสด", "เงินฝากธนาคาร", "เงินฝากส่วนราชการ\nผู้เบิก"]),
    ("หมายเหตุ", []),
]
# ช่อง "ค่าตอบแทน ค่าใช้สอย..." กับ "เงินฝากส่วนราชการ" กว้างพอให้หัวตารางเหลือ 2 บรรทัด
# รวม 27.2 ซม. พอดีหน้า A4 แนวนอนขอบซ้าย-ขวา 1.2 ซม. (_MR_MARGIN)
_MR_W = [Cm(2.3), Cm(1.8), Cm(4.8), Cm(2.1),
         Cm(1.8), Cm(3.9), Cm(2.0),
         Cm(1.8), Cm(2.3), Cm(2.8), Cm(1.6)]
_MR_MARGIN = Cm(1.2)
_MR_ALIGN = (["center", "center", "left", "right"]
             + ["right", "right", "right"]
             + ["right", "right", "right"] + ["left"])
# ช่องคงเหลือที่ต้องลงยอด ขึ้นกับว่าบัญชีนี้เก็บเงินไว้ที่ไหน
_BAL_COL = {"cash": 7, "bank": 8, "agency": 9}
# คำที่บ่งว่ารายการจ่ายนั้นเป็น "ลูกหนี้" (เงินยืม) ไม่ใช่ค่าใช้จ่ายปกติ
_DEBTOR_WORDS = ("ลูกหนี้", "ยืม", "เงินยืม")


def _is_debtor(txn) -> bool:
    text = " ".join(str(getattr(txn, f, "") or "") for f in ("category", "note", "ref"))
    return any(w in text for w in _DEBTOR_WORDS)


def render_money_register(school, account, txns, opening, fiscal_year, *,
                          item=None, doc=None) -> str:
    """ทะเบียนคุมเงินของบัญชีหนึ่ง (หรือรายการย่อยหนึ่ง) ตามแบบฟอร์มกลาง

    item = AccountItem -> ออกทะเบียนคุมเฉพาะรายการย่อยนั้น (เช่น ค่าเครื่องแบบนักเรียน
    ซึ่งเป็นรายการย่อยของเงินอุดหนุน) ตามที่แบบฟอร์มระบุ "ประเภทเงิน ..." ไว้ด้านบน
    """
    own = doc is None
    if own:
        doc = _new(landscape=True)
    else:
        _break_land(doc)
    sec = doc.sections[-1]
    sec.left_margin = sec.right_margin = _MR_MARGIN
    kind = (account.fund_type or "").strip() or "เงินนอกงบประมาณ"
    title = ("ทะเบียนคุมเงินนอกงบประมาณ" if kind == "เงินนอกงบประมาณ"
             else f"ทะเบียนคุม{kind}")
    money_name = (item.name if item is not None else account.name) or ""
    _title(doc, school, title, fiscal_year,
           sub=f"ประเภทเงิน {money_name.strip()}"
               + (f"  (บัญชี{account.name})" if item is not None else ""))

    t = _head2(doc, _MR_GROUPS, _MR_W)
    # คงเหลือแยก 3 ช่องตามที่เก็บเงินจริงของแต่ละรายการ (เงินสด / ธนาคาร / ส่วนราชการผู้เบิก)
    from app.services.money_place import account_place, item_place, opening_by_place, txn_place
    home = item_place(item) if item is not None else account_place(account)
    if item is not None:
        place_bal = {home: float(opening or 0)}
    else:
        place_bal = opening_by_place(account, fiscal_year)
        place_bal[home] += float(opening or 0) - sum(place_bal.values())
    used = {k for k, v in place_bal.items() if abs(v) > 0.005} | {home}
    used |= {txn_place(x) for x in txns}

    def put_bal(vals):
        for k in used:
            vals[_BAL_COL[k]] = _money(place_bal.get(k, 0)) or "0.00"

    def blank_row():
        return [""] * len(_MR_W)

    # ---- แถวยอดยกมา ----
    running = float(opening or 0)
    first = blank_row()
    first[0] = thai_date_short(datetime(fiscal_year - 544, 10, 1))
    first[2] = f"ยอดยกมาจากปีงบประมาณ {fiscal_year - 1}"
    first[3] = _money(running)
    put_bal(first)
    _row(t, first, _MR_W, _MR_ALIGN)

    total_in = total_out = total_debtor = total_exp = 0.0
    rows = sorted(txns, key=lambda x: (x.date or datetime.min, x.id or 0))
    for x in rows:
        amount = float(x.amount or 0)
        vals = blank_row()
        vals[0] = thai_date_short(x.date) if x.date else ""
        vals[1] = (x.ref or "").strip()
        vals[2] = (x.note or x.category or "").strip()
        vals[10] = (getattr(x, "remark", "") or "").strip()
        p = txn_place(x)
        place_bal[p] = place_bal.get(p, 0) + (amount if (x.kind or "in") == "in" else -amount)
        if (x.kind or "in") == "in":
            running += amount
            total_in += amount
            vals[3] = _money(amount)
        else:
            running -= amount
            total_out += amount
            if _is_debtor(x):
                vals[4] = _money(amount)
                total_debtor += amount
            else:
                vals[5] = _money(amount)
                total_exp += amount
            vals[6] = _money(amount)
        put_bal(vals)
        _row(t, vals, _MR_W, _MR_ALIGN)

    _blank_rows(t, max(0, 8 - len(rows)), _MR_W, len(_MR_W))
    sums = blank_row()
    sums[3] = _money(float(opening or 0) + total_in)
    sums[4] = _money(total_debtor)
    sums[5] = _money(total_exp)
    sums[6] = _money(total_out)
    put_bal(sums)
    _merge_label(_row(t, sums, _MR_W, _MR_ALIGN, bold=True), 0, 2, "รวม")   # วันที่+ที่เอกสาร+รายการ = ช่องเดียว
    _sign_finance(doc, school)
    if not own:
        return doc
    tag = f"{money_name.strip()}" if money_name.strip() else account.name
    return _save(doc, f"ทะเบียนคุมเงิน_{tag}_ปีงบ{fiscal_year}")


def _break_land(doc):
    """ขึ้นหน้าใหม่แนวนอนสำหรับทะเบียนฉบับถัดไปในชุดรวม"""
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from app.services.doc_page import A4_H, A4_W
    if not (doc.paragraphs or doc.tables):
        return
    from app.services.doc_page import strip_tail
    strip_tail(doc)
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = A4_H, A4_W
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin = Cm(1.5); sec.bottom_margin = Cm(1.2)
    return sec


def render_all_registers(school, rows, fiscal_year) -> str:
    """ทะเบียนคุมทุกบัญชี (และทุกรายการย่อย) รวมเป็นไฟล์เดียว

    rows = [(account, item_or_None, txns, opening), ...] เรียงมาแล้วจากผู้เรียก
    """
    if not rows:
        raise ValueError("ยังไม่มีบัญชีให้ออกทะเบียนคุม")
    doc = _new(landscape=True)
    for account, item, txns, opening in rows:
        render_money_register(school, account, txns, opening, fiscal_year,
                              item=item, doc=doc)
    return _save(doc, f"ทะเบียนคุมเงินทุกบัญชี_ปีงบ{fiscal_year}")


# --------------------------------------- ทะเบียนคุมใบเสร็จรับเงิน/ใบสำคัญ
def render_receipt_register(school, fiscal_year, rows) -> str:
    """ทะเบียนคุมใบเสร็จรับเงินและใบสำคัญรับเงิน (แยกยอดด้านรับ-ด้านจ่าย)"""
    doc = _new(landscape=True)
    _title(doc, school, "ทะเบียนคุมใบเสร็จรับเงินและใบสำคัญรับเงิน", fiscal_year)
    headers = ["วัน เดือน ปี", "เลขที่", "ประเภท", "ผู้รับ/ผู้จ่ายเงิน", "บัญชี/ประเภทเงิน",
               "จำนวนเงิน", "หมายเหตุ"]
    widths = [Cm(2.6), Cm(3.0), Cm(2.2), Cm(6.4), Cm(5.0), Cm(3.0), Cm(4.5)]
    t = _grid(doc, headers, widths)
    tot_in = tot_out = 0.0
    for r, acct_name in rows:
        amount = float(r.amount or 0)
        if (r.kind or "รับ") == "รับ":
            tot_in += amount
        else:
            tot_out += amount
        _row(t, [thai_date(r.date) if r.date else "", r.receipt_no or "", r.kind or "รับ",
                 r.party or "", acct_name or "", _money(amount), r.note or ""],
             widths, ["center", "center", "center", "left", "left", "right", "left"])
    _blank_rows(t, max(0, 6 - len(rows)), widths, 7)
    _row(t, ["", "", "", "", "รวมด้านรับ", _money(tot_in), ""], widths,
         ["center", "center", "center", "left", "right", "right", "left"], bold=True)
    _row(t, ["", "", "", "", "รวมด้านจ่าย", _money(tot_out), ""], widths,
         ["center", "center", "center", "left", "right", "right", "left"], bold=True)
    _sign_finance(doc, school)
    return _save(doc, f"ทะเบียนคุมใบเสร็จ_ปีงบ{fiscal_year}")
