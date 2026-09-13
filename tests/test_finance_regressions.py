import ast
import asyncio
import io
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace as NS
from unittest.mock import patch

from docx import Document
from fastapi import Depends, Form, Request
from fastapi.responses import RedirectResponse
from openpyxl import Workbook
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session

from app.database import Base, _repair_finance_fund_types
from app.models import FinanceAccount, FinanceTxn, MoneyLoan, LoanReturn
from app.services.asset_utils import opening_for
from app.services.finance_types import FUND_TYPES, FUND_DEFAULT, resolve_fund_type
from app.services.finance_io import import_finance_workbook
from app.services.finance_forms_doc import render_loan_contract, render_loan_returns
from app.services import lunch_ingredient_doc as lunch
from app.thai_utils import parse_be_date

ROOT = Path(__file__).resolve().parents[1]

def route_functions():
    tree = ast.parse((ROOT/'app/routers/finance.py').read_text(encoding='utf-8'))
    names = {'account_add', 'loan_add', 'account_set_fund_type', '_cashbook_fund_data'}
    funcs = [n for n in tree.body if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)) and n.name in names]
    for n in funcs: n.decorator_list = []
    ns = dict(Session=Session, Depends=Depends, Form=Form, Request=Request, get_db=lambda: None,
              FinanceAccount=FinanceAccount, FinanceTxn=FinanceTxn, MoneyLoan=MoneyLoan, RedirectResponse=RedirectResponse,
              datetime=datetime, opening_for=opening_for, thai_date=lambda d:str(d),
              FUND_TYPES=FUND_TYPES, _FUND_DEFAULT=FUND_DEFAULT, DEPOSIT_TYPES={'cash': '', 'bank': '', 'agency': ''},
              resolve_fund_type=resolve_fund_type, parse_be_date=parse_be_date, timedelta=timedelta,
              current_fiscal_year=lambda:2569, _to_int=lambda x,d:int(x) if x else d,
              _to_float=lambda x,d:float(x) if x else d)
    exec(compile(ast.Module(body=funcs,type_ignores=[]),'<finance-routes>','exec'),ns)
    return ns

class FinanceTests(unittest.TestCase):
    def setUp(self):
        self.engine=create_engine('sqlite://')
        Base.metadata.create_all(self.engine)
        self.db=Session(self.engine)
        self.routes=route_functions()
    def tearDown(self):
        self.db.close(); self.engine.dispose()
    def test_create_all_three_funds(self):
        for name in FUND_TYPES:
            self.routes['account_add'](db=self.db,name=name,opening_balance='100',note='',deposit_type='bank',fund_type='')
        self.assertEqual([a.fund_type for a in self.db.query(FinanceAccount).order_by(FinanceAccount.id)],FUND_TYPES)
        self.assertEqual(resolve_fund_type('บัญชีโครงการ','เงินงบประมาณ'),'เงินงบประมาณ')
    def test_import_classifies_accounts_and_preserves_existing(self):
        wb=Workbook(); ws=wb.active; ws.title='บัญชีเงิน'
        ws.append(['คำแนะนำ']); ws.append(['ชื่อ','ยอด','หมายเหตุ','หมวด'])
        for name in FUND_TYPES: ws.append([name,100,''])
        ws.append(['บัญชีทดสอบ',200,'','เงินรายได้แผ่นดิน'])
        buf=io.BytesIO(); wb.save(buf)
        import_finance_workbook(buf.getvalue(),self.db)
        self.assertEqual([a.fund_type for a in self.db.query(FinanceAccount).order_by(FinanceAccount.id)],FUND_TYPES+['เงินรายได้แผ่นดิน'])
    def test_repair_is_one_time_and_leaves_custom_accounts(self):
        self.db.add_all([FinanceAccount(name='เงินงบประมาณ',fund_type=FUND_DEFAULT),FinanceAccount(name='เงินรายได้แผ่นดิน',fund_type=FUND_DEFAULT),FinanceAccount(name='ทุนกิจกรรม',fund_type=FUND_DEFAULT)])
        self.db.commit(); _repair_finance_fund_types(self.engine); self.db.expire_all()
        rows=self.db.query(FinanceAccount).order_by(FinanceAccount.id).all()
        self.assertEqual([a.fund_type for a in rows],FUND_TYPES)
        rows[0].fund_type=FUND_DEFAULT; self.db.commit()
        _repair_finance_fund_types(self.engine); self.db.expire_all()
        self.assertEqual(rows[0].fund_type,FUND_DEFAULT)
    def test_cashbook_separates_receipts_and_payments(self):
        for index, fund in enumerate(FUND_TYPES, 1):
            account=FinanceAccount(name=fund,fund_type=fund,opening_balance=100*index)
            self.db.add(account); self.db.flush()
            self.db.add_all([FinanceTxn(account_id=account.id,fiscal_year=2569,kind='in',amount=10*index),
                             FinanceTxn(account_id=account.id,fiscal_year=2569,kind='out',amount=index)])
        self.db.commit()
        scope, opening, receipts, payments=self.routes['_cashbook_fund_data'](self.db,2569,None)
        self.assertEqual([row['fund'] for row in receipts],FUND_TYPES)
        self.assertEqual([row['fund'] for row in payments],FUND_TYPES)
        self.assertEqual(list(opening.values()),[100,200,300])

    def test_due_date_crosses_year_and_uses_receipt_date(self):
        class Req:
            async def form(self):
                return {'fiscal_year':'2569','borrower':'ครูทดสอบ','date':'30/12/2569','receive_date':'31/12/2569','within_days':'15','amount':'1000','due_date':'01/01/2500'}
        asyncio.run(self.routes['loan_add'](Req(),self.db))
        self.assertEqual(self.db.query(MoneyLoan).one().due_date,datetime(2027,1,15))
    def test_documents_and_lunch_regression(self):
        school=NS(name='โรงเรียนทดสอบ',address='อำเภอเมือง จังหวัดทดสอบ',director_name='นายทดสอบ ใจดี',finance_officer_name='นางการเงิน ทดสอบ',officer_name='ครูพัสดุ')
        loan=MoneyLoan(id=1,fiscal_year=2569,contract_no='1/2569',borrower='นายผู้ยืม ทดสอบ',position='ครู',amount=1000,date=datetime(2026,9,13),receive_date=datetime(2026,9,13),due_date=datetime(2026,9,28),within_days=15,purpose='เดินทางไปราชการ',fund_from='เงินงบประมาณ')
        loan.returns=[LoanReturn(date=datetime(2026,9,15),kind='ใบสำคัญ',amount=600,receipt_no='V1'),LoanReturn(date=datetime(2026,9,16),kind='เงินสด',amount=100,receipt_no='C1')]
        with tempfile.TemporaryDirectory() as tmp, patch('app.services.lunch_doc.get_data_dir',return_value=Path(tmp)), patch('app.services.finance_forms_doc.get_data_dir',return_value=Path(tmp)):
            contract=Document(render_loan_contract(school,loan))
            self.assertEqual(len(contract.tables[0].rows),6)
            xml=contract.element.xml
            self.assertIn('เดินทางไปราชการ',xml); self.assertNotIn('อาหารกลางวัน',xml)
            self.assertIn('V1',xml); self.assertIn('15 วัน',xml)
            repayment=Document(render_loan_returns(school,loan))
            xml=repayment.element.xml
            self.assertIn('คงค้าง 300',xml); self.assertIn('600',xml); self.assertIn('100',xml)
            self.assertNotIn('อาหารกลางวัน',xml)
            loan.returns.append(LoanReturn(date=datetime(2026,9,17),kind='เงินสด',amount=300))
            self.assertIn('คงค้าง 0',Document(render_loan_returns(school,loan)).element.xml)
            # Lunch and finance use the same renderer but retain their own business wording.
            rnd=NS(program=NS(funding_org='เทศบาล',total_students=50,rate_per_head=22,year=2569,lunch_officer='ครูอาหาร'),days=5,amount=5500,start_date=datetime(2026,9,1),end_date=datetime(2026,9,5),order_no='1/2569',order_date=datetime(2026,9,1),seq=1,doc_nos='{}')
            for name in ('render_loan_contract','render_repay_memo'):
                actual=Document(getattr(lunch,name)(rnd,school)).element.xml
                self.assertIn("อาหารกลางวัน",actual)
                self.assertIn("ครูอาหาร",actual)
                if name == "render_loan_contract": self.assertIn("30 วัน",actual)

if __name__=='__main__': unittest.main()
