"""Fiscal-year/term subsidy regressions, isolated from real school data."""
import importlib.util
import pathlib
import re
import sys
import tempfile

import pytest

ROOT = pathlib.Path(__file__).resolve().parents[1]
AY = 2569


@pytest.fixture()
def env(monkeypatch):
    tmp = pathlib.Path(tempfile.mkdtemp())
    import app.accounts as ac
    import app.database as dbm
    import app.tenancy as tn
    for m in (dbm, tn, ac):
        monkeypatch.setattr(m, "get_data_dir", lambda t=tmp: t)
    monkeypatch.setattr(tn, "_engines", type(tn._engines)())
    monkeypatch.setattr(ac, "_engine", None)
    monkeypatch.setattr(ac, "_Session", None)
    import app.main as main_mod
    import app.routers.auth as auth_mod
    monkeypatch.setattr(auth_mod, "authenticate", ac.authenticate)
    for key in ("tenant_state", "can_use_module", "get_account_access"):
        monkeypatch.setattr(main_mod, key, getattr(ac, key))

    spec = importlib.util.spec_from_file_location("seed_demo", ROOT / "tools" / "seed_demo.py")
    seed = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(seed)
    sys.argv = ["seed_demo.py", "--password", "Demo!2569"]
    seed.main()

    from fastapi.testclient import TestClient

    from app.main import app
    c = TestClient(app, raise_server_exceptions=False)
    c.post("/login", data={"username": "demo", "password": "Demo!2569"}, follow_redirects=False)
    return c


from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from starlette.datastructures import FormData
from app.database import Base
from app.models import (School, FinanceAccount, FinanceTxn, AccountItem, SubsidySnapshot,
                        SubsidyRate, SubsidyCensus, SubsidyReceiptLink, SubsidyCensusRevision)
from app.services import subsidy as sub
import json


@pytest.fixture
def db():
    engine = create_engine('sqlite://')
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        session.add(School(name='โรงเรียนทดสอบ'))
        session.commit()
        yield session
    engine.dispose()


def form_for(db, ay=2568, term=1, **overrides):
    s = sub.state(db, ay, term)
    d = {'academic_year': str(ay), 'term': str(term), 'token': s['token'], 'levels': 'ป.1'}
    for k in sub.keys_for(term):
        d['r_ป.1_'+k] = '1000'
    for scan in s['census']:
        d['n_'+scan['key']+'_ป.1'] = '100' if scan == s['census'][0] else '120'
        d['source_'+scan['key']] = 'รายงาน DMC ทดสอบ'
        d['confirm_'+scan['key']] = '1'
    d.update(overrides)
    return FormData(d)


def prepared(db, ay=2568, term=1, **overrides):
    state = sub.save(db, ay, term, form_for(db, ay, term, **overrides))
    db.commit()
    return state


def confirmed(db, **overrides):
    s = prepared(db, **overrides)
    snap = sub.confirm(db, s['academic_year'], s['term'], s['token'])
    db.commit()
    return snap


def test_topup_is_full_latest_minus_first_not_thirty_percent(db):
    s = prepared(db)
    row = s['result']['rows'][0]
    assert (row['first_estimate'],row['full'],row['remaining']) == (70000,120000,50000)
    assert s['result']['ready']
    s = prepared(db, first_teach='65000', first_ref='หนังสือ 1')
    assert s['result']['rows'][0]['remaining'] == 55000


def test_declining_roll_retains_negative_adjustment(db):
    s = prepared(db, **{'n_jun2568_ป.1':'50'})
    assert s['result']['rows'][0]['remaining'] == -20000


def test_term_two_uses_new_fiscal_year_independent_rates(db):
    old = confirmed(db)
    frozen = old.payload
    s = prepared(db, term=2, **{'r_ป.1_teach':'1500'})
    assert s['fiscal_year'] == 2569
    assert len(s['result']['rows']) == 3
    assert s['result']['rows'][0]['full'] == 180000
    assert sub.state(db,2568,1)['config']['rates']['ป.1']['teach'] == 1000
    assert db.get(SubsidySnapshot,old.id).payload == frozen


def test_promotion_uses_destination_grade_rate_and_entry_grade_stays():
    cfg={'levels':['ป.1','ป.2','ม.4','ม.5','อ.2','อ.3'], 'rates':{}, 'rates_confirmed':True,'rate_source':'x'}
    for lv in cfg['levels']:
        cfg['rates'][lv]={k:10 for k in sub.keys_for(1)}
    scans=[{'counts':{'ป.1':10,'ป.2':99,'ม.4':20,'อ.2':30},'confirmed':True,'source':'x','label':'old'},
           {'counts':{lv:1 for lv in cfg['levels']},'confirmed':True,'source':'x','label':'new'}]
    r=sub.calculate(cfg,scans,1)
    assert [b['advance'] for b in r['basis']] == [10,10,20,20,30,30]
    assert r['rows'][0]['first_estimate'] == 840


def test_missing_is_not_zero_and_future_census_cannot_confirm(db):
    s=prepared(db, **{'n_jun2568_ป.1':''})
    assert s['result']['total'] is None and not s['result']['ready']
    with pytest.raises(ValueError,match='ยังยืนยันไม่ได้'): sub.confirm(db,2568,1,s['token'])
    db.rollback()
    s=prepared(db, **{'n_jun2568_ป.1':'0'})
    assert s['result']['total'] == 0 and s['result']['ready']
    future = sub.save(db,2699,2,form_for(db,2699,2))
    assert not future['result']['ready']
    with pytest.raises(ValueError,match='ยังไม่ถึงวันสำรวจ'):
        sub.confirm(db,2699,2,future['token'])
    db.rollback()


@pytest.mark.parametrize('raw',['-1','NaN','Infinity','1000000001','abc'])
def test_invalid_money_rejected(raw):
    with pytest.raises(ValueError): sub.number(raw,'amount')


def test_dmc_fraction_rejected():
    with pytest.raises(ValueError): sub.number('1.5','heads',True)


def test_snapshot_immutable_and_stale_form_rejected(db):
    snap=confirmed(db)
    frozen=snap.payload
    stale=form_for(db)
    prepared(db, **{'r_ป.1_teach':'1100'})
    assert sub.state(db,2568,1)['changed']
    assert db.get(SubsidySnapshot,snap.id).payload==frozen
    with pytest.raises(ValueError,match='หน้าอื่น'): sub.save(db,2568,1,stale)
    db.rollback()
    assert sub.state(db,2568,1)['config']['rates']['ป.1']['teach']==1100


def test_reconfirm_idempotent_and_census_revisions_not_duplicated(db):
    snap=confirmed(db)
    s=sub.state(db,2568,1)
    assert sub.confirm(db,2568,1,s['token']).id==snap.id
    prepared(db)
    assert db.query(SubsidyCensusRevision).count()==2


def test_legacy_data_preserved_unconfirmed(db):
    db.add(SubsidyRate(academic_year=2568,level='ป.1',item_key='teach',amount=2280))
    db.add(SubsidyCensus(academic_year=2567,round='nov',level='ป.1',count=9))
    db.commit()
    s=sub.state(db,2568,1)
    assert s['config']['rates']['ป.1']['teach']==1140
    assert not s['result']['ready']
    prepared(db)
    assert db.query(SubsidyRate).one().amount==2280
    assert db.query(SubsidyCensus).one().count==9


def account(db):
    a=FinanceAccount(name='เงินอุดหนุน')
    db.add(a);db.flush()
    return a


def ensure_mapping(db, fy, aid, extra_keys=()):
    current = sub.mappings(db,fy)
    choices = {}
    for key in sub.keys_for(1)+list(extra_keys):
        matches = db.query(AccountItem).filter_by(account_id=aid,fiscal_year=fy,name=sub.NAMES[key]).all()
        choices[key] = str(current[key].id) if current.get(key) else (str(matches[0].id) if len(matches)==1 else 'new')
    for key,item in current.items():
        if item:choices.setdefault(key,str(item.id))
    sub.save_mapping(db,fy,aid,choices,sub.mapping_token(db,fy))
    db.commit()


def mapped_preview(db,snap,aid):
    if not sub.mappings(db,snap.fiscal_year):
        ensure_mapping(db,snap.fiscal_year,aid,[r['key'] for r in json.loads(snap.payload)['result']['rows'] if r.get('extra')])
    return sub.budget_preview(db,snap,aid)


def mapped_receipt(db,ay,term,tid,key,rnd,amount):
    txn=db.get(FinanceTxn,tid)
    fy=sub.fiscal_year(ay,term)
    if not sub.mappings(db,fy):ensure_mapping(db,fy,txn.account_id)
    if txn.item_id is None:
        txn.item_id=sub.mappings(db,fy)[key].id;db.commit()
    return sub.link_receipt(db,ay,term,tid,key,rnd,amount)


def test_receipt_split_cannot_double_count_between_terms_or_repeat_click(db):
    a=account(db)
    t=FinanceTxn(account_id=a.id,fiscal_year=2569,kind='in',amount=1000)
    db.add(t);db.commit()
    mapped_receipt(db,2568,2,t.id,'teach','first','600');db.commit()
    with pytest.raises(ValueError,match='เชื่อมรายการ'): mapped_receipt(db,2568,2,t.id,'teach','first','100')
    db.rollback()
    with pytest.raises(ValueError,match='เกินเงินรับจริง'): mapped_receipt(db,2569,1,t.id,'teach','first','500')
    db.rollback()
    mapped_receipt(db,2569,1,t.id,'teach','first','400');db.commit()
    assert sub.receipt_summary(db,2568,2)[0]['teach']==600
    assert sub.receipt_summary(db,2569,1)[0]['teach']==400
    assert sub.receipt_summary(db,2570,1)[0]=={}
    t.amount=500;db.commit()
    assert sub.receipt_summary(db,2568,2)[0]=={}
    assert not sub.receipt_summary(db,2569,1)[1][0]['valid']


def test_deleted_receipt_cannot_attach_to_reused_id(db):
    a=account(db)
    t=FinanceTxn(account_id=a.id,fiscal_year=2568,kind='in',amount=1000)
    db.add(t);db.commit();tid=t.id
    mapped_receipt(db,2568,1,tid,'teach','first','500');db.commit()
    db.delete(t);db.commit()
    db.add(FinanceTxn(id=tid,account_id=a.id,fiscal_year=2568,kind='in',amount=1000));db.commit()
    assert sub.receipt_summary(db,2568,1)[0]=={}
    assert db.query(SubsidyReceiptLink).count()==0


def test_deleted_budget_item_does_not_reuse_old_contribution(db):
    from app.models import SubsidyBudgetContribution
    snap=confirmed(db);a=account(db);db.commit()
    p=mapped_preview(db,snap,a.id);sub.apply_budget(db,snap,a.id,sub.fingerprint(p));db.commit()
    item=db.query(AccountItem).filter_by(name=sub.NAMES['teach']).one();iid=item.id
    db.delete(item);db.commit()
    db.add(AccountItem(id=iid,account_id=a.id,fiscal_year=2568,name='งบอื่น',budget=200));db.commit()
    assert db.query(SubsidyBudgetContribution).filter_by(item_key='teach').count()==0
    with pytest.raises(ValueError,match='จับคู่'):
        sub.budget_preview(db,snap,a.id)


def test_budget_preserves_other_money_and_uses_correct_fy(db):
    snap=confirmed(db)
    a=account(db)
    item=AccountItem(account_id=a.id,fiscal_year=2568,name=sub.NAMES['teach'],budget=123)
    db.add(item);db.commit()
    p=mapped_preview(db,snap,a.id)
    token=sub.fingerprint(p)
    sub.apply_budget(db,snap,a.id,token);db.commit()
    assert item.budget==120123
    assert db.query(AccountItem).filter_by(fiscal_year=2569).count()==0
    with pytest.raises(ValueError,match='ยอดงบปลายทางเปลี่ยน'): sub.apply_budget(db,snap,a.id,token)
    db.rollback()
    p=mapped_preview(db,snap,a.id)
    sub.apply_budget(db,snap,a.id,sub.fingerprint(p));db.commit()
    assert item.budget==120123 and db.query(AccountItem).count()==5


def test_budget_rejects_stale_snapshot_and_destination_change(db):
    snap=confirmed(db); a=account(db);db.commit()
    p=mapped_preview(db,snap,a.id)
    sub.mappings(db,2568)['teach'].budget=100;db.commit()
    with pytest.raises(ValueError,match='ยอดงบปลายทางเปลี่ยน'): sub.apply_budget(db,snap,a.id,sub.fingerprint(p))
    db.rollback()
    prepared(db, **{'r_ป.1_teach':'900'})
    with pytest.raises(ValueError,match='ฉบับใหม่'): sub.apply_budget(db,snap,a.id,sub.fingerprint(p))
    db.rollback()


def test_removed_extra_subtracts_only_previous_contribution(db):
    snap=confirmed(db,extra_small='500',extra_ref_small='หนังสือเพิ่มเติม')
    a=account(db);db.commit()
    p=mapped_preview(db,snap,a.id);sub.apply_budget(db,snap,a.id,sub.fingerprint(p));db.commit()
    item=db.query(AccountItem).filter_by(name=sub.NAMES['small']).one()
    item.budget+=50;db.commit()
    snap2=confirmed(db)
    p=mapped_preview(db,snap2,a.id);sub.apply_budget(db,snap2,a.id,sub.fingerprint(p));db.commit()
    assert item.budget==50


def test_two_academic_years_contribute_to_same_fiscal_budget(db):
    a=account(db);db.commit()
    prior=confirmed(db,ay=2567,term=2)
    p=mapped_preview(db,prior,a.id);sub.apply_budget(db,prior,a.id,sub.fingerprint(p));db.commit()
    current=confirmed(db,ay=2568,term=1)
    p=mapped_preview(db,current,a.id);sub.apply_budget(db,current,a.id,sub.fingerprint(p));db.commit()
    item=db.query(AccountItem).filter_by(fiscal_year=2568,name=sub.NAMES['teach']).one()
    assert item.budget==240000
    amended=confirmed(db,ay=2568,term=1,**{'r_ป.1_teach':'900'})
    p=mapped_preview(db,amended,a.id);sub.apply_budget(db,amended,a.id,sub.fingerprint(p));db.commit()
    assert item.budget==228000  # Other term's 120,000 remains intact.


def actual_form(db, **changes):
    data=dict(form_for(db,budget_basis='allocated',first_ref='งวดแรก',second_ref='ปรับยอด'))
    for k in sub.keys_for(1):
        data['first_'+k]='70000'
        data['second_'+k]='50000'
    data.update(changes)
    return FormData(data)


def test_actual_allocation_replaces_estimate_and_only_applies_delta(db):
    snap=confirmed(db,**{'n_jun2568_ป.1':'100'})
    a=account(db);db.commit()
    p=mapped_preview(db,snap,a.id);sub.apply_budget(db,snap,a.id,sub.fingerprint(p));db.commit()
    old_payload=snap.payload
    s=sub.save(db,2568,1,actual_form(db));db.commit()
    snap2=sub.confirm(db,2568,1,s['token']);db.commit()
    p=mapped_preview(db,snap2,a.id)
    teach=next(r for r in p['rows'] if r['key']=='teach')
    assert (teach['old'],teach['contribution'],teach['delta'],teach['new'])==(100000,120000,20000,120000)
    sub.apply_budget(db,snap2,a.id,sub.fingerprint(p));db.commit()
    repeated=mapped_preview(db,snap2,a.id)
    assert all(r['delta']==0 for r in repeated['rows'])
    assert snap.payload==old_payload
    s=sub.save(db,2568,1,actual_form(db,second_teach='10000'));db.commit()
    snap3=sub.confirm(db,2568,1,s['token']);db.commit()
    p=mapped_preview(db,snap3,a.id)
    assert next(r for r in p['rows'] if r['key']=='teach')['delta']==-40000


def test_actual_mode_needs_both_installments_but_not_estimate_inputs(db):
    d=dict(actual_form(db,second_teach=''))
    for scan in sub.state(db,2568,1)['census']:
        d['n_'+scan['key']+'_ป.1']=''
        d['confirm_'+scan['key']]=''
        d['source_'+scan['key']]=''
    for k in sub.keys_for(1):d['r_ป.1_'+k]=''
    s=sub.save(db,2568,1,FormData(d));db.commit()
    assert s['result']['budget_total'] is None and not s['result']['ready']
    with pytest.raises(ValueError,match='สองงวด'):sub.confirm(db,2568,1,s['token'])
    db.rollback()
    d.update(token=sub.state(db,2568,1)['token'],second_teach='0')
    s=sub.save(db,2568,1,FormData(d));db.commit()
    assert s['result']['total'] is None
    assert s['result']['budget_total']==550000 and s['result']['ready']
    assert sub.confirm(db,2568,1,s['token']).id


def test_supplementary_entries_sum_and_preserve_each_award(db):
    d=list(form_for(db).multi_items())+[('extra_small','10000'),('extra_ref_small','ครั้งแรก'),('extra_small','8000'),('extra_ref_small','ครั้งสอง')]
    s=sub.save(db,2568,1,FormData(d));db.commit()
    extra=s['config']['extras']['small']
    assert extra['amount']==18000 and len(extra['entries'])==2
    snap=sub.confirm(db,2568,1,s['token']);db.commit()
    a=account(db);db.commit()
    p=mapped_preview(db,snap,a.id);sub.apply_budget(db,snap,a.id,sub.fingerprint(p));db.commit()
    # Editing/removing an award adjusts only this category's previous contribution.
    s=prepared(db,extra_small='10000',extra_ref_small='ครั้งแรก')
    snap2=sub.confirm(db,2568,1,s['token']);db.commit()
    p=mapped_preview(db,snap2,a.id)
    assert next(r for r in p['rows'] if r['key']=='small')['delta']==-8000


def test_legacy_snapshot_budget_defaults_to_original_estimate(db):
    snap=confirmed(db)
    data=json.loads(snap.payload)
    for r in data['result']['rows']:r.pop('budget_amount')
    for k in ('budget_label','budget_basis','budget_total'):data['result'].pop(k)
    snap.payload=sub.dumps(data);db.commit()
    a=account(db);db.commit()
    assert mapped_preview(db,snap,a.id)['rows'][0]['contribution']==120000


def test_confirmed_report_uses_frozen_amounts_and_sources(db,monkeypatch,tmp_path):
    from app.services import finance_forms_doc, subsidy_report
    from docx import Document
    monkeypatch.setattr(finance_forms_doc,'get_data_dir',lambda:tmp_path)
    snap=confirmed(db)
    prepared(db,rate_source='หนังสือใหม่',**{'r_ป.1_teach':'900'})
    path=subsidy_report.render(db.query(School).first(),json.loads(snap.payload),snap.id)
    doc=Document(path)
    assert doc.tables[0].rows[1].cells[3].text=='120,000.00'
    assert not any('อ้างอิงอัตรา:' in p.text for p in doc.paragraphs)
    assert not any('หนังสือใหม่' in p.text for p in doc.paragraphs)


def test_routes_save_validation_snapshot_and_budget(env):
    from app.tenancy import session_for
    c=env
    r=c.get('/finance/subsidy?year=2568&term=1')
    assert r.status_code==200, r.text[:500]
    assert 'บาท/คน/ภาคเรียน' in r.text
    assert 'name="rate_source"' not in r.text and 'name="rates_confirmed"' not in r.text
    db=session_for(1)
    data=dict(form_for(db))
    a=account(db);db.commit();aid=a.id;ensure_mapping(db,2568,aid);db.close()
    invalid=dict(data, **{'r_ป.1_teach':'-1'})
    r=c.post('/finance/subsidy',data=invalid,follow_redirects=False)
    assert r.status_code==422 and 'ยังบันทึกไม่ได้' in r.text
    assert 'value="1000"' in r.text
    r=c.post('/finance/subsidy',data=data,follow_redirects=False)
    assert r.status_code==303,r.text[:500]
    db=session_for(1);s=sub.state(db,2568,1);db.close()
    r=c.post('/finance/subsidy/confirm',data={'academic_year':2568,'term':1,'token':s['token']},follow_redirects=False)
    assert r.status_code==303,r.text
    assert c.get(r.headers['location']).status_code==200
    sid=int(r.headers['location'].split('/')[-1])
    r=c.get(f'/finance/subsidy/budget-preview?snapshot_id={sid}&account_id={aid}')
    assert r.status_code==200,r.text
    token=re.search(r'name="token" value="([^"]+)"',r.text).group(1)
    post={'snapshot_id':sid,'account_id':aid,'token':token}
    assert c.post('/finance/subsidy/to-budget',data=post).status_code==409
    post['reviewed']='1'
    assert c.post('/finance/subsidy/to-budget',data=post,follow_redirects=False).status_code==303
    assert c.get('/finance/subsidy?year=2568&term=3').status_code==400


def test_dmc_save_without_source_or_checkbox_and_allocations_preserved(db):
    data = dict(form_for(db))
    data = {k:v for k,v in data.items() if not k.startswith(('source_', 'confirm_'))}
    data.update(first_teach='70000', first_ref='งวดแรก', second_teach='50000', second_ref='งวดสอง')
    s = sub.save(db,2568,1,FormData(data))
    assert s['result']['ready']
    assert all(scan['saved'] and not scan['confirmed'] and not scan['source'] for scan in s['census'])
    assert s['config']['first']['teach'] == 70000
    data.update(token=s['token'], budget_basis='estimate')
    s = sub.save(db,2568,1,FormData(data))
    assert s['config']['second']['teach'] == 50000


def test_explicit_mapping_custom_names_rename_and_wrong_receipt(db):
    a=account(db)
    custom=AccountItem(account_id=a.id,fiscal_year=2568,name='ค่ารายหัว',budget=10)
    db.add(custom);db.commit()
    choices={k:'new' for k in sub.keys_for(1)};choices['teach']=str(custom.id)
    sub.save_mapping(db,2568,a.id,choices,sub.mapping_token(db,2568));db.commit()
    snap=confirmed(db)
    p=sub.budget_preview(db,snap,a.id)
    assert p['rows'][0]['item_id']==custom.id and p['rows'][0]['new']==120010
    custom.name='รายหัวใหม่';db.commit()
    assert sub.budget_preview(db,snap,a.id)['rows'][0]['destination']=='รายหัวใหม่'
    t=FinanceTxn(account_id=a.id,item_id=sub.mappings(db,2568)['book'].id,fiscal_year=2568,kind='in',amount=500)
    db.add(t);db.commit()
    with pytest.raises(ValueError,match='ไม่ตรง'):sub.link_receipt(db,2568,1,t.id,'teach','first','100')
    db.rollback()
    sub.link_receipt(db,2568,1,t.id,'book','first','100');db.commit()
    choices={k:str(i.id) for k,i in sub.mappings(db,2568).items()}
    choices['teach']=choices['book']
    with pytest.raises(ValueError,match='คนละหมวด'):sub.save_mapping(db,2568,a.id,choices,sub.mapping_token(db,2568))
    db.rollback()


def test_mapping_setup_idempotent_stale_and_applied_cannot_remap(db):
    choices={k:'new' for k in sub.keys_for(1)}
    token=sub.mapping_token(db,2568)
    a=sub.save_mapping(db,2568,0,choices,token,'เงินอุดหนุนใหม่');db.commit()
    assert db.query(AccountItem).count()==5
    with pytest.raises(ValueError,match='เปลี่ยนแล้ว'):sub.save_mapping(db,2568,0,choices,token,'เงินอุดหนุนใหม่')
    db.rollback()
    snap=confirmed(db);p=sub.budget_preview(db,snap,a.id)
    sub.apply_budget(db,snap,a.id,sub.fingerprint(p));db.commit()
    other=AccountItem(account_id=a.id,fiscal_year=2568,name='อื่น',budget=0);db.add(other);db.commit()
    choices={k:str(i.id) for k,i in sub.mappings(db,2568).items()};choices['teach']=str(other.id)
    with pytest.raises(ValueError,match='เคยตั้งงบ'):sub.save_mapping(db,2568,a.id,choices,sub.mapping_token(db,2568))
    db.rollback()
    assert sub.mappings(db,2568)['teach'].id != other.id
    assert sub.mappings(db,2569)=={}


def test_receipt_changed_category_excluded_and_wrong_year_mapping_rejected(db):
    a=account(db);db.commit();ensure_mapping(db,2568,a.id)
    m=sub.mappings(db,2568)
    t=FinanceTxn(account_id=a.id,item_id=m['teach'].id,fiscal_year=2568,kind='in',amount=1000)
    db.add(t);db.commit()
    sub.link_receipt(db,2568,1,t.id,'teach','first','500');db.commit()
    assert sub.receipt_summary(db,2568,1)[0]['teach']==500
    t.item_id=m['book'].id;db.commit()
    assert sub.receipt_summary(db,2568,1)[0]=={}
    choices={k:str(i.id) for k,i in m.items()}
    with pytest.raises(ValueError,match='ปีงบ'):sub.save_mapping(db,2569,a.id,choices,sub.mapping_token(db,2569))
    db.rollback()


def test_mapping_routes_create_five_and_validation_keeps_inputs(env):
    from app.tenancy import session_for
    c=env
    r=c.get('/finance/subsidy/mapping?year=2568&term=1&account_id=0')
    assert r.status_code==200
    token=re.search(r'name="token" value="([^"]+)"',r.text).group(1)
    data={'academic_year':'2568','term':'1','account_id':'0','new_name':'บัญชีทดสอบ','token':token}
    data.update({'map_'+k:'new' for k in sub.keys_for(1)})
    invalid=dict(data,map_teach='')
    r=c.post('/finance/subsidy/mapping',data=invalid)
    assert r.status_code==422 and 'บัญชีทดสอบ' in r.text
    db=session_for(1)
    assert db.query(FinanceAccount).filter_by(name='บัญชีทดสอบ').count()==0
    db.close()
    r=c.post('/finance/subsidy/mapping',data=data,follow_redirects=False)
    assert r.status_code==303
    db=session_for(1)
    assert len(sub.mappings(db,2568))==5
    assert db.query(AccountItem).count()==5
    assert all(i.budget==0 for i in sub.mappings(db,2568).values())
    db.close()
    assert c.post('/finance/subsidy/mapping',data=data).status_code==422
