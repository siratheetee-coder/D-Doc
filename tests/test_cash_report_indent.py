from datetime import datetime
from types import SimpleNamespace
from docx import Document
from docx.shared import Cm
from app.services.cash_report import render_cash_report


def test_long_category_uses_same_paragraph_indent_as_siblings(tmp_path, monkeypatch):
    import app.services.cash_report as report
    monkeypatch.setattr(report, 'get_data_dir', lambda: tmp_path)
    school=SimpleNamespace(name='โรงเรียนตัวอย่าง',finance_officer_name='',officer_name='',director_name='')
    names=['เงินนอกงบประมาณ','เงินอุดหนุน','เงินอุดหนุนทั่วไป','ค่าจัดการเรียนการสอน','ค่ากิจกรรมพัฒนาคุณภาพผู้เรียน']
    levels=[0,1,2,3,3]
    rows=[dict(name=name,level=level,kind='leaf',bank=100,total=100) for name,level in zip(names,levels)]
    path=render_cash_report(school,rows,dict(bank=500,total=500),datetime(2026,10,11))
    doc=Document(path)
    for row,name,level in zip(doc.tables[0].rows[1:],names,levels):
        p=row.cells[0].paragraphs[0]
        assert p.text==name
        assert abs(p.paragraph_format.left_indent-Cm(.45*level))<1000
        assert p.paragraph_format.first_line_indent==0
