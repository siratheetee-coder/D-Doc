# -*- coding: utf-8 -*-
"""ชุดเอกสารจำหน่ายพัสดุ (ระเบียบกระทรวงการคลังฯ 2560 ข้อ 214-218)

อ้างแบบฟอร์มตามแนวทางการตรวจสอบพัสดุประจำปีและการจำหน่ายพัสดุสำหรับสถานศึกษา
(ฉบับปรับปรุง 2563) ซึ่งกำหนดเลขข้อและวงเงินที่ต้องอ้างในเอกสาร
"""
import datetime as dt
import json

import pytest
from docx import Document
from docx.shared import Emu
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

import app.database as dbm
from app.models import Asset, AssetDisposal, AssetDisposalItem, School
from app.services import asset_dispose_set as ds


@pytest.fixture()
def db(tmp_path, monkeypatch):
    monkeypatch.setattr(dbm, "get_data_dir", lambda: tmp_path)
    monkeypatch.setattr(ds, "get_data_dir", lambda: tmp_path)
    eng = create_engine(f"sqlite:///{tmp_path/'t.db'}")
    dbm.init_school_db(eng)
    with Session(bind=eng) as s:
        yield s


@pytest.fixture()
def school():
    return School(name="โรงเรียนบ้านตัวอย่าง", address="ต.ในเมือง อ.เมือง จ.สมมติ",
                  area_office="สำนักงานเขตพื้นที่การศึกษาประถมศึกษาสมมติ เขต 1",
                  director_name="นายเอนก ทดสอบ", officer_name="นางสาวพัสดุ ใจดี",
                  head_officer_name="นายหัวหน้า พัสดุ")


def _make(db, actions=("ขาย", "ทำลาย", "จำหน่ายเป็นสูญ"), cost=25000.0):
    dp = AssetDisposal(year=2569, stage="fact", sao_region="สมมติ",
                       fact_order_no="45/2569", fact_order_date=dt.datetime(2026, 10, 3),
                       order_no="48/2569", order_date=dt.datetime(2026, 10, 12),
                       sale_total=9500.0, buyer_name="ร้านรับซื้อของเก่าสมชาย",
                       written_off_date=dt.datetime(2026, 10, 31),
                       members=json.dumps({k: [{"name": f"ครู{k}", "position": "ครู",
                                                "role": "ประธานกรรมการ"}]
                                           for k in ("fact", "price", "sale", "destroy")},
                                          ensure_ascii=False))
    db.add(dp)
    db.flush()
    for i, act in enumerate(actions, 1):
        a = Asset(asset_code=f"7440-001-{i:04d}/2569", name=f"ครุภัณฑ์ {i}",
                  category="ครุภัณฑ์สำนักงาน", cost=cost, quantity=1, unit="เครื่อง",
                  useful_life=8, status="ชำรุด", acquired_date=dt.datetime(2023, 8, 31))
        db.add(a)
        db.flush()
        db.add(AssetDisposalItem(disposal_id=dp.id, asset_id=a.id, action=act,
                                 cause="ใช้งานมานาน", price_mid=4500.0))
    db.commit()
    return dp


def _text(path):
    d = Document(path)
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            out += [c.text for c in r.cells]
    return "\n".join(out)


def test_every_doc_renders(db, school):
    """ทั้ง 21 ฉบับต้องออกได้โดยไม่พัง แม้ข้อมูลกรอกไม่ครบ"""
    dp = _make(db)
    assert len(ds.DOCS) == 21
    for key, (label, fn) in ds.DOCS.items():
        path = fn(school, dp)
        assert path.endswith(".docx"), key
        assert len(_text(path)) > 200, key


def test_clause_numbers_match_2560(db, school):
    """เลขข้อที่อ้างต้องตรงระเบียบฯ 2560 ไม่ใช่ระเบียบสำนักนายกฯ 2535"""
    dp = _make(db)
    txt = _text(ds.render_dispose_request(school, dp))
    assert "ข้อ 215 (1)" in txt and "ข้อ 215 (4)" in txt
    assert "ข้อ 217 (1)" in txt                      # มีรายการจำหน่ายเป็นสูญ
    assert "2535" not in txt and "ข้อ 157" not in txt   # ห้ามหลงเหลือของเก่า
    # คำสั่งมอบอำนาจ สพฐ. อ้างในคำสั่งแต่งตั้ง (ตามแบบฟอร์มจริง ไม่ได้อ้างในบันทึก)
    cmd = _text(ds.render_committee_order(school, dp))
    assert "1340/2560" in cmd and "24 สิงหาคม 2560" in cmd


def test_fact_finding_cites_214(db, school):
    dp = _make(db)
    for fn in (ds.render_fact_memo, ds.render_fact_order):
        txt = _text(fn(school, dp))
        assert "ข้อ 214" in txt, fn.__name__


def test_writeoff_limit_switches_authority(db, school):
    """จำหน่ายเป็นสูญ: ไม่เกิน 1 ล้าน ผอ. อนุมัติเอง · เกินกว่านั้นเป็นอำนาจกระทรวงการคลัง"""
    dp = _make(db, actions=("จำหน่ายเป็นสูญ",), cost=900_000.0)
    txt = _text(ds.render_writeoff_memo(school, dp))
    assert "อนุมัติให้จำหน่ายพัสดุเป็นสูญ" in txt
    # ไม่ต้องส่งเรื่องให้กระทรวงการคลังพิจารณา (แต่ยังต้อง "รายงาน" ตามข้อ 218 ได้)
    assert "เสนอกระทรวงการคลังพิจารณาอนุมัติ" not in txt
    # (ข้อความระเบียบที่ยกมาอ้างมีคำว่า "ไม่เกิน 1,000,000 บาท" อยู่เสมอ จึงเช็กที่ประโยคยกระดับ)
    assert "มีราคาซื้อหรือได้มารวมกันเกิน 1,000,000 บาท" not in txt

    dp2 = _make(db, actions=("จำหน่ายเป็นสูญ",), cost=1_500_000.0)
    txt2 = _text(ds.render_writeoff_memo(school, dp2))
    assert "เกิน 1,000,000 บาท" in txt2 and "ข้อ 217 (2) (ก)" in txt2


def test_close_bundle_targets(db, school):
    """ข้อ 218 ให้รายงาน 3 ทาง · แจ้งกระทรวงการคลังเฉพาะกรณีจำหน่ายเป็นสูญ"""
    dp = _make(db, actions=("ขาย", "จำหน่ายเป็นสูญ"))
    assert ds._keys_for(dp, "close") == ["area", "sao", "mof", "remit"]

    dp2 = _make(db, actions=("ขาย",))               # ไม่มีจำหน่ายเป็นสูญ -> ไม่ต้องแจ้งคลัง
    assert "mof" not in ds._keys_for(dp2, "close")

    dp3 = _make(db, actions=("ทำลาย",))             # ไม่มีเงินเข้า -> ไม่ต้องมีหนังสือนำส่งเงิน
    dp3.sale_total = 0.0
    assert "remit" not in ds._keys_for(dp3, "close")


def test_sale_mode_selects_bundle(db, school):
    """เลือกเฉพาะเจาะจง -> ไม่ออกชุดทอดตลาด และกลับกัน"""
    dp = _make(db, actions=("ขาย",))
    dp.sale_mode = "specific"
    assert ds._keys_for(dp, "sell") and not ds._keys_for(dp, "auction")
    dp.sale_mode = "auction"
    assert ds._keys_for(dp, "auction") and not ds._keys_for(dp, "sell")


def test_no_table_exceeds_paper(db, school):
    """ตารางทุกใบต้องไม่ล้นพื้นที่พิมพ์ของ section ที่กว้างที่สุดในไฟล์นั้น"""
    dp = _make(db)
    for key, (label, fn) in ds.DOCS.items():
        doc = Document(fn(school, dp))
        usable = max(s.page_width - s.left_margin - s.right_margin for s in doc.sections)
        for ti, t in enumerate(doc.tables):
            w = sum(c.width for c in t.rows[0].cells if c.width)
            assert w <= usable + Emu(20000), f"{key} ตาราง {ti}: {w/360000:.2f} ซม."


def test_bundle_has_no_blank_page_markers(db, school):
    """ชุดรวมต้องไม่ลงท้ายด้วยย่อหน้าว่าง และไม่มีย่อหน้าที่มีแต่ page break
    (ย่อหน้าแบบนั้นทำให้เกิดหน้าเปล่าคั่นเมื่อหน้าก่อนหน้าเต็มพอดี)"""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    dp = _make(db)
    doc = Document(ds.render_full_set(school, dp))
    body = doc.element.body
    for para in body.findall(qn('w:p')):
        runs = para.findall(qn('w:r'))
        brs = [b for r in runs for b in r.findall(qn('w:br'))
               if b.get(qn('w:type')) == 'page']
        has_text = any(r.findall(qn('w:t')) for r in runs)
        assert not (brs and not has_text), "เหลือย่อหน้าที่มีแต่ page break"
    kids = [e for e in body.iterchildren() if not e.tag.endswith('sectPr')]
    assert not (kids[-1].tag.endswith('}p') and not Paragraph(kids[-1], doc).text.strip())


def test_full_set_orders_stages(db, school):
    """ชุดทั้งสำนวนต้องเรียงตามลำดับขั้นของระเบียบ"""
    dp = _make(db, actions=("ขาย", "ทำลาย", "จำหน่ายเป็นสูญ"))
    dp.sale_mode = "specific"
    txt = _text(ds.render_full_set(school, dp))
    order = ["แต่งตั้งคณะกรรมการสอบหาข้อเท็จจริง", "ขอจำหน่ายพัสดุ",
             "การประเมินราคากลางพัสดุ", "รายงานผลการทำลายพัสดุ",
             "ขออนุมัติจำหน่ายพัสดุเป็นสูญ", "ปลัดกระทรวงการคลัง"]
    pos = [txt.find(x) for x in order]
    assert all(p >= 0 for p in pos), list(zip(order, pos))
    assert pos == sorted(pos), list(zip(order, pos))


def test_empty_disposal_refuses_bundle(db, school):
    """สำนวนที่ยังไม่มีรายการ -> บอกให้ชัดว่าไม่มีอะไรให้ออก ไม่ใช่ไฟล์เปล่า"""
    dp = AssetDisposal(year=2569)
    db.add(dp)
    db.commit()
    with pytest.raises(ValueError):
        ds.render_full_set(school, dp)
    with pytest.raises(ValueError):
        ds.render_bundle(school, dp, "destroy")
