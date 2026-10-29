"""จ้างแม่ครัว/ซื้อวัตถุดิบ: ยืมเงินที่หน้ารอบ -> สัญญาในทะเบียนคุมเงินยืม + บัญชีอาหารกลางวัน/การเงินไม่นับซ้ำ
(คู่มือ สพฐ. 2569 ภาคผนวก ค: ยืม -> ส่งใช้ด้วยใบสำคัญ + เงินสดเหลือจ่าย, ค่าจ้างแม่ครัวไม่ต้องยืม)"""
import io
from datetime import datetime

from tests.test_subsidy import env  # noqa: F401


def _db():
    import app.tenancy as tn
    return tn.session_for(next(iter(tn._engines)))


def _setup():
    from app.models import LunchProgram, LunchClass, LunchHireRound, LunchIngredient, Vendor
    with _db() as s:
        v = Vendor(name="นางแม่ครัว ทดสอบ"); s.add(v); s.flush()
        p = LunchProgram(year=2570, rate_per_head=27, operate_mode="person", lunch_officer="นางสาวผู้ยืม เงินดี")
        s.add(p); s.flush()
        s.add(LunchClass(program_id=p.id, seq=1, level="ป.1", num_students=20)); s.flush()
        # 20 คน x 27 x 10 วัน = 5,400 หักค่าจ้างแม่ครัว 1,500 = ยืม 3,900
        r = LunchHireRound(program_id=p.id, seq=1, vendor_id=v.id, start_date=datetime(2027, 6, 1),
                           end_date=datetime(2027, 6, 14), days=10, cook_wage=1500, fuel_cost=200)
        s.add(r); s.flush()
        s.add(LunchIngredient(program_id=p.id, date=datetime(2027, 6, 2), name="ข้าวสาร", quantity=10,
                              unit="กก.", unit_price=300))          # 3,000 + เชื้อเพลิง 200 = ใช้ 3,200
        s.commit()
        return r.id, p.id


def test_loan_links_finance_without_double_count(env):
    from docx import Document
    from app.models import LunchHireRound, LunchLedger, MoneyLoan, FinanceTxn
    rid, pid = _setup()

    page = env.get(f"/lunch/round/{rid}/plan").text
    assert 'action="/lunch/round/%d/loan"' % rid in page and "3,900.00" in page

    # ยืม (ยังไม่ได้รับเงิน) -> มีสัญญาในทะเบียนคุมเงินยืม แต่ยังไม่มีรายการเงิน
    env.post(f"/lunch/round/{rid}/loan", data={"contract_no": "5/2570", "date": "01/06/2570",
                                               "receive_date": "", "amount": "3900", "borrower": "นางสาวผู้ยืม เงินดี"})
    with _db() as s:
        rnd = s.get(LunchHireRound, rid)
        loan = s.get(MoneyLoan, rnd.loan_id)
        assert loan.amount == 3900 and loan.contract_no == "5/2570" and loan.position
        assert s.query(LunchLedger).filter_by(loan_id=loan.id).count() == 0
        lid = loan.id
    assert "5/2570" in env.get("/finance/loans?year=2570").text

    # ได้รับเงิน -> จ่ายเงินยืม 3,900 ทั้งบัญชีอาหารกลางวันและการเงิน (ผูก loan_id ให้ทะเบียนคุมเงินยืม)
    env.post(f"/lunch/round/{rid}/loan", data={"contract_no": "5/2570", "date": "01/06/2570",
                                               "receive_date": "02/06/2570", "amount": "3900", "borrower": "นางสาวผู้ยืม เงินดี"})
    with _db() as s:
        led = s.query(LunchLedger).filter_by(loan_id=lid).one()
        assert led.kind == "out" and led.amount == 3900
        txn = s.get(FinanceTxn, led.finance_txn_id)
        assert txn.loan_id == lid and txn.loan_return_id is None and txn.amount == 3900
        assert s.get(MoneyLoan, lid).due_date == datetime(2027, 7, 2)

    # ส่งใช้: ใบสำคัญ 3,200 เงินสดคืน 700
    env.post(f"/lunch/round/{rid}/loan/repay", data={"date": "15/06/2570"})
    with _db() as s:
        loan = s.get(MoneyLoan, lid)
        got = sorted((r.kind, r.amount) for r in loan.returns)
        assert dict(got) == {"ใบสำคัญ": 3200, "เงินสด": 700}
        back = s.query(LunchLedger).filter_by(loan_id=lid, kind="in").one()
        assert back.amount == 700 and s.get(FinanceTxn, back.finance_txn_id).loan_return_id == back.loan_return_id

    # ปิดรอบ -> ลงเฉพาะค่าจ้างแม่ครัว 1,500 (ค่าวัตถุดิบ/เชื้อเพลิงออกจากเงินยืมแล้ว)
    env.post(f"/lunch/round/{rid}/status", data={"status": "จ่ายแล้ว"})
    with _db() as s:
        close = s.query(LunchLedger).filter_by(round_id=rid).one()
        assert close.amount == 1500
        out = sum(l.amount for l in s.query(LunchLedger).filter_by(program_id=pid, kind="out"))
        back = sum(l.amount for l in s.query(LunchLedger).filter_by(program_id=pid, kind="in"))
        assert round(out - back, 2) == 3000 + 200 + 1500        # จ่ายจริงทั้งรอบ ไม่ซ้ำ

    # เอกสารยืม/ส่งใช้ใช้ยอดจริง (เดิมได้ 0 บาท เพราะใช้วงเงินรอบที่ไม่ได้กรอก)
    doc = Document(io.BytesIO(env.get(f"/lunch/round/{rid}/ingredient-doc/borrow").content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "เป็นเงิน 3,900 บาท" in text and "หักค่าจ้างแม่ครัว 1,500 บาท" in text
    doc = Document(io.BytesIO(env.get(f"/lunch/round/{rid}/ingredient-doc/repay").content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "จำนวน 3,200 บาท" in text and "จำนวน 700 บาท" in text

    # ยกเลิกการส่งใช้ -> รายการรับคืนหาย ยอดปิดรอบคิดใหม่ (ยังไม่ส่งใช้ = เงินยืมจ่ายวัตถุดิบไป 3,200)
    env.post(f"/lunch/round/{rid}/loan/repay", data={"action": "undo"})
    with _db() as s:
        assert s.query(LunchLedger).filter_by(loan_id=lid, kind="in").count() == 0
        assert s.query(LunchLedger).filter_by(round_id=rid).one().amount == 1500
    # ลบรอบ -> สัญญายืมและรายการเงินที่ระบบลงให้หายด้วย
    env.post(f"/lunch/round/{rid}/delete")
    with _db() as s:
        assert s.get(MoneyLoan, lid) is None
        assert s.query(LunchLedger).filter_by(loan_id=lid).count() == 0


def test_old_round_without_loan_unchanged(env):
    """รอบเดิมที่ไม่ได้ยืมผ่านการ์ด ปิดรอบแล้วลงยอดเต็มเหมือนเดิม"""
    from app.models import LunchLedger
    rid, _ = _setup()
    env.post(f"/lunch/round/{rid}/status", data={"status": "จ่ายแล้ว"})
    with _db() as s:
        assert s.query(LunchLedger).filter_by(round_id=rid).one().amount == 3000 + 200 + 1500
