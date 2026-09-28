# -*- coding: utf-8 -*-
"""ทะเบียนคุมเงิน (ฟอร์มกลาง) · รายบัญชี · รายการย่อย · รวมทุกบัญชี

ยึดตามแบบฟอร์มทะเบียนคุมเงินนอกงบประมาณที่โรงเรียนใช้จริง
คอลัมน์: วัน เดือน ปี | ที่เอกสาร | รายการ | รับ |
         จ่าย(ลูกหนี้ · ค่าตอบแทนฯ · รวมจ่าย) | คงเหลือ(เงินสด · ธนาคาร · ส่วนราชการ) | หมายเหตุ
"""
import pathlib
import sys
import tempfile
from datetime import datetime

import pytest
from docx import Document
from docx.oxml.ns import qn

ROOT = pathlib.Path(__file__).resolve().parents[1]
FY = 2569


@pytest.fixture()
def env(monkeypatch):
    import importlib.util

    tmp = pathlib.Path(tempfile.mkdtemp())
    import app.accounts as ac
    import app.database as dbm
    import app.tenancy as tn
    for m in (dbm, tn, ac):
        monkeypatch.setattr(m, "get_data_dir", lambda t=tmp: t)
    monkeypatch.setattr(tn, "_engines", type(tn._engines)())
    monkeypatch.setattr(ac, "_engine", None)
    monkeypatch.setattr(ac, "_Session", None)
    import app.main as main_mod
    import app.routers.auth as auth_mod
    monkeypatch.setattr(auth_mod, "authenticate", ac.authenticate)
    monkeypatch.setattr(main_mod, "tenant_state", ac.tenant_state)
    monkeypatch.setattr(main_mod, "can_use_module", ac.can_use_module)
    monkeypatch.setattr(main_mod, "get_account_access", ac.get_account_access)

    spec = importlib.util.spec_from_file_location("seed_demo", ROOT / "tools" / "seed_demo.py")
    seed = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(seed)
    sys.argv = ["seed_demo.py", "--password", "Demo!2569"]
    seed.main()

    from app.accounts import Tenant, acc_session
    from app.models import AccountItem, FinanceAccount, FinanceTxn
    from app.tenancy import session_for
    with acc_session() as s:
        tid = s.query(Tenant).filter_by(slug="demo").one().id
    db = session_for(tid)
    acct = FinanceAccount(name="เงินอุดหนุน", fund_type="เงินนอกงบประมาณ",
                          deposit_type="bank", opening_balance=5000.0)
    db.add(acct)
    db.flush()
    ids = {}
    for nm, budget in [("ค่าจัดการเรียนการสอน", 120000), ("ค่าเครื่องแบบนักเรียน", 18000),
                       ("ค่ากิจกรรมพัฒนาคุณภาพผู้เรียน", 31000)]:
        it = AccountItem(account_id=acct.id, fiscal_year=FY, name=nm, budget=budget,
                         deposit_type="bank")
        db.add(it)
        db.flush()
        ids[nm] = it.id
    for item, kind, amt, note, ref, d in [
        ("ค่าจัดการเรียนการสอน", "in", 120000, "รับจัดสรรงวดที่ 1", "ฎ.1/2569",
         datetime(2025, 10, 15)),
        ("ค่าเครื่องแบบนักเรียน", "in", 18000, "รับจัดสรรค่าเครื่องแบบ", "ฎ.3/2569",
         datetime(2025, 10, 15)),
        ("ค่าเครื่องแบบนักเรียน", "out", 18000, "จ่ายค่าเครื่องแบบนักเรียน", "ฎ.9/2569",
         datetime(2025, 11, 20)),
        ("ค่ากิจกรรมพัฒนาคุณภาพผู้เรียน", "out", 15000, "ลูกหนี้เงินยืม - ครูสมศรี",
         "ย.2/2569", datetime(2026, 1, 9)),
    ]:
        db.add(FinanceTxn(account_id=acct.id, fiscal_year=FY, date=d, kind=kind,
                          amount=amt, ref=ref, note=note, item_id=ids[item]))
    other = FinanceAccount(name="รายได้สถานศึกษา", fund_type="เงินนอกงบประมาณ",
                           deposit_type="cash", opening_balance=2500.0)
    db.add(other)
    db.commit()
    aid, uniform = acct.id, ids["ค่าเครื่องแบบนักเรียน"]
    db.close()

    from fastapi.testclient import TestClient

    from app.main import app
    c = TestClient(app, raise_server_exceptions=False)
    r = c.post("/login", data={"username": "demo", "password": "Demo!2569"},
               follow_redirects=False)
    assert r.status_code in (302, 303)
    return c, aid, uniform


def _cells(path):
    d = Document(path)
    return [[c.text.strip() for c in row.cells] for t in d.tables for row in t.rows]


def _save(content, tmp_path, name="r.docx"):
    p = tmp_path / name
    p.write_bytes(content)
    return p


def test_account_register_matches_form(env, tmp_path):
    """ทะเบียนคุมรายบัญชี: แถวแรกเป็นยอดยกมา · คงเหลือไหลถูก · แยกลูกหนี้ออกจากค่าใช้จ่าย"""
    c, aid, _ = env
    r = c.get(f"/finance/accounts/{aid}/money-register.docx?year={FY}")
    assert r.status_code == 200
    rows = _cells(_save(r.content, tmp_path))
    flat = "\n".join(" | ".join(x) for x in rows)
    assert "ทะเบียนคุมเงิน" not in flat        # หัวเรื่องอยู่นอกตาราง
    assert f"ยอดยกมาจากปีงบประมาณ {FY - 1}" in flat
    assert "ลูกหนี้" in flat and "รวมจ่าย" in flat
    assert "เงินฝากส่วนราชการ" in flat
    # ยอดคงเหลือสุดท้าย = ยกมา 5,000 + รับ 138,000 - จ่าย 33,000 = 110,000
    assert "110,000.00" in flat, flat[-600:]
    # เงินยืมต้องไปช่องลูกหนี้ ไม่ใช่ช่องค่าใช้จ่าย
    debtor = next(x for x in rows if "ลูกหนี้เงินยืม" in " ".join(x))
    assert debtor[4] == "15,000.00" and debtor[5] == "", debtor


def test_item_register_is_scoped_to_that_item(env, tmp_path):
    """ทะเบียนคุมของรายการย่อย ต้องมีเฉพาะรายการของตัวเอง"""
    c, aid, uniform = env
    r = c.get(f"/finance/accounts/{aid}/money-register.docx?year={FY}&item={uniform}")
    assert r.status_code == 200
    flat = "\n".join(" | ".join(x) for x in _cells(_save(r.content, tmp_path)))
    assert "จ่ายค่าเครื่องแบบนักเรียน" in flat
    assert "รับจัดสรรงวดที่ 1" not in flat        # ของรายการย่อยอื่น ห้ามหลุดมา
    assert "ลูกหนี้เงินยืม" not in flat
    d = Document(tmp_path / "r.docx")
    head = "\n".join(p.text for p in d.paragraphs)
    assert "ประเภทเงิน ค่าเครื่องแบบนักเรียน" in head    # แบบฟอร์มระบุประเภทเงินไว้ด้านบน


def test_item_register_404_for_wrong_account(env):
    c, aid, _ = env
    assert c.get(f"/finance/accounts/{aid}/money-register.docx"
                 f"?year={FY}&item=999999").status_code == 404


def test_all_registers_covers_accounts_and_items(env, tmp_path):
    """ออกทะเบียนคุมทั้งหมด = ทุกบัญชี + ทุกรายการย่อยที่มีงบหรือมีรายการ"""
    c, _aid, _u = env
    r = c.get(f"/finance/registers.docx?year={FY}")
    assert r.status_code == 200
    p = _save(r.content, tmp_path, "all.docx")
    d = Document(p)
    heads = [x.text.strip() for x in d.paragraphs if x.text.strip().startswith("ประเภทเงิน")]
    names = {h.replace("ประเภทเงิน", "").split("(")[0].strip() for h in heads}
    assert {"เงินอุดหนุน", "รายได้สถานศึกษา"} <= names, names
    assert {"ค่าจัดการเรียนการสอน", "ค่าเครื่องแบบนักเรียน",
            "ค่ากิจกรรมพัฒนาคุณภาพผู้เรียน"} <= names, names


def test_all_registers_pages_are_landscape_and_clean(env, tmp_path):
    from docx.enum.section import WD_ORIENT
    c, _a, _u = env
    r = c.get(f"/finance/registers.docx?year={FY}")
    d = Document(_save(r.content, tmp_path, "all.docx"))
    assert all(s.orientation == WD_ORIENT.LANDSCAPE for s in d.sections)
    widest = max(s.page_width - s.left_margin - s.right_margin for s in d.sections)
    for i, t in enumerate(d.tables):
        grid = t._tbl.find(qn('w:tblGrid'))
        w = sum(int(g.get(qn('w:w'))) for g in grid.findall(qn('w:gridCol'))
                if g.get(qn('w:w'))) * 635
        assert w <= widest + 20000, f"ตาราง {i} กว้าง {w/360000:.2f} ซม."


def test_accounts_page_shows_sub_items(env):
    """หน้าทะเบียนคุมเงินต้องเห็นรายการย่อย + ปุ่มออกทะเบียนคุมทั้งหมด"""
    c, aid, _ = env
    html = c.get(f"/finance/accounts?year={FY}").text
    assert "ออกทะเบียนคุมทั้งหมด" in html
    assert "/finance/registers.docx" in html
    for nm in ("ค่าจัดการเรียนการสอน", "ค่าเครื่องแบบนักเรียน"):
        assert nm in html, nm
    assert f"/finance/accounts/{aid}/money-register.docx" in html
    assert "รายการย่อย 3" in html


def test_ledger_page_offers_money_register_first(env):
    c, aid, _ = env
    html = c.get(f"/finance/accounts/{aid}?year={FY}").text
    assert "ทะเบียนคุมเงิน (Word)" in html
    assert "money-register.docx" in html
