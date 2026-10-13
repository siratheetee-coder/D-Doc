# -*- coding: utf-8 -*-
"""เอกสารจากบัญชีรับ-จ่ายอาหารกลางวัน: ใบสำคัญรับเงิน (รับและจ่าย) + ใบเสร็จรับเงิน (รายการรับ)

เดิมใบสำคัญรับเงินเป็นย่อหน้าเล่าเรื่อง ไม่มีตารางรายการ/จำนวนเงิน (เจ้าของบอกว่าดูไม่เหมือนใบสำคัญ)
รัน: .venv\\Scripts\\python.exe -m pytest tests/test_lunch_receipts.py
"""
from datetime import datetime
from io import BytesIO

from docx import Document

from tests.test_subsidy import env  # noqa: F401  (fixture)
from app.models import LunchHireRound, LunchLedger, LunchProgram, School, Vendor


def _text(content):
    d = Document(BytesIO(content))
    return ' '.join(p.text for p in d.paragraphs) + ' ' + ' '.join(
        c.text for t in d.tables for r in t.rows for c in r.cells)


def _seed(sid):
    import app.tenancy as tn
    with tn.session_for(sid) as s:
        sc = s.query(School).first(); sc.finance_officer_name = 'นางการเงิน ทดสอบ'
        p = LunchProgram(year=2569, funding_org='อบต.ทดสอบ', lunch_officer='นางอาหาร ทดสอบ'); s.add(p); s.flush()
        v = Vendor(name='ร้านครัวทดสอบ', address='12 หมู่ 3'); s.add(v); s.flush()
        r = LunchHireRound(program_id=p.id, seq=1, vendor_id=v.id); s.add(r); s.flush()
        inc = LunchLedger(program_id=p.id, kind='in', amount=128400, date=datetime(2026, 11, 1),
                          detail='เงินอุดหนุนอาหารกลางวัน รอบ 1', ref='012/2569')
        out = LunchLedger(program_id=p.id, kind='out', amount=44000, date=datetime(2026, 11, 20),
                          detail='ค่าจ้างเหมา งวดที่ 1', round_id=r.id)
        s.add_all([inc, out]); s.commit()
        return p.id, inc.id, out.id


def test_income_voucher_is_a_real_voucher(env):
    import app.tenancy as tn
    pid, inc, out = _seed(next(iter(tn._engines)))
    r = env.get(f'/lunch/ledger/{inc}/receipt.docx')
    assert r.status_code == 200
    t = _text(r.content)
    assert 'ใบสำคัญรับเงิน' in t and 'ได้รับเงินจาก อบต.ทดสอบ' in t
    assert 'รายการ' in t and 'ส.ต.' in t and '128,400' in t and 'หนึ่งแสนสองหมื่นแปดพันสี่ร้อยบาทถ้วน' in t
    assert '( นางการเงิน ทดสอบ )' in t, 'เจ้าหน้าที่การเงินลงชื่อเป็นผู้รับเงิน ในนามโรงเรียน'


def test_payment_voucher_names_the_contractor(env):
    import app.tenancy as tn
    pid, inc, out = _seed(next(iter(tn._engines)))
    t = _text(env.get(f'/lunch/ledger/{out}/receipt.docx').content)
    assert 'ข้าพเจ้า ร้านครัวทดสอบ' in t and '12 หมู่ 3' in t and '44,000' in t
    assert '( นางอาหาร ทดสอบ )' in t, 'ผู้จ่ายเงิน = เจ้าหน้าที่โครงการอาหารกลางวัน'


def test_official_receipt_for_income_only(env):
    import app.tenancy as tn
    pid, inc, out = _seed(next(iter(tn._engines)))
    r = env.get(f'/lunch/ledger/{inc}/official-receipt.docx')
    t = _text(r.content)
    assert 'ใบเสร็จรับเงิน' in t and 'ต้นฉบับ' in t and 'สำเนา' in t
    assert 'เลขที่ 012/2569' in t and 'ได้รับเงินจาก  อบต.ทดสอบ' in t and 'รวมเงิน' in t
    r = env.get(f'/lunch/ledger/{out}/official-receipt.docx', follow_redirects=False)
    assert r.status_code == 303, 'รายการจ่ายไม่มีใบเสร็จ'
    html = env.get(f'/lunch/{pid}').text
    assert html.count('official-receipt.docx') == 1 and html.count('/receipt.docx') == 2
