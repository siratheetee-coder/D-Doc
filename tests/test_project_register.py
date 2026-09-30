# -*- coding: utf-8 -*-
"""ทะเบียนคุมโครงการ ผูกกับเงินที่จ่ายจริง

โจทย์: เดิมหน้าโครงการรู้แค่ "งบที่ตั้งไว้" ส่วนเงินที่จ่ายจริงอยู่ในงานการเงิน
คนละที่กัน ครูต้องไล่บวกเอง · ตอนนี้ยอดใช้ไปนับจากของจริงทุกทาง
  1. เรื่องจัดซื้อ/จัดจ้างที่ผูกโครงการ
  2. บันทึกขอเบิกจ่ายที่ผูกโครงการ
  3. รายการจ่ายตรงในทะเบียนคุมเงิน (ใหม่ - FinanceTxn.project_id)
และต้องไม่นับซ้ำเมื่อรายการหนึ่งเชื่อมกันหลายชั้น
รัน: .venv\\Scripts\\python.exe -m pytest tests/test_project_register.py
"""
import importlib.util
import pathlib
import sys
import tempfile
from datetime import datetime

import pytest
from docx import Document

ROOT = pathlib.Path(__file__).resolve().parents[1]
PY = 2569


@pytest.fixture()
def env(monkeypatch):
    tmp = pathlib.Path(tempfile.mkdtemp())
    import app.accounts as ac
    import app.database as dbm
    import app.tenancy as tn
    for m in (dbm, tn, ac):
        monkeypatch.setattr(m, "get_data_dir", lambda t=tmp: t)
    monkeypatch.setattr(tn, "_engines", type(tn._engines)())
    monkeypatch.setattr(ac, "_engine", None)
    monkeypatch.setattr(ac, "_Session", None)
    import app.main as main_mod
    import app.routers.auth as auth_mod
    monkeypatch.setattr(auth_mod, "authenticate", ac.authenticate)
    monkeypatch.setattr(main_mod, "tenant_state", ac.tenant_state)
    monkeypatch.setattr(main_mod, "can_use_module", ac.can_use_module)
    monkeypatch.setattr(main_mod, "get_account_access", ac.get_account_access)

    spec = importlib.util.spec_from_file_location("seed_demo", ROOT / "tools" / "seed_demo.py")
    seed = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(seed)
    sys.argv = ["seed_demo.py", "--password", "Demo!2569"]
    seed.main()

    from app.accounts import Tenant, acc_session
    from app.models import DisburseMemo, FinanceAccount, FinanceTxn, Project, School
    from app.tenancy import session_for
    with acc_session() as s:
        tid = s.query(Tenant).filter_by(slug="demo").one().id
    db = session_for(tid)
    sc = db.query(School).first()
    sc.academic_year = None
    pj = Project(name="ส่งเสริมการอ่านออกเขียนได้", budget=50000, plan_year=PY,
                 responsible="ฝ่ายวิชาการ")
    other = Project(name="โครงการที่ไม่ได้ใช้เงิน", budget=10000, plan_year=PY)
    db.add_all([pj, other])
    db.flush()
    acct = FinanceAccount(name="เงินอุดหนุน", fund_type="เงินนอกงบประมาณ",
                          deposit_type="bank", opening_balance=0)
    db.add(acct)
    db.flush()
    # 1) บันทึกขอเบิกจ่ายที่ผูกโครงการ
    m = DisburseMemo(fiscal_year=PY, memo_no="12/2569", date=datetime(2026, 1, 5),
                     subject="ค่าหนังสือส่งเสริมการอ่าน", payee="ร้านหนังสือ",
                     amount=8000, project_id=pj.id)
    db.add(m)
    # 2) จ่ายตรงจากทะเบียนคุมเงิน
    db.add(FinanceTxn(account_id=acct.id, fiscal_year=PY, date=datetime(2026, 2, 3),
                      kind="out", amount=2500, ref="ฎ.20/2569", note="ค่าป้ายไวนิล",
                      project_id=pj.id))
    # 3) รายการรับ ไม่นับเป็นค่าใช้จ่าย
    db.add(FinanceTxn(account_id=acct.id, fiscal_year=PY, date=datetime(2026, 2, 4),
                      kind="in", amount=1000, note="รับคืนเงิน", project_id=pj.id))
    db.commit()
    pj_id = pj.id
    db.close()

    from fastapi.testclient import TestClient

    from app.main import app
    c = TestClient(app, raise_server_exceptions=False)
    r = c.post("/login", data={"username": "demo", "password": "Demo!2569"},
               follow_redirects=False)
    assert r.status_code in (302, 303)
    return c, pj_id, tid


def test_spent_counts_every_route_money_leaves(env):
    """ใช้ไป = 8,000 (ขอเบิกจ่าย) + 2,500 (จ่ายตรง) = 10,500 · รายการรับไม่นับ"""
    from app.models import Project
    from app.services.budget import project_spent
    from app.tenancy import session_for
    _c, pj_id, tid = env
    db = session_for(tid)
    try:
        pj = db.get(Project, pj_id)
        assert project_spent(pj) == 10500.0, project_spent(pj)
    finally:
        db.close()


def test_direct_payment_via_disburse_memo_is_not_double_counted(env):
    """รายการที่ผูกกับบันทึกขอเบิกจ่ายอยู่แล้ว (disburse_id) ห้ามนับซ้ำ"""
    from app.models import DisburseMemo, FinanceAccount, FinanceTxn, Project
    from app.services.budget import project_spent
    from app.tenancy import session_for
    _c, pj_id, tid = env
    db = session_for(tid)
    try:
        m = db.query(DisburseMemo).first()
        acct = db.query(FinanceAccount).first()
        db.add(FinanceTxn(account_id=acct.id, fiscal_year=PY, date=datetime(2026, 1, 6),
                          kind="out", amount=8000, note="จ่ายตามบันทึก 12/2569",
                          project_id=pj_id, disburse_id=m.id))
        db.commit()
        assert project_spent(db.get(Project, pj_id)) == 10500.0
    finally:
        db.close()


def test_page_shows_budget_spent_left(env):
    c, _pj, _tid = env
    html = c.get(f"/finance/projects?year={PY}").text
    assert "ทะเบียนคุมโครงการ" in html
    assert "ส่งเสริมการอ่านออกเขียนได้" in html
    assert "50,000.00" in html and "10,500.00" in html and "39,500.00" in html
    assert "ค่าป้ายไวนิล" in html          # เห็นรายการที่จ่ายตรงด้วย
    assert "/finance/projects/register.docx" in html


def test_register_docx_one_page_per_project(env, tmp_path):
    """ทะเบียนคุม: 1 โครงการ 1 หน้า · เดินยอดคงเหลือสะสมลงมา"""
    c, _pj, _tid = env
    r = c.get(f"/finance/projects/register.docx?year={PY}")
    assert r.status_code == 200
    f = tmp_path / "pj.docx"
    f.write_bytes(r.content)
    d = Document(f)
    flat = "\n".join(" | ".join(x.text.strip() for x in row.cells)
                     for t in d.tables for row in t.rows)
    assert "งบที่ได้รับอนุมัติ" in flat and "รวมใช้ไป" in flat
    assert "50,000.00" in flat
    assert "42,000.00" in flat        # 50,000 - 8,000
    assert "39,500.00" in flat        # - 2,500
    heads = [p.text.strip() for p in d.paragraphs if p.text.strip().startswith("โครงการ ")]
    assert len(heads) == 2, heads     # 2 โครงการ = 2 หน้า


def test_register_docx_single_project(env, tmp_path):
    c, pj_id, _tid = env
    r = c.get(f"/finance/projects/register.docx?year={PY}&project={pj_id}")
    assert r.status_code == 200
    f = tmp_path / "one.docx"
    f.write_bytes(r.content)
    d = Document(f)
    heads = [p.text.strip() for p in d.paragraphs if p.text.strip().startswith("โครงการ ")]
    assert len(heads) == 1 and "ส่งเสริมการอ่าน" in heads[0]
    assert c.get(f"/finance/projects/register.docx?year={PY}&project=999999").status_code == 404


def test_ledger_form_offers_project(env):
    """ฟอร์มลงรับ-จ่ายต้องเลือกโครงการได้ ไม่งั้นผูกเงินกับโครงการไม่ได้เลย"""
    from app.models import FinanceAccount
    from app.tenancy import session_for
    c, _pj, tid = env
    db = session_for(tid)
    aid = db.query(FinanceAccount).first().id
    db.close()
    html = c.get(f"/finance/accounts/{aid}?year={PY}").text
    assert 'name="project_id"' in html
    assert "ส่งเสริมการอ่านออกเขียนได้" in html


def test_sidebar_links_the_register(env):
    c, _pj, _tid = env
    assert "/finance/projects" in c.get("/finance").text
