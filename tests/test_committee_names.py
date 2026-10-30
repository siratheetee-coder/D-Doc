"""ชื่อคณะกรรมการต้องลงในเอกสาร: ตรวจสอบพัสดุประจำปี + สำนวนจำหน่าย (สอบหาข้อเท็จจริง)
เจ้าของแจ้ง 2026-10-30: เอกสารออกมามีแต่จุดไข่ปลา ไม่มีชื่อกรรมการที่กรอกไว้"""
import io

from tests.test_subsidy import env  # noqa: F401


def _db():
    import app.tenancy as tn
    return tn.session_for(next(iter(tn._engines)))


def _text(content: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join(parts)


def test_audit_docs_have_committee_names(env):
    from app.thai_utils import current_fiscal_year
    data = {"year": str(current_fiscal_year()), "date": "01/10/2569", "result_date": "15/10/2569",
            "m1_name": "นายประธาน ตรวจพัสดุ", "m1_pos": "ครู", "m1_role": "ประธานกรรมการ",
            "m2_name": "นางกรรมการ สองคน", "m2_pos": "ครู", "m2_role": "กรรมการ",
            "m3_name": "", "m3_pos": "ครู", "m3_role": "กรรมการและเลขานุการ"}
    r = env.post("/assets/audit", data=data)
    assert r.status_code == 200
    text = _text(r.content)
    assert "นายประธาน ตรวจพัสดุ" in text and "นางกรรมการ สองคน" in text
    # ชื่อที่กรอกครั้งก่อน จำไว้ให้ครั้งหน้า (ไม่ต้องกรอกใหม่ทุกปี)
    assert "นายประธาน ตรวจพัสดุ" in env.get("/assets/audit").text


def test_disposal_docs_use_names_typed_before_download(env):
    from app.models import Asset, AssetDisposal
    with _db() as s:
        a = Asset(asset_code="X-1", name="โต๊ะทดสอบกรรมการ", cost=1000, status="ชำรุด")
        s.add(a); s.commit()
        aid = a.id
    r = env.post("/assets/disposal", data={"asset_ids": str(aid), "year": "2570"}, follow_redirects=False)
    did = int(r.headers["location"].split("/")[-1].split("?")[0])
    # หน้าสำนวนมีปุ่มเอกสารที่บันทึกก่อนดาวน์โหลด
    page = env.get(f"/assets/disposal/{did}").text
    assert "data-save-first" in page
    env.post(f"/assets/disposal/{did}", data={"fact1_name": "นายสอบ ข้อเท็จจริง", "fact1_pos": "ครู",
                                               "fact1_role": "ประธานกรรมการ"})
    with _db() as s:
        assert "นายสอบ ข้อเท็จจริง" in (s.get(AssetDisposal, did).members or "")
    for kind in ("fact_memo", "fact_order"):
        text = _text(env.get(f"/assets/disposal/{did}/doc/{kind}.docx").content)
        assert "นายสอบ ข้อเท็จจริง" in text, kind
