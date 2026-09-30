# -*- coding: utf-8 -*-
"""รายงานขอซื้อ/ขอจ้าง ต้องจบหน้าเดียว และต้องเป็น 16 pt เสมอ

ที่มา
  ระเบียบงานสารบรรณกำหนดตัวอักษร 16 pt และบันทึกข้อความควรจบในหน้าเดียว
  แต่ความยาวจริงขึ้นกับข้อมูลของแต่ละโรงเรียน (ชื่อ/ที่อยู่ยาว · กรรมการ 3-5 คน)
  ก่อนแก้ พอมีกรรมการ 3 คน บล็อกลงนาม ผอ. หลุดไปหน้า 2 ทั้งบล็อก
  (ยืนยันด้วย Microsoft Word: ก่อนแก้ 3 หน้า · หลังแก้ 2 หน้า = บันทึก 1 + แนบท้าย 1)

  ที่นี่ไม่มี Word จึงวัดด้วยตัวประมาณใน app/services/page_fit.py
  ซึ่งสอบเทียบกับที่ Word วัดจริงไว้แล้ว (เผื่อสูงไว้ ~1.5%)
รัน: .venv\\Scripts\\python.exe -m pytest tests/test_purchase_request_one_page.py
"""
from datetime import datetime

import pytest
from docx import Document

from app.models import Committee, CommitteeMember, Procurement, ProcurementItem
from app.routers.pages import get_school
from app.services.page_fit import (SAFETY_PT, advance_len, block_height, first_block,
                                   page_budget)
from app.services.render import ONE_PAGE_KINDS, build_context, render_document
from tests.test_assignments import _db

MARK = "ทดสอบหน้าเดียว_"
NAMES = [("นายอมรพรรณ จรนามน", "ครู"), ("นายเนติพงษ์ มาตนาเรียง", "ครู"),
         ("นางสาวสุวิญญา พลชำนิ", "ครู"), ("นางกนกวรรณ ศรีสุวรรณ", "ครู"),
         ("นายประดิษฐ์ วัฒนโชติ", "ครู")]


def _make(db, *, kind="จ้าง", members=3, items=1):
    p = Procurement(
        fiscal_year=2569, proc_type=kind,
        subject=MARK + "รถทัศนศึกษาระดับมัธยมศึกษาปีที่ 1-3",
        project_name="ส่งเสริมผู้เรียนตามกิจกรรมพัฒนาผู้เรียน",
        purpose="จ้างทัศนศึกษาระดับมัธยมศึกษาปีที่ 1-3",
        department="ฝ่ายบริหารงานวิชาการ", budget_source="อุดหนุน",
        method="เฉพาะเจาะจง", memo_no="208/2569", total_amount=15000,
        request_date=datetime(2026, 9, 29), delivery_days=7,
        inspection_mode="committee" if members else "single")
    for i in range(items):
        p.items.append(ProcurementItem(name=f"รายการที่ {i + 1}", quantity=1,
                                       unit_price=15000 / max(1, items), unit="คัน"))
    if members:
        c = Committee(kind="inspect", mode="committee")
        for i, (nm, pos) in enumerate(NAMES[:members], 1):
            c.members.append(CommitteeMember(name=nm, position=pos, seq=i, role="กรรมการ"))
        p.committees.append(c)
    db.add(p)
    db.commit()
    db.refresh(p)
    return p


def _render(db, **kw):
    p = _make(db, **kw)
    try:
        return Document(render_document("รายงานขอซื้อ", p, get_school(db)))
    finally:
        from app.models import Document as Doc
        db.query(Doc).filter_by(procurement_id=p.id).delete(synchronize_session=False)
        db.delete(p)
        db.commit()


def test_advance_len_skips_thai_marks():
    """สระบน/สระล่าง/วรรณยุกต์ ซ้อนอยู่บนพยัญชนะ ไม่กินความกว้างบรรทัด"""
    assert advance_len("กิ") == 1 and advance_len("กี่") == 1
    assert advance_len("ที่") == 1          # ท + สระอี + ไม้เอก ซ้อนกันอยู่ที่เดียว
    assert advance_len("กา") == 2 and advance_len("") == 0
    assert advance_len("ทำงาน") == 5        # ำ กินความกว้าง (ไม่ใช่สระซ้อน)


@pytest.mark.parametrize("members", [0, 1, 2, 3, 4, 5])
def test_memo_fits_one_page(members):
    """บันทึกข้อความจบหน้าเดียว ตั้งแต่ผู้ตรวจรับคนเดียวจนถึงกรรมการ 5 คน"""
    db = _db()
    d = _render(db, members=members)
    # page_budget กันชนไว้ ~0.8 บรรทัดสำหรับ "ตอนตัดสินใจบีบ"
    # ส่วนการตรวจว่าจบหน้าเดียวจริง เทียบกับพื้นที่พิมพ์เต็ม
    h, printable = block_height(d), page_budget(d) + SAFETY_PT
    assert h <= printable, f"กรรมการ {members} คน สูง {h:.0f}pt เกินพื้นที่ {printable:.0f}pt"


def test_font_stays_16pt_after_fitting():
    """บีบได้แต่ห้ามลดขนาดตัวอักษร (ระเบียบงานสารบรรณกำหนด 16 pt)"""
    db = _db()
    d = _render(db, members=5)
    body = {r.font.size.pt for b in first_block(d) if hasattr(b, "runs")
            for r in b.runs if r.font.size}
    assert 16.0 in body, body
    assert not [s for s in body if s < 16.0], f"มีตัวอักษรเล็กกว่า 16 pt: {sorted(body)}"


def test_short_document_is_left_alone():
    """เอกสารสั้น ๆ ที่พอดีอยู่แล้ว ต้องไม่ถูกบีบเลย (หน้าตาเหมือนเดิมเป๊ะ)"""
    from docx import Document as NewDoc

    from app.services.doc_page import set_a4
    from app.services.page_fit import fit_one_page
    d = NewDoc()
    set_a4(d)
    for i in range(5):
        d.add_paragraph(f"บรรทัดที่ {i + 1}")
    out = fit_one_page(d)
    assert out["fitted"] and out["spacing"] is None and out["gap_scale"] == 1.0
    assert all(p.paragraph_format.line_spacing is None for p in d.paragraphs)


def test_krut_paragraph_never_squeezed():
    """ย่อหน้าตราครุฑ/หัวเรื่องตัวใหญ่ ห้ามบีบระยะบรรทัด (รูปจะถูกตัดขอบ)"""
    db = _db()
    d = _render(db, members=5)
    from docx.oxml.ns import qn
    for p in d.paragraphs:
        big = any(r.font.size and r.font.size.pt > 18 for r in p.runs)
        if p._p.findall(".//" + qn("w:drawing")) or big:
            assert p.paragraph_format.line_spacing in (None, 1.0), p.text[:30]


@pytest.mark.parametrize("kind,want", [("ซื้อ", "สำหรับจัดซื้อพัสดุ"), ("จ้าง", "สำหรับงานจัดจ้าง")])
def test_attachment_heading_wording(kind, want):
    """หัวรายละเอียดแนบท้าย: งานจ้างไม่ใช่ 'จัดจ้างพัสดุ' เพราะไม่ได้ซื้อตัวพัสดุ"""
    db = _db()
    d = _render(db, kind=kind, members=3)
    heads = [p.text.strip() for p in d.paragraphs if p.text.strip().startswith("สำหรับ")]
    assert heads and heads[0].startswith(want), heads
    assert "จัดจ้างพัสดุ" not in "\n".join(heads)


def test_context_has_attach_for():
    db = _db()
    p = _make(db, kind="จ้าง", members=0)
    try:
        assert build_context(p, get_school(db))["attach_for"] == "งานจัดจ้าง"
    finally:
        db.delete(p)
        db.commit()


def test_purchase_request_is_registered_as_one_page_doc():
    assert "รายงานขอซื้อ" in ONE_PAGE_KINDS


def _sign_gaps(d):
    """ช่องว่างเหนือทุกบรรทัด 'ลงชื่อ' (pt) ทั้งย่อหน้าปกติและในตารางลงนาม"""
    out = []
    for p in d.paragraphs:
        if p.text.strip().lstrip("(").startswith("ลงชื่อ"):
            sb = p.paragraph_format.space_before
            out.append(sb.pt if sb else 0.0)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    if p.text.strip().lstrip("(").startswith("ลงชื่อ"):
                        sb = p.paragraph_format.space_before
                        out.append(sb.pt if sb else 0.0)
    return out


@pytest.mark.parametrize("members", [0, 2, 3])
def test_room_to_sign_above_every_signature(members):
    """ต้องมีที่ว่างเหนือบรรทัด 'ลงชื่อ' ให้เซ็นจริงได้ ทั้งเจ้าหน้าที่/หัวหน้า/ผอ."""
    db = _db()
    d = _render(db, members=members)
    gaps = _sign_gaps(d)
    assert len(gaps) >= 3, f"หาบรรทัดลงชื่อไม่ครบ: {gaps}"
    assert all(g > 0 for g in gaps), f"ไม่มีที่เซ็น: {gaps}"


def test_full_sign_gap_when_there_is_room():
    """เอกสารที่ไม่แน่น ต้องได้ช่องเซ็นเต็มขนาดที่ตั้งไว้"""
    from app.services.build_templates import SIGN_GAP
    db = _db()
    d = _render(db, members=0)
    assert max(_sign_gaps(d)) == SIGN_GAP


def test_sign_gap_is_given_up_last():
    """บีบระยะบรรทัดจนสุดก่อน ค่อยยอมลดช่องเซ็น (ที่เซ็นสำคัญกว่าความโปร่ง)"""
    from app.services.page_fit import _STEPS
    first_gap_cut = next(i for i, (gap, _) in enumerate(_STEPS) if gap < 1.0)
    tightest_before = min(sp for gap, sp in _STEPS[:first_gap_cut] if sp)
    assert tightest_before == min(sp for _, sp in _STEPS if sp), \
        "ต้องลดระยะบรรทัดจนสุดก่อน จึงค่อยลดช่องเซ็น"


def test_no_date_line_under_director():
    """ใต้ชื่อ ผอ. ไม่มีบรรทัด 'วันที่' (วันที่อยู่หัวบันทึกแล้ว) - บรรทัดนั้นเอาไปเป็นที่เซ็น"""
    db = _db()
    d = _render(db, members=3)
    texts = [p.text.strip() for p in first_block(d) if hasattr(p, "text")]
    tail = texts[-4:]
    assert not [t for t in tail if t.startswith("วันที่")], tail


# ---------------- ใบอนุมัติเบิกจ่าย (รายงานผลตรวจรับและเบิกจ่าย) ----------------
def _render_kind(db, kind, **kw):
    from app.models import Vendor
    v = db.query(Vendor).filter_by(name=MARK + "ร้านทดสอบ").first()
    if v is None:
        v = Vendor(name=MARK + "ร้านทดสอบ", owner_name="นายทดสอบ ใจดี")
        db.add(v)
        db.flush()
    p = _make(db, **kw)
    p.vendor_id = v.id
    p.inspect_memo_no = "209/2569"
    p.order_no = "จ22/2569"
    p.order_date = datetime(2026, 9, 29)
    p.inspect_date = datetime(2026, 10, 6)
    db.commit()
    try:
        return Document(render_document(kind, p, get_school(db)))
    finally:
        from app.models import Document as Doc
        db.query(Doc).filter_by(procurement_id=p.id).delete(synchronize_session=False)
        db.delete(p)
        db.query(Vendor).filter_by(id=v.id).delete(synchronize_session=False)
        db.commit()


def test_disbursement_body_is_16pt():
    """ใบอนุมัติเบิกจ่ายเคยย่อเหลือ 14 pt เพื่อให้จบหน้าเดียว
    ตอนนี้คง 16 pt ตามระเบียบงานสารบรรณ แล้วให้ page_fit จัดการเรื่องหน้าแทน"""
    db = _db()
    d = _render_kind(db, "รายงานผลตรวจรับและเบิกจ่าย", members=0)
    sizes = {r.font.size.pt for b in first_block(d) if hasattr(b, "runs")
             for r in b.runs if r.font.size}
    assert 16.0 in sizes, sizes
    assert not [s for s in sizes if s < 16.0], f"ยังมีตัวอักษรเล็กกว่า 16 pt: {sorted(sizes)}"


def test_disbursement_fits_one_page():
    db = _db()
    d = _render_kind(db, "รายงานผลตรวจรับและเบิกจ่าย", members=0)
    h, printable = block_height(d), page_budget(d) + SAFETY_PT
    assert h <= printable, f"สูง {h:.0f}pt เกินพื้นที่ {printable:.0f}pt"


def test_disbursement_has_room_to_sign():
    """4 ช่องลงนาม (เจ้าหน้าที่ · หัวหน้าเจ้าหน้าที่ · เจ้าหน้าที่การเงิน · ผอ.)"""
    db = _db()
    d = _render_kind(db, "รายงานผลตรวจรับและเบิกจ่าย", members=0)
    gaps = _sign_gaps(d)
    assert len(gaps) >= 4, f"หาบรรทัดลงชื่อไม่ครบ: {gaps}"
    assert all(g > 0 for g in gaps), f"ไม่มีที่เซ็น: {gaps}"


def test_disbursement_no_longer_shrinks_font():
    """กันเผลอใส่ _shrink_body_font กลับเข้าไปในแม่แบบเบิกจ่ายอีก"""
    import inspect

    from app.services import build_templates as bt
    assert "_shrink_body_font" not in inspect.getsource(bt.build_disbursement)
    assert "รายงานผลตรวจรับและเบิกจ่าย" in ONE_PAGE_KINDS
