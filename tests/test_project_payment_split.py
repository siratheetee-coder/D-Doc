from tests.test_subsidy import db, env
from app.models import Project, Procurement, DisburseMemo, FinanceAccount, FinanceTxn, LunchProgram, LunchLedger
from app.services.budget import project_money_summary, project_spent
from app.services.project_summary import build_rows


def test_partial_payment_does_not_release_reserved_budget(db):
    p=Project(name='project',plan_year=2570,budget=10000,spent_before=1000)
    a=FinanceAccount(name='account');db.add_all([p,a]);db.flush()
    proc=Procurement(subject='purchase',project_id=p.id,fiscal_year=2570,total_amount=4000)
    db.add(proc);db.flush()
    m=DisburseMemo(project_id=p.id,procurement_id=proc.id,fiscal_year=2570,amount=1500)
    db.add(m);db.flush()
    first=project_money_summary(p)
    assert (first['paid'],first['waiting'],first['remaining'])==(0,4000,5000)
    t=FinanceTxn(account_id=a.id,fiscal_year=2570,kind='out',amount=1500,disburse_id=m.id)
    db.add(t);db.flush()
    now=project_money_summary(p)
    assert (now['paid'],now['waiting'],now['remaining'])==(1500,2500,5000)
    assert project_spent(p)==5000
    row=build_rows(db,[p])[0]
    assert row['paid']==1500 and row['waiting']==2500
    assert sum(w['amount'] for w in row['works'])==row['spent']


def test_overpaid_source_cannot_reduce_another_sources_pending(db):
    p=Project(name='project',budget=10000,plan_year=2570);a=FinanceAccount(name='a');db.add_all([p,a]);db.flush()
    procs=[Procurement(subject=str(i),project_id=p.id,fiscal_year=2570,total_amount=1000) for i in range(2)]
    db.add_all(procs);db.flush()
    m=DisburseMemo(procurement_id=procs[0].id,fiscal_year=2570,amount=1500);db.add(m);db.flush()
    db.add(FinanceTxn(account_id=a.id,fiscal_year=2570,kind='out',amount=1500,disburse_id=m.id));db.flush()
    money=project_money_summary(p)
    assert money['paid']==1500 and money['waiting']==1000 and money['used']==2500


def test_lunch_mirror_and_memo_count_once(db):
    p=Project(name='project',budget=5000,plan_year=2570);a=FinanceAccount(name='a');prog=LunchProgram(year=2570)
    db.add_all([p,a,prog]);db.flush()
    proc=Procurement(subject='lunch',project_id=p.id,fiscal_year=2570,total_amount=1000);db.add(proc);db.flush()
    t=FinanceTxn(account_id=a.id,fiscal_year=2570,kind='out',amount=1000);db.add(t);db.flush()
    db.add(LunchLedger(program_id=prog.id,procurement_id=proc.id,finance_txn_id=t.id,amount=1000,kind='out'));db.flush()
    assert project_money_summary(p)['paid']==1000
    m=DisburseMemo(procurement_id=proc.id,fiscal_year=2570,amount=1000);db.add(m);db.flush()
    t.disburse_id=m.id;db.flush()
    info=project_money_summary(p)
    assert info['paid']==1000 and info['waiting']==0


def test_pages_and_exports_show_separate_values(env):
    import app.tenancy as tn
    import io
    from docx import Document
    from openpyxl import load_workbook
    sid=next(iter(tn._engines))
    with tn.session_for(sid) as db:
        p=Project(name='split-test',plan_year=2570,budget=10000,spent_before=1000)
        db.add(p);db.flush()
        db.add(Procurement(subject='pending',project_id=p.id,fiscal_year=2570,total_amount=4000))
        db.commit();pid=p.id
    for url in [f'/projects/{pid}','/projects?year=2570','/finance/projects?year=2570']:
        res=env.get(url)
        assert res.status_code==200 and 'split-test' in res.text
        assert 'จ่ายจริง' in res.text and ('รอจ่าย' in res.text)
    res=env.get('/finance/projects/register.docx?year=2570')
    assert res.status_code==200
    doc=Document(io.BytesIO(res.content))
    assert 'ทำเรื่องแล้วรอจ่าย' in doc.element.xml
    with tn.session_for(sid) as db:
        from app.services.project_summary import export_project_summary,render_project_summary
        from app.models import School
        rows=build_rows(db,[db.get(Project,pid)])
        book=load_workbook(export_project_summary(db.query(School).first(),2570,'ปีงบประมาณ',rows))
        assert list(book['จ่ายจริงและรอจ่าย'].values)[1][1:]==(1000,0,4000,5000)
        doc=Document(render_project_summary(db.query(School).first(),2570,'ปีงบประมาณ',rows,detail=False))
        assert 'จ่ายจริงในระบบ' in doc.element.xml
