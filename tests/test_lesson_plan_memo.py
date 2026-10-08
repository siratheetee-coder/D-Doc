# -*- coding: utf-8 -*-
"""บันทึกข้อความขออนุมัติใช้แผนการจัดการเรียนรู้

ไฟล์แผนเป็นรูปแบบของแต่ละโรงเรียน ระบบแก้ไม่ได้ ลายเซ็นผู้บริหารจึงต้องไปอยู่
บนบันทึกนำหน้าที่ระบบออกเอง เทสต์นี้คุมสามเรื่อง
  1) แปลงรายการหน่วยที่ครูพิมพ์อิสระให้เป็นตารางได้ถูกต้อง
  2) เอกสารที่ออกมามีเนื้อความราชการ ตารางหน่วย และช่องลงนามครบสามฝ่าย
  3) ลายเซ็นลอย "อยู่หน้าข้อความ" (wrapNone) จุดไข่ปลาจึงไม่เลื่อน
รัน: .venv\\Scripts\\python.exe -m pytest tests/test_lesson_plan_memo.py
"""
import pathlib
from datetime import datetime

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.services.lesson_plan_doc import parse_units, total_hours, render_plan_memo

ROOT = pathlib.Path(__file__).resolve().parents[1]


class O:
    def __init__(self, **kw):
        self.__dict__.update(kw)


def _school():
    return O(name="บ้านหินลาด", address="ตำบลหินลาด อำเภอเมือง จังหวัดขอนแก่น",
             director_name="นายสมชาย ใจดี", director_position="ผู้อำนวยการโรงเรียน",
             academic_head_name="นางสาวมาลี เรียนเก่ง")


def _plan(**over):
    base = dict(id=1, title="รายวิชาคณิตศาสตร์ ค14101 ชั้นประถมศึกษาปีที่ 4",
                term=1, year=2569, note="", status="approved",
                submitted_at=datetime(2026, 5, 12), reviewed_at=datetime(2026, 5, 14),
                director_at=datetime(2026, 5, 15), comment="", director_comment="",
                teacher=O(name="นางสาวสุดา ขยันสอน", position="ครู"),
                units="หน่วยที่ 1 จำนวนนับ (12 ชั่วโมง)\nหน่วยที่ 2 การบวก การลบ (15 ชั่วโมง)")
    base.update(over)
    return O(**base)


def _text(path):
    doc = Document(path)
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            out += [c.text for c in row.cells]
    return "\n".join(out)


# ---------------- 1) อ่านรายการหน่วยที่ครูพิมพ์เอง ----------------
@pytest.mark.parametrize("line, name, hours", [
    ("หน่วยที่ 1 จำนวนนับ (12 ชั่วโมง)", "จำนวนนับ", 12),
    ("2. การบวก การลบจำนวนนับ 15 ชม.", "การบวก การลบจำนวนนับ", 15),
    ("หน่วย 3 เรขาคณิต เวลา 10 ชั่วโมง", "เรขาคณิต", 10),
    ("4) การวัดความยาว (8 ชม)", "การวัดความยาว", 8),
    ("สถิติเบื้องต้น", "สถิติเบื้องต้น", None),
])
def test_unit_line_formats_teachers_actually_type(line, name, hours):
    u = parse_units(line)[0]
    assert u["name"] == name
    assert u["hours"] == (float(hours) if hours is not None else None)


def test_numbering_falls_back_to_line_order_and_blank_lines_are_skipped():
    us = parse_units("จำนวนนับ (2 ชั่วโมง)\n\n   \nเรขาคณิต")
    assert [u["no"] for u in us] == [1, 2]
    assert total_hours(us) == 2


def test_empty_units_is_not_an_error():
    assert parse_units("") == [] and parse_units(None) == []


# ---------------- 2) เนื้อเอกสาร ----------------
def test_memo_has_official_header_body_units_and_three_signature_blocks(tmp_path):
    path = render_plan_memo(_plan(), _school())
    t = _text(path)
    # หัวบันทึกข้อความตามระเบียบสารบรรณ
    for want in ("ส่วนราชการ", "เรื่อง", "เรียน", "ผู้อำนวยการโรงเรียนบ้านหินลาด"):
        assert want in t, want
    # เนื้อความราชการ + ข้อเสนอ
    assert "หลักสูตรแกนกลางการศึกษาขั้นพื้นฐาน พุทธศักราช 2551" in t
    assert "จึงเรียนมาเพื่อโปรดพิจารณาอนุมัติ" in t
    # ตารางหน่วย + เวลารวม
    assert "ชื่อหน่วยการเรียนรู้" in t and "จำนวนนับ" in t
    assert "รวมเวลาเรียน" in t and "27" in t
    # ช่องลงนามสามฝ่าย
    assert "นางสาวสุดา ขยันสอน" in t
    assert "ความเห็นหัวหน้ากลุ่มบริหารงานวิชาการ" in t and "นางสาวมาลี เรียนเก่ง" in t
    assert "คำสั่งผู้อำนวยการ" in t and "นายสมชาย ใจดี" in t
    assert "อนุมัติให้ใช้แผนการจัดการเรียนรู้ตามเสนอ" in t


def test_unapproved_plan_leaves_the_decision_blank_for_hand_signing():
    """ยังไม่ถึงคิวลงนาม ต้องไม่พิมพ์ชื่อผู้ลงนามหรือคำว่าอนุมัติไว้ล่วงหน้า"""
    t = _text(render_plan_memo(_plan(id=2, status="pending", reviewed_at=None, director_at=None),
                               _school()))
    assert "อนุมัติให้ใช้แผนการจัดการเรียนรู้ตามเสนอ" not in t
    assert "อนุมัติ / ไม่อนุมัติ" in t
    assert "นายสมชาย ใจดี" not in t, "ยังไม่ได้อนุมัติ ห้ามพิมพ์ชื่อ ผอ. ลงในช่องลงนาม"
    assert "นางสาวมาลี เรียนเก่ง" not in t


def test_comments_from_both_reviewers_appear():
    t = _text(render_plan_memo(_plan(id=3, comment="ปรับสื่อหน่วยที่ 2", director_comment="เห็นชอบ"),
                               _school()))
    assert "ปรับสื่อหน่วยที่ 2" in t and "เห็นชอบ" in t


def test_plan_without_units_still_produces_a_usable_memo():
    t = _text(render_plan_memo(_plan(id=4, units=""), _school()))
    assert "จึงเรียนมาเพื่อโปรดพิจารณาอนุมัติ" in t
    assert "ชื่อหน่วยการเรียนรู้" not in t
    assert "หน่วยการเรียนรู้ ดังนี้" not in t


def test_non_school_agency_is_not_called_a_school():
    """หน่วยงานอื่นที่ใช้ระบบ ต้องไม่โดนเติมคำว่าโรงเรียนให้เอง (ใช้กติกากลาง org_names)"""
    sc = _school()
    sc.name = "เทศบาลตำบลหินลาด"
    sc.director_position = "นายกเทศมนตรีตำบลหินลาด"
    t = _text(render_plan_memo(_plan(id=5), sc))
    assert "โรงเรียนเทศบาล" not in t
    assert "นายกเทศมนตรีตำบลหินลาด" in t


# ---------------- 3) ลายเซ็นต้องลอยหน้าข้อความ ----------------
def test_signature_is_anchored_in_front_of_text(monkeypatch, tmp_path):
    """ถ้าแปะเป็นรูปแบบ inline เส้นไข่ปลาจะถูกดันจนเละ ต้องเป็น anchor + wrapNone"""
    from PIL import Image
    png = tmp_path / "sig.png"
    Image.new("RGBA", (300, 100), (0, 0, 0, 255)).save(png)
    import app.services.signature as S
    monkeypatch.setattr(S, "signature_path_for_current", lambda name: str(png))

    doc = Document(render_plan_memo(_plan(id=6), _school()))
    anchors = doc.element.body.findall(".//" + qn("wp:anchor"))
    inlines = doc.element.body.findall(".//" + qn("wp:inline"))
    assert len(anchors) == 3, f"ต้องมีลายเซ็นลอย 3 จุด (ครู/วิชาการ/ผอ.) ได้ {len(anchors)}"
    # ครุฑบนหัวกระดาษยังเป็น inline ได้ แต่ลายเซ็นต้องไม่เหลือ inline เกินนั้น
    assert len(inlines) <= 1, "ลายเซ็นยังเป็น inline อยู่ จุดไข่ปลาจะเลื่อน"
    for a in anchors:
        assert a.find(qn("wp:wrapNone")) is not None, "ลายเซ็นต้องเป็นแบบ wrapNone (หน้าข้อความ)"
        assert a.get("behindDoc") == "0", "ต้องอยู่หน้าข้อความ ไม่ใช่หลังข้อความ"


# ---------------- การต่อสาย ----------------
def test_route_and_form_are_wired():
    acad = (ROOT / "app/routers/academic.py").read_text(encoding="utf-8")
    assert '"/academic/lesson-plans/{plan_id}/memo"' in acad, "ยังไม่มีเส้นทางออกบันทึก"
    assert "units=(units or \"\").strip()" in acad, "ตอนส่งแผนยังไม่เก็บรายการหน่วย"
    html = (ROOT / "app/templates/academic_lesson_plans.html").read_text(encoding="utf-8")
    assert 'name="units"' in html, "ฟอร์มส่งแผนยังไม่มีช่องหน่วยการเรียนรู้"
    detail = (ROOT / "app/templates/lesson_plan_detail.html").read_text(encoding="utf-8")
    assert "/memo" in detail, "หน้ารายละเอียดยังไม่มีปุ่มดาวน์โหลดบันทึก"


def test_units_column_is_migrated_for_existing_schools():
    db = (ROOT / "app/database.py").read_text(encoding="utf-8")
    assert '("lesson_plan", "units"' in db, "โรงเรียนเดิมจะพังเพราะไม่มีคอลัมน์ units"


# ---------------- เดินผ่าน HTTP จริง ----------------
def _login_client(person_id=None):
    from fastapi.testclient import TestClient
    import app.routers.auth as auth_mod
    import app.main as main_mod
    from app.main import app
    auth_mod.authenticate = lambda u, p: {
        "uid": 1, "username": "tester", "role": "owner", "tenant_id": 1,
        "display_name": "ผู้ทดสอบ", "must_change": False, "person_id": person_id,
    }
    main_mod.can_use_module = lambda tid, mod: True
    main_mod.get_account_access = lambda uid: {
        "is_owner": True, "modules": "", "active": True, "welcomed": True,
        "person_id": person_id}
    c = TestClient(app)
    r = c.post("/login", data={"username": "tester", "password": "x"}, follow_redirects=False)
    assert r.status_code in (302, 303), f"login ไม่ผ่าน: {r.status_code}"
    return c


def test_submit_keeps_units_and_memo_downloads_over_http():
    """เทสต์ระดับฟังก์ชันผ่านได้ทั้งที่เส้นทางจริงพังก็มี จึงต้องยิง HTTP ด้วย"""
    from app.tenancy import session_for
    from app.models import LessonPlan
    from app.models import Person
    mark = "ทดสอบบันทึกแผน_"
    db = session_for(1)
    teacher = Person(name=mark + "ครูทดสอบ", position="ครู")
    db.add(teacher); db.commit()
    c = _login_client(person_id=teacher.id)

    def _clean(people=False):
        for p in db.query(LessonPlan).filter(LessonPlan.title.like(mark + "%")).all():
            db.delete(p)
        if people:
            for q in db.query(Person).filter(Person.name.like(mark + "%")).all():
                db.delete(q)
        db.commit()

    _clean()
    try:
        r = c.post("/academic/lesson-plans/submit",
                   data={"title": mark + "คณิตศาสตร์ ป.4", "term": "1",
                         "units": "หน่วยที่ 1 จำนวนนับ (12 ชั่วโมง)\nหน่วยที่ 2 การวัด (8 ชั่วโมง)",
                         "note": ""},
                   files={"file": ("plan.pdf", b"%PDF-1.4\ntrailer<</Root 1 0 R>>\n%%EOF\n",
                                   "application/pdf")},
                   follow_redirects=False)
        assert r.status_code in (302, 303), r.status_code
        assert "err=" not in r.headers.get("location", ""), r.headers.get("location")
        p = db.query(LessonPlan).filter(LessonPlan.title.like(mark + "%")).first()
        assert p is not None, "ส่งแผนไม่สำเร็จ"
        assert "จำนวนนับ" in (p.units or ""), "ช่องหน่วยการเรียนรู้ไม่ถูกบันทึก"

        r = c.get(f"/academic/lesson-plans/{p.id}/memo")
        assert r.status_code == 200, r.status_code
        assert r.content[:2] == b"PK", "ไฟล์ที่ได้ไม่ใช่ .docx"
        assert len(r.content) > 5000, "ไฟล์เล็กผิดปกติ"
    finally:
        _clean(people=True)
        db.close()
