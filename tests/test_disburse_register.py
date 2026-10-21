# -*- coding: utf-8 -*-
"""ทะเบียนคุมหลักฐานขอเบิก: คอลัมน์เลขที่บันทึก + ชื่อผู้รับเงิน + หมายเหตุ (กรอกได้ แยกจากรายละเอียดเพิ่มเติม)
รัน: .venv\\Scripts\\python.exe -m pytest tests/test_disburse_register.py
"""
import io

from docx import Document

from tests.test_subsidy import env  # noqa: F401  (fixture)

FY = 2570


def _memo(note):
    import app.tenancy as tn
    from app.models import DisburseMemo
    db = tn.session_for(next(iter(tn._engines)))
    try:
        m = db.query(DisburseMemo).filter_by(note=note).one()
        return m.id, m.receiver_name, m.register_note, m.memo_no
    finally:
        db.close()


def test_fields_saved_and_register_columns(env):
    form = env.get(f"/finance/disburse?year={FY}").text
    assert 'name="receiver_name"' in form and 'name="register_note"' in form
    assert "จ่ายให้ (เจ้าหนี้/ผู้ขาย)" in form and "จ่ายให้ (ผู้รับเงิน)" not in form
    env.post("/finance/disburse", data={"fiscal_year": FY, "memo_no": "77/2570", "date": "15/11/2569",
             "subject": "ซื้อวัสดุ", "payee": "ร้านศึกษาภัณฑ์", "amount": "1250",
             "budget_source": "ค่าวัสดุการเรียนการสอน", "note": "รายละเอียดสำหรับบันทึกขออนุมัติ",
             "receiver_name": "  นางสาวสมหญิง ใจดี ", "register_note": "จ่ายเช็คเลขที่ 0012345"})
    mid, recv, rnote, no = _memo("รายละเอียดสำหรับบันทึกขออนุมัติ")
    assert (recv, rnote, no) == ("นางสาวสมหญิง ใจดี", "จ่ายเช็คเลขที่ 0012345", "77/2570")
    detail = env.get(f"/finance/disburse/{mid}").text
    assert 'value="นางสาวสมหญิง ใจดี"' in detail and 'value="จ่ายเช็คเลขที่ 0012345"' in detail

    r = env.get(f"/finance/disburse-register.docx?year={FY}")
    assert r.status_code == 200
    t = Document(io.BytesIO(r.content)).tables[0]
    head = [c.text.replace("\n", " ") for c in t.rows[0].cells]
    assert head == ["วัน เดือน ปี", "เลขที่บันทึก", "เจ้าหนี้หรือผู้ขอเบิก", "ประเภทรายจ่าย", "จำนวนเงิน",
                    "ชื่อผู้รับเงิน", "ลายมือชื่อ ผู้รับหลักฐาน", "วัน เดือน ปี ที่ส่ง ส่วนราชการผู้เบิก", "หมายเหตุ"]
    row = [c.text for c in t.rows[1].cells]
    assert row == ["15 พ.ย. 2569", "77/2570", "ร้านศึกษาภัณฑ์", "ค่าวัสดุการเรียนการสอน", "1,250.00",
                   "นางสาวสมหญิง ใจดี", "", "", "จ่ายเช็คเลขที่ 0012345"]
    assert "รายละเอียดสำหรับบันทึกขออนุมัติ" not in [c.text for r_ in t.rows for c in r_.cells], \
        "รายละเอียดเพิ่มเติมไม่ปนเข้าช่องหมายเหตุ"
    total = t.rows[-1].cells
    assert total[3].text == "รวม" and total[4].text == "1,250.00"
    widths = sum(c.width for c in t.rows[0].cells)
    assert abs(widths / 360000 - 26.7) < 0.05, "รวมความกว้างพอดีกระดาษ A4 แนวนอน"

    # แก้ไขภายหลังได้ · เว้นว่าง = ช่องว่างไว้เขียนมือ
    env.post(f"/finance/disburse/{mid}/update", data={"memo_no": "77/2570", "date": "15/11/2569", "subject": "ซื้อวัสดุ",
             "payee": "ร้านศึกษาภัณฑ์", "amount": "1250", "note": "รายละเอียดสำหรับบันทึกขออนุมัติ",
             "receiver_name": "", "register_note": "โอนเข้าบัญชี"})
    assert _memo("รายละเอียดสำหรับบันทึกขออนุมัติ")[1:3] == ("", "โอนเข้าบัญชี")
    t = Document(io.BytesIO(env.get(f"/finance/disburse-register.docx?year={FY}").content)).tables[0]
    row = [c.text for c in t.rows[1].cells]
    assert row[5] == "" and row[8] == "โอนเข้าบัญชี"


def test_name_suggestions_and_project_combo(env):
    """ช่องจ่ายให้/ชื่อผู้รับเงิน มีรายชื่อแนะนำจากผู้ขายในระบบ + ชื่อที่เคยใช้ · ช่องโครงการพิมพ์ค้นหาได้"""
    import json
    import re
    import app.tenancy as tn
    from app.models import Vendor
    db = tn.session_for(next(iter(tn._engines)))
    try:
        db.add(Vendor(name="ร้านรุ่งเรืองพาณิชย์", owner_name="นายสมชาย รุ่งเรือง")); db.commit()
    finally:
        db.close()
    env.post("/finance/disburse", data={"fiscal_year": FY, "subject": "ซื้อวัสดุ", "payee": "ร้านศึกษาภัณฑ์",
             "amount": "100", "receiver_name": "นางสาวสมหญิง ใจดี", "note": "x1"})
    for url in (f"/finance/disburse?year={FY}", f"/finance/disburse/{_memo('x1')[0]}"):
        html = env.get(url).text
        assert 'name="payee" autocomplete="off" data-suggest="payee"' in html
        assert 'data-suggest="receiver"' in html and 'id="projSel" data-combo' in html
        assert "/static/combo-pick.js" in html
        names = json.loads(re.search(r"window.SUGGEST_NAMES = (\{.*?\});</script>", html).group(1))
        assert {"ร้านรุ่งเรืองพาณิชย์", "ร้านศึกษาภัณฑ์"} <= set(names["payee"])
        assert {"นายสมชาย รุ่งเรือง", "นางสาวสมหญิง ใจดี"} <= set(names["receiver"])
        assert "" not in names["payee"] and len(names["payee"]) == len(set(names["payee"]))
