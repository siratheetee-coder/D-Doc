"""Period attendance regressions; all data is isolated from school databases."""
from datetime import date
from io import BytesIO
import re
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from app.database import Base
from app.models import (AcadClass, AcadSubject, AcadStudent, AcadCalendar, AcadPeriod,
                        AcadTimetable, AcadLesson, AcadLessonMark, AcadAttendance,
                        AcadClassMonth, School)
from app.services import academic_lessons as svc
from tests.test_subsidy import env


def seed(db):
    c = AcadClass(year=2569, level="ป.1", room="99")
    subject = AcadSubject(year=2569, level="ป.1", name="ภาษาไทยทดสอบ", code="ท11101", hours=80, term=0)
    db.add_all([c, subject]); db.flush()
    student = AcadStudent(class_id=c.id, name="นักเรียนทดสอบ", seq=1)
    db.add(student); db.flush()
    db.add(AcadAttendance(acad_student_id=student.id, subject_id=subject.id, month=5, present=1, marks="/"))
    db.add(AcadClassMonth(class_id=c.id, month=5, days_open=20))
    db.commit()
    return c, subject, student


@pytest.fixture
def db():
    engine=create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        yield db
    engine.dispose()


def test_periods_same_day_duration_blank_and_legacy_are_separate(db):
    c, subject, student=seed(db)
    assert svc.add_lesson(db,c,subject,1,date(2026,5,18),"คาบ 1",50)
    assert svc.add_lesson(db,c,subject,1,date(2026,5,18),"คาบ 2",40)
    assert not svc.add_lesson(db,c,subject,1,date(2026,5,18),"คาบ 1",50)
    rows=svc.lessons_for(db,c.id,subject.id)
    db.add(AcadLessonMark(lesson_id=rows[0].id,acad_student_id=student.id,mark="/"));db.commit()
    state=svc.snapshot(db,c,subject,1)
    assert state['totals']=={1:90,2:0}
    assert state['checked']==50 and state['target'] is None
    assert state['summary'][student.id]=={'/':50,'ป':0,'ล':0,'ข':0,'':40}
    assert db.query(AcadAttendance).one().present==1
    assert db.query(AcadClassMonth).one().days_open==20


def test_annual_targets_remain_optional_and_can_be_unequal(db):
    c, subject, _=seed(db)
    assert svc.snapshot(db,c,subject,1)['target'] is None
    plan=svc.plan_for(db,c.id,subject.id,True)
    plan.term1_minutes=svc.target_minutes('38')
    plan.term2_minutes=svc.target_minutes('42')
    db.commit()
    assert svc.snapshot(db,c,subject,2)['target']==2520
    assert subject.hours==80
    assert svc.target_minutes('') is None
    subject.term=2; subject.hours=30
    assert svc.snapshot(db,c,subject,2)['target']==1800
    with pytest.raises(ValueError):svc.check_term(subject,1)


@pytest.mark.parametrize('value',['-1','nan','inf','x','0.00001'])
def test_invalid_targets(value):
    with pytest.raises(ValueError):svc.target_minutes(value)


def test_import_calendar_periods_preview_and_snapshot(db):
    c,subject,_=seed(db)
    # Monday 18 is open, Monday 25 is closed. Two distinct periods on 18.
    db.add(AcadCalendar(year=2569,month=5,days_csv="18,19"))
    p1=AcadPeriod(year=2569,seq=1,name="คาบ 1",time_label="08:30-09:20",is_break=False)
    p2=AcadPeriod(year=2569,seq=2,name="คาบ 2",time_label="09:20-10:00",is_break=False)
    db.add_all([p1,p2]);db.flush()
    for p in (p1,p2):db.add(AcadTimetable(class_id=c.id,subject_id=subject.id,period_id=p.id,day=1))
    db.commit()
    preview=svc.import_preview(db,c,subject,1,'2026-05-18','2026-05-31')
    assert len(preview)==2 and [r['minutes'] for r in preview]==[50,40]
    assert db.query(AcadLesson).count()==0
    for r in preview:svc.add_lesson(db,c,subject,1,date.fromisoformat(r['date']),r['slot'],r['minutes'],'timetable')
    db.commit()
    assert all(r['duplicate'] for r in svc.import_preview(db,c,subject,1,'2026-05-18','2026-05-31'))
    p1.time_label='08:30-09:30';db.commit()
    assert svc.lessons_for(db,c.id,subject.id)[0].minutes==50
    with pytest.raises(ValueError,match='ปฏิทิน'):svc.import_preview(db,c,subject,1,'2026-05-18','2026-06-01')


def test_dates_cross_calendar_year(db):
    c,_,_=seed(db)
    assert svc.check_date(c,'2027-01-05')==date(2027,1,5)
    assert svc.check_date(c,'18052569')==date(2026,5,18)
    assert svc.check_date(c,'18/05/2569')==date(2026,5,18)
    with pytest.raises(ValueError):svc.check_date(c,'2026-01-05')
    with pytest.raises(ValueError):svc.check_date(c,'2027-05-01')
    assert svc.period_minutes('13.00 – 13.50')==50
    assert svc.period_minutes('พัก') is None


def test_routes_marks_export_permissions_and_legacy(env,monkeypatch):
    import app.tenancy as tn
    from app.routers import academic
    # The demo seed creates one tenant.
    tid=next(iter(tn._engines))
    with tn.session_for(tid) as db:
        c,subject,student=seed(db)
        cid,sid,student_id=c.id,subject.id,student.id
    args=dict(cid=cid,sid=sid,term=1)
    url=f'/academic/lessons?cid={cid}&sid={sid}&term=1'
    response=env.get(url)
    assert response.status_code==200, response.text
    assert 'ยังไม่กำหนดเป้าหมายรายเทอม' in response.text
    for slot,minutes in [('คาบ 1',50),('คาบ 2',40)]:
        r=env.post('/academic/lessons/add',data={**args,'date':'2026-05-18','slot':slot,'minutes':minutes})
        assert r.status_code==200 and 'บันทึกแล้ว' in r.text
    with tn.session_for(tid) as db:
        lessons=svc.lessons_for(db,cid,sid);lid=lessons[0].id
        db.add(AcadCalendar(year=2569,month=5,days_csv='18,19,20,21,22,25'))
        period=AcadPeriod(year=2569,seq=1,name='คาบ 1',time_label='08:30-09:20',is_break=False)
        db.add(period);db.flush()
        db.add(AcadTimetable(class_id=cid,subject_id=sid,period_id=period.id,day=1));db.commit()
    preview=env.post('/academic/lessons/preview',data={**args,'start':'2026-05-18','end':'2026-05-18'})
    assert preview.status_code==200 and 'มีอยู่แล้ว · ข้าม' in preview.text
    token=re.search(r'name="token" value="([^"]+)"',preview.text).group(1)
    for _ in range(2):
        imported=env.post('/academic/lessons/import',data={**args,'token':token})
        assert 'เพิ่มแล้ว 0 คาบ' in imported.text
    bad=env.post('/academic/lessons/import',data={**args,'token':'broken'})
    assert 'ตัวอย่างหมดอายุ' in bad.text
    r=env.post('/academic/lessons/marks',data={**args,'lesson_id':lid,f'mark_{student_id}':'/'})
    assert r.status_code==200
    r=env.post('/academic/lessons/remove',data={**args,'lesson_id':lid})
    assert 'มีผลเช็คชื่อแล้ว' in r.text
    legacy=env.get(f'/academic/attendance?mode=subject&cid={cid}&sid={sid}&legacy=1')
    assert legacy.status_code==200 and 'อ่านอย่างเดียว' in legacy.text
    assert '/academic/attendance/day-save' not in legacy.text
    for endpoint in ['day-save','save','fill-year']:
        assert env.post('/academic/attendance/'+endpoint,data={**args,'mode':'subject','month':5}).status_code==409
    home=env.post('/academic/attendance/day-save',data={'cid':cid,'mode':'overall','month':5,f'd_{student_id}_18':'/'})
    assert home.status_code==200
    r=env.get(f'/academic/attendance?mode=subject&cid={cid}&sid={sid}',follow_redirects=False)
    assert r.status_code==303 and '/academic/lessons?' in r.headers['location']
    from docx import Document
    r=env.get(f'/academic/lessons/report.docx?cid={cid}&sid={sid}&term=1')
    assert r.status_code==200,r.text[:300] if r.status_code!=200 else ''
    doc=Document(BytesIO(r.content))
    text=' '.join(p.text for p in doc.paragraphs)+' '.join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    assert '1.5 ชั่วโมง' in text and 'ยังไม่เช็ค' in text and '50 นาที' in text
    with tn.session_for(tid) as db:
        assert db.query(AcadAttendance).filter_by(acad_student_id=student_id,subject_id=sid).one().present==1
        assert db.query(AcadClassMonth).filter_by(class_id=cid).one().days_open==6  # only homeroom save updates this
    class Denied:
        is_teacher=True
        def can_teach(self,*args):return False
    monkeypatch.setattr(academic,'_scope',lambda *args:Denied())
    assert env.get(url).status_code==403
    assert env.post('/academic/lessons/add',data={**args,'date':'2026-05-19','slot':'คาบ 1','minutes':50}).status_code==403


def test_pp5_uses_period_marks_and_removes_orphans(db):
    from app.services.acad_doc import _doc, _pp5_attendance_daily
    c,subject,student=seed(db)
    svc.add_lesson(db,c,subject,1,date(2026,5,18),'คาบ 1',50)
    row=svc.lessons_for(db,c.id,subject.id)[0]
    db.add(AcadLessonMark(lesson_id=row.id,acad_student_id=student.id,mark='/'));db.commit()
    doc=_doc()
    _pp5_attendance_daily(doc,c,[student],db,subject_id=subject.id,page_break=False)
    text=' '.join(p.text for p in doc.paragraphs)
    assert 'รายคาบ' in text and '0.83 ชั่วโมง' in text
    db.delete(c);db.commit()
    assert db.query(AcadLesson).count()==0
    assert db.query(AcadLessonMark).count()==0
