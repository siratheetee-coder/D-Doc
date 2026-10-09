# -*- coding: utf-8 -*-
"""ลายเซ็นในเอกสารงานไปราชการ

เอกสารชุดนี้ไม่เคยมีเทสต์คุมลายเซ็นมาก่อน จึงไม่มีอะไรจับได้เลยว่า
  - บันทึกขออนุญาตไปราชการ: ลายเซ็นผู้ขอถูกลบทิ้งทันทีหลังแปะ
    (โค้ดสั่ง P[21].runs[7].text = "" เพื่อล้างจุดไข่ปลา แต่หลังแปะรูปแล้ว
     ดัชนี 7 คือ "รันรูป" การตั้ง text บน run = clear_content() ลบ w:drawing ไปด้วย)
  - คำสั่งไปราชการ: ไม่เคยแปะลายเซ็น ผอ. เลย ทั้งที่ระบบรู้ว่าใครลงนาม
รัน: .venv\\Scripts\\python.exe -m pytest tests/test_travel_signature.py
"""
import pathlib
import re
from datetime import date, datetime

import pytest
from docx import Document
from docx.oxml.ns import qn

ROOT = pathlib.Path(__file__).resolve().parents[1]


class O:
    def __init__(self, **kw):
        self.__dict__.update(kw)


@pytest.fixture
def sig(tmp_path, monkeypatch):
    from PIL import Image
    png = tmp_path / "sig.png"
    Image.new("RGBA", (300, 100), (0, 0, 0, 255)).save(png)
    import app.services.signature as S
    monkeypatch.setattr(S, "signature_path_for", lambda db, name: str(png) if name else None)
    monkeypatch.setattr(S, "signature_path_for_current", lambda name: str(png) if name else None)
    return str(png)


def _school():
    return O(name="โรงเรียนบ้านหินลาด", area_office="สพป.ขอนแก่น เขต 1",
             director_name="นายสมชาย ใจดี", director_position="ผู้อำนวยการโรงเรียน")


def _person():
    return O(name="นางสาวสุดา ขยันสอน", position="ครู", rank="ครู คศ.1")


def _record():
    return O(id=1, subject=" อบรมเชิงปฏิบัติการ", place="สพป.ขอนแก่น เขต 1", days=2,
             start_date=date(2026, 8, 17), end_date=date(2026, 8, 18),
             created_at=datetime(2026, 8, 10), purpose_type="อบรม", doc_ref="", doc_date=None,
             substitute_person_id=None, budget=0, budget_words="", reimburse="no",
             doc_no="123/2569", destination="สพป.ขอนแก่น เขต 1")


def _anchors(path):
    doc = Document(path)
    return doc.element.body.findall(".//" + qn("wp:anchor"))


def _text(path):
    doc = Document(path)
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            out += [c.text for c in row.cells]
    return "\n".join(out)


def _check_float(anchors):
    for a in anchors:
        assert a.find(qn("wp:wrapNone")) is not None, "ลายเซ็นต้องลอยหน้าข้อความ"
        assert a.get("behindDoc") == "0", "ต้องอยู่หน้าข้อความ ไม่ใช่หลัง"


# ---------------- บันทึกขออนุญาตไปราชการ (แบบฟอร์มจริง) ----------------
def test_official_travel_form_signs_both_the_teacher_and_the_director(sig):
    from app.services.gov_forms import render_travel_official
    path = render_travel_official(_school(), _person(), _record(), db=object(),
                                  approver=O(name="นายสมชาย ใจดี"), approve_date=date(2026, 8, 12))
    anchors = _anchors(path)
    assert len(anchors) == 2, f"ต้องมีลายเซ็นผู้ขอ + ผอ. = 2 จุด ได้ {len(anchors)}"
    _check_float(anchors)


def test_dotted_lines_survive_the_stamping(sig):
    """ของเดิมล้างจุดไข่ปลาทิ้ง ทำให้ข้อความท้ายบรรทัดเลื่อนมาทับลายเซ็น"""
    from app.services.gov_forms import render_travel_official
    path = render_travel_official(_school(), _person(), _record(), db=object(),
                                  approver=O(name="นายสมชาย ใจดี"), approve_date=date(2026, 8, 12))
    doc = Document(path)
    lines = [p.text for p in doc.paragraphs if "ลงชื่อ" in p.text]
    assert len(lines) >= 2, lines
    for ln in lines:
        assert re.search(r"\.{5,}|…{3,}", ln), f"เส้นไข่ปลาหายไปจากบรรทัด: {ln!r}"


def test_no_signature_before_the_director_approves(sig):
    """ยังไม่อนุมัติ ต้องมีแค่ลายเซ็นผู้ขอ ไม่มีของ ผอ."""
    from app.services.gov_forms import render_travel_official
    path = render_travel_official(_school(), _person(), _record(), db=object(), approver=None)
    assert len(_anchors(path)) == 1


def test_no_signature_in_the_registry_leaves_the_form_blank(monkeypatch):
    from app.services.gov_forms import render_travel_official
    import app.services.signature as S
    monkeypatch.setattr(S, "signature_path_for", lambda db, name: None)
    path = render_travel_official(_school(), _person(), _record(), db=object(),
                                  approver=O(name="นายสมชาย ใจดี"), approve_date=date(2026, 8, 12))
    assert _anchors(path) == []
    assert "นายสมชาย ใจดี" in _text(path), "ไม่มีลายเซ็นก็ต้องยังพิมพ์ชื่อผู้ลงนาม"


# ---------------- คำสั่งไปราชการ ----------------
def test_travel_order_is_signed_by_the_director(sig):
    from app.services.hr_doc import render_travel_order
    path = render_travel_order(_school(), _person(), _record())
    anchors = _anchors(path)
    assert len(anchors) == 1, f"คำสั่งไปราชการต้องมีลายเซ็น ผอ. ได้ {len(anchors)}"
    _check_float(anchors)


# ---------------- บันทึกขออนุญาตแบบย่อ (สำรอง) ----------------
def test_simple_travel_request_signs_the_director_when_approved(sig):
    from app.services.hr_doc import render_travel_request
    path = render_travel_request(_school(), _person(), _record(), approver=O(name="นายสมชาย ใจดี"))
    assert len(_anchors(path)) == 1
    path = render_travel_request(_school(), _person(), _record(), approver=None)
    assert _anchors(path) == [], "ยังไม่อนุมัติ ต้องไม่มีลายเซ็น"


# ---------------- กันบัคเดิมกลับมา ----------------
def test_no_document_sets_run_text_right_after_stamping_a_signature():
    """ตั้ง .text บน run หลังแปะรูป = ลบรูปทิ้ง (python-docx clear_content)

    เป็นต้นเหตุของบัคลายเซ็นหายในใบลาและในบันทึกไปราชการ กันไม่ให้ใครเผลอเขียนแบบนี้อีก
    """
    bad = []
    for f in sorted((ROOT / "app/services").glob("*.py")):
        lines = f.read_text(encoding="utf-8").splitlines()
        for i, ln in enumerate(lines):
            if "_stamp_sig(" not in ln and "_float_signature(" not in ln:
                continue
            for nxt in lines[i + 1:i + 4]:
                if re.search(r"runs\[[^\]]+\]\.text\s*=", nxt):
                    bad.append(f"{f.name}:{i + 1} -> {nxt.strip()}")
    assert not bad, "แปะลายเซ็นแล้วตั้ง .text บนรันถัดไปทันที จะลบรูปทิ้ง:\n" + "\n".join(bad)
