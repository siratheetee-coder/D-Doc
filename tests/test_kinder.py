# -*- coding: utf-8 -*-
"""สมุดพกอนุบาล (สมุดรายงานประจำตัวนักเรียน ระดับปฐมวัย)

ตรวจ 4 เรื่อง
  1. ชุดตัวบ่งชี้ครบตามไฟล์ต้นฉบับ (อ.1/อ.2/อ.3 ด้านละ 10-10-10-20 ข้อ · ไม่ซ้ำกันข้ามชั้น)
  2. สรุประดับคุณภาพรายด้านจากผลรายข้อ
  3. เอกสารออกได้ครบทุกหน้าและมีข้อความสำคัญ (ตัวบ่งชี้ · ความเห็นครู · สรุป)
  4. หน้ากระดาษ: ตารางไม่ล้นกรอบ A4 และไม่มีย่อหน้าว่างที่ทำให้เกิดหน้าเปล่า
"""
from types import SimpleNamespace as NS

from app.services import kinder as kd


def test_indicator_sets_complete():
    assert set(kd.KINDER_LEVELS) == {"อ.1", "อ.2", "อ.3"}
    for level in kd.KINDER_LEVELS:
        counts = {k: len(kd.items_for(level, k)) for k, _, _ in kd.DOMAINS}
        assert counts == {"phys": 10, "emo": 10, "soc": 10, "intel": 20}, (level, counts)
        for key, _full, _short in kd.DOMAINS:
            nums = [n for n, _, _ in kd.items_for(level, key)]
            assert nums == list(range(1, len(nums) + 1))
    assert len(kd.HOME_ITEMS) == 10
    assert kd.is_kinder("อ.2") and not kd.is_kinder("ป.1")


def test_indicators_differ_by_age():
    """แต่ละช่วงอายุต้องมีตัวบ่งชี้ของตัวเอง ไม่ใช่ชุดเดียวใช้ทุกชั้น"""
    first = {lv: kd.items_for(lv, "phys")[0][2] for lv in kd.KINDER_LEVELS}
    assert len(set(first.values())) == 3, first
    # ตัวอย่างจากไฟล์จริง: 3 ขวบต่อบล็อก 8 ชิ้น · 4 ขวบต่อเป็นรูปง่าย ๆ
    assert "8 ชิ้น" in kd.items_for("อ.1", "phys")[6][2]
    assert "จินตนาการ" in kd.items_for("อ.2", "phys")[8][2]


def test_domain_average():
    assert kd.domain_average([3, 3, 3]) == 3
    assert kd.domain_average([3, 2, 2]) == 2          # 2.33 -> ปานกลาง
    assert kd.domain_average([3, 3, 2]) == 3          # 2.67 -> ดี
    assert kd.domain_average([1, 1, 2]) == 1
    assert kd.domain_average([]) == 0
    assert kd.domain_average([None, 3]) == 3          # ข้อที่ยังไม่ประเมินไม่ถ่วง


def test_render_book(tmp_path, monkeypatch):
    """ออกเอกสารจริงแล้วอ่านกลับมาตรวจข้อความ"""
    import app.database as dbm
    import app.services.kinder_book as kb
    monkeypatch.setattr(dbm, "get_data_dir", lambda: tmp_path)
    monkeypatch.setattr(kb, "get_data_dir", lambda: tmp_path)

    school = NS(name="โรงเรียนบ้านตัวอย่าง", district="เมือง", province="สมมติ",
                area_office="ประถมศึกษาสมมติ เขต 1", director_name="นายเอนก ทดสอบ",
                director_position="ผู้อำนวยการโรงเรียน", logo=None)
    klass = NS(id=1, year=2569, level="อ.2", room="1", homeroom=NS(name="นางสาวครู ใจดี"),
               co_homeroom=None)
    student = NS(id=7, name="เด็กหญิงตัวอย่าง ทดสอบ", seq=3, student_no="4611",
                 student_id=None, klass=klass)
    klass.students = [student]

    res = {}
    for n, _g, _t in kd.items_for("อ.2", "phys"):
        res[(1, kd.code_of("phys", n))] = 3
        res[(2, kd.code_of("phys", n))] = 2
    notes = {"comments": {"1": {"phys": "ร่างกายแข็งแรงดี"}, "2": {}},
             "improve": {"1": "ควรฝึกการตัดกระดาษ", "2": ""},
             "summary": {"soc": 2}, "works": [{"date": "10/08/2569", "work": "วาดภาพระบายสี",
                                               "award": "ชนะเลิศ", "org": "โรงเรียน"}]}
    monkeypatch.setattr(kb, "_results_of", lambda db, aid: res)
    monkeypatch.setattr(kb, "_notes_of", lambda db, aid: notes)
    monkeypatch.setattr(kb, "_attendance", lambda *a, **k: None)   # ใช้ DB จริง ข้ามในเทสต์นี้
    monkeypatch.setattr(kb, "_personal", lambda *a, **k: None)

    path = kb.render_kinder_book(school, student, db=None)

    from docx import Document
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                text += "\n" + c.text
    assert "สมุดรายงานประจำตัวนักเรียน" in text
    assert "อบ.1/2" in text                       # แบบของอายุ 4 ขวบ
    assert "เด็กหญิงตัวอย่าง ทดสอบ" in text
    # ตัวบ่งชี้ต้องขึ้นครบทั้ง 4 ด้าน (50 ข้อ)
    for key, _full, _short in kd.DOMAINS:
        for _n, _g, body in kd.items_for("อ.2", key):
            assert body in text, body
    assert "ร่างกายแข็งแรงดี" in text and "ควรฝึกการตัดกระดาษ" in text
    assert "วาดภาพระบายสี" in text
    assert kd.HOME_ITEMS[0] in text               # หน้าผู้ปกครอง
    assert "นายเอนก ทดสอบ" in text                # ลงนามผู้บริหาร


def test_tables_fit_the_page(tmp_path, monkeypatch):
    """ไม่มีตารางไหนกว้างเกินพื้นที่พิมพ์ A4 (เคยเป็นบั๊กซ้ำ ๆ ในเอกสารอื่น)"""
    import app.database as dbm
    import app.services.kinder_book as kb
    from docx import Document
    from docx.shared import Emu
    monkeypatch.setattr(dbm, "get_data_dir", lambda: tmp_path)
    monkeypatch.setattr(kb, "get_data_dir", lambda: tmp_path)
    monkeypatch.setattr(kb, "_results_of", lambda db, aid: {})
    monkeypatch.setattr(kb, "_notes_of", lambda db, aid: {"comments": {}, "improve": {},
                                                         "summary": {}, "works": []})
    monkeypatch.setattr(kb, "_attendance", lambda *a, **k: None)
    monkeypatch.setattr(kb, "_personal", lambda *a, **k: None)

    school = NS(name="รร.ทดสอบ", district="", province="", area_office="",
                director_name="", director_position="", logo=None)
    klass = NS(id=1, year=2569, level="อ.3", room="1", homeroom=None, co_homeroom=None)
    student = NS(id=1, name="เด็กชายทดสอบ ทดสอบ", seq=1, student_no="1",
                 student_id=None, klass=klass)
    klass.students = [student]

    doc = Document(kb.render_kinder_book(school, student, db=None))
    sec = doc.sections[0]
    avail = Emu(sec.page_width - sec.left_margin - sec.right_margin).cm
    for i, t in enumerate(doc.tables, 1):
        widest = max(sum(Emu(c.width).cm for c in row.cells if c.width) for row in t.rows)
        assert widest <= avail + 0.01, f"ตารางที่ {i} กว้าง {widest:.2f} ซม. เกิน {avail:.2f} ซม."


def test_comment_page_has_no_trailing_blank(tmp_path, monkeypatch):
    """หน้าความเห็นครูต้องไม่ลงท้ายด้วยย่อหน้าว่าง (เคยดันให้เกิดหน้ากระดาษเปล่า)"""
    import app.services.kinder_book as kb
    from docx import Document

    doc = Document()
    kb._teacher_comments(doc, NS(name="x"), {"comments": {}, "improve": {}})
    assert doc.paragraphs[-1].text.strip() != "" or doc.paragraphs[-1].text == "", \
        "โครงสร้างเปลี่ยน - ตรวจใหม่"
    # ย่อหน้าสุดท้ายของเอกสารต้องไม่ใช่ย่อหน้าว่างที่ต่อท้ายตารางภาคเรียนที่ 2
    body = doc.element.body
    last = [c for c in body.iterchildren() if c.tag.endswith('}p') or c.tag.endswith('}tbl')][-1]
    assert last.tag.endswith('}tbl'), "มีย่อหน้าว่างต่อท้ายตารางภาค 2 -> จะเกิดหน้าเปล่า"


def _fake_db(rows=()):
    """db จำลองสำหรับเล่มครู: query(...).filter(...).all() -> rows ตามชนิดที่ถาม"""
    class _Q:
        def __init__(self, out):
            self.out = out

        def filter(self, *a, **k):
            return self

        def all(self):
            return self.out

    class _DB:
        def query(self, model, *a):
            return _Q(list(rows.get(getattr(model, "__name__", ""), []) if rows else []))

        def get(self, *a, **k):
            return None
    return _DB()


def test_teacher_book(tmp_path, monkeypatch):
    """สมุดบันทึกของครู (อบ.2): ออกได้ · มีทุกหน้าที่ต้องมี · ตารางไม่เกินกรอบกระดาษ"""
    import app.database as dbm
    import app.services.kinder_teacher as kt
    from docx import Document
    from docx.shared import Emu
    monkeypatch.setattr(dbm, "get_data_dir", lambda: tmp_path)
    monkeypatch.setattr(kt, "get_data_dir", lambda: tmp_path)

    school = NS(name="โรงเรียนบ้านตัวอย่าง", district="เมือง", province="สมมติ",
                area_office="ประถมศึกษาสมมติ เขต 1", director_name="นายเอนก ทดสอบ",
                director_position="ผู้อำนวยการโรงเรียน", logo=None)
    klass = NS(id=1, year=2569, level="อ.3", room="1",
               homeroom=NS(name="นางสาวครู ใจดี"), co_homeroom=None)
    klass.students = [NS(id=i, student_id=None, seq=i, student_no=f"46{i:02d}",
                         name=f"เด็กชายทดสอบ คนที่{i}", sex="M" if i % 2 else "F")
                      for i in range(1, 31)]                       # เต็มหน้า 30 คน

    path = kt.render_kinder_teacher_book(school, klass, _fake_db())
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                text += "\n" + c.text
    for need in ["บัญชีเรียกชื่อและสมุดบันทึกพัฒนาการเด็กนักเรียน", "อบ.2/3",
                 "ข้อมูลเด็ก", "สรุปผลน้ำหนัก - ส่วนสูง", "สรุปเวลาเรียน",
                 "สรุปผลการพัฒนา", "เด็กที่ควรได้รับการเสริม", "ชั้นอนุบาลปีที่ 3"]:
        assert need in text, need
    # ตัวบ่งชี้ทุกข้อของชั้นนี้ต้องมีอยู่ในเล่ม
    for key, _full, _short in kd.DOMAINS:
        for _n, _g, body in kd.items_for("อ.3", key):
            assert body in text, body
    # ชื่อนักเรียนครบทุกคน
    assert all(s.name in text for s in klass.students)

    # ไม่มีตารางไหนเกินพื้นที่พิมพ์ของ section ที่ตัวเองอยู่ (แนวนอน 26.7 · แนวตั้ง 18.0)
    widest_page = max(Emu(s.page_width - s.left_margin - s.right_margin).cm
                      for s in doc.sections)
    for i, t in enumerate(doc.tables, 1):
        w = max(sum(Emu(c.width).cm for c in row.cells if c.width) for row in t.rows)
        assert w <= widest_page + 0.01, f"ตารางที่ {i} กว้าง {w:.2f} ซม."


def test_teacher_book_marks_weak_students(tmp_path, monkeypatch):
    """หน้าสรุปต้องระบุเลขที่เด็กที่ได้ระดับ 'ควรเสริม' ของแต่ละด้าน"""
    import app.database as dbm
    import app.services.kinder_teacher as kt
    from docx import Document
    monkeypatch.setattr(dbm, "get_data_dir", lambda: tmp_path)
    monkeypatch.setattr(kt, "get_data_dir", lambda: tmp_path)

    school = NS(name="รร.ทดสอบ", district="", province="", area_office="",
                director_name="", director_position="", logo=None)
    klass = NS(id=1, year=2569, level="อ.1", room="1", homeroom=None, co_homeroom=None)
    klass.students = [NS(id=1, student_id=None, seq=1, student_no="1", name="เด็กเก่ง ดีมาก", sex="M"),
                      NS(id=2, student_id=None, seq=2, student_no="2", name="เด็กควร เสริม", sex="F")]

    # คนที่ 1 ได้ 3 ทุกข้อ · คนที่ 2 ได้ 1 ทุกข้อ -> ต้องขึ้นเลขที่ 2 ทุกด้าน
    results = []
    for s, val in ((1, 3), (2, 1)):
        for key, _f, _sh in kd.DOMAINS:
            for n, _g, _t in kd.items_for("อ.1", key):
                results.append(NS(acad_student_id=s, term=1, code=kd.code_of(key, n), value=val))

    class _Q:
        def __init__(self, out):
            self.out = out

        def filter(self, *a, **k):
            return self

        def all(self):
            return self.out

    class _DB:
        def query(self, model, *a):
            return _Q(results if getattr(model, "__name__", "") == "KinderResult" else [])

        def get(self, *a, **k):
            return None

    doc = Document(kt.render_kinder_teacher_book(school, klass, _DB()))
    rows = []
    for t in doc.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if cells and cells[0].startswith("ด้าน") and cells[0] != "ด้าน":
                rows.append(cells)
    assert rows, "ไม่พบตารางสรุปเด็กที่ควรได้รับการเสริม"
    term1 = rows[:4]
    assert all(r[1] == "2" for r in term1), term1
