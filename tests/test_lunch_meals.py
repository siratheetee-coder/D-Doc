"""จ้างเหมาปรุงสำเร็จ: บันทึกมื้อจริงรายวัน รายชั้น -> ยอดงวด = มื้อจริง x ราคาต่อหัว + ค่าปรับ + เอกสาร
(คู่มือ สพฐ. 2569: ชำระเงินตามที่ได้ประกอบอาหารจริงในแต่ละงวด, ค่าปรับร้อยละ 0.2 ต่อวัน)"""
from datetime import datetime

from tests.test_subsidy import env  # noqa: F401


def _db():
    import app.tenancy as tn
    return tn.session_for(next(iter(tn._engines)))


def test_meals_set_installment_amount_and_docs(env):
    from docx import Document
    from app.models import LunchProgram, LunchClass, LunchHireRound, LunchInstallment, Vendor
    with _db() as s:
        v = Vendor(name="นางทดสอบ ปรุงอร่อย"); s.add(v); s.flush()
        p = LunchProgram(year=2569, rate_per_head=22, operate_mode="hire"); s.add(p); s.flush()
        a = LunchClass(program_id=p.id, seq=1, level="ป.1", num_students=10)
        b = LunchClass(program_id=p.id, seq=2, level="ป.2", num_students=12)
        s.add_all([a, b]); s.flush()
        r = LunchHireRound(program_id=p.id, seq=1, vendor_id=v.id, start_date=datetime(2026, 6, 8),
                           end_date=datetime(2026, 6, 12), days=5, amount=2420)
        s.add(r); s.flush()
        i = LunchInstallment(round_id=r.id, seq=1, start_date=datetime(2026, 6, 8), end_date=datetime(2026, 6, 12),
                             days=5, amount=5 * 22 * 22)        # ตั้งต้นตามสิทธิ์เต็ม 2,420
        s.add(i); s.commit()
        iid, ca, cb, rid = i.id, a.id, b.id, r.id

    page = env.get(f"/lunch/installment/{iid}/meals").text
    assert 'name="m_2026-06-08_%d"' % ca in page and "ยังไม่บันทึก" in page

    # บันทึก 4 วัน -> ยังไม่ครบ ยอดงวดไม่เปลี่ยน
    days = ["2026-06-08", "2026-06-09", "2026-06-10", "2026-06-11", "2026-06-12"]
    data = {f"m_{d}_{ca}": "10" for d in days[:4]} | {f"m_{d}_{cb}": "12" for d in days[:4]}
    env.post(f"/lunch/installment/{iid}/meals", data=data)
    with _db() as s:
        assert s.get(LunchInstallment, iid).amount == 2420

    # ครบ 5 วัน โดยวันศุกร์ ป.2 มา 9 คน -> มื้อจริง 107 x 22 = 2,354
    data = {f"m_{days[4]}_{ca}": "10", f"m_{days[4]}_{cb}": "9", "penalty_days": "1"}
    r = env.post(f"/lunch/installment/{iid}/meals", data=data, follow_redirects=False)
    assert "amt=1" in r.headers["location"]
    from app.services.lunch_meals import summary
    with _db() as s:
        inst = s.get(LunchInstallment, iid)
        assert inst.amount == 2354 and inst.penalty_days == 1
        sm = summary(s, inst)
        assert sm["meals"] == 107 and sm["complete"]
        assert sm["penalty"] == round(1 * 0.002 * 2354 / 5, 2)
        assert sm["net"] == round(2354 - 23.54 - sm["penalty"], 2)

    # จ่ายแล้ว: แก้มื้อไม่เปลี่ยนยอด
    env.post(f"/lunch/installment/{iid}/update", data={"status": "จ่ายแล้ว"})
    env.post(f"/lunch/installment/{iid}/meals", data={f"m_{days[0]}_{ca}": "1"})
    with _db() as s:
        assert s.get(LunchInstallment, iid).amount == 2354

    # เอกสารงวดมีบัญชีจำนวนนักเรียนรายวัน
    import io
    doc = Document(io.BytesIO(env.get(f"/lunch/installment/{iid}/doc").content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "บัญชีจำนวนนักเรียนที่รับประทานอาหารกลางวัน" in text
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert "ป.2" in cells
    # บันทึกขอเบิกจ่ายมีค่าปรับ
    doc = Document(io.BytesIO(env.get(f"/lunch/installment/{iid}/disburse-doc").content))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert any("0.94 บาท" in c for c in cells)
    # หน้ารอบแสดงสถานะมื้อจริงในแถวงวด
    assert "ยอดนักเรียนรายวัน" in env.get(f"/lunch/round/{rid}/plan").text
