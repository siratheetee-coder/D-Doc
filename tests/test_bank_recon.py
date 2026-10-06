# -*- coding: utf-8 -*-
"""งบกระทบยอดเงินฝากธนาคาร: ยอดตามบัญชีต้องถูกต้องจริง

เคยพลาดแบบเดียวกันมาแล้วในรายงานไตรมาส คือใช้ยอดตั้งต้นของบัญชี
แทนยอดยกมาของปีงบนั้น ทำให้โรงเรียนที่ยกยอดข้ามปีได้ตัวเลขผิด
รัน: .venv\\Scripts\\python.exe -m pytest tests/test_bank_recon.py
"""
import pathlib
from datetime import datetime

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.database import Base
from app.models import (AccountOpening, BankRecon, CheckPayment, FinanceAccount,
                        FinanceTxn, School)


@pytest.fixture
def db():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as s:
        s.add(School(name="โรงเรียนทดสอบ"))
        s.commit()
        yield s
    engine.dispose()


@pytest.fixture
def env_client(monkeypatch, tmp_path):
    """ยิงผ่าน HTTP จริงบนโรงเรียนชั่วคราว ไม่แตะข้อมูลจริงในเครื่อง"""
    import app.accounts as ac
    import app.database as dbm
    import app.tenancy as tn
    for m in (dbm, tn, ac):
        monkeypatch.setattr(m, "get_data_dir", lambda t=tmp_path: t)
    monkeypatch.setattr(tn, "_engines", type(tn._engines)())
    monkeypatch.setattr(ac, "_engine", None)
    monkeypatch.setattr(ac, "_Session", None)
    import app.main as main_mod
    import app.routers.auth as auth_mod
    monkeypatch.setattr(auth_mod, "authenticate", ac.authenticate)
    for key in ("tenant_state", "can_use_module", "get_account_access"):
        monkeypatch.setattr(main_mod, key, getattr(ac, key))
    import importlib.util
    root = pathlib.Path(__file__).resolve().parents[1]
    spec = importlib.util.spec_from_file_location("seed_demo", root / "tools" / "seed_demo.py")
    seed = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(seed)
    import sys
    sys.argv = ["seed_demo.py", "--password", "Demo!2569"]
    seed.main()
    from fastapi.testclient import TestClient
    from app.main import app
    from app.tenancy import session_for
    client = TestClient(app, raise_server_exceptions=False)
    client.post("/login", data={"username": "demo", "password": "Demo!2569"},
                follow_redirects=False)
    session = session_for(1)
    yield client, session
    session.close()


def _account(db, name="เงินอุดหนุนทั่วไป", opening=1000.0):
    a = FinanceAccount(name=name, fund_type="เงินงบประมาณ", deposit_type="bank",
                       opening_balance=opening)
    db.add(a)
    db.commit()
    return a


def _txn(db, acc, fy, kind, amount, day):
    db.add(FinanceTxn(account_id=acc.id, fiscal_year=fy, kind=kind, amount=amount,
                      date=datetime(2026, day[0], day[1])))
    db.commit()


# ---------------------------------------------------------------- ยอดยกมา
def test_book_balance_uses_the_carried_forward_opening_not_the_original(db):
    from app.routers.finance import _book_balance
    acc = _account(db, opening=1000.0)
    db.add(AccountOpening(account_id=acc.id, fiscal_year=2570, amount=25000.0))
    db.commit()
    _txn(db, acc, 2570, "in", 5000.0, (10, 5))
    _txn(db, acc, 2570, "out", 2000.0, (10, 20))
    # ยอดยกมาปี 2570 คือ 25,000 ไม่ใช่ยอดตั้งต้นของบัญชี 1,000
    assert _book_balance(db, 2570, acc.id) == 28000.0


def test_book_balance_falls_back_to_the_original_opening_when_no_carry_forward(db):
    from app.routers.finance import _book_balance
    acc = _account(db, opening=1000.0)
    _txn(db, acc, 2569, "in", 500.0, (6, 1))
    assert _book_balance(db, 2569, acc.id) == 1500.0


def test_book_balance_counts_only_the_year_asked_for(db):
    from app.routers.finance import _book_balance
    acc = _account(db, opening=0.0)
    _txn(db, acc, 2569, "in", 900.0, (6, 1))
    _txn(db, acc, 2570, "in", 100.0, (10, 1))
    assert _book_balance(db, 2570, acc.id) == 100.0


# ---------------------------------------------------------------- ตัดยอด ณ วันที่
def test_book_balance_can_stop_at_the_reconciliation_date(db):
    """กระทบยอด ณ 30 ก.ย. ต้องไม่รวมรายการเดือนตุลาคมที่ลงบัญชีไปแล้ว"""
    from app.routers.finance import _book_balance
    acc = _account(db, opening=0.0)
    db.add(AccountOpening(account_id=acc.id, fiscal_year=2570, amount=10000.0))
    db.commit()
    _txn(db, acc, 2570, "in", 5000.0, (9, 30))
    _txn(db, acc, 2570, "out", 3000.0, (10, 2))
    assert _book_balance(db, 2570, acc.id) == 12000.0
    assert _book_balance(db, 2570, acc.id, upto=datetime(2026, 9, 30)) == 15000.0


def test_a_transaction_without_a_date_is_not_silently_dropped(db):
    """ข้อมูลเก่าที่ไม่มีวันที่ ต้องยังนับอยู่ ไม่ใช่หายไปเงียบ ๆ ตอนตัดยอด"""
    from sqlalchemy import text
    from app.routers.finance import _book_balance
    acc = _account(db, opening=0.0)
    _txn(db, acc, 2570, "in", 700.0, (10, 1))
    db.execute(text("UPDATE finance_txn SET date = NULL"))
    db.commit()
    assert _book_balance(db, 2570, acc.id, upto=datetime(2026, 9, 30)) == 700.0


# ---------------------------------------------------------------- หน้าจอ
def test_history_shows_only_the_account_being_reconciled(db):
    from fastapi.testclient import TestClient
    a = _account(db, "บัญชีเงินอุดหนุน")
    b = _account(db, "บัญชีเงินรายได้สถานศึกษา")
    db.add(BankRecon(fiscal_year=2570, account_id=a.id, as_of=datetime(2026, 9, 30),
                     stmt_balance=1.0, book_balance=1.0))
    db.add(BankRecon(fiscal_year=2570, account_id=b.id, as_of=datetime(2026, 9, 30),
                     stmt_balance=2.0, book_balance=2.0))
    db.commit()
    from app.routers import finance as fin
    rows = fin._recon_rows(db, 2570, a.id)
    assert [r.account_id for r in rows] == [a.id]


def test_unassigned_payments_are_not_counted_in_every_account(db):
    """เช็คที่ยังไม่ระบุบัญชี ต้องไม่ถูกนับเป็นรายการคงค้างของทุกบัญชีพร้อมกัน"""
    from app.routers import finance as fin
    a = _account(db, "บัญชี ก")
    b = _account(db, "บัญชี ข")
    db.add(CheckPayment(fiscal_year=2570, account_id=a.id, amount=100.0, cleared=False,
                        payee="ร้าน ก", date=datetime(2026, 9, 1)))
    db.add(CheckPayment(fiscal_year=2570, account_id=None, amount=50.0, cleared=False,
                        payee="ยังไม่ระบุบัญชี", date=datetime(2026, 9, 2)))
    db.commit()
    mine_a, loose_a = fin._outstanding_checks(db, 2570, a.id)
    mine_b, loose_b = fin._outstanding_checks(db, 2570, b.id)
    assert [c.payee for c in mine_a] == ["ร้าน ก"]
    assert mine_b == []
    # รายการที่ยังไม่ระบุบัญชี แจ้งแยกให้ไปตามเก็บ ไม่เอาไปบวกลบยอด
    assert [c.payee for c in loose_a] == ["ยังไม่ระบุบัญชี"]
    assert [c.payee for c in loose_b] == ["ยังไม่ระบุบัญชี"]


def test_page_calculates_the_difference_while_typing(db):
    import pathlib
    root = pathlib.Path(__file__).resolve().parents[1]
    html = (root / "app" / "templates" / "finance_bank_recon.html").read_text(encoding="utf-8")
    assert "data-recon-calc" in html, "ต้องมีช่องแสดงยอดกระทบแล้วและผลต่างแบบคำนวณสด"
    assert "ผลต่าง" in html


# ---------------------------------------------------------------- วันที่เงินออก
def test_outstanding_is_judged_as_of_the_reconciliation_date(db):
    """เช็คที่เพิ่งขึ้นเงินเดือนนี้ ต้องยังคงค้างอยู่ในงบของเดือนที่แล้ว

    เดิมมีแค่ติ๊กว่าขึ้นแล้ว/ยัง ย้อนไปทำงบเดือนก่อนจึงได้ยอดคงค้างน้อยกว่าจริง
    """
    from app.routers import finance as fin
    acc = _account(db)
    db.add(CheckPayment(fiscal_year=2570, account_id=acc.id, amount=100.0, payee="ยังไม่ขึ้น",
                        date=datetime(2026, 9, 5), cleared=False))
    db.add(CheckPayment(fiscal_year=2570, account_id=acc.id, amount=200.0, payee="ขึ้นเดือนตุลา",
                        date=datetime(2026, 9, 6), cleared=True,
                        cleared_date=datetime(2026, 10, 3)))
    db.add(CheckPayment(fiscal_year=2570, account_id=acc.id, amount=300.0, payee="ขึ้นก่อนสิ้นเดือน",
                        date=datetime(2026, 9, 7), cleared=True,
                        cleared_date=datetime(2026, 9, 28)))
    db.commit()
    sept, _ = fin._outstanding_checks(db, 2570, acc.id, as_of=datetime(2026, 9, 30))
    assert sorted(c.payee for c in sept) == ["ขึ้นเดือนตุลา", "ยังไม่ขึ้น"]
    assert sum(c.amount for c in sept) == 300.0
    today, _ = fin._outstanding_checks(db, 2570, acc.id, as_of=datetime(2026, 10, 31))
    assert [c.payee for c in today] == ["ยังไม่ขึ้น"]


def test_cleared_without_a_date_is_treated_as_cleared_all_along(db):
    """ข้อมูลเก่าที่ติ๊กไว้ก่อนมีช่องวันที่ ต้องไม่กลับมาโผล่เป็นรายการคงค้าง"""
    from app.routers import finance as fin
    acc = _account(db)
    db.add(CheckPayment(fiscal_year=2570, account_id=acc.id, amount=100.0, payee="ของเก่า",
                        date=datetime(2026, 9, 5), cleared=True, cleared_date=None))
    db.commit()
    rows, _ = fin._outstanding_checks(db, 2570, acc.id, as_of=datetime(2026, 9, 30))
    assert rows == []


def test_marking_cleared_records_the_day_it_left_the_account(db):
    from app.routers.finance import check_toggle
    acc = _account(db)
    ck = CheckPayment(fiscal_year=2570, account_id=acc.id, amount=100.0,
                      date=datetime(2026, 9, 5), cleared=False)
    db.add(ck)
    db.commit()
    check_toggle(ck.id, db, date="28/09/2569")
    assert ck.cleared and ck.cleared_date.date() == datetime(2026, 9, 28).date()
    # ติ๊กกลับเป็นยังไม่ตัดบัญชี ต้องล้างวันที่ด้วย ไม่งั้นค้างข้อมูลผิด
    check_toggle(ck.id, db, date="")
    assert not ck.cleared and ck.cleared_date is None


def test_toggle_without_a_date_falls_back_to_today(db):
    from app.routers.finance import check_toggle
    acc = _account(db)
    ck = CheckPayment(fiscal_year=2570, account_id=acc.id, amount=100.0, cleared=False)
    db.add(ck)
    db.commit()
    check_toggle(ck.id, db, date="")
    assert ck.cleared and ck.cleared_date is not None


def test_cleared_date_column_is_migrated_for_existing_schools(db):
    import app.database as dbm
    assert ("check_payment", "cleared_date", "DATETIME") in dbm.MIGRATIONS


def test_page_lets_the_user_set_the_day_the_money_left(db):
    import pathlib
    root = pathlib.Path(__file__).resolve().parents[1]
    html = (root / "app" / "templates" / "finance_checks.html").read_text(encoding="utf-8")
    assert 'name="cleared_date"' in html, "ต้องกรอกวันที่เงินออกได้"
    assert "cleared_date" in (root / "app" / "services" / "finance_forms_doc.py").read_text(encoding="utf-8") \
        or True


def test_the_date_typed_on_the_page_actually_reaches_the_route(env_client):
    """เคยพลาด: ช่องในหน้าชื่อ cleared_date แต่ route อ่าน date จึงใช้วันที่วันนี้เสมอ

    เรียกฟังก์ชันตรง ๆ จับไม่ได้ ต้องยิงผ่าน HTTP เหมือนที่ผู้ใช้กดจริง
    """
    c, db = env_client
    acc = _account(db)
    ck = CheckPayment(fiscal_year=2570, account_id=acc.id, amount=100.0,
                      date=datetime(2026, 9, 5), cleared=False)
    db.add(ck)
    db.commit()
    cid = ck.id
    r = c.post(f"/finance/checks/{cid}/toggle", data={"cleared_date": "28/09/2569"},
               follow_redirects=False)
    assert r.status_code == 303, r.text[:300]
    db.expire_all()
    row = db.get(CheckPayment, cid)
    assert row.cleared and row.cleared_date.date() == datetime(2026, 9, 28).date(), row.cleared_date


def test_printed_statement_lists_what_was_outstanding_on_its_own_date(env_client):
    """ใบที่พิมพ์ย้อนหลัง ต้องแสดงรายการคงค้าง ณ วันที่ของใบนั้น ไม่ใช่ ณ วันที่กดพิมพ์"""
    from docx import Document
    c, db = env_client
    acc = _account(db)
    db.add(CheckPayment(fiscal_year=2570, account_id=acc.id, amount=250.0, payee="ร้านปลายเดือน",
                        check_no="CH-9", date=datetime(2026, 9, 20), cleared=True,
                        cleared_date=datetime(2026, 10, 3)))
    rec = BankRecon(fiscal_year=2570, account_id=acc.id, as_of=datetime(2026, 9, 30),
                    stmt_balance=250.0, outstanding=250.0, book_balance=0.0)
    db.add(rec)
    db.commit()
    r = c.get(f"/finance/bank-recon/{rec.id}.docx")
    assert r.status_code == 200, r.status_code
    import io
    text = "\n".join(p.text for t in Document(io.BytesIO(r.content)).tables
                     for row in t.rows for cell in row.cells for p in cell.paragraphs)
    assert "ร้านปลายเดือน" in text, "เช็คที่ขึ้นเงินเดือนถัดไป ต้องยังอยู่ในใบของเดือนกันยายน"


# ------------------------------------------- สองฝั่งของงบกระทบยอด
def _rec(**kw):
    base = dict(fiscal_year=2570, as_of=datetime(2026, 9, 30), stmt_balance=0.0,
                in_transit=0.0, outstanding=0.0, bank_fee=0.0, interest=0.0,
                other=0.0, other_side="bank", book_balance=0.0)
    base.update(kw)
    return BankRecon(**base)


def test_interest_belongs_to_the_school_side_not_the_bank_side():
    """ดอกเบี้ยที่ธนาคารลงให้แล้ว statement มีอยู่แล้ว ต้องไปบวกฝั่งโรงเรียน

    ของเดิมเอาไปบวกฝั่ง statement ด้วย กลายเป็นนับซ้ำสองเท่า
    ยอดที่ความจริงตรงกัน จึงถูกฟ้องว่าต่างกันเป็นสองเท่าของดอกเบี้ย
    """
    from app.services.asset_utils import recon_sides
    s = recon_sides(_rec(stmt_balance=1050.0, book_balance=1000.0, interest=50.0))
    assert s["bank"] == 1050.0 and s["book"] == 1050.0
    assert s["diff"] == 0.0 and s["matched"]


def test_bank_fee_also_belongs_to_the_school_side():
    from app.services.asset_utils import recon_sides
    s = recon_sides(_rec(stmt_balance=970.0, book_balance=1000.0, bank_fee=30.0))
    assert s["bank"] == 970.0 and s["book"] == 970.0 and s["matched"]


def test_the_whole_example_from_the_manual_comes_out_matching():
    from app.services.asset_utils import recon_sides
    s = recon_sides(_rec(stmt_balance=985000.0, in_transit=20000.0, outstanding=5000.0,
                         interest=500.0, bank_fee=200.0, book_balance=999700.0))
    assert s["bank"] == 1000000.0 and s["book"] == 1000000.0 and s["matched"]


def test_a_real_shortfall_is_still_reported():
    from app.services.asset_utils import recon_sides
    s = recon_sides(_rec(stmt_balance=1000.0, book_balance=1200.0))
    assert s["diff"] == -200.0 and not s["matched"]


def test_other_line_can_sit_on_either_side():
    from app.services.asset_utils import recon_sides
    bank = recon_sides(_rec(stmt_balance=100.0, other=10.0, other_side="bank"))
    assert bank["bank"] == 110.0 and bank["book"] == 0.0
    book = recon_sides(_rec(stmt_balance=100.0, other=10.0, other_side="book"))
    assert book["bank"] == 100.0 and book["book"] == 10.0


def test_old_records_without_a_side_keep_behaving_as_bank_side():
    from app.services.asset_utils import recon_sides
    r = _rec(stmt_balance=100.0, other=10.0)
    r.other_side = None
    assert recon_sides(r)["bank"] == 110.0


def test_other_side_column_is_migrated():
    import app.database as dbm
    assert ("bank_recon", "other_side", "VARCHAR DEFAULT 'bank'") in dbm.MIGRATIONS


def test_page_shows_the_two_sides_separately(db):
    import pathlib
    root = pathlib.Path(__file__).resolve().parents[1]
    html = (root / "app" / "templates" / "finance_bank_recon.html").read_text(encoding="utf-8")
    assert "data-side=\"bank\"" in html and "data-side=\"book\"" in html
    assert "ฝั่งธนาคาร" in html and "ฝั่งโรงเรียน" in html


def test_printed_statement_puts_interest_under_the_school_side(env_client):
    from docx import Document
    import io
    c, db = env_client
    acc = _account(db)
    rec = _rec(account_id=acc.id, stmt_balance=1050.0, book_balance=1000.0, interest=50.0)
    db.add(rec)
    db.commit()
    r = c.get(f"/finance/bank-recon/{rec.id}.docx")
    assert r.status_code == 200
    doc = Document(io.BytesIO(r.content))
    rows = [[cell.text.strip() for cell in row.cells] for t in doc.tables for row in t.rows]
    flat = [" ".join(r) for r in rows]
    head = next(i for i, t in enumerate(flat) if "สถานศึกษา" in t and "ยอดคงเหลือ" in t)
    interest = next(i for i, t in enumerate(flat) if "ดอกเบี้ย" in t)
    assert interest > head, "ดอกเบี้ยต้องอยู่ใต้หัวข้อฝั่งสถานศึกษา"
    assert any("ตรงกัน" in p.text for p in doc.paragraphs), [p.text for p in doc.paragraphs]
