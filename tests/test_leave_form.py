# -*- coding: utf-8 -*-
"""ใบลา (แบบ ก.พ.) ที่ออกจากระบบ ต้องหน้าตาเหมือนแบบฟอร์มเปล่า

เคยพลาด: ตารางสถิติการลาในแบบฟอร์มอยู่ใน "กล่องข้อความ" ซึ่งผูกกับ run ว่าง
ตัวแรกของบรรทัด 'วันที่' พอโค้ดจัดแนวช่องลงนามสั่งล้าง run ว่างข้างหน้าทิ้ง
python-docx ล้างลูกของ run ทั้งหมด ตารางสถิติจึงหายไปทั้งตารางแบบเงียบ ๆ
และเส้นไข่ปลาของช่อง 'ตำแหน่ง' อีกคอลัมน์ก็หายไปด้วย
รัน: .venv\\Scripts\\python.exe -m pytest tests/test_leave_form.py
"""
import datetime
import pathlib
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.models import Person, School

ROOT = pathlib.Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "app" / "data" / "forms" / "leave_form.docx"
ALT = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"


class _Rec:
    id = 3
    leave_type = "personal"
    start_date = datetime.datetime(2026, 5, 18)
    end_date = datetime.datetime(2026, 7, 31)
    days = 75
    reason = "ไปทำธุระที่ต่างจังหวัด"
    contact = "0961066910"
    created_at = datetime.datetime(2026, 9, 1)
    year = None


@pytest.fixture
def render(monkeypatch, tmp_path):
    import app.services.gov_forms as gf
    monkeypatch.setattr(gf, "get_data_dir", lambda: tmp_path)

    def go(**kw):
        school = School(name="โรงเรียนบ้านหินลาด", director_name="นายสมชาย ใจดี",
                        director_position="ผู้อำนวยการโรงเรียนบ้านหินลาด")
        person = Person(name="นายสิรธีร์  ตีเมืองซ้าย", position="ครู")
        person.id = 11
        kw.setdefault("db", None)
        return pathlib.Path(gf.render_leave_official(
            school, person, _Rec(), work_group="กลุ่มสาระการเรียนรู้ภาษาไทย", **kw))
    return go


def _counts(path):
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")
    return {"tbl": xml.count("w:tbl>"), "txbx": xml.count("w:txbxContent"),
            "alt": xml.count("mc:AlternateContent")}


# ------------------------------------------------- ตารางสถิติการลา
def test_the_template_really_keeps_the_stats_table_in_a_text_box():
    """กันไว้: ถ้าวันหน้าเปลี่ยนแม่แบบ เทสต์ข้างล่างจะได้ไม่ผ่านแบบไร้ความหมาย"""
    c = _counts(TEMPLATE)
    assert c["tbl"] >= 1 and c["txbx"] >= 1, c


def test_generated_form_still_has_the_stats_table(render):
    out = _counts(render())
    assert out["tbl"] >= 1, "ตารางสถิติการลาหายไปจากใบที่ออก"
    assert out["txbx"] >= 1 and out["alt"] >= 1


def test_stats_table_is_filled_with_this_leave(render):
    doc = Document(str(render()))
    cells = [tc for tbl in doc.element.body.iter(qn("w:tbl"))
             for tr in tbl.findall(qn("w:tr")) for tc in tr.findall(qn("w:tc"))]
    text = " ".join("".join(t.text or "" for t in tc.iter(qn("w:t"))) for tc in cells)
    assert "ประเภทการลา" in text and "กิจส่วนตัว" in text
    assert "75" in text, text


# ------------------------------------------------- ช่องลงนามท้ายใบ
def _para_texts(path):
    return [p.text for p in Document(str(path)).paragraphs]


def test_unsigned_form_keeps_the_dotted_fields_of_the_boss_column(render):
    """ยังไม่มีใครลงนาม ท้ายใบต้องเหมือนแบบฟอร์มเปล่า ไม่ใช่เส้นประลอยพาดหน้า"""
    doc = Document(str(render()))
    tail = [p for p in doc.paragraphs if p.text.strip()][-4:]
    joined = " | ".join(p.text.strip() for p in tail)
    assert "ตำแหน่ง" in joined and "วันที่" in joined, joined
    # ต้องยังมี run ขีดเส้นใต้ (เส้นไข่ปลา) เหลืออยู่ในบรรทัดท้าย ๆ
    dotted = sum(1 for p in tail for r in p.runs
                 if r._element.rPr is not None and r._element.rPr.findall(qn("w:u")))
    assert dotted >= 2, f"เส้นไข่ปลาท้ายใบหายไปหมด ({dotted} เส้น)"


def test_a_signed_form_fills_both_checker_and_director(render):
    checker = Person(name="นางสาวศศิธร เทียกมา", position="หัวหน้ากลุ่มงานบุคคล")
    checker.id = 6
    boss = Person(name="นายสมชาย ใจดี", position="ผู้อำนวยการโรงเรียน")
    boss.id = 1
    texts = " ".join(_para_texts(render(
        checker=checker, approver=boss,
        checker_date=datetime.datetime(2026, 9, 2),
        approve_date=datetime.datetime(2026, 9, 3))))
    assert "นางสาวศศิธร เทียกมา" in texts
    assert "นายสมชาย ใจดี" in texts
    assert "ผู้อำนวยการโรงเรียนบ้านหินลาด" in texts
    assert "3 กันยายน 2569" in texts


def test_signing_does_not_cost_the_stats_table_either(render):
    boss = Person(name="นายสมชาย ใจดี", position="ผู้อำนวยการโรงเรียน")
    boss.id = 1
    assert _counts(render(approver=boss, approve_date=datetime.datetime(2026, 9, 3)))["tbl"] >= 1


def test_form_stays_on_one_page_worth_of_paragraphs(render):
    """แม่แบบพอดีหน้าเดียว ถ้าย่อหน้าบวมขึ้นแปลว่ามีอะไรแทรกเกิน"""
    assert len(Document(str(render())).paragraphs) <= len(Document(str(TEMPLATE)).paragraphs)


# ------------------------------------------------- ตัวช่วยระดับ run
def test_tab_before_never_deletes_a_picture_or_text_box():
    from app.services.gov_forms import _tab_before
    doc = Document(str(TEMPLATE))
    p = next(p for p in doc.paragraphs
             if p._element.findall(".//" + ALT) or p._element.findall(".//" + qn("w:drawing")))
    before = len(p._element.findall(".//" + ALT)) + len(p._element.findall(".//" + qn("w:drawing")))
    assert before > 0
    i = next(k for k, r in enumerate(p.runs) if r.text.strip())
    _tab_before(p, i)
    after = len(p._element.findall(".//" + ALT)) + len(p._element.findall(".//" + qn("w:drawing")))
    assert after == before, "จัดแนวแล้วรูป/กล่องข้อความหายไป"


def test_tab_before_keeps_the_dotted_field_of_the_other_column():
    from app.services.gov_forms import _tab_before, _find_run
    doc = Document(str(TEMPLATE))
    p = doc.paragraphs[28]                       # 'ตำแหน่ง....(ลงชื่อ)' สองคอลัมน์บรรทัดเดียว
    i = _find_run(p, "(ลงชื่อ)")
    assert i is not None
    before = sum(1 for r in p.runs
                 if r._element.rPr is not None and r._element.rPr.findall(qn("w:u")))
    _tab_before(p, i)
    after = sum(1 for r in p.runs
                if r._element.rPr is not None and r._element.rPr.findall(qn("w:u")))
    assert after == before, "เส้นไข่ปลาของ 'ตำแหน่ง' หายตอนจัดแนวคอลัมน์ขวา"


# ------------------------------------------------- ลายเซ็นต้องลอยหน้าข้อความ
def test_signature_floats_in_front_of_text_so_dots_do_not_move(render, monkeypatch, tmp_path):
    """ลายเซ็นแบบไหลตามข้อความจะดันเส้นไข่ปลาเลื่อน ต้องลอยทับแทน"""
    from PIL import Image
    import app.services.signature as sigmod
    png = tmp_path / "sig.png"
    Image.new("RGBA", (300, 100), (0, 0, 0, 0)).save(png)
    monkeypatch.setattr(sigmod, "signature_path_for", lambda db, name: str(png))

    class _Q:
        def filter_by(self, **k): return self
        def filter(self, *a, **k): return self
        def order_by(self, *a, **k): return self
        def all(self): return []
        def first(self): return None

    class _DB:
        def query(self, *a, **k): return _Q()

    boss = Person(name="นายอัครพงศ์ ศรีวงศ์", position="ผู้อำนวยการโรงเรียน")
    boss.id = 1
    path = render(db=_DB(), approver=boss, approve_date=datetime.datetime(2026, 10, 7))
    import zipfile
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")
    assert "wp:anchor" in xml, "ลายเซ็นยังเป็นรูปแบบไหลตามข้อความ"
    assert 'behindDoc="0"' in xml, "ต้องอยู่หน้าข้อความ ไม่ใช่หลังข้อความ"
    assert "wp:wrapNone" in xml, "ต้องไม่ตัดข้อความรอบรูป"


def test_approved_form_has_no_blank_line_before_the_director_date(render):
    """วันที่ของ ผอ. ต้องอยู่ติดใต้บรรทัดตำแหน่ง ไม่เว้นบรรทัดลอย"""
    boss = Person(name="นายอัครพงศ์ ศรีวงศ์", position="ผู้อำนวยการโรงเรียน")
    boss.id = 1
    doc = Document(str(render(approver=boss, approve_date=datetime.datetime(2026, 10, 7))))
    texts = [p.text.strip() for p in doc.paragraphs]
    pos = max(i for i, t in enumerate(texts) if "ตำแหน่ง" in t and "ผู้อำนวยการ" in t)
    nxt = next(i for i in range(pos + 1, len(texts)) if texts[i])
    assert "วันที่" in texts[nxt], texts[pos:pos + 4]
