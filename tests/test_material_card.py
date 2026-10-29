"""บัญชีวัสดุตามแบบ: หัวบัญชี (ขนาด/รหัส/ที่เก็บ/จำนวนอย่างสูง-ต่ำ) รับจาก/จ่ายให้ ราคาต่อหน่วยทุกแถว พิมพ์ Word"""
from datetime import datetime

from docx import Document

from tests.test_subsidy import env  # noqa: F401


def _db():
    import app.tenancy as tn
    return tn.session_for(next(iter(tn._engines)))


def test_material_card_fields_party_price_and_docx(env, tmp_path, monkeypatch):
    from app.models import MaterialItem, MaterialTxn, Requisition, RequisitionItem
    env.post("/materials", data={"name": "กระดาษ A4", "unit": "รีม", "opening": "10", "category": "วัสดุสำนักงาน",
                                 "spec": "A4 80 แกรม", "code": "01/03/3", "location": "ห้องพัสดุ",
                                 "max_stock": "20", "min_stock": "3"})
    with _db() as s:
        it = s.query(MaterialItem).filter_by(name="กระดาษ A4").one()
        mid = it.id
        assert (it.spec, it.code, it.location, it.max_stock) == ("A4 80 แกรม", "01/03/3", "ห้องพัสดุ", 20)
    # รับจากร้าน ราคา 140
    r = env.post(f"/materials/{mid}/txn", data={"kind": "in", "qty": "5", "unit_price": "140", "date": "20/03/2569",
                                                "party": "ร้านเชียงรายไอทีดี", "ref": "รร53/2569"}, follow_redirects=False)
    assert "over" not in r.headers["location"]
    # รับเกินจำนวนอย่างสูง -> เตือน
    r = env.post(f"/materials/{mid}/txn", data={"kind": "in", "qty": "8", "unit_price": "140", "date": "21/03/2569",
                                                "party": "ร้านเชียงรายไอทีดี"}, follow_redirects=False)
    assert r.headers["location"].endswith("?over=23")
    # จ่ายผ่านใบเบิก: จ่ายให้ = ผู้เบิก, ราคา = ราคารับล่าสุด
    with _db() as s:
        req = Requisition(req_no="จ1/2569", requester="ครูสมศรี ใจดี", date=datetime(2026, 3, 22), status="รออนุมัติ")
        s.add(req); s.flush()
        s.add(RequisitionItem(requisition_id=req.id, material_id=mid, name="กระดาษ A4", qty=2, unit="รีม"))
        s.commit(); rid = req.id
    env.post(f"/requisitions/{rid}/issue")
    with _db() as s:
        out = s.query(MaterialTxn).filter_by(material_id=mid, kind="out").one()
        assert out.party == "ครูสมศรี ใจดี" and out.unit_price == 140 and not out.note

    page = env.get(f"/materials/{mid}").text
    for txt in ("รับจาก / จ่ายให้", "ราคาต่อหน่วย", "ร้านเชียงรายไอทีดี", "ครูสมศรี ใจดี", "01/03/3", "ห้องพัสดุ",
                "รร53/2569", "เกินจำนวนอย่างสูง"):
        assert txt in page, txt

    # แก้หัวบัญชี
    env.post(f"/materials/{mid}/update", data={"name": "กระดาษ A4", "unit": "รีม", "category": "วัสดุสำนักงาน",
                                               "spec": "A4 70 แกรม", "code": "01/03/3", "location": "ห้องธุรการ",
                                               "max_stock": "30", "min_stock": "3"})
    with _db() as s:
        it = s.get(MaterialItem, mid)
        assert (it.spec, it.location, it.max_stock) == ("A4 70 แกรม", "ห้องธุรการ", 30)

    # Word ตามแบบ
    import app.database as dbm
    monkeypatch.setattr(dbm, "get_data_dir", lambda: tmp_path)
    import app.services.material_card as mc
    monkeypatch.setattr(mc, "get_data_dir", lambda: tmp_path)
    r = env.get(f"/materials/{mid}/card.docx")
    assert r.status_code == 200
    p = tmp_path / "card.docx"; p.write_bytes(r.content)
    d = Document(str(p))
    text = "\n".join(x.text for x in d.paragraphs)
    assert "บัญชีวัสดุ" in text and "จำนวนอย่างสูง  30" in text and "ที่เก็บ  ห้องธุรการ" in text
    cells = [c.text for row in d.tables[0].rows for c in row.cells]
    assert "ยอดยกมา" in cells and "รับจาก ร้านเชียงรายไอทีดี" in cells and "จ่ายให้ ครูสมศรี ใจดี" in cells
    assert "140.00" in cells and "21" in cells          # คงเหลือหลังจ่าย 23 - 2
    assert env.get("/materials-cards.docx?category=วัสดุสำนักงาน").status_code == 200


def test_backfill_party_from_requisition_and_procurement():
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from app.database import Base, _backfill_material_party
    from app import models as m
    eng = create_engine("sqlite://")
    Base.metadata.create_all(eng)
    with Session(eng) as s:
        v = m.Vendor(name="ร้านศึกษาภัณฑ์"); s.add(v); s.flush()
        p = m.Procurement(subject="ซื้อวัสดุ", fiscal_year=2569, memo_no="12/2569", vendor_id=v.id); s.add(p)
        it = m.MaterialItem(name="ปากกา", unit="ด้าม"); s.add(it); s.flush()
        req = m.Requisition(req_no="1", requester="ครูเอ"); s.add(req); s.flush()
        s.add_all([m.MaterialTxn(material_id=it.id, kind="in", qty=5, ref="เรื่องจัดซื้อ 12/2569"),
                   m.MaterialTxn(material_id=it.id, kind="out", qty=1, requisition_id=req.id, note="เบิกโดย ครูเอ")])
        s.commit()
    _backfill_material_party(eng)
    with Session(eng) as s:
        rows = {t.kind: t for t in s.query(m.MaterialTxn)}
        assert rows["in"].party == "ร้านศึกษาภัณฑ์"
        assert rows["out"].party == "ครูเอ" and rows["out"].note == ""


def test_sheets_split_with_carry_and_full_blank_rows(env):
    """วัสดุ 1 รายการ = 1 ชุดแผ่น แผ่นละ ROWS_PER_SHEET แถว แผ่นต่อไปมีแถวยกมา + เติมบรรทัดว่างจนเต็มแผ่น"""
    from datetime import datetime, timedelta
    from types import SimpleNamespace as N
    from docx import Document
    from app.services import material_card as mc
    R = mc.ROWS_PER_SHEET
    assert [len(c) for c in mc.sheets(list(range(1)))] == [1]
    assert [len(c) for c in mc.sheets(list(range(R)))] == [R]
    assert [len(c) for c in mc.sheets(list(range(R + 1)))] == [R, 1]
    assert [len(c) for c in mc.sheets(list(range(R + 2 * (R - 1))))] == [R, R - 1, R - 1]

    txns = [N(id=i, kind="in" if i % 2 == 0 else "out", qty=2 if i % 2 == 0 else 1, date=datetime(2026, 5, 1) + timedelta(days=i),
              unit_price=10, party="ร้าน", ref="", note="") for i in range(R + 3)]
    item = N(name="ปากกา", category="", code="", spec="", unit="ด้าม", location="", max_stock=0, min_stock=0, txns=txns)
    path = mc.render_material_cards(N(name="ทดสอบ", affiliation=""), [item])
    doc = Document(path)
    assert len(doc.tables) == 2
    body = [len(t.rows) - 2 for t in doc.tables]          # ไม่นับหัวตาราง 2 แถว
    assert body == [R, R]                                 # เต็มแผ่นทั้งสองแผ่น (แผ่นสองเติมบรรทัดว่าง)
    second = [c.text for c in doc.tables[1].rows[2].cells]
    assert second[1] == "ยกมาจากแผ่นที่ 1"
    assert any("แผ่นที่  2" in p.text for p in doc.paragraphs)


def test_delete_many_materials(env):
    from app.models import MaterialItem, MaterialTxn
    import app.tenancy as tn
    with tn.session_for(next(iter(tn._engines))) as s:
        a, b, c = (MaterialItem(name=n, unit="ชิ้น") for n in ("ลบหลายก", "ลบหลายข", "เก็บไว้"))
        s.add_all([a, b, c]); s.flush()
        s.add(MaterialTxn(material_id=a.id, kind="in", qty=3, note="ยอดยกมา"))
        s.commit(); ids = (a.id, b.id, c.id)
    r = env.post("/materials/delete-many", data={"material_id": [str(ids[0]), str(ids[1])]}, follow_redirects=False)
    assert "deleted=2" in r.headers["location"]
    with tn.session_for(next(iter(tn._engines))) as s:
        assert s.get(MaterialItem, ids[0]) is None and s.get(MaterialItem, ids[1]) is None
        assert s.get(MaterialItem, ids[2]) is not None
        assert s.query(MaterialTxn).filter_by(material_id=ids[0]).count() == 0
    page = env.get("/materials").text
    assert 'id="delete-selected"' in page and 'class="group-pick"' in page
