# -*- coding: utf-8 -*-
"""สมุดเงินสดรายเดือน (ยอดยกมา = ก่อนวันที่ 1 ของเดือน) + Word/Excel/พิมพ์ PDF ตามเดือนที่เลือก
รัน: .venv\Scripts\python.exe -m pytest tests/test_cashbook_monthly.py
"""
import io
from datetime import datetime

from tests.test_subsidy import env  # noqa: F401  (fixture)

FY = 2570


def _seed():
    import app.tenancy as tn
    from app.models import AccountOpening, FinanceAccount, FinanceTxn
    db = tn.session_for(next(iter(tn._engines)))
    try:
        a = FinanceAccount(name="เงินรายได้สถานศึกษา", deposit_type="bank", fund_type="เงินนอกงบประมาณ")
        b = FinanceAccount(name="เงินอุดหนุน", deposit_type="bank", fund_type="เงินงบประมาณ")
        db.add_all([a, b]); db.flush()
        db.add(AccountOpening(account_id=a.id, fiscal_year=FY, amount=1000))
        db.add_all([
            FinanceTxn(account_id=a.id, kind="in", amount=500, date=datetime(2026, 10, 15), fiscal_year=FY, note="รับเดือนตุลา"),
            FinanceTxn(account_id=a.id, kind="out", amount=200, date=datetime(2026, 10, 31, 15), fiscal_year=FY, note="จ่ายสิ้นตุลา"),
            FinanceTxn(account_id=a.id, kind="in", amount=70, date=datetime(2026, 11, 1), fiscal_year=FY, note="รับต้นพฤศจิกา"),
            FinanceTxn(account_id=b.id, kind="out", amount=30, date=datetime(2026, 11, 30, 23), fiscal_year=FY, note="จ่ายสิ้นพฤศจิกา"),
            FinanceTxn(account_id=a.id, kind="in", amount=9, date=datetime(2026, 12, 1), fiscal_year=FY, note="รับธันวา"),
            FinanceTxn(account_id=a.id, kind="in", amount=4, date=None, fiscal_year=FY, note="ไม่มีวันที่"),
            FinanceTxn(account_id=a.id, kind="in", amount=5, date=datetime(2027, 1, 10), fiscal_year=FY, note="รับมกรา"),
        ])
        db.commit()
        db.query(FinanceTxn).filter_by(note="ไม่มีวันที่").update({"date": None})   # default=now ต้องล้างทีหลัง
        db.commit()
        return a.id
    finally:
        db.close()


def test_month_range_crosses_calendar_year():
    from app.routers.finance import _fy_month_range, _month_options
    assert _fy_month_range(2570, 10)[:2] == (datetime(2026, 10, 1), datetime(2026, 11, 1))
    assert _fy_month_range(2570, 12)[:2] == (datetime(2026, 12, 1), datetime(2027, 1, 1))
    assert _fy_month_range(2570, 1)[2] == "ประจำเดือน มกราคม 2570"
    assert _fy_month_range(2570, 11)[2] == "ประจำเดือน พฤศจิกายน 2569"
    assert [m for m, _ in _month_options(2570)][:3] == [10, 11, 12]


def test_monthly_opening_rows_and_exports(env):
    aid = _seed()
    from app.routers.finance import _cashbook_fund_data
    import app.tenancy as tn
    db = tn.session_for(next(iter(tn._engines)))
    try:
        _s, op, rec, pay = _cashbook_fund_data(db, FY, None, 11)
        assert op["เงินนอกงบประมาณ"] == 1300 and op["เงินงบประมาณ"] == 0, "ยกมา = 1000 + 500 - 200 (ทั้งวัน 31 ต.ค.)"
        assert [r["desc"].split("] ")[-1] for r in rec] == ["รับต้นพฤศจิกา"]
        assert [r["desc"].split("] ")[-1] for r in pay] == ["จ่ายสิ้นพฤศจิกา"], "รวมรายการวันสุดท้ายของเดือน"
        _s, op, rec, pay = _cashbook_fund_data(db, FY, None, 1)
        assert op["เงินนอกงบประมาณ"] == 1379 and op["เงินงบประมาณ"] == -30, "ข้ามปี ค.ศ. ยกมาถูก"
        assert [r["amount"] for r in rec] == [5]
        _s, op, rec, pay = _cashbook_fund_data(db, FY, None, 0)
        assert len(rec) == 5 and op["เงินนอกงบประมาณ"] == 1000, "ทั้งปีเหมือนเดิม (รวมรายการไม่มีวันที่)"
    finally:
        db.close()

    html = env.get(f"/finance/cashbook?year={FY}&month=11").text
    assert "ประจำเดือน พฤศจิกายน 2569" in html and "รับต้นพฤศจิกา" in html and "รับเดือนตุลา" not in html
    assert '<option value="11" selected>เดือน พฤศจิกายน 2569</option>' in html
    assert "มี 1 รายการที่ไม่ได้ระบุวันที่" in html
    assert f'/finance/cashbook.xlsx?year={FY}&amp;month=11' in html and f'/finance/cashbook/print?year={FY}&amp;month=11' in html
    html = env.get(f"/finance/cashbook?year={FY}&account={aid}&month=11").text
    assert f'/finance/cashbook.docx?year={FY}&amp;account={aid}&amp;month=11' in html
    assert "ไม่ได้ระบุวันที่" not in env.get(f"/finance/cashbook?year={FY}").text, "แบบทั้งปีไม่ต้องเตือน"
    assert "ทั้งปีงบประมาณ" in env.get(f"/finance/cashbook?year={FY}&month=99").text, "เดือนผิด = ทั้งปี"

    p = env.get(f"/finance/cashbook/print?year={FY}&month=11").text
    assert "สมุดเงินสด" in p and "ประจำเดือน พฤศจิกายน 2569" in p and "window.print()" in p
    assert "1,300.00" in p and "รับต้นพฤศจิกา" in p and "รับธันวา" not in p
    assert "เงินสดคงเหลือยกไป 1,340.00 บาท" in p

    from openpyxl import load_workbook
    r = env.get(f"/finance/cashbook.xlsx?year={FY}&month=11")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    assert "ประจำเดือน พฤศจิกายน 2569" in ws["A1"].value
    vals = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "รับธันวา" not in " ".join(str(v) for v in vals) and 1300 in vals
    from docx import Document
    r = env.get(f"/finance/cashbook.docx?year={FY}&month=11")
    assert r.status_code == 200
    text = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "ประจำเดือน พฤศจิกายน 2569" in text
