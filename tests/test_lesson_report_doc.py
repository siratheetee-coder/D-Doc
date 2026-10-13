# -*- coding: utf-8 -*-
"""แบบบันทึกเวลาเรียนรายคาบ (Word) หน้าตาเดียวกับแบบรายวัน

เจ้าของขอ: แนวนอน หน้าละเดือน วันไหนสอน 2 คาบ ใส่เลขวันที่ซ้ำได้
รัน: .venv\\Scripts\\python.exe -m pytest tests/test_lesson_report_doc.py
"""
from datetime import date, timedelta

from docx import Document
from docx.enum.section import WD_ORIENT

import app.services.acad_doc as acad_doc
from app.models import AcadLessonMark, School
from app.services import academic_lessons as svc
from tests.test_academic_lessons import db, seed  # noqa: F401  (fixture)


def _render(db, tmp_path, monkeypatch, **kw):
    monkeypatch.setattr(acad_doc, 'get_data_dir', lambda: tmp_path)
    c, subject, student = seed(db)
    return c, subject, student, lambda: Document(acad_doc.render_lesson_attendance(School(name='โรงเรียนทดสอบ'),
                                                                                    c, subject, db, **kw))


def _day_header(table):
    return [cell.text for cell in table.rows[1].cells[2:] if cell.text.isdigit()]   # ตัดช่องหัวที่ผสานแนวตั้ง (มา ป่วย ...)


def test_landscape_one_page_per_month_and_repeated_dates(db, tmp_path, monkeypatch):
    c, subject, student, make = _render(db, tmp_path, monkeypatch, term=1)
    for d in (18, 19):
        svc.add_lesson(db, c, subject, 1, date(2026, 5, d), 'คาบ 1', 60)
    svc.add_lesson(db, c, subject, 1, date(2026, 5, 19), 'คาบ 2', 60)
    svc.add_lesson(db, c, subject, 1, date(2026, 6, 1), 'คาบ 1', 60)
    db.commit()
    rows = svc.lessons_for(db, c.id, subject.id)
    db.add(AcadLessonMark(lesson_id=rows[0].id, acad_student_id=student.id, mark='/'))
    db.add(AcadLessonMark(lesson_id=rows[1].id, acad_student_id=student.id, mark='ข'))
    db.commit()
    doc = make()
    assert doc.sections[0].orientation == WD_ORIENT.LANDSCAPE
    texts = [p.text for p in doc.paragraphs]
    assert any('ประจำเดือนพฤษภาคม' in t for t in texts) and any('ประจำเดือนมิถุนายน' in t for t in texts)
    may, june, summary = doc.tables
    assert _day_header(may)[:3] == ['18', '19', '19'], 'วันที่สอน 2 คาบ เลขวันซ้ำ 2 ช่อง'
    assert [c.text for c in may.rows[0].cells[-5:]] == ['มา', 'ป่วย', 'ลา', 'ขาด', 'ร้อยละ']
    assert [c.text for c in may.rows[2].cells[-5:]] == ['1', '0', '0', '1', '33.3']
    assert may.rows[-1].cells[0].text == 'รวมมาเรียน (คน/คาบ)'
    assert summary.rows[0].cells[1].text == 'ชื่อ-นามสกุล' and summary.rows[1].cells[2].text == '1'


def test_single_month_has_no_summary_page(db, tmp_path, monkeypatch):
    c, subject, _, make = _render(db, tmp_path, monkeypatch, month=5)
    svc.add_lesson(db, c, subject, 1, date(2026, 5, 18), 'คาบ 1', 60)
    svc.add_lesson(db, c, subject, 1, date(2026, 6, 1), 'คาบ 1', 60)
    db.commit()
    doc = make()
    assert len(doc.tables) == 1 and _day_header(doc.tables[0]) == ['18']


def test_busy_month_splits_and_counts_the_whole_month_on_the_last_page(db, tmp_path, monkeypatch):
    c, subject, student, make = _render(db, tmp_path, monkeypatch, month=5)
    day = date(2026, 5, 4)
    while day.month == 5:
        if day.weekday() < 5:
            for slot in ('คาบ 1', 'คาบ 2'):
                svc.add_lesson(db, c, subject, 1, day, slot, 50)
        day += timedelta(days=1)
    db.commit()
    for r in svc.lessons_for(db, c.id, subject.id):
        db.add(AcadLessonMark(lesson_id=r.id, acad_student_id=student.id, mark='/'))
    db.commit()
    doc = make()
    first, second = doc.tables
    assert len(_day_header(first)) == acad_doc._LESSON_COLS and first.rows[0].cells[-1].text != 'ร้อยละ'
    assert any('(ต่อ)' in p.text for p in doc.paragraphs)
    assert second.rows[2].cells[-5].text == '40', 'หน้าสุดท้ายนับทั้งเดือน 20 วัน × 2 คาบ'
