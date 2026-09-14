import ast
import asyncio
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace as NS
from unittest.mock import patch

from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy import MetaData, Table, Column
from sqlalchemy.orm import Session
from starlette.datastructures import FormData
from docx import Document
from app.database import Base, init_school_db
from app.models import TextBook, TextbookPurchase, Student, AcadClass, School
from app.services.textbook_people import read_people, read_parties, PARTY_LABELS
from app.services.textbook_selection import load_selection, parse_selection, selection_groups
from app.services.textbook_catalog import parse_results
from app.thai_utils import SCHOOL_LEVELS, parse_be_date


def routes():
    source = ast.parse((Path(__file__).parents[1] / 'app/routers/textbooks.py').read_text(encoding='utf-8'))
    names = {'_purchase_for', '_book_groups', '_student_counts', '_estimate_rows', '_survey_groups',
             '_legacy_selection', 'book_purchase_save', 'selected_to_register'}
    nodes = [node for node in source.body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and node.name in names]
    for node in nodes:
        node.decorator_list = []
    from fastapi import Request, Depends, Form
    scope = dict(globals(), Request=Request, Depends=Depends, Form=Form, get_db=lambda: None,
                 current_academic_year=lambda:2569, _to_int=lambda x,d:int(x) if x else d,
                 _to_float=lambda x,d:float(x) if x else d)
    exec(compile(ast.Module(body=nodes,type_ignores=[]), 'textbook_routes', 'exec'),scope)
    return scope


class TextbookWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.engine=create_engine('sqlite://')
        Base.metadata.create_all(self.engine)
        self.db=Session(self.engine)
        self.routes=routes()

    def tearDown(self):
        self.db.close(); self.engine.dispose()

    def save(self, fields):
        class Request:
            headers={'X-Requested-With':'fetch'}
            async def form(self, **kwargs): return FormData(fields)
        return asyncio.run(self.routes['book_purchase_save'](Request(),self.db))

    def test_dynamic_members_and_combined_parties(self):
        form={f'm{i}_name':f'ครู {i}' for i in range(1,22)}
        self.assertEqual(len(read_people(form,'m')),21)
        form={'parties_editor':'combined','party1_name':'ผู้ปกครอง','party1_kind':'parent',
              'party9_name':'นักเรียน','party9_kind':'student'}
        self.assertEqual(read_parties(form)['parent'][0]['name'],'ผู้ปกครอง')
        self.assertEqual(read_parties(form)['student'][0]['name'],'นักเรียน')
        self.assertEqual(read_parties({'pt8_name':'ครูเดิม'})['teacher'][0]['name'],'ครูเดิม')

    def test_existing_school_migration_twice(self):
        engine=create_engine('sqlite://')
        metadata=MetaData()
        for model, exclude in [(TextBook,{'selection_key'}),(TextbookPurchase,{'selection_items','contact_phone'})]:
            Table(model.__tablename__, metadata, *[Column(c.name,c.type,primary_key=c.primary_key,nullable=c.nullable)
                  for c in model.__table__.columns if c.name not in exclude])
        metadata.create_all(engine)
        with engine.begin() as conn:
            conn.exec_driver_sql("INSERT INTO textbook (id,year,title,qty_received) VALUES (1,2569,'old',5)")
        init_school_db(engine);init_school_db(engine)
        with Session(engine) as db:
            book=db.get(TextBook,1)
            self.assertEqual(book.qty_received,5)
            self.assertIsNone(book.selection_key)
        engine.dispose()

    def test_selection_save_and_explicit_idempotent_transfer(self):
        fields={'year':'2569','selection_editor':'1','book1_title':'เลือกแล้ว','book1_level':'ป.1',
                'book1_price':'80','book1_qty':'10','book1_selected':'1',
                'book2_title':'ไม่เลือก','book2_level':'ป.2','book2_price':'90','book2_qty':'12'}
        saved=self.save(fields)
        self.assertEqual(self.db.query(TextBook).count(),0)
        self.assertEqual(len(self.routes['_book_groups'](self.db,2569)),1)
        self.assertEqual(len(self.routes['_survey_groups'](self.db,2569)),2)
        fields.update({f'book{i}_key':row['key'] for i,row in enumerate(saved['selection'],1)})
        again=self.save(fields)
        self.assertEqual(saved['selection'],again['selection'])
        transfer=self.routes['selected_to_register']
        self.assertEqual(transfer(self.db,2569),{'added':1,'skipped':0})
        self.assertEqual(transfer(self.db,2569),{'added':0,'skipped':1})
        book=self.db.query(TextBook).one()
        self.assertEqual(book.qty_received,0)
        book.qty_received=7; self.db.commit()
        transfer(self.db,2569)
        self.assertEqual(book.qty_received,7)
        self.assertEqual(transfer(self.db,2570),{'added':0,'skipped':0})

    def test_legacy_register_is_preserved(self):
        book=TextBook(year=2569,title='เดิม',level='ป.1',qty_received=20,unit_price=10)
        self.db.add(book); self.db.commit()
        previous=self.routes['_legacy_selection'](self.db,2569)
        self.save({'year':'2569','selection_editor':'1','book1_key':previous[0]['key'],
                   'book1_title':'เดิม','book1_level':'ป.1','book1_selected':'1','book1_qty':'20'})
        self.assertEqual(self.routes['selected_to_register'](self.db,2569)['added'],0)
        self.assertEqual(self.db.query(TextBook).count(),1)
        self.assertEqual(book.qty_received,20)

    def test_all_school_levels_and_negative_input(self):
        self.db.add_all([AcadClass(year=2569,level='ป.3'),Student(name='นักเรียน',level='ป.2')]); self.db.commit()
        rows=self.routes['_estimate_rows'](self.db,NS(year=2569,rates='{}'),[('ป.1',[{'price':80}])])
        self.assertEqual([row['level'] for row in rows],['ป.1','ป.2','ป.3'])
        with self.assertRaises(ValueError):
            parse_selection({'book1_title':'ผิด','book1_level':'ป.1','book1_price':'nan'})

    def test_catalogue_fixture(self):
        fixture=Path(__file__).with_name('fixtures')/'textbook_catalog.html'
        result=parse_results(fixture.read_text(encoding='utf-8'))
        self.assertEqual(result['items'][0]['title'],'ภาษาพาที')
        self.assertEqual(result['items'][0]['price'],79)
        self.assertEqual(result['items'][0]['level'],'ป.4')
        self.assertEqual(result['pages'],14)
        with self.assertRaises(ValueError): parse_results('<html>unavailable</html>')

    def test_tor_standard_text_and_selected_appendix(self):
        from app.services.book_tor_doc import render_book_tor
        school=School(name='โรงเรียนทดสอบ',address='ที่อยู่ทดสอบ',officer_name='ครูทดสอบ')
        tp=TextbookPurchase(year=2569,fiscal_year=2569,period_text='เมษายน 2569',
                            members=json.dumps([{'name':'ผู้จัดทำคนเดียว'}]),
                            purpose='ข้อความเก่าที่ไม่ใช้',conditions='เงื่อนไขเก่า',contact='ติดต่อเก่า')
        with tempfile.TemporaryDirectory() as directory, patch('app.services.book_tor_doc.get_data_dir',return_value=Path(directory)):
            doc=Document(render_book_tor(school,tp,[('ป.1',[{'title':'เลือกแล้ว','price':80,'qty':10}])]))
            text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            self.assertIn('รายการหนังสือเรียนแนบท้าย TOR',text)
            self.assertIn('เลือกแล้ว',text)
            self.assertIn('800.00',text)
            self.assertNotIn('ข้อความเก่าที่ไม่ใช้',text)
            self.assertIn('( ผู้จัดทำคนเดียว )',text)
            self.assertNotIn('( ครูทดสอบ )',text)
            self.assertEqual(len(doc.tables[0].rows),1)


if __name__=='__main__': unittest.main()
